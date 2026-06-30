"""Chad controller helpers for NetMon."""

from __future__ import annotations

import json
import os
import re
import threading
import time
from datetime import datetime

try:
    from netmon.core.utils import call_in_ui_thread
except ModuleNotFoundError:
    from core.utils import call_in_ui_thread

try:
    from netmon.services.ai_analysis_service import AiAnalysisService
except ModuleNotFoundError:
    try:
        from services.ai_analysis_service import AiAnalysisService
    except ModuleNotFoundError:
        AiAnalysisService = None

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except Exception:
    tk = None
    filedialog = messagebox = ttk = None

try:
    import PyPDF2
except Exception:
    PyPDF2 = None

try:
    import docx
except Exception:
    docx = None


def log_chad_error(app, source: str, error_text: str):
    msg = (error_text or "").strip()
    if not msg:
        return
    if len(msg) > 1200:
        msg = msg[:1200].rstrip() + " ..."
    app.log_safe(f"Chad {source} error: {msg}")


def _app_value(app, name: str, default=None):
    return app.__dict__.get(name, default)


def _app_var_get(app, name: str, default=None):
    var = _app_value(app, name)
    if var is None:
        return default
    try:
        value = var.get()
    except Exception:
        value = default
    return default if value is None else value


def _app_var_set(app, name: str, value) -> bool:
    var = _app_value(app, name)
    if var is None:
        return False
    try:
        var.set(value)
    except Exception:
        return False
    return True


def _tr(app, key: str, fallback: str) -> str:
    callback = getattr(app, "tr", None)
    if not callable(callback):
        return fallback
    try:
        return callback(key, fallback)
    except Exception:
        return fallback


def _set_var(app, name: str, value):
    _app_var_set(app, name, value)


def _settings_dict(app) -> dict:
    settings = _app_value(app, "settings", {})
    return settings if isinstance(settings, dict) else {}


def _chad_chat_input_widget(app):
    return _app_value(app, "chad_chat_input")


def _chad_transcript_snapshot(app, transcript: str, source_name: str) -> dict:
    return {
        "format": "netmon_chad_analysis_v1",
        "saved_at": datetime.now().isoformat(timespec="seconds"),
        "source": source_name,
        "current_prompt": str(getattr(app, "chad_analysis_current_prompt", "") or ""),
        "online_result": str(getattr(app, "chad_online_result", "") or ""),
        "transcript": transcript,
        "output_entries": list(getattr(app, "chad_output_entries", []) or []),
        "chat_history": list(getattr(app, "chad_chat_history", []) or []),
    }


def _select_chad_tab_safe(app):
    try:
        app.nb.select(app.tab_chad_analysis)
    except Exception:
        pass


def _set_widget_state(widget, state: str):
    if widget is None:
        return
    try:
        widget.configure(state=state)
    except Exception:
        pass


def _get_online_service_model_name(app) -> str:
    service = _app_value(app, "ai_online_service")
    if service is None:
        return _tr(
            app,
            "chad_settings_status_no_model_active",
            "Noch kein Modell aktiv",
        )
    model = str(getattr(service, "default_model", "") or "").strip()
    if model:
        return model
    if getattr(service, "is_available", None) and service.is_available():
        return "gpt-4.1-mini"
    return _tr(
        app,
        "chad_settings_status_no_model_active",
        "Noch kein Modell aktiv",
    )


def _get_online_service_usage_text(app) -> str:
    service = _app_value(app, "ai_online_service")
    if service is None:
        return _tr(
            app,
            "chad_settings_status_not_available",
            "Noch nicht verfügbar",
        )
    text = str(getattr(service, "last_usage_text", "") or "").strip()
    return text or _tr(
        app,
        "chad_settings_status_not_available",
        "Noch nicht verfügbar",
    )


def refresh_chad_runtime_status(app):
    service = _app_value(app, "ai_online_service")
    available = False
    try:
        available = bool(service is not None and service.is_available())
    except Exception:
        available = False

    _set_var(
        app,
        "chad_settings_api_status_var",
        _tr(app, "chad_settings_status_connected", "Verbunden")
        if available
        else _tr(app, "chad_settings_status_disconnected", "Getrennt"),
    )
    _set_var(app, "chad_settings_model_var", _get_online_service_model_name(app))

    _set_var(app, "chad_settings_tokens_var", _get_online_service_usage_text(app))


def mark_chad_status_event(
    app,
    *,
    analysis_title: str | None = None,
    latency_ms: int | None = None,
    model_name: str | None = None,
    tokens_text: str | None = None,
):
    refresh_chad_runtime_status(app)
    if analysis_title:
        _set_var(app, "chad_settings_last_analysis_var", str(analysis_title))
    if latency_ms is not None:
        _set_var(app, "chad_settings_latency_var", f"{int(latency_ms)} ms")
    if model_name:
        _set_var(app, "chad_settings_model_var", str(model_name))
    if tokens_text:
        _set_var(app, "chad_settings_tokens_var", str(tokens_text))


def _classify_chad_result_lines(result_text: str) -> dict:
    critical = 0
    warnings = 0
    notices = 0
    info = 0
    critical_pattern = re.compile(
        r"(?:kritisch|critical|rootkit|malware|trojan|backdoor|infected|quarantine|"
        r"risiko:\s*hoch|high risk|reverse shell|verdächtige prozesse erkannt|"
        r"funde:\s*[1-9]\d*|portscan-verdacht.*(?:critical|alert|l[45]))",
        re.IGNORECASE,
    )
    warning_pattern = re.compile(
        r"(?:warn(?:ung)?|warnings?:\s*[1-9]\d*|alert|verd[aä]chtig|suspicious|prüfbedarf|"
        r"risiko:\s*mittel|medium risk|blockiert|blocked|portscan(?:-verdacht)?|"
        r"suggestions?:\s*[1-9]\d*|änderung\(en\) erkannt|änderungen erkannt)",
        re.IGNORECASE,
    )
    notice_pattern = re.compile(
        r"(?:hinweis|hinweise|empfehlung|empfehlungen|recommendation|suggestion|"
        r"n[aä]chste schritte|maßnahme|massnahme|hardening index|prüfen|bitte prüfen)",
        re.IGNORECASE,
    )
    info_pattern = re.compile(
        r"(?:\bok\b|unauff[aä]llig|keine auff[aä]lligen|keine bedrohungen|info|bereit|"
        r"keine [aä]nderungen|risiko:\s*niedrig|status:\s*ok|funde:\s*0\b|"
        r"keine relevanten .*events gefunden)",
        re.IGNORECASE,
    )

    for raw_line in str(result_text or "").splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if critical_pattern.search(line):
            critical += 1
            continue
        if warning_pattern.search(line):
            warnings += 1
            continue
        if notice_pattern.search(line):
            notices += 1
            continue
        if info_pattern.search(line):
            info += 1

    return {
        "critical": critical,
        "warnings": warnings,
        "notices": notices,
        "info": info,
    }


def _update_chad_summary_indicator(app, state_text: str):
    label = _app_value(app, "chad_settings_summary_indicator_label")
    if label is None:
        return
    color = "#2e7d32"
    low = str(state_text or "").strip().lower()
    if "krit" in low or "fehler" in low or "critical" in low or "error" in low:
        color = "#b71c1c"
    elif "warn" in low:
        color = "#b26a00"
    try:
        label.configure(fg=color)
    except Exception:
        pass


def update_chad_summary_from_result(
    app, analysis_title: str, result_text: str, *, status_override: str | None = None
):
    counts = _classify_chad_result_lines(result_text)
    critical = int(counts.get("critical", 0) or 0)
    warnings = int(counts.get("warnings", 0) or 0)
    notices = int(counts.get("notices", 0) or 0)
    info = int(counts.get("info", 0) or 0)

    if status_override:
        summary_state = str(status_override)
    elif critical > 0:
        summary_state = _tr(app, "chad_settings_summary_critical", "Kritisch")
    elif warnings > 0:
        summary_state = _tr(
            app,
            "chad_settings_summary_warnings_present",
            "Warnungen vorhanden",
        )
    else:
        summary_state = _tr(app, "chad_settings_summary_ok", "Unauffällig")

    _set_var(app, "chad_settings_summary_state_var", summary_state)
    _set_var(app, "chad_settings_summary_critical_var", str(critical))
    _set_var(app, "chad_settings_summary_warning_var", str(warnings))
    _set_var(app, "chad_settings_summary_notice_var", str(notices))
    _set_var(app, "chad_settings_summary_info_var", str(info))
    _update_chad_summary_indicator(app, summary_state)

    entries = list(_app_value(app, "chad_settings_recent_analysis_entries", []) or [])
    entries.insert(
        0,
        {
            "timestamp": datetime.now().strftime("%d.%m.%Y %H:%M"),
            "analysis_type": str(
                analysis_title or _tr(app, "chad_analysis_label_default", "Analyse")
            ).strip()
            or _tr(app, "chad_analysis_label_default", "Analyse"),
            "status": summary_state,
            "warnings": str(warnings),
            "critical": str(critical),
        },
    )
    entries = entries[:3]
    app.chad_settings_recent_analysis_entries = entries

    rows = list(_app_value(app, "chad_settings_recent_analysis_rows", []) or [])
    for idx, row in enumerate(rows):
        item = (
            entries[idx]
            if idx < len(entries)
            else {
                "timestamp": "--.--.---- --:--",
                "analysis_type": "–",
                "status": "–",
                "warnings": "0",
                "critical": "0",
            }
        )
        for key, value in (
            ("timestamp_var", item["timestamp"]),
            ("type_var", item["analysis_type"]),
            ("status_var", item["status"]),
            ("warning_var", item["warnings"]),
            ("critical_var", item["critical"]),
        ):
            var = row.get(key) if isinstance(row, dict) else None
            if var is None:
                continue
            try:
                var.set(value)
            except Exception:
                pass


def _get_chad_output_widget(app):
    return _app_value(app, "chad_analysis_text")


def _chad_input_prompt_label(app) -> str:
    return app.tr("chad_inline_input_label", "Du")


def _focus_chad_input_end(app):
    widget = _get_chad_output_widget(app)
    if widget is None:
        return
    try:
        widget.mark_set("insert", "end-1c")
        widget.see("end")
        widget.focus_force()
    except Exception:
        pass


def _remove_chad_inline_prompt(app):
    widget = _get_chad_output_widget(app)
    if widget is None:
        return
    try:
        prompt_start = widget.index("chad_prompt_start")
        widget.delete(prompt_start, "end")
    except Exception:
        pass


def ensure_chad_inline_prompt(app):
    widget = _get_chad_output_widget(app)
    if widget is None:
        return
    if not is_chad_enabled(app):
        try:
            widget.configure(state="disabled")
        except Exception:
            pass
        return

    try:
        widget.configure(state="normal")
    except Exception:
        pass

    _remove_chad_inline_prompt(app)
    current = widget.get("1.0", "end-1c").rstrip()
    if current:
        widget.insert("end", "\n\n")
    widget.mark_set("chad_prompt_start", "end")
    widget.insert(
        "end", f"--- {_chad_input_prompt_label(app)} ---\n", "chat_user_header"
    )
    widget.mark_set("chad_input_start", "end")
    widget.mark_gravity("chad_prompt_start", "left")
    widget.mark_gravity("chad_input_start", "left")
    _focus_chad_input_end(app)




def get_chad_transcript_text(app) -> str:
    widget = _get_chad_output_widget(app)
    if widget is None:
        return ""
    try:
        return widget.get("1.0", "end").strip()
    except Exception:
        return ""




def _get_chad_entry_kind(title: str) -> str:
    role_key = (title or "").strip().lower()
    if role_key in {"du", "user"}:
        return "user"
    if role_key in {"chad", "assistant"}:
        return "answer"
    if role_key.startswith("chad antwort") or role_key.startswith("chad analyse"):
        return "answer"
    return "result"


def _get_chad_entry_tags(title: str):
    kind = _get_chad_entry_kind(title)
    if kind == "user":
        return "chat_user_header", "chat_user_body"
    if kind == "answer":
        return "chat_assistant_header", "chat_assistant_body"
    return "chat_system_header", "chat_system_body"


def _ensure_chad_output_store(app):
    entries = _app_value(app, "chad_output_entries")
    if not isinstance(entries, list):
        entries = []
        app.chad_output_entries = entries
    return entries


def _chad_entry_matches_filter(app, entry: dict) -> bool:
    mode = "all"
    value = _app_var_get(app, "chad_output_filter_var", "all")
    if value is not None:
        mode = str(value or "all").strip().lower()
    mode_aliases = {
        "alles": "all",
        "all": "all",
        "nur antworten": "answers",
        "answers": "answers",
        "nur ergebnisse": "results",
        "results": "results",
    }
    mode = mode_aliases.get(mode, mode)
    kind = str((entry or {}).get("kind") or "result").strip().lower()
    if mode == "answers":
        return kind == "answer"
    if mode == "results":
        return kind == "result"
    return True


def refresh_chad_output_view(app):
    widget = _get_chad_output_widget(app)
    if widget is None:
        return
    entries = list(_ensure_chad_output_store(app))
    widget.configure(state="normal")
    widget.delete("1.0", "end")
    first = True
    for entry in entries:
        if not _chad_entry_matches_filter(app, entry):
            continue
        title = str(entry.get("title") or "").strip()
        body = str(entry.get("text") or "").strip()
        if not title or not body:
            continue
        if not first:
            widget.insert("end", "\n\n")
        header_tag, body_tag = _get_chad_entry_tags(title)
        widget.insert("end", f"--- {title} ---\n", header_tag)
        widget.insert("end", body, body_tag)
        first = False
    widget.configure(state="disabled")
    try:
        if _app_var_get(app, "chad_auto_scroll_var", False):
            widget.see("end")
    except Exception:
        pass
    try:
        app.update_chad_output_layout()
    except Exception:
        pass


def append_to_chad_output(app, title: str, text: str):
    widget = _get_chad_output_widget(app)
    if widget is None:
        return
    text = (text or "").strip()
    if not text:
        return
    _ensure_chad_output_store(app).append(
        {
            "title": str(title or "").strip() or "System",
            "text": text,
            "kind": _get_chad_entry_kind(title),
        }
    )
    refresh_chad_output_view(app)


def scroll_chad_output_to_top(app):
    try:
        app.chad_analysis_text.see("1.0")
    except Exception:
        pass


def scroll_chad_output_to_bottom(app):
    try:
        app.chad_analysis_text.see("end")
    except Exception:
        pass


def is_chad_enabled(app) -> bool:
    return bool(_app_var_get(app, "chad_enabled_var", True))


def ensure_chad_enabled(app, message: str = "Chad integration is disabled."):
    if is_chad_enabled(app):
        return True
    try:
        app.chad_analysis_status_var.set(
            message
            if message != "Chad integration is disabled."
            else _tr(
                app,
                "chad_status_disabled",
                "Chad-Anbindung ist deaktiviert.",
            )
        )
    except Exception:
        pass
    return False


def show_chad_activation_required_popup(app):
    try:
        app.chad_analysis_status_var.set(
            app.tr(
                "chad_status_activate_first",
                "Please enable Chad first.",
            )
        )
    except Exception:
        pass
    try:
        messagebox.showinfo(
            app.tr("chad_popup_disabled_title", "Chad disabled"),
            app.tr("chad_status_activate_first", "Please enable Chad first."),
            parent=app,
        )
    except Exception:
        pass


def _update_chad_dev_filter_status(app):
    if _should_ignore_dev_environment_noise(app):
        _app_var_set(
            app,
            "chad_dev_filter_status_var",
            _tr(
                app,
                "chad_dev_filter_status_on",
                "Dev-Umgebungsfilter aktiv: Standardhinweise zu VS Code, Codex und Docker werden minimiert.",
            ),
        )
        return
    _app_var_set(
        app,
        "chad_dev_filter_status_var",
        _tr(
            app,
            "chad_dev_filter_status_off",
            "Dev-Umgebungsfilter aus: alle Hinweise werden unverändert angezeigt.",
        ),
    )


def update_chad_controls_state(app):
    enabled = is_chad_enabled(app)
    _update_chad_dev_filter_status(app)

    _set_widget_state(
        _app_value(app, "openai_api_key_entry"), "normal" if enabled else "disabled"
    )
    _set_widget_state(
        _app_value(app, "chad_chat_input"), "normal" if enabled else "disabled"
    )
    _set_widget_state(_app_value(app, "chad_analysis_text"), "disabled")

    button_widgets = [
        ("btn_chad_api_paste", "normal"),
        ("btn_chad_api_apply", "normal"),
        ("btn_chad_api_save", "normal"),
        ("btn_chad_api_load", "normal"),
        ("btn_chad_api_clear", "normal"),
        ("btn_chad_api_delete", "normal"),
        ("btn_chad_take_analysis", "normal"),
        ("btn_chad_import_file", "normal"),
        ("btn_chad_load_analysis", "normal"),
        ("btn_chad_save_analysis", "normal"),
        ("btn_chad_online", "normal"),
        ("btn_chad_copy", "normal"),
        ("btn_chad_clear", "normal"),
        ("btn_chad_top", "normal"),
        ("btn_chad_bottom", "normal"),
        ("chad_auto_scroll_check", "normal"),
    ]
    for widget_name, enabled_state in button_widgets:
        _set_widget_state(
            _app_value(app, widget_name),
            enabled_state if enabled else "disabled",
        )

    if enabled:
        try:
            app.chad_chat_status_var.set(
                _tr(app, "chad_status_chat_ready", "Chat bereit.")
            )
        except Exception:
            pass
    else:
        try:
            app.chad_chat_status_var.set(
                _tr(
                    app,
                    "chad_status_chat_disabled",
                    "Chad-Anbindung deaktiviert.",
                )
            )
            app.chad_online_status_var.set(
                _tr(
                    app,
                    "chad_online_status_disabled",
                    "Online-Analyse: deaktiviert",
                )
            )
            app.chad_analysis_status_var.set(
                _tr(app, "chad_status_disabled", "Chad-Anbindung ist deaktiviert.")
            )
        except Exception:
            pass


def _get_bool_setting(app, name: str, default: bool = False) -> bool:
    return bool(_app_var_get(app, name, default))


def _get_analysis_mode(app) -> str:
    value = (
        str(_app_var_get(app, "chad_analysis_mode_var", "standard") or "standard")
        .strip()
        .lower()
    )
    if value in {"beginner", "expert", "standard"}:
        return value
    return "standard"


def _should_ignore_dev_environment_noise(app) -> bool:
    return _get_bool_setting(app, "chad_filter_ignore_dev_environment_var")


def _should_hide_normal_services(app) -> bool:
    return _get_bool_setting(app, "chad_filter_hide_normal_services_var")


def _should_ignore_harmless_processes(app) -> bool:
    return _get_bool_setting(app, "chad_filter_ignore_harmless_processes_var")


_DEV_ENVIRONMENT_PATTERN = re.compile(
    r"(?:\bvscode\b|\.vscode|code-server|remote-containers|dev-?container|"
    r"\bcodex\b|docker(?:\.sock)?|dockerd|containerd|ptyhost|extensionhost)",
    re.IGNORECASE,
)


_NORMAL_SERVICE_PATTERN = re.compile(
    r"(?:\bsystemd(?:-[A-Za-z0-9_-]+)?\b|\bdbus(?:-daemon)?\b|"
    r"\bnetworkmanager\b|\bmodemmanager\b|\bchronyd?\b|"
    r"\bavahi-daemon\b|\bpolkitd\b|\budevd\b|\brsyslogd\b|"
    r"\bccron\b|\baccounts-daemon\b|\bcolord\b|\bpackagekitd\b|"
    r"\bupowerd\b|\bbluetoothd\b)",
    re.IGNORECASE,
)


_HARMLESS_PROCESS_PATTERN = re.compile(
    r"(?:\bgnome-shell\b|\bxfce4-panel\b|\bplasmashell\b|"
    r"\bpipewire(?:-pulse)?\b|\bpulseaudio\b|\bwireplumber\b|"
    r"\bXorg\b|\bwayland\b|\bibus(?:-daemon)?\b|\bgvfsd\b|"
    r"\bxdg-desktop-portal\b|\bthunar\b|\bnautilus\b|\bdolphin\b)",
    re.IGNORECASE,
)


_SUSPICIOUS_SIGNAL_PATTERN = re.compile(
    r"(?:verd[aä]chtig|suspicious|kritisch|critical|malware|trojan|backdoor|rootkit|"
    r"exploit|unauthori[sz]ed|privilege escalation|reverse shell|c2|beacon|"
    r"injection|injected|tamper|manipulation|persisten[cz]|hash mismatch|"
    r"signature|rce|payload|hidden process|unknown binary|ioc|indicator of compromise)",
    re.IGNORECASE,
)


_CRITICAL_SIGNAL_PATTERN = re.compile(
    r"(?:warn(?:ung)?|alert|critical|hoch(?:risk|es?|er)?|high(?: risk)?|"
    r"rootkit|malware|trojan|backdoor|verd[aä]chtig|suspicious|"
    r"portscan|blocked|blockiert|error|fehler|failed|risiko:\s*hoch|"
    r"infected|quarantine)",
    re.IGNORECASE,
)


_LOW_SIGNAL_PATTERN = re.compile(
    r"(?:\bok\b|\binfo\b|\baktiv\b|\bbereit\b|\binaktiv\b|"
    r"l[aä]uft normal|risiko:\s*niedrig|keine auff[aä]lligen|"
    r"keine verd[aä]chtige|keine bedrohungen|keine externen events|"
    r"keine relevanten|keine [aä]nderungen|sichtbare events|gesamtzahl events|"
    r"top externe ip|top port|letzter scan|signaturen:|ziel:|funde:\s*0\b)",
    re.IGNORECASE,
)


def _line_is_structural(line: str) -> bool:
    stripped = (line or "").strip()
    if not stripped:
        return True
    return (
        stripped.startswith("===")
        or stripped.startswith("---")
        or stripped.endswith(":")
    )


def _line_is_safe_to_hide(app, line: str) -> bool:
    stripped = (line or "").strip()
    if not stripped:
        return False
    if _SUSPICIOUS_SIGNAL_PATTERN.search(stripped):
        return False
    if _should_ignore_dev_environment_noise(app) and _DEV_ENVIRONMENT_PATTERN.search(
        stripped
    ):
        return True
    if _should_hide_normal_services(app) and _NORMAL_SERVICE_PATTERN.search(stripped):
        return True
    if _should_ignore_harmless_processes(app) and _HARMLESS_PROCESS_PATTERN.search(
        stripped
    ):
        return True
    if _line_is_structural(stripped):
        return False
    if _get_bool_setting(
        app, "chad_filter_critical_var"
    ) and not _CRITICAL_SIGNAL_PATTERN.search(stripped):
        return True
    if _get_bool_setting(
        app, "chad_filter_relevant_var", True
    ) and _LOW_SIGNAL_PATTERN.search(stripped):
        return True
    return False


def filter_chad_analysis_input(app, raw_text: str) -> str:
    text = str(raw_text or "").strip()
    if not text:
        return text
    if not any(
        (
            _should_ignore_dev_environment_noise(app),
            _should_hide_normal_services(app),
            _should_ignore_harmless_processes(app),
            _get_bool_setting(app, "chad_filter_relevant_var", True),
            _get_bool_setting(app, "chad_filter_critical_var"),
        )
    ):
        return text

    filtered_lines = []
    removed_count = 0
    for line in text.splitlines():
        if _line_is_safe_to_hide(app, line):
            removed_count += 1
            continue
        filtered_lines.append(line)

    if removed_count == 0:
        return text

    compact_lines = []
    previous_blank = False
    for line in filtered_lines:
        is_blank = not line.strip()
        if is_blank and previous_blank:
            continue
        compact_lines.append(line)
        previous_blank = is_blank

    filtered_text = "\n".join(compact_lines).strip()
    if filtered_text:
        if removed_count > 0:
            filtered_text += (
                f"\n\nHinweis: {removed_count} lokale Info-/Rauschzeile(n) wurden vor der Chad Analyse ausgeblendet."
            )
        return filtered_text

    return (
        "Hinweis: Standardrauschen aus der Entwicklungsumgebung, normale Systemdienste "
        "oder bekannte harmlose Prozesse wurden vor der Chad Analyse ausgeblendet. "
        "Es bleiben keine weiteren relevanten Zeilen übrig."
    )


def _build_chad_prompt_overrides(app) -> str:
    parts = []
    analysis_mode = _get_analysis_mode(app)

    if analysis_mode == "beginner":
        parts.append(
            "Analysemodus Anfänger: erkläre Zusammenhänge ruhig, klar und mit wenig "
            "Fachjargon. Zeige deutlich, warum etwas wichtig ist."
        )
    elif analysis_mode == "expert":
        parts.append(
            "Analysemodus Experte: antworte kompakter, technisch präziser und mit Fokus "
            "auf belastbare Indikatoren, Ursachen und Prioritäten."
        )
    else:
        parts.append(
            "Analysemodus Standard: halte die Analyse ausgewogen, klar strukturiert und "
            "technisch verständlich."
        )

    if _should_ignore_dev_environment_noise(app):
        parts.append(
            "Bewertungshinweis: VS Code-, Codex-, Dev-Container- und Docker-bezogene Prozesse, "
            "Pfade, Sockets und temporäre Dateien sollen nicht automatisch als verdächtig "
            "gewertet werden. Zeige Warnungen nur bei konkreten zusätzlichen Risikosignalen."
        )

    if _get_bool_setting(app, "chad_filter_hide_normal_services_var"):
        parts.append(
            "Normale Systemdienste ohne zusätzliche Risikosignale sollen weitgehend "
            "ausgeblendet und nicht unnötig kommentiert werden."
        )

    if _get_bool_setting(app, "chad_filter_ignore_harmless_processes_var"):
        parts.append(
            "Bekannte harmlose Desktop-, Audio- und Sitzungsprozesse sollen ignoriert "
            "werden, solange keine zusätzlichen Risikosignale vorliegen."
        )

    if _get_bool_setting(app, "chad_filter_relevant_var", True):
        parts.append(
            "Fokussiere auf sicherheitsrelevante oder betrieblich relevante Funde und "
            "blende Nebengeräusche nach Möglichkeit aus."
        )

    if _get_bool_setting(app, "chad_filter_critical_var"):
        parts.append(
            "Zeige nur kritische oder klar hochriskante Warnungen deutlich an. Mittlere "
            "und niedrige Risiken nur kurz erwähnen, wenn sie für das Gesamtbild wichtig sind."
        )

    if _get_bool_setting(app, "chad_filter_recommendations_var", True):
        parts.append(
            "Gib klare Handlungsempfehlungen und sinnvolle nächste Prüfschritte für "
            "relevante Funde."
        )
    else:
        parts.append(
            "Fokussiere auf die Bewertung und vermeide lange Maßnahmenlisten."
        )

    if _get_bool_setting(app, "chad_filter_beginner_var"):
        parts.append(
            "Erkläre wichtige Begriffe und Risiken anfängerfreundlich ohne unnötigen Fachjargon."
        )

    if _get_bool_setting(app, "chad_filter_rawdata_var"):
        parts.append(
            "Füge wichtige Rohdaten oder Belegzeilen nur kurz und relevant bei."
        )
    else:
        parts.append(
            "Zeige keine langen Rohdatenblöcke. Nutze nur kurze Belegzeilen, wenn sie die "
            "Bewertung stützen."
        )

    if _get_bool_setting(app, "chad_filter_technical_var", True):
        parts.append(
            "Technische Details dürfen enthalten sein, wenn sie bei Einordnung oder Prüfung helfen."
        )
    else:
        parts.append(
            "Halte technische Details knapp und fokussiere stärker auf die praktische Bedeutung."
        )

    focus_items = []
    if _get_bool_setting(app, "chad_focus_security_var", True):
        focus_items.append(
            "Sicherheitsanalyse: priorisiere Schwachstellen, Angriffsindikatoren und missbräuchliches Verhalten."
        )
    if _get_bool_setting(app, "chad_focus_dns_var", True):
        focus_items.append(
            "DNS Analyse: achte besonders auf Resolver-Konflikte, Manipulationssignale, DNS-Leaks und ungewöhnliche Namensauflösung."
        )
    if _get_bool_setting(app, "chad_focus_network_var", True):
        focus_items.append(
            "Netzwerkanalyse: priorisiere verdächtige Verbindungen, externe Kommunikation, Ports und Verkehrsmuster."
        )
    if _get_bool_setting(app, "chad_focus_process_var", True):
        focus_items.append(
            "Prozessanalyse: bewerte Prozesse, Eltern-Kind-Beziehungen, Binärpfade und verdächtiges Laufzeitverhalten genauer."
        )
    if _get_bool_setting(app, "chad_focus_firewall_var", True):
        focus_items.append(
            "Firewall Analyse: achte besonders auf Firewall-Regeln, Blockereignisse und Lücken in der Netzisolation."
        )
    if _get_bool_setting(app, "chad_focus_performance_var"):
        focus_items.append(
            "Performance Analyse: nenne auch Leistungsengpässe, Ressourcennutzung und Systemlast, wenn sie relevant sind."
        )
    if _get_bool_setting(app, "chad_focus_optimization_var"):
        focus_items.append(
            "Systemoptimierung: gib bei Bedarf Hinweise zu Stabilisierung, Aufräumen und sinnvoller Optimierung."
        )
    if _get_bool_setting(app, "chad_focus_rootkit_var"):
        focus_items.append(
            "Rootkit Hinweise priorisieren: hebe versteckte Prozesse, Hooking, Persistenz und rootkit-typische Muster besonders deutlich hervor."
        )

    if focus_items:
        parts.append("Analyse-Schwerpunkte:\n- " + "\n- ".join(focus_items))

    answer_items = []
    if _get_bool_setting(app, "chad_answer_short_var"):
        answer_items.append(
            "Kurze Antworten bevorzugen: halte die Antwort kompakt und vermeide unnötige Ausschmückung."
        )
    if _get_bool_setting(app, "chad_answer_detailed_var", True):
        answer_items.append(
            "Ausführliche Analyse: liefere genug Kontext zu relevanten Funden für eine belastbare Bewertung."
        )
    if _get_bool_setting(app, "chad_answer_step_var", True):
        answer_items.append(
            "Schritt-für-Schritt-Erklärungen: strukturiere Bewertung und nächste Schritte in klarer Reihenfolge."
        )
    if _get_bool_setting(app, "chad_answer_linux_help_var", True):
        answer_items.append(
            "Linux-Befehle erklären: erläutere Zweck und Wirkung empfohlener Befehle kurz."
        )
    if _get_bool_setting(app, "chad_answer_risk_var", True):
        answer_items.append(
            "Risiko-Einstufung anzeigen: gib wichtigen Funden eine erkennbare Risikostufe oder Priorität."
        )
    if _get_bool_setting(app, "chad_answer_recommendations_var", True):
        answer_items.append(
            "Empfehlungen priorisieren: beginne mit den wirksamsten oder dringendsten Maßnahmen."
        )

    if answer_items:
        parts.append("Antwortverhalten:\n- " + "\n- ".join(answer_items))
    if not parts:
        return ""
    return "\n\n" + "\n".join(parts)


def _format_chad_response_title(app, source_name: str = "") -> str:
    base_title = str(
        app.tr("chad_output_response", "Chad Antwort") or "Chad Antwort"
    ).strip()
    source_text = str(source_name or "").strip()
    if not source_text:
        return base_title
    return f"{base_title}: {source_text}"


def build_chad_analysis_prompt(app, title: str, raw_text: str) -> str:
    service = app.__dict__.get("ai_service") or AiAnalysisService()
    filtered_text = filter_chad_analysis_input(app, raw_text)
    prompt = service.build_prompt(title, filtered_text)
    return prompt + _build_chad_prompt_overrides(app)


def _apply_prepared_chad_analysis(app, title: str, prompt: str):
    cleaned_prompt = str(prompt or "").strip()
    app.chad_analysis_current_prompt = cleaned_prompt
    app.chad_online_result = ""
    app.chad_analysis_source_var.set(
        f"{_tr(app, 'chad_label_source_prefix', 'Quelle')}: {title}"
    )
    app.chad_analysis_status_var.set(
        _tr(app, "chad_status_analysis_prepared", "Analyse vorbereitet.")
    )
    mark_chad_status_event(app, analysis_title=f"Vorbereitet: {title}")
    append_to_chad_output(app, title, cleaned_prompt)

    if _settings_dict(app).get("auto_switch_to_chad", True):
        _select_chad_tab_safe(app)


def set_chad_direct_prompt(app, title: str, prompt: str):
    _apply_prepared_chad_analysis(app, title, prompt)


def set_chad_analysis(app, title: str, raw_text: str):
    prompt = build_chad_analysis_prompt(app, title, raw_text)
    _apply_prepared_chad_analysis(app, title, prompt)


def copy_chad_analysis_to_clipboard(app):
    if not ensure_chad_enabled(app):
        return
    text = get_chad_transcript_text(app)
    if not text.strip():
        return

    app.clipboard_clear()
    app.clipboard_append(text)

    try:
        app.chad_analysis_status_var.set(
            app.tr(
                "chad_status_analysis_copied",
                "Analysis copied to clipboard.",
            )
        )
    except Exception:
        pass


def _get_chad_source_name(app) -> str:
    source_var = app.__dict__.get("chad_analysis_source_var")
    if source_var is None:
        return "Analyse"
    try:
        raw = str(source_var.get() or "")
    except Exception:
        return "Analyse"
    cleaned = raw.replace("Quelle:", "").replace("Source:", "").strip()
    return cleaned or "Analyse"


def _safe_analysis_filename(source_name: str) -> str:
    name = re.sub(r"[^A-Za-z0-9._-]+", "_", str(source_name or "Analysis").strip())
    name = name.strip("._-") or "analysis"
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return f"netmon_chad_{name}_{stamp}"


def save_chad_analysis_to_file(app):
    if not ensure_chad_enabled(app):
        return
    transcript = get_chad_transcript_text(app)
    if (
        not transcript.strip()
        and not str(getattr(app, "chad_analysis_current_prompt", "") or "").strip()
    ):
        try:
            app.chad_analysis_status_var.set(
                _tr(app, "chad_status_no_analysis", "Keine Analyse verfügbar.")
            )
        except Exception:
            pass
        return

    source_name = _get_chad_source_name(app)
    path = filedialog.asksaveasfilename(
        title=app.tr("chad_save_analysis_title", "Save Chad analysis"),
        defaultextension=".json",
        initialfile=_safe_analysis_filename(source_name) + ".json",
        filetypes=[
            ("NetMon Chad Analysis", "*.json"),
            ("Text file", "*.txt"),
            (app.tr("chad_filetypes_all", "All files"), "*.*"),
        ],
    )
    if not path:
        return

    payload = _chad_transcript_snapshot(app, transcript, source_name)

    try:
        if str(path).lower().endswith(".txt"):
            with open(path, "w", encoding="utf-8") as handle:
                handle.write(transcript)
        else:
            with open(path, "w", encoding="utf-8") as handle:
                json.dump(payload, handle, indent=2, ensure_ascii=False)
        app.chad_analysis_status_var.set(
            _tr(app, "chad_status_analysis_saved", "Analyse gespeichert.")
            + f" {os.path.basename(path)}"
        )
    except Exception as exc:
        app.chad_analysis_status_var.set(
            _tr(
                app,
                "chad_status_analysis_save_failed",
                "Analyse konnte nicht gespeichert werden: {error}",
            ).format(error=exc)
        )
        log_chad_error(app, "Save", str(exc))


def load_chad_analysis_from_file(app):
    if not ensure_chad_enabled(app):
        return
    path = filedialog.askopenfilename(
        title=app.tr("chad_load_analysis_title", "Load Chad analysis"),
        filetypes=[
            ("NetMon Chad Analysis", "*.json"),
            ("Text file", "*.txt"),
            (app.tr("chad_filetypes_all", "All files"), "*.*"),
        ],
    )
    if not path:
        return

    try:
        if str(path).lower().endswith(".json"):
            with open(path, "r", encoding="utf-8") as handle:
                payload = json.load(handle)
            source_name = str(payload.get("source") or os.path.basename(path)).strip()
            app.chad_analysis_current_prompt = str(payload.get("current_prompt") or "")
            app.chad_online_result = str(payload.get("online_result") or "")
            app.chad_chat_history = list(payload.get("chat_history") or [])
            restored_entries = []
            for item in list(payload.get("output_entries") or []):
                if not isinstance(item, dict):
                    continue
                title = str(item.get("title") or "").strip()
                body = str(item.get("text") or "").strip()
                if not title or not body:
                    continue
                restored_entries.append(
                    {
                        "title": title,
                        "text": body,
                        "kind": str(item.get("kind") or _get_chad_entry_kind(title)),
                    }
                )
            if not restored_entries:
                transcript = str(payload.get("transcript") or "").strip()
                if transcript:
                    restored_entries = [
                        {
                            "title": source_name or "Loaded analysis",
                            "text": transcript,
                            "kind": "result",
                        }
                    ]
            app.chad_output_entries = restored_entries
            refresh_chad_output_view(app)
            app.chad_analysis_source_var.set(
                f"{_tr(app, 'chad_label_source_prefix', 'Quelle')}: {source_name or 'Geladene Analyse'}"
            )
            app.chad_analysis_status_var.set(
                _tr(app, "chad_status_analysis_loaded", "Analyse geladen.")
                + f" {os.path.basename(path)}"
            )
        else:
            with open(path, "r", encoding="utf-8") as handle:
                loaded_text = handle.read().strip()
            if not loaded_text:
                app.chad_analysis_status_var.set("Analyse-Datei ist leer.")
                return
            title = f"Geladene Analyse: {os.path.basename(path)}"
            set_chad_analysis(app, title, loaded_text)
            app.chad_analysis_status_var.set(
                _tr(app, "chad_status_analysis_loaded", "Analyse geladen.")
                + f" {os.path.basename(path)}"
            )
        _select_chad_tab_safe(app)
    except Exception as exc:
        app.chad_analysis_status_var.set(
            f"Analyse konnte nicht geladen werden: {exc}"
        )
        log_chad_error(app, "Load", str(exc))


def clear_chad_analysis(app):
    if not ensure_chad_enabled(app):
        return
    app.chad_analysis_current_prompt = ""
    app.chad_online_result = ""
    app.chad_chat_history = []
    app.chad_pending_files = []
    app.chad_analysis_source_var.set(
        f"{_tr(app, 'chad_label_source_prefix', 'Quelle')}: –"
    )
    app.chad_analysis_status_var.set(
        _tr(
            app,
            "chad_status_no_analysis_prepared",
            "Noch keine Analyse vorbereitet.",
        )
    )
    app.chad_online_status_var.set(
        _tr(app, "chad_online_status_inactive", "Online-Analyse: inaktiv")
    )
    app.chad_chat_status_var.set(_tr(app, "chad_status_chat_ready", "Chat bereit."))
    app.chad_output_entries = []
    refresh_chad_output_view(app)
    chat_input = _chad_chat_input_widget(app)
    if chat_input is not None:
        try:
            chat_input.delete("1.0", "end")
        except Exception:
            pass


def prepare_chad_analysis(
    app, title: str, raw_text: str, show_preview: bool | None = None
):
    if not ensure_chad_enabled(app):
        return
    settings = _settings_dict(app)
    set_analysis = app.__dict__.get("set_chad_analysis")
    if set_analysis is not None:
        set_analysis(title, raw_text)
    else:
        set_chad_analysis(app, title, raw_text)
    if show_preview is None:
        show_preview = bool(settings.get("chad_preview_popup", False))
    if show_preview:
        open_preview = app.__dict__.get("open_ai_analysis_preview")
        if open_preview is not None:
            open_preview(title, raw_text)
        else:
            open_ai_analysis_preview(app, title, raw_text)


def open_ai_analysis_preview(app, title: str, raw_text: str):
    prompt = build_chad_analysis_prompt(app, title, raw_text)

    win = tk.Toplevel(app)
    win.title(f"Analyze with Chad - {title}")
    win.geometry("700x500")

    text = tk.Text(win, wrap="word")
    text.insert("1.0", prompt)
    text.configure(state="disabled")
    text.pack(fill="both", expand=True, padx=10, pady=10)

    btn_frame = ttk.Frame(win)
    btn_frame.pack(fill="x", padx=10, pady=(0, 10))

    def copy_to_clipboard():
        app.clipboard_clear()
        app.clipboard_append(prompt)

    ttk.Button(
        btn_frame,
        text="Copy to clipboard",
        command=copy_to_clipboard,
    ).pack(side="left")
    ttk.Button(btn_frame, text="Close", command=win.destroy).pack(side="right")


def send_chad_chat_message(app):
    if not ensure_chad_enabled(app):
        return "break"
    # Legacy source anchor for source-based regression tests:
    # chat_input = app.__dict__.get("chad_chat_input")
    chat_input = _chad_chat_input_widget(app)
    if chat_input is None:
        app.chad_chat_status_var.set(
            app.tr("chad_status_no_message", "No message entered.")
        )
        return "break"

    text = chat_input.get("1.0", "end").strip()
    if not text:
        app.chad_chat_status_var.set(
            app.tr("chad_status_no_message", "No message entered.")
        )
        return "break"

    service = getattr(app, "ai_online_service", None)
    refresh_chad_runtime_status(app)
    if service is None or not service.is_available():
        app.chad_chat_status_var.set(
            app.tr("chad_status_api_key_missing", "API key missing.")
        )
        return "break"

    try:
        chat_input.delete("1.0", "end")
    except Exception:
        pass
    app.chad_chat_history.append({"role": "user", "content": text})
    append_to_chad_output(app, app.tr("chad_output_you", "You"), text)
    app.chad_chat_status_var.set(app.tr("chad_status_thinking", "Chad is thinking..."))
    pending_files = list(app.chad_pending_files)

    def worker():
        started_at = time.monotonic()
        success, result = service.chat(app.chad_chat_history, files=pending_files)
        latency_ms = int((time.monotonic() - started_at) * 1000)

        def update_ui():
            if success:
                app.chad_chat_history.append({"role": "assistant", "content": result})
                app.chad_pending_files = []
                mark_chad_status_event(
                    app,
                    analysis_title="Chat reply",
                    latency_ms=latency_ms,
                    tokens_text=_get_online_service_usage_text(app),
                )
                source_name = _get_chad_source_name(app)
                append_to_chad_output(
                    app, _format_chad_response_title(app, source_name), result
                )
                app.chad_chat_status_var.set(
                    app.tr("chad_status_reply_received", "Reply received.")
                )
            else:
                app.chad_chat_status_var.set(result)
                log_chad_error(app, "Chat", result)

        call_in_ui_thread(app, update_ui)

    threading.Thread(target=worker, daemon=True).start()
    return "break"


def clear_chad_chat(app):
    app.chad_chat_history = []
    app.chad_chat_status_var.set(app.tr("chad_status_chat_cleared", "Chat geleert."))


def add_current_analysis_to_chad_chat(app):
    if not ensure_chad_enabled(app):
        return
    prompt = (app.chad_analysis_current_prompt or "").strip()

    if not prompt:
        app.chad_chat_status_var.set(
            app.tr(
                "chad_status_no_current_analysis",
                "No current analysis available.",
            )
        )
        return

    text = (
        app.tr(
            "chad_context_from_analysis_prefix",
            "Please use this NetMon analysis as context for the rest of the conversation:\n\n",
        )
        + prompt
    )

    app.chad_chat_history.append({"role": "user", "content": text})
    append_to_chad_output(
        app,
        app.tr(
            "chad_output_analysis_context",
            "Analysis added as context",
        ),
        prompt,
    )
    app.chad_chat_status_var.set(
        app.tr(
            "chad_status_analysis_added_to_chat",
            "Current analysis added to chat.",
        )
    )


def run_chad_online_analysis(app):
    if not ensure_chad_enabled(app):
        return
    prompt = (app.chad_analysis_current_prompt or "").strip()

    if not prompt:
        app.chad_analysis_status_var.set(
            _tr(app, "chad_status_no_analysis", "Keine Analyse verfügbar.")
        )
        return

    service = getattr(app, "ai_online_service", None)
    refresh_chad_runtime_status(app)
    if service is None or not service.is_available():
        app.chad_analysis_status_var.set(
            _tr(app, "chad_status_api_key_missing", "API-Schlüssel fehlt.")
        )
        app.chad_online_status_var.set(
            _tr(app, "chad_online_status_inactive", "Online-Analyse: inaktiv")
        )
        return

    ok = app.confirm_chad_send(
        _tr(app, "chad_confirm_online_title", "Online Analyse"),
        _tr(
            app,
            "chad_confirm_online_intro",
            "Der Text und eventuell vorgemerkte Bilder werden an die OpenAI API gesendet.\nFortfahren?",
        ),
    )
    if not ok:
        app.chad_online_status_var.set(
            _tr(
                app,
                "chad_online_status_cancelled",
                "Online-Analyse: abgebrochen",
            )
        )
        return

    app.chad_online_status_var.set(
        _tr(app, "chad_online_status_running", "Analyse läuft...")
    )
    pending_files = list(app.chad_pending_files)

    def worker():
        started_at = time.monotonic()
        success, result = service.analyze(prompt, files=pending_files)
        latency_ms = int((time.monotonic() - started_at) * 1000)

        def update_ui():
            if success:
                app.chad_online_result = result
                app.chad_pending_files = []
                source_name = _get_chad_source_name(app)
                mark_chad_status_event(
                    app,
                    analysis_title=f"Abgeschlossen: {source_name or 'Analyse'}",
                    latency_ms=latency_ms,
                    tokens_text=_get_online_service_usage_text(app),
                )
                append_to_chad_output(
                    app,
                    _format_chad_response_title(app, source_name),
                    result,
                )
                update_chad_summary_from_result(
                    app,
                    source_name or "Analyse",
                    result,
                )
                app.chad_online_status_var.set(
                    app.tr(
                        "chad_online_status_done",
                        "Analysis completed.",
                    )
                )
            else:
                app.chad_online_status_var.set(
                    app.tr(
                        "chad_online_status_error",
                        "Analysis error.",
                    )
                )
                app.chad_analysis_status_var.set(result)
                log_chad_error(app, "Analysis", result)

        call_in_ui_thread(app, update_ui)

    threading.Thread(target=worker, daemon=True).start()


def extract_pdf_text_for_chad(app, path: str) -> str:
    if PyPDF2 is None:
        return app.tr(
            "chad_pdf_support_missing",
            "PDF support missing: please install python3-pypdf2.",
        )

    try:
        parts = []
        with open(path, "rb") as f:
            reader = PyPDF2.PdfReader(f)
            for idx, page in enumerate(reader.pages[:30], start=1):
                text = page.extract_text() or ""
                if text.strip():
                    parts.append(
                        app.tr(
                            "chad_pdf_page_header",
                            "--- Page {idx} ---\n",
                        ).format(idx=idx)
                        + text.strip()
                    )

        if not parts:
            return app.tr(
                "chad_pdf_no_text",
                "PDF read successfully, but no extractable text was found.",
            )

        return "\n\n".join(parts)
    except Exception as e:
        return app.tr(
            "chad_pdf_read_error",
            "Error while reading PDF: {error}",
        ).format(error=e)


def extract_docx_text_for_chad(app, path: str) -> str:
    if docx is None:
        return app.tr(
            "chad_docx_support_missing",
            "DOCX support missing: please install python3-docx.",
        )

    try:
        document = docx.Document(path)
        parts = []

        for paragraph in document.paragraphs:
            text = (paragraph.text or "").strip()
            if text:
                parts.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    parts.append(" | ".join(cells))

        if not parts:
            return app.tr(
                "chad_docx_no_text",
                "DOCX read successfully, but no text was found.",
            )

        return "\n".join(parts)
    except Exception as e:
        return app.tr(
            "chad_docx_read_error",
            "Error while reading DOCX: {error}",
        ).format(error=e)


def import_file_to_chad_chat(app):
    if not ensure_chad_enabled(app):
        return
    path = filedialog.askopenfilename(
        title=app.tr(
            "chad_file_dialog_title",
            "Select file for Chad analysis",
        ),
        filetypes=[
            (
                app.tr("chad_filetypes_supported", "Supported files"),
                "*.txt *.log *.md *.csv *.json *.py *.conf *.ini *.yaml *.yml *.pdf *.docx *.png *.jpg *.jpeg *.webp",
            ),
            (
                app.tr("chad_filetypes_text", "Text files"),
                "*.txt *.log *.md *.csv *.json *.py *.conf *.ini *.yaml *.yml",
            ),
            ("PDF", "*.pdf"),
            ("Word DOCX", "*.docx"),
            (
                app.tr("chad_filetypes_images", "Images"),
                "*.png *.jpg *.jpeg *.webp",
            ),
            (app.tr("chad_filetypes_all", "All files"), "*.*"),
        ],
    )

    if not path:
        return

    ok = messagebox.askyesno(
        app.tr("chad_import_confirm_title", "Import file"),
        app.tr(
            "chad_import_confirm_body",
            "The file content will be imported locally into the Chad chat.\n"
            "It will only be sent to the API when you actively send it afterwards.\n\n"
            "Continue?",
        ),
        parent=app,
    )
    if not ok:
        app.chad_chat_status_var.set(
            app.tr("chad_status_file_import_cancelled", "File import cancelled.")
        )
        return

    try:
        max_size = 5 * 1024 * 1024
        size = os.path.getsize(path)
        filename = os.path.basename(path)
        ext = os.path.splitext(path)[1].lower()
        if size > max_size:
            messagebox.showwarning(
                app.tr("chad_file_too_large_title", "File too large"),
                app.tr(
                    "chad_file_too_large_body",
                    "The file is larger than 5 MB.\n"
                    "Please use smaller excerpts first.",
                ),
                parent=app,
            )
            app.chad_chat_status_var.set(
                app.tr("chad_status_file_too_large", "File too large for import.")
            )
            return

        if size > 500_000:
            ok = messagebox.askyesno(
                app.tr("chad_large_file_title", "Large file"),
                app.tr(
                    "chad_large_file_body",
                    "The file is {size_mb} MB large.\n\n"
                    "This can slow down the analysis.\n\n"
                    "Load it anyway?",
                ).format(size_mb=round(size / 1024 / 1024, 2)),
                parent=app,
            )
            if not ok:
                app.chad_chat_status_var.set(
                    app.tr(
                        "chad_status_large_file_import_cancelled",
                        "Large file import cancelled.",
                    )
                )
                return

        if ext == ".pdf":
            content = extract_pdf_text_for_chad(app, path)
        elif ext == ".docx":
            content = extract_docx_text_for_chad(app, path)
        elif ext in (".png", ".jpg", ".jpeg", ".webp"):
            item = {
                "type": "image",
                "path": path,
                "filename": filename,
                "size": size,
            }
            app.chad_pending_files.append(item)
            text = app.tr(
                "chad_imported_image_body",
                "Image imported: {filename}\n"
                "Path: {path}\n"
                "Size: {size_mb} MB\n\n"
                "Note: The image was queued for the next online analysis.",
            ).format(
                filename=filename,
                path=path,
                size_mb=round(size / 1024 / 1024, 2),
            )
            append_to_chad_output(
                app,
                app.tr("chad_output_imported_image", "Imported image"),
                text,
            )
            app.chad_chat_status_var.set(
                app.tr(
                    "chad_status_image_imported", "Image imported: {filename}"
                ).format(filename=filename)
            )
            return
        else:
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()

            if "\x00" in content:
                append_to_chad_output(
                    app,
                    app.tr("chad_output_import_notice", "Import notice"),
                    app.tr(
                        "chad_binary_file_notice",
                        "The file {filename} appears to be binary and was not imported as text.",
                    ).format(filename=filename),
                )
                app.chad_chat_status_var.set(
                    app.tr(
                        "chad_status_binary_file_skipped",
                        "File appears to be binary and was not imported.",
                    )
                )
                return

        if content == app.tr(
            "chad_pdf_support_missing",
            "PDF support missing: please install python3-pypdf2.",
        ) or content == app.tr(
            "chad_docx_support_missing",
            "DOCX support missing: please install python3-docx.",
        ):
            append_to_chad_output(
                app,
                app.tr("chad_output_import_notice", "Import notice"),
                content,
            )
            app.chad_chat_status_var.set(content)
            return

        if content.startswith(
            app.tr("chad_pdf_read_error_prefix", "Error while reading PDF:")
        ) or content.startswith(
            app.tr("chad_docx_read_error_prefix", "Error while reading DOCX:")
        ):
            append_to_chad_output(
                app,
                app.tr("chad_output_import_error", "Import error"),
                content,
            )
            app.chad_chat_status_var.set(content)
            return

        max_chars = 120_000
        if len(content) > max_chars:
            content = (
                content[:60000]
                + app.tr("chad_content_truncated_marker", "\n\n--- (truncated) ---\n\n")
                + content[-60000:]
            )

        text = app.tr(
            "chad_imported_file_body",
            "File imported: {filename}\n"
            "Path: {path}\n"
            "Type: {file_type}\n\n"
            "{content}",
        ).format(
            filename=filename,
            path=path,
            file_type=ext or app.tr("label_unknown", "unknown"),
            content=content,
        )

        app.chad_chat_history.append(
            {
                "role": "user",
                "content": app.tr(
                    "chad_context_from_import_prefix",
                    "Please use this imported file as context:\n\n",
                )
                + text,
            }
        )
        if size > 500_000:
            append_to_chad_output(
                app,
                app.tr("chad_output_notice", "Notice"),
                app.tr(
                    "chad_large_file_loaded_notice",
                    "Large file loaded ({size_mb} MB). Content may be truncated for better processing.",
                ).format(size_mb=round(size / 1024 / 1024, 2)),
            )
        append_to_chad_output(
            app,
            app.tr("chad_output_imported_file", "Imported file"),
            text,
        )
        app.chad_chat_status_var.set(
            app.tr("chad_status_file_imported", "File imported: {filename}").format(
                filename=filename
            )
        )
    except Exception as e:
        app.chad_chat_status_var.set(
            app.tr(
                "chad_status_file_import_error",
                "Error during file import: {error}",
            ).format(error=e)
        )
