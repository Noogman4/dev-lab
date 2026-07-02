"""Linux command database controller helpers for NetMon."""

from __future__ import annotations

import csv
import json
import os
import re
import tkinter as tk
from tkinter import ttk

try:
    from tkinter import filedialog, messagebox
except Exception:
    filedialog = messagebox = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

try:
    from netmon.core.utils import safe_write_json
except ModuleNotFoundError:
    from core.utils import safe_write_json


NETMON_REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
LINUX_HELP_EXPORT_FIELDS = [
    "category",
    "name",
    "info",
    "cmd",
    "syntax",
    "example",
    "explanation",
    "root_required",
    "tags",
    "risk",
]
RISK_UI = {
    "safe": ("linux_help_risk_safe", "🟢 Safe", "green"),
    "caution": ("linux_help_risk_caution", "🟡 Admin", "#d8a106"),
    "danger": ("linux_help_risk_danger", "🔴 Dangerous", "red"),
    "unknown": ("linux_help_risk_unknown_label", "⚪ Unknown", "gray"),
}
SOURCE_FIELD_ALIASES = {
    "category": ("category", "kategorie"),
    "name": ("name", "befehl", "command_name", "command name", "title"),
    "info": (
        "info",
        "description",
        "beschreibung",
        "kurzbeschreibung",
        "short_description",
        "short description",
        "summary",
    ),
    "cmd": ("cmd", "command", "kommando", "befehl", "command_line", "command line"),
    "syntax": ("syntax", "usage"),
    "example": ("example", "beispiel", "examples"),
    "explanation": (
        "explanation",
        "ausgabe/erklaerung",
        "erklaerung",
        "details",
        "notes",
    ),
    "warning_text": ("warning_text", "warning", "warnung", "hinweis"),
    "tags": ("tags",),
    "root_required": (
        "root_required",
        "root noetig",
        "root nötig",
        "requires_root",
        "requires root",
        "root_needed",
        "root needed",
    ),
    "risk": ("risk", "risiko", "risk_level", "risk level"),
}
RISK_DIALOG_CHOICES = (
    ("linux_help_risk_safe", "🟢 Safe", "safe"),
    ("linux_help_risk_caution", "🟡 Admin", "caution"),
    ("linux_help_risk_danger", "🔴 Dangerous", "danger"),
)


def _tr(app, key: str, fallback: str) -> str:
    translate = getattr(app, "tr", None)
    if callable(translate):
        try:
            return str(translate(key, fallback))
        except Exception:
            return fallback
    return fallback


def _linux_help_all_categories_label(app) -> str:
    return _tr(app, "linux_help_all_categories", "All categories")


def _linux_help_all_sources_label(app) -> str:
    return _tr(app, "linux_help_source_both", "Both")


def _linux_help_all_risks_label(app) -> str:
    return _tr(app, "linux_help_all_risks", "All risks")


def _risk_label(app, risk_key: str) -> str:
    key, fallback, _color = RISK_UI.get(risk_key, RISK_UI["unknown"])
    return _tr(app, key, fallback)


def _source_label(app, source_key: str) -> str:
    mapping = {
        "command": ("linux_help_source_commands", "Commands"),
        "glossary": ("linux_help_source_glossary", "Glossary"),
    }
    key, fallback = mapping.get(source_key, ("linux_help_source_unknown", "Unknown"))
    return _tr(app, key, fallback)


def _risk_dialog_choices_for_app(app):
    return [
        (_tr(app, key, fallback), risk_key)
        for key, fallback, risk_key in RISK_DIALOG_CHOICES
    ]


def _app_value(app, name: str, default=None):
    return app.__dict__.get(name, default)


def _app_var_get(app, name: str, default=""):
    var = _app_value(app, name)
    if var is None:
        return default
    try:
        value = var.get()
    except Exception:
        value = default
    return value if value is not None else default


def _app_var_set(app, name: str, value) -> bool:
    var = _app_value(app, name)
    if var is None:
        return False
    try:
        var.set(value)
    except Exception:
        return False
    return True


def _set_text_widget_content(widget, text: str):
    if widget is None:
        return
    widget.configure(state="normal")
    widget.delete("1.0", "end")
    if text:
        widget.insert("1.0", text)
    widget.configure(state="disabled")


def set_linux_help_dirty(app, dirty=True):
    app.linux_help_dirty = bool(dirty)


def linux_help_csv_candidates():
    return [
        os.path.join(
            NETMON_REPO_ROOT, "tools", "netmon-linux-befehlsdatenbank-v1.csv"
        ),
        os.path.join(NETMON_REPO_ROOT, "netmon", "linux-befehle.csv"),
        os.path.join(NETMON_REPO_ROOT, "netmon", "linux-befehle.xlsx"),
        os.path.join(NETMON_REPO_ROOT, "netmon", "netmon-linux-befehlsdatenbank-v1.csv"),
        os.path.join(NETMON_REPO_ROOT, "linux-befehle.csv"),
        os.path.join(NETMON_REPO_ROOT, "linux-befehle.xlsx"),
    ]


def linux_help_csv_path():
    for path in linux_help_csv_candidates():
        if path and os.path.isfile(path):
            return path
    return linux_help_csv_candidates()[0]


def linux_help_slugify(text: str) -> str:
    slug = re.sub(r"[^a-zA-Z0-9]+", "_", (text or "").strip().lower()).strip("_")
    return slug or "linux_help"


def linux_help_make_id(raw, existing_ids=None):
    entry = raw or {}
    base = linux_help_slugify(
        entry.get("id")
        or entry.get("name")
        or entry.get("cmd")
        or entry.get("category")
        or "linux_help"
    )
    used = set(existing_ids or [])
    candidate = base
    idx = 2
    while candidate in used:
        candidate = f"{base}_{idx}"
        idx += 1
    return candidate


def linux_help_get_risk_level(cmd: str) -> str:
    text = (cmd or "").strip().lower()
    if not text:
        return "unknown"

    danger_patterns = [
        " rm ",
        "rm ",
        " dd ",
        "dd if=",
        "mkfs",
        " fdisk",
        " parted",
        "chmod ",
        "chown ",
        "systemctl restart",
        "systemctl stop",
        "systemctl disable",
        "apt remove",
        "apt purge",
        "kill -9",
        "reboot",
        "shutdown",
        "userdel",
        "groupdel",
    ]
    caution_patterns = [
        "sudo ",
        "apt install",
        "apt update",
        "systemctl status",
        "systemctl start",
        "mount ",
        "umount ",
        "cp ",
        "mv ",
        "tar ",
        "find ",
        "grep ",
        "journalctl",
        "ping ",
        "curl ",
        "wget ",
        "ssh ",
        "ufw ",
        "iptables",
        "docker ",
    ]
    safe_prefixes = (
        "ls",
        "pwd",
        "ip a",
        "ip addr",
        "df -h",
        "free -h",
        "uname -a",
        "whoami",
        "id",
        "cat ",
        "head ",
        "tail ",
        "ps ",
        "top",
        "du ",
        "which ",
        "history",
    )

    padded = f" {text} "
    if any(
        pattern in padded or text.startswith(pattern) for pattern in danger_patterns
    ):
        return "danger"
    if any(pattern in text for pattern in caution_patterns):
        return "caution"
    if text.startswith(safe_prefixes):
        return "safe"
    return "caution"


def linux_help_requires_root(cmd: str, risk: str) -> bool:
    text = (cmd or "").strip().lower()
    if text.startswith("sudo "):
        return True
    if risk == "danger":
        return True
    root_patterns = (
        "apt ",
        "systemctl ",
        "mount ",
        "umount ",
        "ufw ",
        "iptables",
        "mkfs",
        "fdisk",
        "parted",
        "chown ",
        "chmod ",
        "journalctl -u",
        "useradd",
        "usermod",
        "userdel",
    )
    return any(pattern in text for pattern in root_patterns)


def linux_help_tag_text(value) -> str:
    if isinstance(value, (list, tuple, set)):
        parts = [str(part).strip() for part in value if str(part).strip()]
    else:
        text = str(value or "").strip()
        if not text:
            parts = []
        else:
            parts = [part.strip() for part in re.split(r"[;,|]", text) if part.strip()]
    deduped = []
    seen = set()
    for part in parts:
        lowered = part.lower()
        if lowered in seen:
            continue
        seen.add(lowered)
        deduped.append(part)
    return ", ".join(deduped)


def _derive_linux_help_tags(item) -> str:
    explicit = linux_help_tag_text(item.get("tags"))
    if explicit:
        return explicit
    derived = []
    for value in (item.get("category"), item.get("name")):
        text = str(value or "").strip()
        if text:
            derived.append(text)
    cmd = str(item.get("cmd") or "").strip()
    if cmd:
        first_token = cmd.split()[0]
        if first_token:
            derived.append(first_token)
    return linux_help_tag_text(derived)


def _glossary_category(item) -> str:
    tags = item.get("tags")
    if isinstance(tags, (list, tuple, set)):
        tag_values = [str(tag).strip() for tag in tags if str(tag).strip()]
    else:
        tag_values = [
            part.strip() for part in re.split(r"[;,|]", str(tags or "")) if part.strip()
        ]
    if tag_values:
        return tag_values[0].title()
    return "Glossary"


def _bool_from_raw(value, default=True):
    if value is None or value == "":
        return bool(default)
    if isinstance(value, str):
        text = value.strip().lower()
        false_tokens = {"0", "false", "no", "off", "nein", "n", "selten"}
        true_tokens = {"1", "true", "yes", "on", "ja", "j", "immer"}
        parts = [part.strip() for part in re.split(r"[/,;|]", text) if part.strip()]
        if parts:
            if any(part in false_tokens for part in parts):
                return False
            if any(part in true_tokens for part in parts):
                return True
        return text not in false_tokens
    return bool(value)


def _normalize_source_row(row):
    normalized = {}
    source = {
        str(key or "").strip().lower(): value for key, value in dict(row or {}).items()
    }
    for target_key, aliases in SOURCE_FIELD_ALIASES.items():
        for alias in aliases:
            if alias in source:
                normalized[target_key] = source.get(alias)
                break
    return normalized


def _normalize_risk_value(value):
    text = str(value or "").strip().lower()
    if not text:
        return ""
    if any(token in text for token in ("rot", "red", "danger", "gefähr", "gefaehr")):
        return "danger"
    if any(
        token in text for token in ("gelb", "yellow", "admin", "vorsicht", "caution")
    ):
        return "caution"
    if any(token in text for token in ("gruen", "grün", "green", "safe", "sicher")):
        return "safe"
    return text


def _risk_storage_to_dialog_label(app, value):
    normalized = _normalize_risk_value(value)
    for label, key in _risk_dialog_choices_for_app(app):
        if key == normalized:
            return label
    return _risk_dialog_choices_for_app(app)[1][0]


def _risk_dialog_label_to_storage(app, value):
    label = str(value or "").strip()
    for option_label, key in _risk_dialog_choices_for_app(app):
        if label == option_label:
            return key
    return _normalize_risk_value(value) or "caution"


def _selected_linux_help_id(app) -> str:
    tree = _app_value(app, "linux_help_tree")
    if tree is not None:
        try:
            selection = tree.selection()
        except Exception:
            selection = ()
        if selection:
            return str(selection[0] or "").strip()
    listbox = _app_value(app, "linux_help_listbox")
    if listbox is not None:
        try:
            idx = int(listbox.curselection()[0])
        except Exception:
            idx = -1
        visible_ids = list(_app_value(app, "linux_help_visible_command_ids", []) or [])
        if 0 <= idx < len(visible_ids):
            return str(visible_ids[idx] or "").strip()
    command_id = str(_app_value(app, "linux_help_selected_id", "") or "").strip()
    if command_id:
        return command_id
    return ""


def _linux_help_command_resolver(app):
    resolver = _app_value(app, "get_linux_help_command_by_id")
    return resolver if callable(resolver) else None


def _resolve_linux_help_item(app, command_id=None):
    target_id = str(command_id or _selected_linux_help_id(app) or "").strip()
    if not target_id:
        return None
    resolver = _linux_help_command_resolver(app)
    if resolver is not None:
        return resolver(target_id)
    return get_linux_help_command_by_id(app, target_id)


def _set_chad_direct_prompt_safe(app, title, prompt):
    direct_prompt_setter = getattr(app, "set_chad_direct_prompt", None)
    if callable(direct_prompt_setter):
        direct_prompt_setter(title, prompt)
        return True
    return False


def _get_selected_linux_help_item(app):
    return _resolve_linux_help_item(app)


def _is_glossary_item(item) -> bool:
    return str((item or {}).get("entry_type") or "").strip().lower() == "glossary"


def _linux_help_source_value(app):
    label = str(_app_var_get(app, "linux_help_source_var", "") or "").strip()
    return {
        _tr(app, "linux_help_source_commands", "Commands"): "command",
        _tr(app, "linux_help_source_glossary", "Glossary"): "glossary",
        _tr(app, "linux_help_source_both", "Both"): "both",
    }.get(label, "both")


def _linux_help_chad_title(item) -> str:
    entry = dict(item or {})
    if _is_glossary_item(entry):
        return f"Explain glossary term: {entry.get('name') or 'Unknown'}"
    return (
        f"Explain Linux command: {entry.get('name') or entry.get('cmd') or 'Unknown'}"
    )


def build_linux_help_chad_prompt(app, item):
    entry = dict(item or {})
    if _is_glossary_item(entry):
        lines = [
            "Please explain the following IT/Linux glossary term for beginners in clear, well-structured English.",
            "",
            "Important requirements for your answer:",
            "- Start with a simple definition.",
            "- Explain why the term matters in Linux or IT practice.",
            "- Use the stored example and expand it in plain language.",
            "- Mention 2 or 3 related concepts beginners should learn next.",
            "",
            "Glossary data from NetMon:",
            f"Term: {entry.get('name') or '–'}",
            f"Category: {entry.get('category') or '–'}",
            f"Short description: {entry.get('info') or '–'}",
            f"Example: {entry.get('example') or '–'}",
            f"Tags: {str(entry.get('tags') or '–').strip() or '–'}",
            f"Related terms: {str(entry.get('related') or '–').strip() or '–'}",
        ]
        return "\n".join(lines).strip()
    _risk_key, risk_label, _fill = RISK_UI.get(
        str(entry.get("risk") or "").strip().lower(),
        RISK_UI["unknown"],
    )
    root_text = (
        app.tr("label_yes", "Yes")
        if entry.get("root_required")
        else app.tr("label_no", "No")
    )
    tags = str(entry.get("tags") or "–").strip() or "–"
    explanation = (
        str(entry.get("explanation") or "").strip()
        or "No additional explanation stored."
    )
    warning_text = (
        str(entry.get("warning_text") or "").strip() or "No additional warning stored."
    )
    lines = [
        "Please explain the following Linux command for beginners in clear, well-structured English.",
        "",
        "Important requirements for your answer:",
        "- Write in a beginner-friendly way without assuming prior knowledge.",
        "- Start with a short explanation of what the command is used for.",
        "- Then break down the syntax and the most important parts.",
        "- Explain the example step by step.",
        "- Mention typical real-world use cases.",
        "- Clearly point out risks, root requirements, and common mistakes.",
        "- End with 3 safe beginner tips or variants.",
        "- Use calm, clear language instead of unnecessary jargon.",
        "",
        "Command data from NetMon:",
        f"Command name: {entry.get('name') or entry.get('cmd') or '–'}",
        f"Category: {entry.get('category') or '–'}",
        f"Command: {entry.get('cmd') or '–'}",
        f"Short description: {entry.get('info') or '–'}",
        f"Syntax: {entry.get('syntax') or entry.get('cmd') or '–'}",
        f"Example: {entry.get('example') or entry.get('cmd') or '–'}",
        f"Database explanation/note: {explanation}",
        f"Additional warning: {warning_text}",
        f"Root required: {root_text}",
        f"Risk level: {risk_label}",
        f"Tags: {tags}",
    ]
    return "\n".join(lines).strip()


def _selected_linux_help_chad_payload(app):
    item = _get_selected_linux_help_item(app)
    if item is None:
        return None, "", ""
    prompt = build_linux_help_chad_prompt(app, item)
    return item, _linux_help_chad_title(item), prompt


def _iter_source_rows(path):
    extension = os.path.splitext(str(path or ""))[1].lower()
    if extension == ".xlsx":
        if load_workbook is None:
            raise RuntimeError("XLSX import requires openpyxl.")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = list(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
        if not rows:
            return []
        headers = [str(cell or "").strip() for cell in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(row)]
    with open(path, "r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return [dict(row) for row in reader if row]


def normalize_linux_help_entry(app, raw, existing_ids=None):
    item = dict(raw or {})
    cmd = str(item.get("cmd") or "").strip()
    info = str(item.get("info") or item.get("description") or "").strip()
    category = str(item.get("category") or "Allgemein").strip() or "Allgemein"
    name = str(item.get("name") or cmd or "Unbenannt").strip() or "Unbenannt"
    syntax = str(item.get("syntax") or cmd).strip() or cmd
    example = str(item.get("example") or cmd).strip() or cmd
    explanation = str(item.get("explanation") or info).strip()
    warning_text = str(item.get("warning_text") or item.get("warning") or "").strip()
    risk = _normalize_risk_value(item.get("risk"))
    if risk not in {"safe", "caution", "danger"}:
        risk = linux_help_get_risk_level(cmd)
    enabled = _bool_from_raw(item.get("enabled"), default=True)
    root_required = _bool_from_raw(
        item.get("root_required"),
        default=linux_help_requires_root(cmd, risk),
    )
    normalized = {
        "id": linux_help_make_id(item, existing_ids=existing_ids),
        "entry_type": "command",
        "category": category,
        "name": name,
        "info": info,
        "cmd": cmd,
        "syntax": syntax or cmd,
        "example": example or cmd,
        "explanation": explanation,
        "warning_text": warning_text,
        "tags": _derive_linux_help_tags(item),
        "risk": risk,
        "root_required": root_required,
        "enabled": enabled,
    }
    return normalized


def normalize_linux_help_glossary_entry(app, raw, existing_ids=None):
    item = dict(raw or {})
    term = str(item.get("term") or item.get("name") or "").strip() or "Glossary entry"
    short = str(item.get("short") or item.get("info") or "").strip()
    explanation = str(item.get("explanation") or short).strip()
    warning_text = str(item.get("warning_text") or "").strip()
    related = item.get("related") or []
    related_text = (
        linux_help_tag_text(related)
        if isinstance(related, (list, tuple, set))
        else linux_help_tag_text(str(related or ""))
    )
    return {
        "id": linux_help_make_id(
            {"id": item.get("id") or term, "name": term}, existing_ids=existing_ids
        ),
        "entry_type": "glossary",
        "category": str(item.get("category") or "").strip() or _glossary_category(item),
        "name": term,
        "info": short,
        "cmd": "",
        "syntax": "",
        "example": str(item.get("example") or "").strip(),
        "explanation": explanation,
        "warning_text": warning_text,
        "tags": linux_help_tag_text(item.get("tags")),
        "risk": "unknown",
        "root_required": False,
        "enabled": _bool_from_raw(item.get("enabled"), default=True),
        "related": related_text,
    }


def _serialize_linux_help_glossary_entry(item):
    entry = dict(item or {})
    return {
        "id": str(entry.get("id") or "").strip(),
        "term": str(entry.get("name") or entry.get("term") or "").strip(),
        "short": str(entry.get("info") or entry.get("short") or "").strip(),
        "example": str(entry.get("example") or "").strip(),
        "tags": [
            part.strip()
            for part in str(entry.get("tags") or "").split(",")
            if part.strip()
        ],
        "related": [
            part.strip()
            for part in str(entry.get("related") or "").split(",")
            if part.strip()
        ],
        "explanation": str(entry.get("explanation") or "").strip(),
        "warning_text": str(entry.get("warning_text") or "").strip(),
        "enabled": bool(entry.get("enabled", True)),
        "category": str(entry.get("category") or "").strip(),
    }


def rebuild_linux_help_indexes(app):
    active_entries = [item for item in app.linux_help_data if item.get("enabled", True)]
    active_entries.sort(
        key=lambda item: (
            item.get("category", "").lower(),
            item.get("name", "").lower(),
        )
    )
    categories = sorted(
        {item.get("category", "Allgemein") for item in active_entries},
        key=str.lower,
    )
    commands_by_category = {category: [] for category in categories}
    for item in active_entries:
        commands_by_category.setdefault(item["category"], []).append(item)
    app.linux_help_commands = active_entries
    app.linux_help_categories = categories
    app.linux_help_commands_by_category = commands_by_category
    glossary_categories = set(
        (_app_value(app, "linux_help_glossary_by_category", {}) or {}).keys()
    )
    app.linux_help_combined_categories = sorted(
        set(categories).union(glossary_categories),
        key=str.lower,
    )


def rebuild_linux_help_glossary_indexes(app):
    glossary_entries = list(_app_value(app, "linux_help_glossary", []) or [])
    glossary_entries.sort(
        key=lambda item: (
            item.get("category", "").lower(),
            item.get("name", "").lower(),
        )
    )
    categories = sorted(
        {item.get("category", "Glossary") for item in glossary_entries},
        key=str.lower,
    )
    glossary_by_category = {category: [] for category in categories}
    for item in glossary_entries:
        glossary_by_category.setdefault(item["category"], []).append(item)
    app.linux_help_glossary = glossary_entries
    app.linux_help_glossary_by_category = glossary_by_category
    app.linux_help_combined_categories = sorted(
        set(_app_value(app, "linux_help_categories", []) or []).union(categories),
        key=str.lower,
    )


def load_linux_help_glossary(app):
    path = str(_app_value(app, "linux_help_glossary_path", "") or "").strip()
    seed_path = str(_app_value(app, "linux_help_glossary_seed_path", "") or "").strip()
    source_path = path if path and os.path.isfile(path) else seed_path
    if not source_path or not os.path.isfile(source_path):
        app.linux_help_glossary = []
        app.linux_help_glossary_by_category = {}
        app.linux_help_combined_categories = list(
            _app_value(app, "linux_help_categories", []) or []
        )
        return []
    with open(source_path, "r", encoding="utf-8") as handle:
        payload = json.load(handle)
    entries = payload.get("entries") if isinstance(payload, dict) else []
    seen_ids = set(
        item.get("id") for item in (_app_value(app, "linux_help_data", []) or [])
    )
    glossary = []
    for raw in entries if isinstance(entries, list) else []:
        normalized = normalize_linux_help_glossary_entry(
            app, raw, existing_ids=seen_ids
        )
        seen_ids.add(normalized["id"])
        glossary.append(normalized)
    app.linux_help_glossary = glossary
    rebuild_linux_help_glossary_indexes(app)
    if source_path == seed_path and path:
        save_linux_help_glossary(app, update_status=False)
    return glossary


def save_linux_help_glossary(app, update_status=True):
    payload = {
        "name": "IT-Lexikon fuer Linux Hilfe",
        "version": "0.1.0",
        "language": "de",
        "entries": [
            _serialize_linux_help_glossary_entry(item)
            for item in (_app_value(app, "linux_help_glossary", []) or [])
            if item.get("enabled", True)
        ],
    }
    ok, err = safe_write_json(
        app.linux_help_glossary_path,
        payload,
        indent=2,
        ensure_ascii=False,
    )
    if not ok:
        if update_status:
            app.linux_help_status_var.set(
                app.tr(
                    "linux_help_status_glossary_save_failed",
                    "Linux help: glossary save failed.",
                )
            )
        update_linux_help_notes(app, None, extra_error=str(err))
        return False
    if update_status:
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_glossary_saved_count",
                "Linux help: {count} glossary entries saved.",
            ).format(count=len(_app_value(app, "linux_help_glossary", []) or []))
        )
    return True


def import_linux_help_csv(app, path):
    commands = []
    seen_ids = set()
    for row in _iter_source_rows(path):
        source_row = _normalize_source_row(row)
        reduced = {
            "category": source_row.get("category"),
            "name": source_row.get("name"),
            "info": source_row.get("info"),
            "cmd": source_row.get("cmd"),
            "syntax": source_row.get("syntax"),
            "example": source_row.get("example"),
            "explanation": source_row.get("explanation"),
            "warning_text": source_row.get("warning_text"),
            "tags": source_row.get("tags"),
            "root_required": source_row.get("root_required"),
            "risk": source_row.get("risk"),
        }
        if not (
            str(reduced.get("category") or "").strip()
            and str(reduced.get("name") or "").strip()
            and str(reduced.get("cmd") or "").strip()
        ):
            continue
        item = normalize_linux_help_entry(app, reduced, existing_ids=seen_ids)
        seen_ids.add(item["id"])
        commands.append(item)
    commands.sort(
        key=lambda item: (
            item.get("category", "").lower(),
            item.get("name", "").lower(),
        )
    )
    app.linux_help_data = commands
    rebuild_linux_help_indexes(app)
    set_linux_help_dirty(app, True)
    return commands


def export_linux_help_csv(app, path):
    with open(path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=LINUX_HELP_EXPORT_FIELDS)
        writer.writeheader()
        for item in app.linux_help_data:
            row = dict(item)
            row["root_required"] = "yes" if item.get("root_required") else "no"
            writer.writerow(
                {field: row.get(field, "") for field in LINUX_HELP_EXPORT_FIELDS}
            )


def load_linux_help_data(app):
    data_path = app.linux_help_data_path
    if os.path.isfile(data_path):
        with open(data_path, "r", encoding="utf-8") as handle:
            raw_items = json.load(handle)
        items = raw_items if isinstance(raw_items, list) else []
        seen_ids = set()
        app.linux_help_data = []
        for item in items:
            normalized = normalize_linux_help_entry(app, item, existing_ids=seen_ids)
            seen_ids.add(normalized["id"])
            app.linux_help_data.append(normalized)
        rebuild_linux_help_indexes(app)
        load_linux_help_glossary(app)
        set_linux_help_dirty(app, False)
        return app.linux_help_data

    commands = import_linux_help_csv(app, linux_help_csv_path())
    save_linux_help_data(app, update_status=False)
    load_linux_help_glossary(app)
    set_linux_help_dirty(app, False)
    return commands


def save_linux_help_data(app, update_status=True):
    ok, err = safe_write_json(
        app.linux_help_data_path,
        app.linux_help_data,
        indent=2,
        ensure_ascii=False,
    )
    if not ok:
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_save_failed",
                "Linux help: save failed.",
            )
        )
        update_linux_help_notes(app, None, extra_error=str(err))
        return False
    set_linux_help_dirty(app, False)
    if update_status:
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_saved_count",
                "Linux help: {count} entries saved.",
            ).format(count=len(app.linux_help_data))
        )
    return True


def get_linux_help_categories(app):
    combined = _app_value(app, "linux_help_combined_categories")
    if combined:
        return list(combined)
    return list(app.linux_help_categories)


def get_linux_help_commands_for_category(app, category):
    name = str(category or "").strip()
    source_mode = _linux_help_source_value(app)
    command_items = []
    glossary_items = []
    if source_mode in {"command", "both"}:
        if not name or name == _linux_help_all_categories_label(app):
            command_items = list(app.linux_help_commands)
        else:
            command_items = list(app.linux_help_commands_by_category.get(name, []))
    if source_mode in {"glossary", "both"}:
        glossary_by_category = (
            _app_value(app, "linux_help_glossary_by_category", {}) or {}
        )
        if not name or name == _linux_help_all_categories_label(app):
            glossary_items = list(_app_value(app, "linux_help_glossary", []) or [])
        else:
            glossary_items = list(glossary_by_category.get(name, []))
    items = [*command_items, *glossary_items]
    items.sort(
        key=lambda item: (
            item.get("category", "").lower(),
            item.get("name", "").lower(),
        )
    )
    return items


def get_linux_help_command_by_id(app, command_id):
    target = str(command_id or "").strip()
    data = list(_app_value(app, "linux_help_data", []) or []) + list(
        _app_value(app, "linux_help_glossary", []) or []
    )
    for item in data:
        if item.get("id") == target:
            return item
    return None


def sync_linux_help_form_categories(app):
    categories = [
        _linux_help_all_categories_label(app),
        *get_linux_help_categories(app),
    ]
    category_combo = _app_value(app, "linux_help_category_combo")
    if category_combo is not None:
        category_combo["values"] = categories
    form_category_combo = _app_value(app, "linux_help_form_category_combo")
    if form_category_combo is not None:
        form_category_combo["values"] = categories[1:]


def set_linux_help_form_from_entry(app, item=None):
    entry = item or {}
    _app_var_set(app, "linux_help_form_category_var", entry.get("category", ""))
    _app_var_set(app, "linux_help_form_name_var", entry.get("name", ""))
    _app_var_set(app, "linux_help_form_info_var", entry.get("info", ""))
    _app_var_set(app, "linux_help_form_cmd_var", entry.get("cmd", ""))
    _app_var_set(app, "linux_help_form_risk_var", entry.get("risk", "caution"))
    _app_var_set(app, "linux_help_form_explanation_var", entry.get("explanation", ""))
    _app_var_set(app, "linux_help_form_warning_var", entry.get("warning_text", ""))
    _app_var_set(app, "linux_help_form_enabled_var", bool(entry.get("enabled", True)))


def set_linux_help_risk_light(app, risk):
    canvas = _app_value(app, "linux_help_risk_canvas")
    dot = _app_value(app, "linux_help_risk_dot")
    if canvas is None or dot is None:
        return
    _risk_key, _label, fill = RISK_UI.get(
        (risk or "").strip().lower(),
        RISK_UI["unknown"],
    )
    canvas.itemconfigure(dot, fill=fill)


def update_linux_help_notes(app, item=None, extra_error=""):
    widget = _app_value(app, "linux_help_text")
    if widget is None:
        return
    if not item:
        body = app.tr(
            "linux_help_notes_empty",
            "Select a Linux command or glossary term above to view explanations, notes, and examples here.",
        )
        if extra_error:
            body = f"{body}\n\nError: {extra_error}"
        _set_text_widget_content(widget, body)
        return

    if _is_glossary_item(item):
        lines = [
            f"{app.tr('linux_help_label_term', 'Term')}: {item.get('name') or '–'}",
            f"{app.tr('label_category', 'Category')}: {item.get('category') or '–'}",
            f"{app.tr('linux_help_col_source', 'Source')}: {_source_label(app, 'glossary')}",
            f"{app.tr('linux_help_label_tags', 'Tags')}: {item.get('tags') or '–'}",
            "",
            f"{app.tr('linux_help_label_description', 'Description')}:",
            item.get("info") or "–",
            "",
            f"{app.tr('linux_help_label_example', 'Example')}:",
            item.get("example") or "–",
            "",
            f"{app.tr('linux_help_label_related', 'Related terms')}:",
            item.get("related") or "–",
            "",
            f"{app.tr('linux_help_label_note_readonly', 'Note')}:",
            app.tr(
                "linux_help_note_glossary_readonly",
                "This glossary entry is read-only reference content and cannot be sent to the terminal.",
            ),
        ]
    else:
        _risk_key, risk_label, _fill = RISK_UI.get(
            str(item.get("risk") or "").strip().lower(),
            RISK_UI["unknown"],
        )
        lines = [
            f"{app.tr('linux_help_label_command_name', 'Command name')}: {item.get('name') or item.get('cmd') or '–'}",
            f"{app.tr('label_category', 'Category')}: {item.get('category') or '–'}",
            f"{app.tr('linux_help_col_source', 'Source')}: {_source_label(app, 'command')}",
            f"{app.tr('linux_help_col_risk', 'Risk')}: {risk_label}",
            f"{app.tr('linux_help_label_root_required', 'Root required')}: "
            + (
                app.tr("label_yes", "Yes")
                if item.get("root_required")
                else app.tr("label_no", "No")
            ),
            f"{app.tr('linux_help_label_tags', 'Tags')}: {item.get('tags') or '–'}",
            "",
            f"{app.tr('linux_help_col_command', 'Command')}: {item.get('cmd') or '–'}",
            f"{app.tr('linux_help_label_syntax', 'Syntax')}: {item.get('syntax') or item.get('cmd') or '–'}",
            f"{app.tr('linux_help_label_example', 'Example')}: {item.get('example') or item.get('cmd') or '–'}",
            "",
            f"{app.tr('linux_help_label_description', 'Description')}:",
            item.get("info") or "–",
            "",
            f"{app.tr('linux_help_label_explanation_text', 'Explanation')}:",
            item.get("explanation") or item.get("info") or "–",
            "",
            f"{app.tr('linux_help_label_warning_text', 'Note')}:",
            item.get("warning_text")
            or app.tr(
                "linux_help_note_display_only",
                "This entry acts as NetMon reference data and can be sent to the terminal when needed.",
            ),
        ]
    _set_text_widget_content(widget, "\n".join(lines))


def _risk_filter_value(app):
    label = str(_app_var_get(app, "linux_help_risk_filter_var", "") or "").strip()
    return {
        _tr(app, "linux_help_risk_safe", "🟢 Safe"): "safe",
        _tr(app, "linux_help_risk_caution", "🟡 Admin"): "caution",
        _tr(app, "linux_help_risk_danger", "🔴 Dangerous"): "danger",
    }.get(label, "")


def _matches_linux_help_query(item, query):
    text = " ".join(
        [
            str(item.get("category") or ""),
            str(item.get("name") or ""),
            str(item.get("info") or ""),
            str(item.get("cmd") or ""),
            str(item.get("syntax") or ""),
            str(item.get("example") or ""),
            str(item.get("explanation") or ""),
            str(item.get("tags") or ""),
        ]
    ).lower()
    return query in text


def clear_linux_help_filters(app):
    _app_var_set(app, "linux_help_search_var", "")
    _app_var_set(app, "linux_help_category_var", _linux_help_all_categories_label(app))
    _app_var_set(app, "linux_help_source_var", _linux_help_all_sources_label(app))
    _app_var_set(app, "linux_help_risk_filter_var", _linux_help_all_risks_label(app))
    refresh_linux_help_list(app)


def refresh_linux_help_list(app, category=None):
    category_name = str(
        category or _app_var_get(app, "linux_help_category_var", "") or ""
    ).strip()
    if not category_name:
        category_name = _linux_help_all_categories_label(app)
    if category_name != _linux_help_all_categories_label(
        app
    ) and category_name not in get_linux_help_categories(app):
        category_name = _linux_help_all_categories_label(app)
    _app_var_set(app, "linux_help_category_var", category_name)

    commands = list(get_linux_help_commands_for_category(app, category_name))
    query = str(_app_var_get(app, "linux_help_search_var", "") or "").strip().lower()
    if query:
        commands = [item for item in commands if _matches_linux_help_query(item, query)]
    risk_filter = _risk_filter_value(app)
    if risk_filter:
        commands = [item for item in commands if item.get("risk") == risk_filter]

    app.linux_help_visible_commands = commands
    app.linux_help_visible_command_ids = [item.get("id", "") for item in commands]

    tree = _app_value(app, "linux_help_tree")
    listbox = _app_value(app, "linux_help_listbox")
    if tree is not None:
        tree.delete(*tree.get_children())
        for item in commands:
            _risk_key, risk_label, _fill = RISK_UI.get(
                item.get("risk"),
                RISK_UI["unknown"],
            )
            tree.insert(
                "",
                "end",
                iid=item.get("id", ""),
                text=item.get("name") or item.get("cmd") or "–",
                values=(
                    item.get("category") or "–",
                    _source_label(
                        app, "glossary" if _is_glossary_item(item) else "command"
                    ),
                    risk_label,
                    (
                        _tr(app, "label_yes", "Yes")
                        if item.get("root_required")
                        else _tr(app, "label_no", "No")
                    ),
                    item.get("tags") or "–",
                    item.get("info") or "–",
                ),
            )
    elif listbox is not None:
        listbox.delete(0, "end")
        for item in commands:
            listbox.insert("end", item.get("name") or item.get("cmd") or "–")

    app.linux_help_selected_id = ""
    _app_var_set(app, "linux_help_detail_name_var", "–")
    _app_var_set(app, "linux_help_detail_category_var", "–")
    _app_var_set(
        app,
        "linux_help_command_var",
        app.tr("linux_help_no_command_selected", "No command selected."),
    )
    _app_var_set(
        app,
        "linux_help_desc_var",
        app.tr(
            "linux_help_select_left_hint",
            "Select a command above to view syntax, description, tags and risks.",
        ),
    )
    _app_var_set(app, "linux_help_detail_syntax_var", "–")
    _app_var_set(app, "linux_help_detail_example_var", "–")
    _app_var_set(app, "linux_help_detail_root_var", "–")
    _app_var_set(app, "linux_help_detail_tags_var", "–")
    _app_var_set(
        app, "linux_help_risk_var", app.tr("linux_help_risk_unknown", "Risk: -")
    )
    _app_var_set(
        app,
        "linux_help_risk_badge_var",
        app.tr("linux_help_risk_unknown_label", "⚪ Unknown"),
    )
    set_linux_help_risk_light(app, "unknown")
    set_linux_help_form_from_entry(app, None)
    update_linux_help_notes(app, None)

    if commands:
        first_id = commands[0].get("id", "")
        if tree is not None:
            tree.selection_set(first_id)
            tree.focus(first_id)
        elif listbox is not None:
            listbox.selection_clear(0, "end")
            listbox.selection_set(0)
            listbox.activate(0)
        on_linux_help_select(app)


def on_linux_help_select(app, _event=None):
    command_id = _selected_linux_help_id(app)
    if not command_id:
        return
    item = _resolve_linux_help_item(app, command_id)
    if not item:
        return

    app.linux_help_selected_id = command_id
    risk = str(item.get("risk") or "").strip().lower()
    _risk_key, risk_label, _fill = RISK_UI.get(risk, RISK_UI["unknown"])

    _app_var_set(app, "linux_help_detail_name_var", item.get("name") or "–")
    _app_var_set(app, "linux_help_detail_category_var", item.get("category") or "–")
    _app_var_set(
        app,
        "linux_help_command_var",
        item.get("cmd")
        or (
            app.tr("linux_help_glossary_no_command", "No terminal command")
            if _is_glossary_item(item)
            else "–"
        ),
    )
    _app_var_set(
        app,
        "linux_help_desc_var",
        item.get("info")
        or app.tr("linux_help_no_description", "No description available."),
    )
    _app_var_set(
        app,
        "linux_help_detail_syntax_var",
        item.get("syntax")
        or item.get("cmd")
        or (
            app.tr("linux_help_glossary_not_applicable", "Not applicable")
            if _is_glossary_item(item)
            else "–"
        ),
    )
    _app_var_set(
        app,
        "linux_help_detail_example_var",
        item.get("example") or item.get("cmd") or "–",
    )
    _app_var_set(
        app,
        "linux_help_detail_root_var",
        (
            app.tr("linux_help_glossary_not_applicable", "Not applicable")
            if _is_glossary_item(item)
            else (
                app.tr("label_yes", "Yes")
                if item.get("root_required")
                else app.tr("label_no", "No")
            )
        ),
    )
    _app_var_set(app, "linux_help_detail_tags_var", item.get("tags") or "–")
    _app_var_set(app, "linux_help_risk_var", risk_label)
    _app_var_set(app, "linux_help_risk_badge_var", risk_label)
    set_linux_help_risk_light(app, risk)
    set_linux_help_form_from_entry(app, None if _is_glossary_item(item) else item)
    update_linux_help_notes(app, item)


def linux_help_apply_form_to_entry(app, mark_dirty=True):
    target = get_linux_help_command_by_id(app, app.linux_help_selected_id)
    if target is None or _app_value(app, "linux_help_form_category_var") is None:
        return False
    target["category"] = (
        str(
            _app_var_get(app, "linux_help_form_category_var", "Allgemein")
            or "Allgemein"
        ).strip()
        or "Allgemein"
    )
    target["name"] = (
        str(_app_var_get(app, "linux_help_form_name_var", "") or "").strip()
        or target["name"]
    )
    target["info"] = str(
        _app_var_get(app, "linux_help_form_info_var", "") or ""
    ).strip()
    target["cmd"] = str(_app_var_get(app, "linux_help_form_cmd_var", "") or "").strip()
    target["explanation"] = str(
        _app_var_get(app, "linux_help_form_explanation_var", "") or ""
    ).strip()
    target["warning_text"] = str(
        _app_var_get(app, "linux_help_form_warning_var", "") or ""
    ).strip()
    manual_risk = (
        str(_app_var_get(app, "linux_help_form_risk_var", "") or "").strip().lower()
    )
    target["risk"] = (
        manual_risk
        if manual_risk in {"safe", "caution", "danger"}
        else linux_help_get_risk_level(target.get("cmd", ""))
    )
    target["enabled"] = bool(_app_var_get(app, "linux_help_form_enabled_var", False))
    normalized = normalize_linux_help_entry(
        app,
        target,
        existing_ids={item["id"] for item in app.linux_help_data if item is not target},
    )
    target.update(normalized)
    rebuild_linux_help_indexes(app)
    sync_linux_help_form_categories(app)
    refresh_linux_help_list(app, category=target.get("category"))
    if mark_dirty:
        set_linux_help_dirty(app, True)
    return True


def linux_help_new_entry(app):
    if _linux_help_source_value(app) == "glossary":
        base = {
            "entry_type": "glossary",
            "category": app.linux_help_category_var.get() or "Glossary",
            "name": "New term",
            "info": "Short definition",
            "example": "",
            "tags": "",
            "related": "",
            "explanation": "",
            "warning_text": "",
            "enabled": True,
        }
    else:
        base = {
            "entry_type": "command",
            "category": app.linux_help_category_var.get() or "General",
            "name": "New command",
            "info": "New short description",
            "cmd": "",
            "syntax": "",
            "example": "",
            "tags": "",
            "root_required": False,
            "explanation": "",
            "warning_text": "",
            "risk": "caution",
            "enabled": True,
        }
    open_linux_help_entry_dialog(app, item=base, is_new=True)


def linux_help_edit_selected(app):
    if not app.linux_help_selected_id:
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_select_entry_first",
                "Linux help: please select an entry first.",
            )
        )
        return
    target = get_linux_help_command_by_id(app, app.linux_help_selected_id)
    if target is None:
        return
    open_linux_help_entry_dialog(app, item=target, is_new=False)


def linux_help_delete_selected(app):
    target = get_linux_help_command_by_id(app, app.linux_help_selected_id)
    if target is None:
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_select_entry_first",
                "Linux help: please select an entry first.",
            )
        )
        return
    if messagebox is not None:
        confirmed = messagebox.askyesno(
            app.tr("linux_help_delete_title", "Delete entry"),
            app.tr(
                (
                    "linux_help_delete_confirm_glossary"
                    if _is_glossary_item(target)
                    else "linux_help_delete_confirm"
                ),
                "Delete the selected glossary entry now?"
                if _is_glossary_item(target)
                else "Delete the selected Linux command now?",
            ),
            parent=app,
        )
        if not confirmed:
            return
    if _is_glossary_item(target):
        app.linux_help_glossary = [
            item
            for item in app.linux_help_glossary
            if item.get("id") != app.linux_help_selected_id
        ]
        rebuild_linux_help_glossary_indexes(app)
        save_linux_help_glossary(app, update_status=False)
    else:
        app.linux_help_data = [
            item
            for item in app.linux_help_data
            if item.get("id") != app.linux_help_selected_id
        ]
        rebuild_linux_help_indexes(app)
        save_linux_help_data(app, update_status=False)
    app.linux_help_selected_id = ""
    sync_linux_help_form_categories(app)
    refresh_linux_help_list(app)
    set_linux_help_dirty(app, True)
    app.linux_help_status_var.set(
        app.tr(
            (
                "linux_help_status_glossary_deleted"
                if _is_glossary_item(target)
                else "linux_help_status_entry_deleted"
            ),
            "Linux help: glossary entry deleted."
            if _is_glossary_item(target)
            else "Linux help: entry deleted.",
        )
    )


def open_linux_help_entry_dialog(app, item=None, is_new=False):
    existing = getattr(app, "linux_help_entry_dialog", None)
    if existing is not None and existing.winfo_exists():
        try:
            existing.destroy()
        except Exception:
            pass

    base = dict(item or {})
    is_glossary = _is_glossary_item(base) or str(
        base.get("entry_type") or ""
    ).strip().lower() == "glossary"
    dialog = tk.Toplevel(app)
    dialog.title(
        app.tr(
            "linux_help_dialog_title_new_glossary"
            if is_new and is_glossary
            else "linux_help_dialog_title_new",
            "Create glossary term" if is_new and is_glossary else "Create Linux command",
        )
        if is_new
        else app.tr(
            "linux_help_dialog_title_edit_glossary"
            if is_glossary
            else "linux_help_dialog_title_edit",
            "Edit glossary term" if is_glossary else "Edit Linux command",
        )
    )
    dialog.transient(app)
    dialog.resizable(True, True)
    dialog.minsize(760, 560)
    dialog.columnconfigure(0, weight=1)
    dialog.rowconfigure(0, weight=1)
    app.linux_help_entry_dialog = dialog

    category_var = tk.StringVar(value=base.get("category") or "General")
    name_var = tk.StringVar(value=base.get("name") or "")
    info_var = tk.StringVar(value=base.get("info") or "")
    cmd_var = tk.StringVar(value=base.get("cmd") or "")
    syntax_var = tk.StringVar(value=base.get("syntax") or base.get("cmd") or "")
    example_var = tk.StringVar(value=base.get("example") or "")
    tags_var = tk.StringVar(value=base.get("tags") or "")
    related_var = tk.StringVar(value=base.get("related") or "")
    risk_var = tk.StringVar(value=_risk_storage_to_dialog_label(app, base.get("risk")))
    root_var = tk.BooleanVar(value=bool(base.get("root_required")))
    enabled_var = tk.BooleanVar(value=bool(base.get("enabled", True)))

    shell = ttk.Frame(dialog, padding=12)
    shell.grid(row=0, column=0, sticky="nsew")
    shell.columnconfigure(0, weight=1)
    shell.rowconfigure(1, weight=1)

    form = ttk.LabelFrame(
        shell,
        text=app.tr("linux_help_dialog_group", "Edit entry"),
        padding=10,
    )
    form.grid(row=0, column=0, sticky="ew")
    form.columnconfigure(1, weight=1)
    form.columnconfigure(3, weight=1)

    ttk.Label(form, text=app.tr("label_category", "Category")).grid(
        row=0, column=0, sticky="w", padx=(0, 8), pady=(0, 6)
    )
    category_combo = ttk.Combobox(
        form,
        textvariable=category_var,
        values=[value for value in get_linux_help_categories(app) if value],
    )
    category_combo.grid(row=0, column=1, sticky="ew", pady=(0, 6))

    ttk.Label(
        form,
        text=app.tr(
            "linux_help_label_term" if is_glossary else "linux_help_label_command_name",
            "Term" if is_glossary else "Command name",
        ),
    ).grid(
        row=0, column=2, sticky="w", padx=(12, 8), pady=(0, 6)
    )
    ttk.Entry(form, textvariable=name_var).grid(
        row=0, column=3, sticky="ew", pady=(0, 6)
    )

    ttk.Label(form, text=app.tr("linux_help_label_description", "Description")).grid(
        row=1, column=0, sticky="w", padx=(0, 8), pady=(0, 6)
    )
    ttk.Entry(form, textvariable=info_var).grid(
        row=1, column=1, columnspan=3, sticky="ew", pady=(0, 6)
    )

    if is_glossary:
        ttk.Label(form, text=app.tr("linux_help_label_example", "Example")).grid(
            row=2, column=0, sticky="w", padx=(0, 8), pady=(0, 6)
        )
        ttk.Entry(form, textvariable=example_var).grid(
            row=2, column=1, columnspan=3, sticky="ew", pady=(0, 6)
        )

        ttk.Label(form, text=app.tr("linux_help_label_tags", "Tags")).grid(
            row=3, column=0, sticky="w", padx=(0, 8), pady=(0, 6)
        )
        ttk.Entry(form, textvariable=tags_var).grid(
            row=3, column=1, sticky="ew", pady=(0, 6)
        )
        ttk.Label(form, text=app.tr("linux_help_label_related", "Related terms")).grid(
            row=3, column=2, sticky="w", padx=(12, 8), pady=(0, 6)
        )
        ttk.Entry(form, textvariable=related_var).grid(
            row=3, column=3, sticky="ew", pady=(0, 6)
        )
    else:
        ttk.Label(form, text=app.tr("linux_help_col_command", "Command")).grid(
            row=2, column=0, sticky="w", padx=(0, 8), pady=(0, 6)
        )
        ttk.Entry(form, textvariable=cmd_var).grid(
            row=2, column=1, columnspan=3, sticky="ew", pady=(0, 6)
        )

        ttk.Label(form, text=app.tr("linux_help_label_syntax", "Syntax")).grid(
            row=3, column=0, sticky="w", padx=(0, 8), pady=(0, 6)
        )
        ttk.Entry(form, textvariable=syntax_var).grid(
            row=3, column=1, columnspan=3, sticky="ew", pady=(0, 6)
        )

        ttk.Label(form, text=app.tr("linux_help_label_example", "Example")).grid(
            row=4, column=0, sticky="w", padx=(0, 8), pady=(0, 6)
        )
        ttk.Entry(form, textvariable=example_var).grid(
            row=4, column=1, columnspan=3, sticky="ew", pady=(0, 6)
        )

        ttk.Label(form, text=app.tr("linux_help_label_tags", "Tags")).grid(
            row=5, column=0, sticky="w", padx=(0, 8), pady=(0, 6)
        )
        ttk.Entry(form, textvariable=tags_var).grid(
            row=5, column=1, sticky="ew", pady=(0, 6)
        )

        ttk.Label(form, text=app.tr("linux_help_label_risk_filter", "Risk")).grid(
            row=5, column=2, sticky="w", padx=(12, 8), pady=(0, 6)
        )
        ttk.Combobox(
            form,
            textvariable=risk_var,
            values=[label for label, _key in _risk_dialog_choices_for_app(app)],
            state="readonly",
        ).grid(row=5, column=3, sticky="ew", pady=(0, 6))

    flags_row = ttk.Frame(form)
    flags_row.grid(
        row=4 if is_glossary else 6, column=0, columnspan=4, sticky="w", pady=(2, 0)
    )
    if not is_glossary:
        ttk.Checkbutton(
            flags_row,
            text=app.tr("linux_help_label_root_required", "Root required"),
            variable=root_var,
        ).pack(side="left")
    ttk.Checkbutton(
        flags_row,
        text=app.tr("label_enabled", "Enabled"),
        variable=enabled_var,
    ).pack(side="left", padx=(14 if not is_glossary else 0, 0))

    editors = ttk.Frame(shell)
    editors.grid(row=1, column=0, sticky="nsew", pady=(10, 0))
    editors.columnconfigure(0, weight=1)
    editors.columnconfigure(1, weight=1)
    editors.rowconfigure(0, weight=1)

    explanation_frame = ttk.LabelFrame(
        editors,
        text=app.tr("linux_help_label_explanation_text", "Explanation"),
        padding=8,
    )
    explanation_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 6))
    explanation_frame.columnconfigure(0, weight=1)
    explanation_frame.rowconfigure(0, weight=1)
    explanation_text = tk.Text(explanation_frame, wrap="word", height=10)
    explanation_text.grid(row=0, column=0, sticky="nsew")
    explanation_text.insert("1.0", base.get("explanation") or "")

    warning_frame = ttk.LabelFrame(
        editors,
        text=app.tr("linux_help_label_warning_text", "Note"),
        padding=8,
    )
    warning_frame.grid(row=0, column=1, sticky="nsew", padx=(6, 0))
    warning_frame.columnconfigure(0, weight=1)
    warning_frame.rowconfigure(0, weight=1)
    warning_text = tk.Text(warning_frame, wrap="word", height=10)
    warning_text.grid(row=0, column=0, sticky="nsew")
    warning_text.insert("1.0", base.get("warning_text") or "")

    buttons = ttk.Frame(shell)
    buttons.grid(row=2, column=0, sticky="ew", pady=(10, 0))
    buttons.columnconfigure(0, weight=1)

    def _close_dialog():
        app.__dict__.pop("linux_help_entry_dialog", None)
        try:
            dialog.destroy()
        except Exception:
            pass

    def _save_dialog():
        raw = {
            "id": base.get("id", ""),
            "entry_type": "glossary" if is_glossary else "command",
            "category": (category_var.get() or "").strip() or "General",
            "name": (name_var.get() or "").strip(),
            "info": (info_var.get() or "").strip(),
            "cmd": (cmd_var.get() or "").strip(),
            "syntax": (syntax_var.get() or "").strip(),
            "example": (example_var.get() or "").strip(),
            "tags": (tags_var.get() or "").strip(),
            "related": (related_var.get() or "").strip(),
            "risk": _risk_dialog_label_to_storage(app, risk_var.get()),
            "root_required": bool(root_var.get()),
            "enabled": bool(enabled_var.get()),
            "explanation": explanation_text.get("1.0", "end").strip(),
            "warning_text": warning_text.get("1.0", "end").strip(),
        }
        if (not is_glossary and (not raw["name"] or not raw["cmd"])) or (
            is_glossary and not raw["name"]
        ):
            if messagebox is not None:
                messagebox.showwarning(
                    app.tr("linux_help_dialog_validation_title", "Incomplete input"),
                    app.tr(
                        (
                            "linux_help_dialog_validation_body_glossary"
                            if is_glossary
                            else "linux_help_dialog_validation_body"
                        ),
                        "Please fill at least the glossary term."
                        if is_glossary
                        else "Please fill at least name and command.",
                    ),
                    parent=dialog,
                )
            return
        target = (
            None if is_new else get_linux_help_command_by_id(app, base.get("id", ""))
        )
        existing_ids = {
            item.get("id")
            for item in [*app.linux_help_data, *app.linux_help_glossary]
            if item is not target and item.get("id")
        }
        normalized = (
            normalize_linux_help_glossary_entry(app, raw, existing_ids=existing_ids)
            if is_glossary
            else normalize_linux_help_entry(app, raw, existing_ids=existing_ids)
        )
        if target is None:
            if is_glossary:
                app.linux_help_glossary.append(normalized)
            else:
                app.linux_help_data.append(normalized)
        else:
            target.clear()
            target.update(normalized)
        if is_glossary:
            rebuild_linux_help_glossary_indexes(app)
            save_linux_help_glossary(app, update_status=False)
        else:
            rebuild_linux_help_indexes(app)
            save_linux_help_data(app, update_status=False)
        sync_linux_help_form_categories(app)
        app.linux_help_category_var.set(_linux_help_all_categories_label(app))
        refresh_linux_help_list(app)
        if hasattr(app, "linux_help_tree") and app.linux_help_tree is not None:
            try:
                app.linux_help_tree.selection_set(normalized["id"])
                app.linux_help_tree.focus(normalized["id"])
            except Exception:
                pass
        app.linux_help_selected_id = normalized["id"]
        on_linux_help_select(app)
        set_linux_help_dirty(app, True)
        app.linux_help_status_var.set(
            app.tr(
                (
                    "linux_help_status_glossary_saved"
                    if is_glossary
                    else "linux_help_status_entry_saved"
                ),
                "Linux help: glossary entry saved."
                if is_glossary
                else "Linux help: entry saved.",
            )
        )
        _close_dialog()

    ttk.Button(
        buttons,
        text=app.tr("btn_cancel", "Cancel"),
        command=_close_dialog,
    ).pack(side="right")
    ttk.Button(
        buttons,
        text=app.tr("btn_save", "Save"),
        command=_save_dialog,
    ).pack(side="right", padx=(0, 8))

    dialog.protocol("WM_DELETE_WINDOW", _close_dialog)
    try:
        dialog.update_idletasks()
        dialog.grab_set()
        dialog.focus_force()
    except Exception:
        pass


def copy_linux_help_chad_prompt(app):
    item, _title, prompt = _selected_linux_help_chad_payload(app)
    if item is None:
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_no_command_selected",
                "Linux help: no command selected.",
            )
        )
        return False
    app.clipboard_clear()
    app.clipboard_append(prompt)
    app.linux_help_status_var.set(
        app.tr(
            "linux_help_status_prompt_copied",
            "Linux help: Chad prompt copied to clipboard.",
        )
    )
    return True


def send_linux_help_to_chad(app):
    item, title, prompt = _selected_linux_help_chad_payload(app)
    if item is None:
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_no_command_selected",
                "Linux help: no command selected.",
            )
        )
        return False
    if not _set_chad_direct_prompt_safe(app, title, prompt):
        app.chad_analysis_current_prompt = prompt
        app.chad_online_result = ""
        try:
            app.chad_analysis_source_var.set(
                f"{app.tr('chad_label_source_prefix', 'Source')}: {title}"
            )
            app.chad_analysis_status_var.set(
                app.tr("chad_status_analysis_prepared", "Analysis prepared.")
            )
        except Exception:
            pass
    app.linux_help_status_var.set(
        app.tr(
            (
                "linux_help_status_sent_to_chad_glossary"
                if _is_glossary_item(item)
                else "linux_help_status_sent_to_chad"
            ),
            "Linux help: entry handed to Chad for explanation.",
        )
    )
    return True


def show_linux_help_tree_menu(app, event):
    tree = getattr(app, "linux_help_tree", None)
    if tree is None:
        return
    try:
        row_id = tree.identify_row(event.y)
    except Exception:
        row_id = ""
    if row_id:
        try:
            tree.selection_set(row_id)
            tree.focus(row_id)
        except Exception:
            pass
        on_linux_help_select(app)
    item = _get_selected_linux_help_item(app)
    if item is None:
        return
    menu = tk.Menu(app, tearoff=0)
    menu.add_command(
        label=app.tr("linux_help_menu_send_chad", "Send to Chad"),
        command=lambda: send_linux_help_to_chad(app),
    )
    menu.add_command(
        label=app.tr("linux_help_menu_copy_prompt", "Copy as prompt"),
        command=lambda: copy_linux_help_chad_prompt(app),
    )
    menu.add_separator()
    menu.add_command(
        label=app.tr("btn_edit", "Edit"),
        command=app.linux_help_edit_selected,
    )
    show_fn = getattr(app, "_show_context_menu", None)
    if callable(show_fn):
        return show_fn(menu, event)
    try:
        menu.tk_popup(event.x_root, event.y_root)
    finally:
        try:
            menu.grab_release()
        except Exception:
            pass
    return "break"


def linux_help_import_csv_dialog(app):
    if filedialog is None:
        return
    path = filedialog.askopenfilename(
        parent=app,
        title=app.tr("linux_help_import_csv_title", "Import Linux command database"),
        filetypes=[
            (app.tr("filetype_data", "Databases"), "*.csv *.xlsx"),
            (app.tr("filetype_csv", "CSV files"), "*.csv"),
            (app.tr("filetype_excel", "Excel files"), "*.xlsx"),
            (app.tr("filetype_all", "All files"), "*.*"),
        ],
        initialfile=os.path.basename(linux_help_csv_path()),
    )
    if not path:
        return
    app.linux_help_import_path = path
    try:
        commands = import_linux_help_csv(app, path)
        sync_linux_help_form_categories(app)
        app.linux_help_category_var.set(_linux_help_all_categories_label(app))
        app.linux_help_source_var.set(_linux_help_all_sources_label(app))
        app.linux_help_risk_filter_var.set(_linux_help_all_risks_label(app))
        refresh_linux_help_list(app)
        save_linux_help_data(app, update_status=False)
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_imported_count",
                "Linux help: {count} entries imported.",
            ).format(count=len(commands))
        )
    except Exception as exc:
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_import_failed",
                "Linux help: import failed.",
            )
        )
        update_linux_help_notes(app, None, extra_error=str(exc))


def linux_help_export_csv_dialog(app):
    if filedialog is None:
        return
    path = filedialog.asksaveasfilename(
        parent=app,
        title=app.tr("linux_help_export_csv_title", "Export Linux database"),
        defaultextension=".csv",
        filetypes=[
            (app.tr("filetype_csv", "CSV files"), "*.csv"),
            (app.tr("filetype_all", "All files"), "*.*"),
        ],
        initialfile="linux_befehlsdatenbank_export.csv",
    )
    if not path:
        return
    export_linux_help_csv(app, path)
    app.linux_help_status_var.set(
        app.tr("linux_help_status_exported", "Linux help: CSV exported.")
    )


def reload_linux_help_data(app):
    try:
        commands = load_linux_help_data(app)
    except Exception as exc:
        app.linux_help_data = []
        app.linux_help_commands = []
        app.linux_help_categories = []
        app.linux_help_commands_by_category = {}
        app.linux_help_visible_commands = []
        app.linux_help_visible_command_ids = []
        app.linux_help_selected_id = ""
        sync_linux_help_form_categories(app)
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_load_failed",
                "Linux help: data could not be loaded.",
            )
        )
        update_linux_help_notes(app, None, extra_error=str(exc))
        return

    sync_linux_help_form_categories(app)
    app.linux_help_category_var.set(_linux_help_all_categories_label(app))
    app.linux_help_source_var.set(_linux_help_all_sources_label(app))
    app.linux_help_risk_filter_var.set(_linux_help_all_risks_label(app))
    refresh_linux_help_list(app)
    app.linux_help_status_var.set(
        app.tr(
            "linux_help_status_loaded_count",
            "Linux help: {count} commands loaded.",
        ).format(count=len(commands))
        + " | "
        + app.tr(
            "linux_help_status_loaded_glossary_count",
            "{count} glossary entries loaded.",
        ).format(count=len(_app_value(app, "linux_help_glossary", []) or []))
    )


def send_linux_help_command_to_terminal(app):
    item = _resolve_linux_help_item(app, app.linux_help_selected_id)
    if item is None:
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_no_command_selected",
                "Linux help: no command selected.",
            )
        )
        return False
    if _is_glossary_item(item):
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_glossary_no_terminal",
                "Linux help: glossary entries cannot be sent to the terminal.",
            )
        )
        return False
    if not item.get("enabled", True):
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_entry_disabled",
                "Linux help: this entry is disabled and cannot be sent.",
            )
        )
        return False

    cmd = item.get("cmd", "")
    ok, reason = app._validate_linux_help_command_for_terminal(cmd)
    if not ok:
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_terminal_not_possible",
                "Linux help: terminal send not possible ({reason})",
            ).format(reason=reason)
        )
        return False

    risk = str(item.get("risk") or "").strip().lower()
    requires_confirmation = bool(item.get("root_required")) or risk in {
        "caution",
        "danger",
    }
    if requires_confirmation and messagebox is not None:
        confirmed = messagebox.askyesno(
            app.tr("linux_help_confirm_send_title", "Send command to terminal"),
            app.tr(
                "linux_help_confirm_send_prompt",
                "This command may change system state.\n\nSend it to the terminal anyway?",
            ),
            parent=app,
        )
        if not confirmed:
            app.linux_help_status_var.set(
                app.tr(
                    "linux_help_status_terminal_cancelled",
                    "Linux help: terminal send canceled.",
                )
            )
            return False

    terminal_name = app._open_command_in_external_terminal(cmd)
    if not terminal_name:
        app.linux_help_status_var.set(
            app.tr(
                "linux_help_status_terminal_failed",
                "Linux help: terminal could not be started.",
            )
        )
        return False

    app.linux_help_status_var.set(
        app.tr(
            "linux_help_status_sent_to_terminal",
            "Linux help: command was sent to an external terminal.",
        )
    )
    if hasattr(app, "log"):
        try:
            app.log(
                app.tr(
                    "linux_help_log_command_sent",
                    "Linux help: command sent to terminal",
                )
            )
        except Exception:
            pass
    if hasattr(app, "linux_help_text") and app.linux_help_text is not None:
        if hasattr(app.linux_help_text, "configure"):
            app.linux_help_text.configure(state="normal")
        app.linux_help_text.delete("1.0", "end")
        app.linux_help_text.insert(
            "1.0",
            "\n".join(
                [
                    f"{app.tr('linux_help_label_name', 'Name')}: {item.get('name') or '–'}",
                    f"{app.tr('linux_help_label_command_name', 'Command name')}: {cmd}",
                    f"Terminal: {terminal_name}",
                    "",
                    app.tr(
                        "linux_help_terminal_note_sent",
                        "Command was sent to an external terminal.",
                    ),
                ]
            ),
        )
        if hasattr(app.linux_help_text, "configure"):
            app.linux_help_text.configure(state="disabled")
    return True
