#!/usr/bin/env python3
import importlib
import importlib.util
import os
import subprocess
import sys
import platform
import socket
import ssl
import getpass
from datetime import datetime
import threading
import time
import re
import shutil
import webbrowser
import shlex
import ipaddress
import json
import tempfile
import signal
import traceback
import warnings
import hashlib

try:
    import keyring
except Exception:
    keyring = None

# Ensure repository root is importable when launching via tools/netmon.py
NETMON_REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if NETMON_REPO_ROOT not in sys.path:
    sys.path.insert(0, NETMON_REPO_ROOT)

KEYRING_SERVICE = "NetMon"
KEYRING_USERNAME = "openai_api_key"
ANSI_RE = re.compile(r"\x1B\[[0-?]*[ -/]*[@-~]")


LIGHT_THEME = {
    "name": "light",
    "bg": "#f4f6f8",
    "panel_bg": "#f4f6f8",
    "fg": "#101418",
    "muted_fg": "#2f3a44",
    "accent": "#2f6feb",
    "entry_bg": "#ffffff",
    "entry_fg": "#111111",
    "text_bg": "#ffffff",
    "text_fg": "#101418",
    "select_bg": "#2f6feb",
    "select_fg": "#ffffff",
    "tree_bg": "#ffffff",
    "tree_fg": "#111111",
    "tab_bg": "#dde3ea",
    "tab_fg": "#111111",
    "tab_selected_bg": "#ffffff",
    "tab_selected_fg": "#000000",
}


DARK_THEME = {
    "name": "dark",
    "bg": "#111418",
    "panel_bg": "#111418",
    "fg": "#e6edf3",
    "muted_fg": "#b8c1cc",
    "accent": "#58a6ff",
    "entry_bg": "#181d24",
    "entry_fg": "#e6edf3",
    "text_bg": "#0d1117",
    "text_fg": "#dce4ee",
    "select_bg": "#264f78",
    "select_fg": "#ffffff",
    "tree_bg": "#0d1117",
    "tree_fg": "#dce4ee",
    "tab_bg": "#1a212b",
    "tab_fg": "#c3cfdb",
    "tab_selected_bg": "#263142",
    "tab_selected_fg": "#ffffff",
}


def strip_ansi_codes(text: str) -> str:
    return ANSI_RE.sub("", text or "")


def normalize_lynis_output(text: str) -> str:
    text = strip_ansi_codes(text)
    lines = []
    for line in text.splitlines():
        line = line.rstrip()
        if line:
            lines.append(line)
    return "\n".join(lines)


try:
    from netmon.netmon_ops import NetworkOperationCoordinator, RollbackTransaction
except ModuleNotFoundError:
    try:
        from tools.netmon_ops import NetworkOperationCoordinator, RollbackTransaction
    except ModuleNotFoundError:
        NetworkOperationCoordinator = None
        RollbackTransaction = None

try:
    import netmon.features.clamav_controller as clamav_controller
    import netmon.features.dns_controller as dns_controller
    import netmon.features.dot_controller as dot_controller
    import netmon.features.chad_controller as chad_controller
    import netmon.features.firewall_controller as firewall_controller
    import netmon.features.hdd_controller as hdd_controller
    import netmon.features.linux_help_controller as linux_help_controller
    import netmon.features.packet_capture_controller as packet_capture_controller
    import netmon.features.security_audit_controller as security_audit_controller
    import netmon.features.security_monitoring_controller as security_monitoring_controller
    import netmon.features.speedtest_controller as speedtest_controller
    import netmon.features.vpn_controller as vpn_controller
    from netmon.services.network_service import NetmonNetworkService
except ModuleNotFoundError:
    try:
        import features.clamav_controller as clamav_controller
        import features.dns_controller as dns_controller
        import features.dot_controller as dot_controller
        import features.chad_controller as chad_controller
        import features.firewall_controller as firewall_controller
        import features.hdd_controller as hdd_controller
        import features.linux_help_controller as linux_help_controller
        import features.packet_capture_controller as packet_capture_controller
        import features.security_audit_controller as security_audit_controller
        import features.security_monitoring_controller as security_monitoring_controller
        import features.speedtest_controller as speedtest_controller
        import features.vpn_controller as vpn_controller
        from services.network_service import NetmonNetworkService
    except ModuleNotFoundError:
        clamav_controller = None
        dns_controller = None
        dot_controller = None
        chad_controller = None
        firewall_controller = None
        hdd_controller = None
        linux_help_controller = None
        packet_capture_controller = None
        security_audit_controller = None
        security_monitoring_controller = None
        speedtest_controller = None
        vpn_controller = None
        NetmonNetworkService = None

try:
    from netmon.services.dns_control_service import DNSControlService
except ModuleNotFoundError:
    try:
        from services.dns_control_service import DNSControlService
    except ModuleNotFoundError:
        DNSControlService = None

try:
    from netmon.services.ai_analysis_service import AiAnalysisService
except ModuleNotFoundError:
    try:
        from services.ai_analysis_service import AiAnalysisService
    except ModuleNotFoundError:
        AiAnalysisService = None

try:
    from netmon.services.ai_online_service import AiOnlineService
except ModuleNotFoundError:
    try:
        from services.ai_online_service import AiOnlineService
    except ModuleNotFoundError:
        AiOnlineService = None

try:
    from netmon.services.startup_guard_service import StartupGuardService
except ModuleNotFoundError:
    try:
        from services.startup_guard_service import StartupGuardService
    except ModuleNotFoundError:
        StartupGuardService = None

try:
    from netmon.services.path_guard_service import PathGuardService
except ModuleNotFoundError:
    try:
        from services.path_guard_service import PathGuardService
    except ModuleNotFoundError:
        PathGuardService = None

try:
    from netmon.services.lang_service import LangService
except ModuleNotFoundError:
    try:
        from services.lang_service import LangService
    except ModuleNotFoundError:
        LangService = None

try:
    from netmon.ui.tabs.connections_tab import build_connections_tab
    from netmon.ui.tabs.chad_tab import build_chad_tab
    from netmon.ui.tabs.dns_tab import build_dns_tab
    from netmon.ui.tabs.hdd_tab import build_hdd_tab
    from netmon.ui.tabs.linux_help_tab import build_linux_help_tab
    from netmon.ui.tabs.links_tab import build_links_tab
    from netmon.ui.tabs.log_tab import build_log_tab
    from netmon.ui.tabs.security_tabs import build_security_tabs
    from netmon.ui.tabs.settings_tab import build_settings_tab as build_settings_tab_ui
    from netmon.ui.tabs.system_tab import build_system_tab
    from netmon.ui.tabs.tools_tabs import build_basic_tools_tabs
    from netmon.ui.tabs.trace_tab import build_trace_tab
    from netmon.ui.tabs.vpn_slots_tab import build_vpn_slots_tab
    from netmon.ui.widgets.lights import make_light as build_light_widget
    from netmon.ui.widgets.lights import set_light as update_light_widget
    from netmon.ui.widgets.tooltip import SimpleTooltip
except ModuleNotFoundError:
    from ui.tabs.connections_tab import build_connections_tab
    from ui.tabs.chad_tab import build_chad_tab
    from ui.tabs.log_tab import build_log_tab
    from ui.tabs.dns_tab import build_dns_tab
    from ui.tabs.hdd_tab import build_hdd_tab
    from ui.tabs.linux_help_tab import build_linux_help_tab
    from ui.tabs.links_tab import build_links_tab
    from ui.tabs.security_tabs import build_security_tabs
    from ui.tabs.settings_tab import build_settings_tab as build_settings_tab_ui
    from ui.tabs.system_tab import build_system_tab
    from ui.tabs.tools_tabs import build_basic_tools_tabs
    from ui.tabs.trace_tab import build_trace_tab
    from ui.tabs.vpn_slots_tab import build_vpn_slots_tab
    from ui.widgets.lights import make_light as build_light_widget
    from ui.widgets.lights import set_light as update_light_widget
    from ui.widgets.tooltip import SimpleTooltip

try:
    from netmon.core.coordinator import NetmonStateCoordinator
except ModuleNotFoundError:
    try:
        from core.coordinator import NetmonStateCoordinator
    except ModuleNotFoundError:
        NetmonStateCoordinator = None

try:
    from netmon.config.constants import (
        CLAMAV_FINDINGS_FILE,
        DNS_MODE_LABELS,
        DNS_PROFILES,
        DOT_DROPIN_DIR,
        DOT_DROPIN_PATH,
        DOT_PROFILES,
        HDD_SLOT_CONFIG_FILE,
        LINK_SLOT_CONFIG_FILE,
        LINK_SLOT_COUNT,
        NETMON_APP_VERSION,
        PORT_REALITY_DISPLAY_MAP,
        PRIMARY_VPN_SLOT_CONFIG_FILE,
        PRIVACY_CONFIG_FILE,
        SECURITY_BASELINE_FILE,
        TRACE_TARGETS,
        VPN_CH,
        VPN_DE,
        VPN_SLOT_CONFIG_FILE,
        VPN_SLOT_COUNT,
        VSCODE_DESKTOP_FILES,
        WG_IPV4_PROFILE_SUFFIX,
    )
except (ImportError, ModuleNotFoundError):
    from config.constants import (
        CLAMAV_FINDINGS_FILE,
        DNS_MODE_LABELS,
        DNS_PROFILES,
        DOT_DROPIN_DIR,
        DOT_DROPIN_PATH,
        DOT_PROFILES,
        HDD_SLOT_CONFIG_FILE,
        LINK_SLOT_CONFIG_FILE,
        LINK_SLOT_COUNT,
        NETMON_APP_VERSION,
        PORT_REALITY_DISPLAY_MAP,
        PRIMARY_VPN_SLOT_CONFIG_FILE,
        PRIVACY_CONFIG_FILE,
        SECURITY_BASELINE_FILE,
        TRACE_TARGETS,
        VPN_CH,
        VPN_DE,
        VPN_SLOT_CONFIG_FILE,
        VPN_SLOT_COUNT,
        VSCODE_DESKTOP_FILES,
        WG_IPV4_PROFILE_SUFFIX,
    )

try:
    from netmon.core.command import (
        check_sudo_nopasswd,
        renew_sudo_timestamp,
        run_cmd,
    )
except (ImportError, ModuleNotFoundError):
    from core.command import (
        check_sudo_nopasswd,
        renew_sudo_timestamp,
        run_cmd,
    )

try:
    from netmon.core.crash import CRASH_LOG_PATH, _append_crash_dump
except (ImportError, ModuleNotFoundError):
    from core.crash import CRASH_LOG_PATH, _append_crash_dump

try:
    from netmon.core.state import get_netmon_config_dir, get_netmon_state_dir
except ModuleNotFoundError:
    try:
        from core.state import get_netmon_config_dir, get_netmon_state_dir
    except ModuleNotFoundError:

        def _xdg_dir(env_name: str, default_suffix: str) -> str:
            base = (os.environ.get(env_name) or "").strip()
            if not base:
                base = os.path.expanduser(default_suffix)
            path = os.path.join(base, "netmon")
            try:
                os.makedirs(path, exist_ok=True)
            except Exception:
                pass
            return path

        def get_netmon_config_dir() -> str:
            return _xdg_dir("XDG_CONFIG_HOME", "~/.config")

        def get_netmon_state_dir() -> str:
            return _xdg_dir("XDG_STATE_HOME", "~/.local/state")


try:
    from netmon.core.utils import human_rate, human_size, safe_write_json, ts
except ModuleNotFoundError:
    try:
        from core.utils import human_rate, human_size, safe_write_json, ts
    except ModuleNotFoundError:
        pass

try:
    from netmon.core import tooling as _tooling
except ModuleNotFoundError:
    try:
        from core import tooling as _tooling
    except ModuleNotFoundError:
        _tooling = None

if _tooling is not None:
    check_speedtest_module_available = _tooling.check_speedtest_module_available
    get_cmd_path = _tooling.get_cmd_path
    get_user_pictures_dir = _tooling.get_user_pictures_dir
    has_cmd = _tooling.has_cmd
    load_speedtest_module_from_file = _tooling.load_speedtest_module_from_file


try:
    from netmon.services import network_helpers as _network_helpers
except ModuleNotFoundError:
    try:
        from services import network_helpers as _network_helpers
    except ModuleNotFoundError:
        _network_helpers = None

if _network_helpers is not None:
    format_traceroute_table = _network_helpers.format_traceroute_table
    is_local_or_lan_ip = _network_helpers.is_local_or_lan_ip
    is_safe_nmap_target = _network_helpers.is_safe_nmap_target
    is_safe_trace_target = _network_helpers.is_safe_trace_target
    normalize_external_url_target = _network_helpers.normalize_external_url_target
    parse_cf_trace = _network_helpers.parse_cf_trace
    parse_nmap_open_tcp_ports = _network_helpers.parse_nmap_open_tcp_ports
    parse_ss_lines = _network_helpers.parse_ss_lines
    parse_traceroute_hops = _network_helpers.parse_traceroute_hops


try:
    from netmon.core.rkhunter_manager import (
        RKHUNTER_BACKUP_PATH,
        detect_update_failure,
        is_rkhunter_installed,
        repair_config as rkhunter_repair_config,
        run_rkhunter_update,
        update_after_repair as rkhunter_update_after_repair,
        backup_config as rkhunter_backup_config,
    )
except ModuleNotFoundError:
    try:
        from core.rkhunter_manager import (
            RKHUNTER_BACKUP_PATH,
            detect_update_failure,
            is_rkhunter_installed,
            repair_config as rkhunter_repair_config,
            run_rkhunter_update,
            update_after_repair as rkhunter_update_after_repair,
            backup_config as rkhunter_backup_config,
        )
    except ModuleNotFoundError:
        RKHUNTER_BACKUP_PATH = "/etc/rkhunter.conf.netmon.bak"

        def is_rkhunter_installed():
            return has_cmd("rkhunter")

        def run_rkhunter_update(timeout=600):
            return run_cmd(["sudo", "-n", "rkhunter", "--update"], timeout=timeout)

        def detect_update_failure(output):
            lowered = str(output or "").strip().lower()
            markers = [
                "mirrors.dat does not exist",
                "unable to determine the latest version number",
                "update failed",
                "invalid web_cmd configuration option",
            ]
            return any(marker in lowered for marker in markers)

        def rkhunter_backup_config():
            return 1, "rkhunter_manager Modul fehlt"

        def rkhunter_repair_config():
            return 1, "rkhunter_manager Modul fehlt"

        def rkhunter_update_after_repair(timeout=600):
            return run_rkhunter_update(timeout=timeout)


try:
    from netmon.modules.port_analysis import check_port_real_service
except ModuleNotFoundError:

    def check_port_real_service(host, port, timeout=3):
        def _safe_close(sock):
            if sock is None:
                return
            try:
                sock.close()
            except Exception:
                pass

        port = int(port)
        result = {
            "host": host,
            "port": port,
            "state": "unknown",
            "reason": "",
        }

        sock = None
        try:
            sock = socket.create_connection((host, port), timeout=timeout)
            sock.settimeout(timeout)

            if port == 22:
                try:
                    banner = sock.recv(1024).decode(errors="ignore").strip()
                except socket.timeout:
                    banner = ""
                    result["state"] = "pseudo-open"
                    result["reason"] = "no_ssh_banner_timeout"
                except OSError:
                    banner = ""
                    result["state"] = "pseudo-open"
                    result["reason"] = "no_ssh_banner_connection_closed"

                if result["state"] == "unknown":
                    if banner.startswith("SSH-"):
                        result["state"] = "open"
                        result["reason"] = "ssh_banner_received"
                        result["banner"] = banner[:100]
                    elif banner:
                        result["state"] = "open"
                        result["reason"] = "unexpected_banner_received"
                        result["banner"] = banner[:100]
                    else:
                        result["state"] = "pseudo-open"
                        result["reason"] = "no_ssh_banner_connection_closed"

            elif port == 80:
                request = (
                    f"HEAD / HTTP/1.1\r\nHost: {host}\r\nConnection: close\r\n\r\n"
                ).encode("ascii", errors="ignore")
                sock.sendall(request)
                try:
                    response = sock.recv(1024).decode(errors="ignore")
                except socket.timeout:
                    response = ""
                    result["state"] = "filtered"
                    result["reason"] = "http_response_timeout"
                except OSError:
                    response = ""
                    result["state"] = "pseudo-open"
                    result["reason"] = "no_http_response_connection_closed"

                if result["state"] == "unknown":
                    text_resp = response.strip()
                    if text_resp:
                        first_line = (
                            text_resp.splitlines()[0].strip()
                            if text_resp.splitlines()
                            else text_resp
                        )
                        if first_line.startswith("HTTP/") or " HTTP/" in first_line:
                            result["state"] = "open"
                            result["reason"] = "http_response_received"
                        else:
                            result["state"] = "open"
                            result["reason"] = "non_http_response_received"
                        result["banner"] = text_resp[:100]
                    else:
                        result["state"] = "pseudo-open"
                        result["reason"] = "no_http_response_connection_closed"

            elif port == 443:
                context = ssl.create_default_context()
                context.check_hostname = False
                context.verify_mode = ssl.CERT_NONE
                tls_sock = None
                try:
                    tls_sock = context.wrap_socket(sock, server_hostname=host)
                    tls_sock.settimeout(timeout)
                    result["state"] = "open"
                    result["reason"] = "tls_handshake_success"
                except socket.timeout:
                    result["state"] = "filtered"
                    result["reason"] = "tls_handshake_timeout"
                except ssl.SSLError as exc:
                    msg = str(exc).lower()
                    if any(
                        token in msg
                        for token in (
                            "wrong version number",
                            "unknown protocol",
                            "http request",
                            "plaintext",
                        )
                    ):
                        result["state"] = "open"
                        result["reason"] = "tls_handshake_protocol_mismatch"
                    elif any(
                        token in msg
                        for token in (
                            "eof occurred",
                            "unexpected eof",
                            "connection reset",
                            "handshake failure",
                        )
                    ):
                        result["state"] = "pseudo-open"
                        result["reason"] = "tls_handshake_closed"
                    else:
                        result["state"] = "filtered"
                        result["reason"] = f"tls_handshake_error:{exc}"
                except (ConnectionResetError, BrokenPipeError):
                    result["state"] = "pseudo-open"
                    result["reason"] = "tls_handshake_closed"
                finally:
                    _safe_close(tls_sock)
                    sock = None

            else:
                try:
                    banner = sock.recv(1024).decode(errors="ignore").strip()
                except socket.timeout:
                    banner = ""
                    result["state"] = "filtered"
                    result["reason"] = "no_service_banner_timeout"
                except OSError:
                    banner = ""
                    result["state"] = "pseudo-open"
                    result["reason"] = "no_banner_connection_closed"

                if result["state"] == "unknown":
                    if banner:
                        result["state"] = "open"
                        result["reason"] = "service_detected"
                        result["banner"] = banner[:100]
                    else:
                        result["state"] = "pseudo-open"
                        result["reason"] = "no_banner_connection_closed"

        except ConnectionRefusedError:
            result["state"] = "closed"
            result["reason"] = "connection_refused"
        except OSError as exc:
            msg = str(exc).lower()
            errno_value = getattr(exc, "errno", None)
            if errno_value == 111 or "connection refused" in msg:
                result["state"] = "closed"
                result["reason"] = "connection_refused"
            else:
                result["state"] = "filtered"
                result["reason"] = str(exc)
        except Exception as exc:
            result["state"] = "filtered"
            result["reason"] = str(exc)
        finally:
            _safe_close(sock)

        return result


try:
    _speedtest_mod, SPEEDTEST_MODULE_FILE = load_speedtest_module_from_file(
        app_file=__file__,
        repo_root=NETMON_REPO_ROOT,
        search_paths=sys.path,
    )
    format_speedtest_summary = _speedtest_mod.format_speedtest_summary
    format_speedtest_table = _speedtest_mod.format_speedtest_table
    run_speedtest = _speedtest_mod.run_speedtest
    SPEEDTEST_MODULE_AVAILABLE = True
except Exception:
    SPEEDTEST_MODULE_FILE = ""
    SPEEDTEST_MODULE_AVAILABLE = False

    def run_speedtest(run_cmd_fn, has_cmd_fn, timeout=120):
        return None, "Kein Speedtest-Modul verfügbar."

    def format_speedtest_summary(result):
        if not isinstance(result, dict):
            return "Speedtest: keine Daten verfügbar."
        ping = result.get("ping_ms")
        jitter = result.get("jitter_ms")
        download = result.get("download_mbps")
        upload = result.get("upload_mbps")
        ping_txt = "-" if ping is None else f"{ping:.0f}"
        jitter_txt = "-" if jitter is None else f"{jitter:.0f}"
        dl_txt = "-" if download is None else f"{download:.1f}"
        ul_txt = "-" if upload is None else f"{upload:.1f}"
        return (
            f"Speedtest abgeschlossen: Ping {ping_txt} ms | "
            f"Jitter {jitter_txt} ms | Download {dl_txt} Mbit/s | Upload {ul_txt} Mbit/s"
        )

    def format_speedtest_table(result):
        if not isinstance(result, dict):
            return "Kein Speedtest-Ergebnis verfügbar."
        ping = result.get("ping_ms")
        jitter = result.get("jitter_ms")
        download = result.get("download_mbps")
        upload = result.get("upload_mbps")
        ping_txt = "-" if ping is None else f"{ping:.0f} ms"
        jitter_txt = "-" if jitter is None else f"{jitter:.0f} ms"
        dl_txt = "-" if download is None else f"{download:.1f} Mbit/s"
        ul_txt = "-" if upload is None else f"{upload:.1f} Mbit/s"
        header = "+--------------+--------+-----------+--------------+-------------+"
        titles = "| Tool         | Ping   | Jitter    | Download     | Upload      |"
        row = (
            f"| {'Speedtest':<12} | {ping_txt:<6} | {jitter_txt:<9} | "
            f"{dl_txt:<12} | {ul_txt:<11} |"
        )
        return "\n".join([header, titles, header, row, header])


try:
    from netmon.netmon_security_parser import NetmonSecurityParser
except ModuleNotFoundError:
    try:
        from tools.netmon_security_parser import NetmonSecurityParser
    except ModuleNotFoundError:
        NetmonSecurityParser = None

try:
    from netmon.netmon_security_store import NetmonSecurityStore
except ModuleNotFoundError:
    try:
        from tools.netmon_security_store import NetmonSecurityStore
    except ModuleNotFoundError:
        NetmonSecurityStore = None

try:
    import tkinter as tk
    from tkinter import ttk, messagebox, filedialog, simpledialog
except ModuleNotFoundError as exc:
    tk = None
    ttk = messagebox = filedialog = simpledialog = None
    TK_IMPORT_ERROR = exc
else:
    TK_IMPORT_ERROR = None

try:
    import psutil
except ModuleNotFoundError:
    psutil = None

try:
    from PIL import Image, ImageTk
except Exception as exc:
    Image = None
    ImageTk = None
    PIL_IMPORT_ERROR = str(exc)
else:
    PIL_IMPORT_ERROR = ""

try:
    import PyPDF2
except Exception as exc:
    PyPDF2 = None
    PYPDF2_IMPORT_ERROR = str(exc)
else:
    PYPDF2_IMPORT_ERROR = ""

try:
    import docx
except Exception as exc:
    docx = None
    DOCX_IMPORT_ERROR = str(exc)
else:
    DOCX_IMPORT_ERROR = ""

TkBase = tk.Tk if tk is not None else object
TkError = tk.TclError if tk is not None else RuntimeError


# ---------- PATH Fix (wichtig bei Start über Desktop/Icon) ----------
DEFAULT_PATH = "/usr/local/sbin:/usr/local/bin:/usr/sbin:/usr/bin:/sbin:/bin"
cur_path = os.environ.get("PATH", "")
if "/usr/sbin" not in cur_path:
    os.environ["PATH"] = DEFAULT_PATH


def _desktop_bool(value):
    return (value or "").strip().lower() in ("1", "true", "yes")


def run_headless_snapshot():
    """Print a compact status snapshot without starting Tk."""
    print("NetMon Headless Snapshot")
    print("=" * 24)

    rc, out = run_cmd(["ip", "-4", "route", "show", "default"], timeout=5)
    default_route = out.splitlines()[0] if rc == 0 and out else "n/a"
    print(f"Default route: {default_route}")

    if has_cmd("wg"):
        rc, out = run_cmd(["wg", "show", "interfaces"], timeout=5)
        wg_ifaces = out.strip() if rc == 0 and out.strip() else "none"
        print(f"WireGuard ifaces: {wg_ifaces}")
    else:
        print("WireGuard ifaces: wg not installed")

    if has_cmd("ufw"):
        rc, out = run_cmd(["ufw", "status"], timeout=8)
        first_line = out.splitlines()[0] if out else "n/a"
        print(f"UFW: {first_line} (rc={rc})")
    else:
        print("UFW: not installed")

    rc, out = run_cmd(["curl", "-s", "--max-time", "3", "ifconfig.me"], timeout=5)
    pub_ip = out.strip() if rc == 0 and out.strip() else "n/a"
    print(f"Public IP: {pub_ip}")


class App(TkBase):
    def __init__(self, start_theme: str | None = None):
        if tk is None:
            raise RuntimeError("tkinter is not available")
        super().__init__()

        self.debug_mode = False

        # ----- Locks gegen "Schaltchaos" -----
        self.ufw_lock = threading.Lock()
        self.network_ops_lock = threading.RLock()
        self.network_active_operation = ""
        self._network_op_owner = None
        self._network_op_depth = 0
        self.network_ops_coordinator = None
        if NetworkOperationCoordinator is not None:
            self.network_ops_coordinator = NetworkOperationCoordinator(
                lock=self.network_ops_lock,
                log_busy_fn=self.log_safe,
            )
            self._sync_network_operation_state_from_coordinator()

        self.network_service = None
        if NetmonNetworkService is not None:
            self.network_service = NetmonNetworkService(run_cmd)

        self.state_coordinator = None
        if NetmonStateCoordinator is not None:
            self.state_coordinator = NetmonStateCoordinator()

        self.security_parser = None
        if NetmonSecurityParser is not None:
            self.security_parser = NetmonSecurityParser()

        self.security_store = None
        if NetmonSecurityStore is not None:
            self.security_store = NetmonSecurityStore()

        self.ai_service = AiAnalysisService() if AiAnalysisService is not None else None
        self.ai_online_service = (
            AiOnlineService() if AiOnlineService is not None else None
        )
        self.startup_guard_service = (
            StartupGuardService() if StartupGuardService is not None else None
        )
        self.path_guard_service = (
            PathGuardService() if PathGuardService is not None else None
        )
        self.openai_api_key_var = tk.StringVar(value="")
        self.openai_api_key_saved = False
        self.chad_analysis_current_prompt = ""
        self.chad_online_result = ""
        self.chad_chat_history = []
        self.chad_pending_files = []
        self.chad_auto_scroll_var = tk.BooleanVar(value=True)
        self.chad_enabled_var = tk.BooleanVar(value=True)
        self._chad_context_menu_attached = set()
        self._selection_copy_shortcuts_attached = set()
        self._active_context_menu = None
        self._active_context_menu_after_id = None
        self.keyring_available = False
        self.keyring_error = ""
        self.keyring_backend_name = ""
        self.chad_analysis_source_var = tk.StringVar(value="Quelle: –")
        self.chad_analysis_status_var = tk.StringVar(
            value="Noch keine Analyse vorbereitet."
        )
        self.chad_online_status_var = tk.StringVar(value="Online-Analyse: inaktiv")
        self.chad_chat_status_var = tk.StringVar(value="Chat bereit.")

        self.isp_stealth_lock = threading.Lock()
        self.isp_stealth_running = False

        # ---------------- Desktop-Scaling / Kompakt ----------------
        try:
            self.tk.call("tk", "scaling", 0.90)
        except Exception:
            pass

        self.style = ttk.Style(self)
        self.style.configure("TButton", padding=(6, 3))
        self.style.configure("LinkSlot.TButton", padding=(12, 10), width=22)
        self.style.configure("TCheckbutton", padding=(4, 2))
        self.style.configure("TNotebook.Tab", foreground="#111111")
        self.style.map(
            "TNotebook.Tab",
            foreground=[("selected", "#000000"), ("!selected", "#1a1a1a")],
        )
        self._init_system_monitor_styles()
        self._app_theme = None
        self._theme_toggle_btn = None
        self._theme_name_var = tk.StringVar(value="")
        self._theme_start_override = (
            str(start_theme).strip().lower() if str(start_theme or "").strip() else ""
        )
        self._preferred_theme_name = "light"

        self.title(
            f"NetMon - Firewall, VPN, DNS, Trace    Version {NETMON_APP_VERSION}"
        )

        # Fenstergröße: 100% vom Screen, maximale Größe fix auf Screen
        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        self.geometry(f"{sw}x{sh}+0+0")
        self.minsize(900, 600)
        self.maxsize(sw, sh)
        self.bind("<Button-1>", self._handle_app_click_for_context_menu, add="+")

        # state
        self.last_dns_consistency_ok = None  # True/False/None
        self.last_dns_consistency_ts = 0.0
        self.dns_consistency_running = False
        self.dns_consistency_enabled = False
        self.dns_consistency_one_shot = False
        self.dns_mode = "auto"
        self.dns_mode_running = False
        self.dns_mode_var = tk.StringVar(value=DNS_MODE_LABELS["auto"])
        self.dns_mode_status_var = tk.StringVar(value="Aktiver DNS-Modus: AUTO")
        self.dns_resolver_status_var = tk.StringVar(value="Aktiver Resolver: –")
        self.dns_conflict_status_var = tk.StringVar(value="Konflikte: keine")
        self.dns_control_hint_var = tk.StringVar(value="DNS-Auswahl aktiv.")
        self.dns_guard_enabled_var = tk.BooleanVar(value=False)
        self.dns_guard_status_var = tk.StringVar(value="DNS Guard: inaktiv")
        self.dns_guard_last_action_var = tk.StringVar(value="Letzte Korrektur: –")
        self.dns_guard_last_fix_ts = 0.0
        self.last_dns_guard_repair_ts = 0.0
        self.dns_guard_allowed_dns = []
        self.dns_guard_check_running = False
        self.last_dns_guard_check_ts = 0.0
        self.dns_guard_transition_until = 0.0
        self.dns_guard_state = "off"
        self.last_dns_guard_state = None
        self.last_vpn_state = False
        self.dot_enabled = False
        self.dot_last_test_ok = None
        self.dot_last_test_ts = 0.0
        self.dot_running = False
        self.dot_vpn_conflict = False
        self.dot_vpn_dns_active = False
        self.dot_status_var = tk.StringVar(value="DoT: unbekannt")
        self.dot_server_var = tk.StringVar(value="Server: –")
        self.dot_mode_var = tk.StringVar(value="Modus: –")
        self.dot_profile_var = tk.StringVar(value="Profil: –")
        self.dot_last_test_var = tk.StringVar(value="Letzter Test: –")
        self.linux_help_status_var = tk.StringVar(
            value=self.tr(
                "linux_help_status_preparing",
                "Linux Hilfe: Vorbereitung aktiv – Integration folgt später.",
            )
        )
        self.linux_help_command_var = tk.StringVar(
            value=self.tr("linux_help_no_command_selected", "Kein Befehl ausgewählt.")
        )
        self.linux_help_risk_var = tk.StringVar(
            value=self.tr("linux_help_risk_unknown", "Risiko: –")
        )
        self.linux_help_desc_var = tk.StringVar(
            value=self.tr(
                "linux_help_desc_placeholder",
                "Hier erscheinen später Erklärungen zu ausgewählten Linux-Befehlen.",
            )
        )
        self.linux_help_category_var = tk.StringVar(value="Allgemein")
        self.linux_help_data = []
        self.linux_help_commands = []
        self.linux_help_categories = []
        self.linux_help_commands_by_category = {}
        self.linux_help_visible_commands = []
        self.linux_help_visible_command_ids = []
        self.linux_help_selected_id = ""
        self.linux_help_dirty = False
        self.linux_help_data_path = ""
        self.linux_help_import_path = ""
        self.linux_help_form_category_var = tk.StringVar(value="")
        self.linux_help_form_name_var = tk.StringVar(value="")
        self.linux_help_form_info_var = tk.StringVar(value="")
        self.linux_help_form_cmd_var = tk.StringVar(value="")
        self.linux_help_form_risk_var = tk.StringVar(value="caution")
        self.linux_help_form_explanation_var = tk.StringVar(value="")
        self.linux_help_form_warning_var = tk.StringVar(value="")
        self.linux_help_form_enabled_var = tk.BooleanVar(value=True)

        self.last_ip_ts = 0.0
        self.ip_country_running = False
        self.public_ip = "–"
        self.country = "–"

        self.last_wg_transfer = None
        self.last_wg_ts = None
        self.wg_rx_rate = 0.0
        self.wg_tx_rate = 0.0

        self.nm_dns_header_text = "NM DNS: –"

        # privacy: external calls toggle (safer default)
        self.auto_external_checks = tk.BooleanVar(value=False)
        self.wg_profiles = []
        self.feature_buttons = {}
        self.dep_status = {}
        self.sudo_nopasswd = False
        self.danger_guard_enabled = True
        self.danger_confirm_cooldown_s = 15
        self._danger_last_confirm = {}
        self.netmon_config_dir = get_netmon_config_dir()
        self.netmon_state_dir = get_netmon_state_dir()
        self.linux_help_data_path = os.path.join(
            self.netmon_config_dir, "linux_help_commands.json"
        )
        try:
            os.makedirs(self.netmon_config_dir, exist_ok=True)
            os.makedirs(self.netmon_state_dir, exist_ok=True)
        except Exception:
            pass

        self.hdd_config_path = os.path.join(
            self.netmon_config_dir, HDD_SLOT_CONFIG_FILE
        )
        self.link_config_path = os.path.join(
            self.netmon_config_dir, LINK_SLOT_CONFIG_FILE
        )
        self.vpn_config_path = os.path.join(
            self.netmon_config_dir, VPN_SLOT_CONFIG_FILE
        )
        self.primary_vpn_config_path = os.path.join(
            self.netmon_config_dir, PRIMARY_VPN_SLOT_CONFIG_FILE
        )
        self.privacy_config_path = os.path.join(
            self.netmon_config_dir, PRIVACY_CONFIG_FILE
        )
        self.netmon_config_path = os.path.join(self.netmon_config_dir, "config.json")
        self.ai_config_path = os.path.join(self.netmon_config_dir, "ai_config.json")
        self.clamav_findings_path = os.path.join(
            self.netmon_state_dir, CLAMAV_FINDINGS_FILE
        )
        self.security_baseline_path = os.path.join(
            self.netmon_state_dir, SECURITY_BASELINE_FILE
        )
        self.startup_guard_baseline_path = os.path.join(
            self.netmon_state_dir, "startup_guard_baseline.json"
        )
        self.startup_guard_trustlist_path = os.path.join(
            self.netmon_state_dir, "startup_guard_trustlist.json"
        )
        self.packet_capture_dir = os.path.join(self.netmon_state_dir, "captures")
        try:
            os.makedirs(self.packet_capture_dir, exist_ok=True)
        except Exception:
            pass

        self.settings = self._default_settings_payload()
        self.load_settings()
        lang_base_path = os.path.join(os.path.dirname(__file__), "lang")
        self.lang_service = (
            LangService(lang_base_path) if LangService is not None else None
        )
        lang_code = str(self.settings.get("language", "de") or "de").strip().lower()
        self.lang = (
            self.lang_service.load(lang_code) if self.lang_service is not None else {}
        )
        self.var_language = tk.StringVar(
            value=str(self.settings.get("language", "de") or "de")
        )
        self.var_startup_guard_autorun = tk.BooleanVar(
            value=bool(self.settings.get("startup_guard_autorun", True))
        )
        self.var_path_guard_autorun = tk.BooleanVar(
            value=bool(self.settings.get("path_guard_autorun", False))
        )
        self.var_dns_guard_enabled = tk.BooleanVar(
            value=bool(self.settings.get("dns_guard_enabled", True))
        )
        self.var_auto_switch_to_chad = tk.BooleanVar(
            value=bool(self.settings.get("auto_switch_to_chad", True))
        )
        self.var_chad_preview_popup = tk.BooleanVar(
            value=bool(self.settings.get("chad_preview_popup", False))
        )
        self.dns_guard_enabled = bool(self.settings.get("dns_guard_enabled", True))
        self.dns_guard_enabled_var.set(self.dns_guard_enabled)

        legacy_root = os.path.dirname(__file__)
        legacy_repo_root = os.path.abspath(os.path.join(legacy_root, ".."))
        migration_map = {
            self.hdd_config_path: os.path.join(legacy_root, HDD_SLOT_CONFIG_FILE),
            self.link_config_path: os.path.join(legacy_root, LINK_SLOT_CONFIG_FILE),
            self.vpn_config_path: os.path.join(legacy_root, VPN_SLOT_CONFIG_FILE),
            self.primary_vpn_config_path: os.path.join(
                legacy_root, PRIMARY_VPN_SLOT_CONFIG_FILE
            ),
            self.privacy_config_path: os.path.join(legacy_root, PRIVACY_CONFIG_FILE),
            self.clamav_findings_path: os.path.join(legacy_root, CLAMAV_FINDINGS_FILE),
            self.security_baseline_path: os.path.join(
                legacy_root, SECURITY_BASELINE_FILE
            ),
            os.path.join(
                self.netmon_state_dir, "logs", "speedtest_history.jsonl"
            ): os.path.join(legacy_repo_root, "logs", "speedtest_history.jsonl"),
        }
        for new_path, old_path in migration_map.items():
            try:
                if os.path.isfile(new_path) or not os.path.isfile(old_path):
                    continue
                os.makedirs(os.path.dirname(new_path), exist_ok=True)
                shutil.copy2(old_path, new_path)
            except Exception:
                pass

        self.detected_disks = []
        self.detected_disks_by_device = {}
        self.hdd_slots = self.load_hdd_slots_config()
        self.link_slots = self.load_link_slots_config()
        self.vpn_slots = self.load_vpn_slots_config()
        self.primary_vpn_slots = self.load_primary_vpn_slots_config()
        self.desktop_app_entries = self._load_desktop_app_entries()
        self.desktop_app_labels = [entry["label"] for entry in self.desktop_app_entries]
        self.desktop_app_by_label = {
            entry["label"]: entry for entry in self.desktop_app_entries
        }
        self.desktop_app_by_id = {
            entry["desktop_id"]: entry
            for entry in self.desktop_app_entries
            if entry.get("desktop_id")
        }
        self.packet_capture_running = False
        self.packet_capture_cancel_requested = False
        self.packet_capture_active_process = None
        self.packet_capture_text_after_id = None
        self.packet_capture_started_at = 0.0
        self.packet_capture_output_path = ""
        self.packet_capture_interfaces = []
        self.packet_capture_interface_var = tk.StringVar(value="")
        self.packet_capture_filter_var = tk.StringVar(value="")
        self.packet_capture_file_var = tk.StringVar(value="")
        self.packet_capture_status_var = tk.StringVar(value="bereit")
        self.packet_capture_runtime_var = tk.StringVar(value="-")
        self.packet_capture_packets_var = tk.StringVar(value="-")
        self.packet_capture_size_var = tk.StringVar(value="-")
        self.lynis_running = False
        self.lynis_active_process = None
        self.lynis_status_var = tk.StringVar(value="bereit")
        self.lynis_summary_var = tk.StringVar(value="Noch kein Audit ausgeführt.")
        self.lynis_last_result = {}
        self.lynis_last_output = ""
        self.rkhunter_running = False
        self.rkhunter_updating = False
        self.rkhunter_active_process = None
        self.rkhunter_status_var = tk.StringVar(value="bereit")
        self.rkhunter_summary_var = tk.StringVar(value="Noch kein Scan ausgeführt.")
        self.rkhunter_last_update_var = tk.StringVar(value="-")
        self.rkhunter_last_result = {}
        self.rkhunter_repair_available = False
        self.startup_guard_running = False
        self.startup_guard_status_var = tk.StringVar(
            value="Startup Guard bereit. Noch keine Baseline geladen."
        )
        self.startup_guard_last_state = "unknown"
        self.startup_guard_last_result = {}
        self.path_guard_status_var = tk.StringVar(
            value="PATH Guard bereit. Noch keine Prüfung ausgeführt."
        )
        self.path_guard_last_state = "unknown"
        self.path_guard_last_result = None
        self.startup_guard_visible_entries = []
        self.startup_guard_selected_entry = None
        self.hdd_slot_lights = []
        self.hdd_slot_states = [None] * len(self.hdd_slots)
        self.hdd_slot_name_buttons = []
        self.link_slot_buttons = []
        self.vpn_slot_buttons = []
        self.vpn_slot_lights = []
        self.hdd_selected_device = None
        self.hdd_luks_mappers = {}
        self.hdd_refresh_running = False
        self.hdd_name_editing = False
        self.hdd_name_dirty = False
        self.hdd_last_loaded_name = ""
        self._hdd_name_programmatic_update = False
        self.security_monitoring = False
        self.security_scan_running = False
        self.security_clamav_running = False
        self.security_clamav_update_running = False
        self.clamav_cancel_requested = False
        self.clamav_active_process = None
        self.nmap_running = False
        self.nmap_cancel_requested = False
        self.nmap_active_process = None
        self.nmap_status_var = tk.StringVar(value="bereit")
        self.nmap_error_var = tk.StringVar(value="")
        self.port_reality_default_ports = (22, 80, 443)
        self.port_reality_last_scan = {"host": "", "results": [], "ts": 0.0}
        self.speedtest_last_result = {}
        self.speedtest_history = []
        self.speedtest_history_limit = 5
        self._load_speedtest_history()
        self.log_overlay_active = False
        self.log_overlay_image_path = ""
        self.log_overlay_source_image = None
        self.log_overlay_source_photo = None
        self.log_overlay_photo = None
        self.log_overlay_picker_open = False
        self.log_overlay_picker_dir = ""
        self.log_overlay_picker_files = []
        self.log_overlay_picker_index = -1
        self.speedtest_ping_var = tk.StringVar(value="-")
        self.speedtest_jitter_var = tk.StringVar(value="-")
        self.speedtest_download_var = tk.StringVar(value="-")
        self.speedtest_upload_var = tk.StringVar(value="-")
        self.speedtest_timestamp_var = tk.StringVar(value="-")

        # LCD clock + alarm (system time based)
        self.lcd_time_var = tk.StringVar(value="--:--:--")
        self.lcd_date_var = tk.StringVar(value="--.--.----")
        self.lcd_tz_var = tk.StringVar(value="TZ: -")
        self.alarm_enabled_var = tk.BooleanVar(value=False)
        self.alarm_time_var = tk.StringVar(value="07:00")
        self.alarm_status_var = tk.StringVar(value="Wecker aus")
        self._clock_after_id = None
        self._alarm_last_fired_key = ""
        self._alarm_invalid_logged = False

        self.clamav_status_var = tk.StringVar(value="nicht installiert")
        self.clamav_signatures_var = tk.StringVar(value="unbekannt")
        self.clamav_last_scan_var = tk.StringVar(value="-")
        self.clamav_last_target_var = tk.StringVar(value="-")
        self.clamav_findings_var = tk.StringVar(value="-")
        self.clamav_result_var = tk.StringVar(value="-")
        self.clamav_runtime_var = tk.StringVar(value="-")
        self.clamav_scanned_files_var = tk.StringVar(value="0")
        self.clamav_speed_var = tk.StringVar(value="-")
        self.clamav_eta_var = tk.StringVar(value="-")
        self.clamav_scan_started_at = 0.0
        self.clamav_scanned_files = 0
        self.clamav_expected_files = None
        self.clamav_progress_after_id = None
        self.clamav_quarantine_dir = os.path.expanduser("~/.netmon_quarantine")
        self.clamav_findings = []
        self._load_clamav_findings()
        self.security_events = []
        self.security_event_keys = set()
        self.security_visible_events = []
        self.security_filter_mode = "all"
        self.security_last_alarm = "–"
        self.security_log_level = "–"
        self.security_last_portscan_alert = {}
        self.security_aggregated = {}  # Für aggregierte interne Events
        # Für Deduplizierung: key -> (count, last_time)
        self.security_dedup_cache = {}

        self.integrity_baseline = {}
        self.device_baseline = {}
        self.input_access_baseline = {}
        self.security_input_access_seen = {}
        self.security_input_access_whitelist = {
            "systemd",
            "kthreadd",
            "kworker",
            "udevd",
            "dbus-daemon",
            "agetty",
            "ssh-agent",
            "gnome-keyring-daemon",
            "xdg-desktop-portal",
            "bash",
            "sh",
            "gnome-shell",
        }
        self.tooltips_enabled = tk.BooleanVar(value=True)
        self.strict_protection_var = tk.BooleanVar(value=True)
        self._tooltips = []
        self.system_monitor_active = True
        self.system_monitor_interval_ms = 1500
        self.system_monitor_psutil_hint_logged = False
        self.sysinfo_static = {}

        # Hop Lookup (Land/Provider) – optional (Privacy)
        self.hop_lookup_enabled = tk.BooleanVar(value=False)
        self.hop_cache = {}  # ip -> (country_code, org)
        self.hop_lookup_warned_ips = set()

        # ISP Stealth Mode
        self.isp_stealth_enabled = tk.BooleanVar(value=False)
        self.isp_stealth_autostart = tk.BooleanVar(value=False)
        self.isp_stealth_disable_ipv6 = tk.BooleanVar(value=False)
        self.isp_stealth_enable_killswitch = tk.BooleanVar(value=False)
        self.isp_stealth_vpn_profile = tk.StringVar(value="")
        self.isp_stealth_vpn_choice = tk.StringVar(value="")
        self.isp_stealth_dns_profile = tk.StringVar(value="System (DHCP/auto)")
        self.privacy_config = self.load_privacy_config()
        self._apply_privacy_config_to_vars(self.privacy_config)
        self._apply_settings_runtime_state()
        self.isp_stealth_runtime_state = None

        # Internet AUS/AN Zustandsspeicherung
        self.internet_suspend_state = self._default_internet_suspend_state()

        self.protocol("WM_DELETE_WINDOW", self.on_close)

        # ---------- Global scroll (whole page) ----------
        outer = ttk.Frame(self)
        outer.pack(fill="both", expand=True)

        self.page_canvas = tk.Canvas(outer, highlightthickness=0)
        self.page_canvas.pack(side="left", fill="both", expand=True)

        page_scroll = ttk.Scrollbar(
            outer, orient="vertical", command=self.page_canvas.yview
        )
        page_scroll.pack(side="right", fill="y")

        self.page_canvas.configure(yscrollcommand=page_scroll.set)

        self.page = ttk.Frame(self.page_canvas, padding=0)
        self.page_window = self.page_canvas.create_window(
            (0, 0), window=self.page, anchor="nw"
        )

        def _on_page_config(_event):
            self.page_canvas.configure(scrollregion=self.page_canvas.bbox("all"))

        def _on_canvas_config(event):
            self.page_canvas.itemconfig(self.page_window, width=event.width)

        self.page.bind("<Configure>", _on_page_config)
        self.page_canvas.bind("<Configure>", _on_canvas_config)

        def _bind_scroll(_e=None):
            self.page_canvas.bind_all("<Button-4>", self._scroll_up)
            self.page_canvas.bind_all("<Button-5>", self._scroll_down)

        def _unbind_scroll(_e=None):
            self.page_canvas.unbind_all("<Button-4>")
            self.page_canvas.unbind_all("<Button-5>")

        self.page_canvas.bind("<Enter>", _bind_scroll)
        self.page_canvas.bind("<Leave>", _unbind_scroll)

        # ---------- Build content inside self.page ----------
        self.page.columnconfigure(1, weight=1)

        header = ttk.Frame(self.page, padding=10)
        header.grid(row=0, column=0, columnspan=2, sticky="ew")
        header.columnconfigure(0, weight=1)
        header.columnconfigure(1, weight=0)
        header.columnconfigure(2, weight=0)

        ttk.Label(
            header,
            text=self.tr("header_status", "Status"),
            font=("TkDefaultFont", 12, "bold"),
        ).grid(row=0, column=0, sticky="w")
        self.status_text = ttk.Label(
            header,
            text="–",
            wraplength=520,
            justify="left",
            anchor="w",
        )
        self.status_text.grid(row=0, column=1, sticky="w", padx=10)
        header.columnconfigure(1, weight=0, minsize=520)

        self.header_clock_label = tk.Label(
            header,
            textvariable=self.lcd_time_var,
            font=("Courier", 14, "bold"),
            bg="#0a0f0a",
            fg="#7CFF6B",
            padx=10,
            pady=3,
            bd=1,
            relief="solid",
        )
        self.header_clock_label._preserve_theme_colors = True
        self.header_clock_label.place(relx=0.5, y=6, anchor="n")

        self.dns_header = ttk.Label(header, text=self.nm_dns_header_text)
        self.dns_header.grid(row=1, column=0, sticky="w", pady=(6, 0))

        systemf = ttk.Frame(header)
        systemf.grid(row=1, column=1, columnspan=2, sticky="e", pady=(4, 0))
        ttk.Label(systemf, text=self.tr("header_system", "System")).pack(
            side="left", padx=(0, 6)
        )
        self.system_action_var = tk.StringVar(value="")
        self.system_action_combo = ttk.Combobox(
            systemf,
            textvariable=self.system_action_var,
            values=[
                "Reboot",
                "Abmelden",
                "Ausschalten",
                "Ausschalten - 15 Min",
                "Ausschalten - 30 Min",
                "Ausschalten - 45 Min",
                "Ausschalten - 60 Min",
                "Shutdown abbrechen",
            ],
            state="readonly",
            width=22,
        )
        self.system_action_combo.pack(side="left")
        self.system_action_button = ttk.Button(
            systemf,
            text=self.tr("btn_execute", "Ausführen"),
            command=self.execute_selected_system_action,
        )
        self.system_action_button.pack(side="left", padx=(6, 0))
        self.add_tooltip(
            self.system_action_button,
            "Führt die gewählte Systemaktion erst nach Sicherheitsabfrage aus.",
        )

        self.copyright_label = tk.Label(
            header,
            text="© by Chad, Codi und Happy  free to use and no warranty",
            fg="red",
            font=("TkDefaultFont", 8),
        )
        self.copyright_label._preserve_theme_colors = True
        self.copyright_label.grid(
            row=2, column=0, columnspan=3, sticky="e", pady=(2, 0)
        )

        # main area
        self.page.columnconfigure(0, weight=0, minsize=260)
        self.page.columnconfigure(1, weight=1)

        left = ttk.Frame(self.page, padding=10, width=260)
        left.grid(row=1, column=0, sticky="nsew")
        left.grid_propagate(False)
        left.columnconfigure(0, weight=1)

        right = ttk.Frame(self.page, padding=10)
        right.grid(row=1, column=1, sticky="nsew")
        # Hauptbereich flexibel, zwei symmetrische Randspalten rechts.
        main_work_width = 860
        side_panel_width = 220
        right.columnconfigure(0, weight=1, minsize=main_work_width)
        right.columnconfigure(1, weight=0, minsize=side_panel_width)
        right.columnconfigure(2, weight=0, minsize=side_panel_width)
        right.rowconfigure(0, weight=1)

        self.main_work_area = ttk.Frame(right, width=main_work_width)
        self.main_work_area.grid(row=0, column=0, sticky="nsew")
        self.main_work_area.columnconfigure(0, weight=1, minsize=main_work_width)
        self.main_work_area.rowconfigure(1, weight=1)

        self.right_button_space = ttk.Frame(right, padding=(12, 0, 0, 0))
        self.right_button_space.grid(row=0, column=1, sticky="nw")

        self.right_reserve_space = ttk.Frame(right, padding=(12, 0, 0, 0))
        self.right_reserve_space.grid(row=0, column=2, sticky="nw")
        self.right_reserve_space.columnconfigure(0, weight=1, minsize=side_panel_width)

        # Tooltip-Steuerung (oben links, über Ampel)
        tipf = ttk.Frame(left)
        tipf.pack(fill="x", pady=(0, 8))
        ttk.Label(tipf, text=self.tr("label_tooltips", "Tooltips:")).pack(side="left")
        ttk.Checkbutton(
            tipf,
            text=self.tr("check_on_off", "AN / AUS"),
            variable=self.tooltips_enabled,
            onvalue=True,
            offvalue=False,
            command=self.on_tooltip_toggle,
        ).pack(side="left", padx=(8, 0))

        # Ampel
        amp = ttk.LabelFrame(left, text=self.tr("group_ampel", "Ampel"), padding=10)
        amp.pack(fill="x")

        inet_row = ttk.Frame(amp)
        inet_row.pack(fill="x", pady=4)
        inet_canvas = tk.Canvas(inet_row, width=18, height=18, highlightthickness=0)
        inet_canvas.pack(side="left")
        inet_dot = inet_canvas.create_oval(2, 2, 16, 16, fill="gray")
        ttk.Label(inet_row, text=self.tr("label_internet", "Internet")).pack(
            side="left", padx=8
        )

        btn_inet_off = ttk.Button(
            inet_row,
            text=self.tr("btn_off", "AUS"),
            command=self.internet_disable,
            width=5,
        )
        btn_inet_off.pack(side="left", padx=2)
        self.register_feature_button("nmcli", btn_inet_off)
        self.add_tooltip(
            btn_inet_off,
            "Deaktiviert die Netzwerkverbindung systemweit über NetworkManager.",
        )

        btn_inet_on = ttk.Button(
            inet_row,
            text=self.tr("btn_on", "AN"),
            command=self.internet_enable,
            width=5,
        )
        btn_inet_on.pack(side="left", padx=2)
        self.register_feature_button("nmcli", btn_inet_on)
        self.add_tooltip(
            btn_inet_on,
            "Aktiviert die Netzwerkverbindung systemweit wieder über NetworkManager.",
        )

        btn_ip_country = ttk.Button(
            inet_row,
            text=self.tr("btn_get_ip", "IP holen"),
            command=lambda: self.ip_country_async(force=True),
        )
        btn_ip_country.pack(side="right")
        self.add_tooltip(
            btn_ip_country,
            "Ermittelt deine öffentliche IP-Adresse und das zugehörige Land.",
        )
        vpn_row = ttk.Frame(amp)
        vpn_row.pack(fill="x", pady=4)
        vpn_canvas = tk.Canvas(vpn_row, width=18, height=18, highlightthickness=0)
        vpn_canvas.pack(side="left")
        vpn_dot = vpn_canvas.create_oval(2, 2, 16, 16, fill="gray")
        ttk.Label(
            vpn_row,
            text=self.tr("label_vpn_wireguard", "VPN (WireGuard)"),
        ).pack(side="left", padx=8)
        btn_whoami = ttk.Button(
            vpn_row,
            text=self.tr("btn_whoami", "Whoami"),
            command=lambda: self.whoami_async(force=True),
        )
        btn_whoami.pack(side="right")
        self.add_tooltip(
            btn_whoami,
            "Zeigt, wie dein Anschluss nach außen im Internet sichtbar ist.",
        )
        btn_terminal = ttk.Button(
            vpn_row,
            text=self.tr("btn_terminal", "Terminal"),
            command=self.open_external_terminal,
        )
        btn_terminal.pack(side="right", padx=(0, 8))
        self.add_tooltip(
            btn_terminal,
            "Öffnet ein externes System-Terminal.",
        )
        self.l_inet = (inet_canvas, inet_dot)
        self.l_vpn = (vpn_canvas, vpn_dot)
        # UI Cleanup inventory (Ampel):
        # Kernstatus: Firewall, DNS Guard, Startup Guard, DoT
        # Zusatzstatus: Killswitch, Security, IPv6, SSH
        amp_core = ttk.LabelFrame(
            amp,
            text=self.tr("group_core_status", "Kernstatus"),
            padding=6,
        )
        amp_core.pack(fill="x", pady=(6, 0))
        amp_extra = ttk.LabelFrame(
            amp,
            text=self.tr("group_extra_status", "Weitere Status"),
            padding=6,
        )
        amp_extra.pack(fill="x", pady=(8, 0))

        self.l_fw = self.make_light(amp_core, "Firewall")
        self.l_dnscons = self.make_light(amp_core, "DNS Guard")
        self.l_startup_guard = self.make_light(amp_core, "Startup Guard")
        self.l_path_guard = self.make_light(amp_extra, "PATH Guard")
        self.add_tooltip(
            self.l_startup_guard[0],
            self.tr(
                "tooltip_startup_guard_light",
                "Startup Guard überwacht Autostart- und systemd-Startup-Einträge.\n"
                "Grün: keine Änderungen\n"
                "Gelb: neue/geänderte/entfernte Einträge\n"
                "Rot: Prüf-/Baseline-Fehler\n"
                "Grau: keine Baseline vorhanden",
            ),
        )
        self.add_tooltip(
            self.l_startup_guard[2],
            self.tr(
                "tooltip_startup_guard_light",
                "Startup Guard überwacht Autostart- und systemd-Startup-Einträge.\n"
                "Grün: keine Änderungen\n"
                "Gelb: neue/geänderte/entfernte Einträge\n"
                "Rot: Prüf-/Baseline-Fehler\n"
                "Grau: keine Baseline vorhanden",
            ),
        )
        self.l_startup_guard[0].bind(
            "<Button-1>", lambda _e: self.open_startup_guard_tab()
        )
        self.l_startup_guard[2].bind(
            "<Button-1>", lambda _e: self.open_startup_guard_tab()
        )
        self.add_tooltip(
            self.l_path_guard[0],
            self.tr(
                "tooltip_path_guard_light",
                "PATH Guard bewertet PATH-Einträge und Shell-Profile.\n"
                "Grün: Risiko niedrig\n"
                "Gelb: Risiko mittel\n"
                "Rot: Risiko hoch\n"
                "Grau: noch nicht geprüft",
            ),
        )
        self.add_tooltip(
            self.l_path_guard[2],
            self.tr(
                "tooltip_path_guard_light",
                "PATH Guard bewertet PATH-Einträge und Shell-Profile.\n"
                "Grün: Risiko niedrig\n"
                "Gelb: Risiko mittel\n"
                "Rot: Risiko hoch\n"
                "Grau: noch nicht geprüft",
            ),
        )
        self.l_path_guard[0].bind(
            "<Button-1>", lambda _e: self._select_security_tab(self.tab_security)
        )
        self.l_path_guard[2].bind(
            "<Button-1>", lambda _e: self._select_security_tab(self.tab_security)
        )
        self.dns_guard_indicator = self.l_dnscons
        self.dns_guard_label = self.l_dnscons[2]
        self.add_tooltip(
            self.dns_guard_label,
            "DNS Guard überwacht die aktiven DNS-Server.\n\n"
            "Grün:\nNur erlaubte DNS aktiv (z. B. VPN-DNS)\n\n"
            "Rot:\nFremde DNS erkannt -> mögliche Leaks\n\n"
            "Grau:\nDNS Guard deaktiviert",
        )
        self.add_tooltip(
            self.dns_guard_indicator[0],
            "DNS Guard überwacht die aktiven DNS-Server.\n\n"
            "Grün:\nNur erlaubte DNS aktiv (z. B. VPN-DNS)\n\n"
            "Rot:\nFremde DNS erkannt -> mögliche Leaks\n\n"
            "Grau:\nDNS Guard deaktiviert",
        )
        self.dns_guard_indicator[0].bind(
            "<Button-1>", lambda _e: self.show_dns_guard_details()
        )
        self.dns_guard_label.bind(
            "<Button-1>", lambda _e: self.show_dns_guard_details()
        )
        self.l_dot = self.make_light(amp_core, "DoT")
        self.l_ks = self.make_light(amp_extra, "Killswitch")
        self.l_sec = self.make_light(amp_extra, "Security")
        self.l_ipv6 = self.make_light(amp_extra, "IPv6")
        self.l_ssh = self.make_light(amp_extra, "SSH")

        # Steuerbereich links: zwei kompakte Spalten
        control_area = ttk.Frame(left)
        control_area.pack(fill="x", pady=10)
        control_area.columnconfigure(0, weight=1, uniform="left_controls")
        control_area.columnconfigure(1, weight=1, uniform="left_controls")

        left_controls = ttk.Frame(control_area, padding=(0, 0, 6, 0))
        left_controls.grid(row=0, column=0, sticky="new")
        left_controls.columnconfigure(0, weight=1)

        right_controls = ttk.Frame(control_area, padding=(6, 0, 0, 0))
        right_controls.grid(row=0, column=1, sticky="new")
        right_controls.columnconfigure(0, weight=1)

        # UI Cleanup inventory (Seitenblock):
        # Standard / Core: Security, DNS
        # Diagnose: Trace, MTR, Nmap
        right_core_group = ttk.LabelFrame(
            right_controls,
            text=self.tr("group_standard_core", "Standard / Core"),
            padding=8,
        )
        right_core_group.grid(row=0, column=0, sticky="new")
        right_core_group.columnconfigure(0, weight=1)

        right_diag_group = ttk.LabelFrame(
            right_controls,
            text=self.tr("group_diagnostics", "Diagnose"),
            padding=8,
        )
        right_diag_group.grid(row=1, column=0, sticky="new", pady=(10, 0))
        right_diag_group.columnconfigure(0, weight=1)

        firewallf = ttk.LabelFrame(
            left_controls,
            text=self.tr("group_firewall", "Firewall"),
            padding=10,
        )
        firewallf.grid(row=0, column=0, sticky="new")

        btn_fw_on = ttk.Button(
            firewallf,
            text=self.tr("btn_firewall_on", "Firewall EIN"),
            command=self.fw_enable,
        )
        btn_fw_on.pack(fill="x")
        self.register_feature_button("ufw", btn_fw_on)
        self.add_tooltip(
            btn_fw_on,
            "Aktiviert die Firewall. Unerwünschte Verbindungen werden dann blockiert.",
        )
        btn_fw_off = ttk.Button(
            firewallf,
            text=self.tr("btn_firewall_off", "Firewall AUS"),
            command=self.fw_disable,
        )
        btn_fw_off.pack(fill="x", pady=5)
        self.register_feature_button("ufw", btn_fw_off)
        self.add_tooltip(
            btn_fw_off,
            "Schaltet die Firewall aus. Verbindungen werden dann weniger streng gefiltert.",
        )
        btn_fw_status = ttk.Button(
            firewallf,
            text=self.tr("btn_firewall_status", "Firewall Status"),
            command=self.fw_status,
        )
        btn_fw_status.pack(fill="x")
        self.register_feature_button("ufw", btn_fw_status)
        self.add_tooltip(
            btn_fw_status,
            "Zeigt den aktuellen Zustand und die Regeln der Firewall an.",
        )

        vpnf = ttk.LabelFrame(
            left_controls, text=self.tr("group_vpn", "VPN"), padding=10
        )
        vpnf.grid(row=1, column=0, sticky="new", pady=(10, 0))

        self.btn_primary_vpn_1 = ttk.Button(
            vpnf,
            text=self.primary_vpn_slots[0].get("custom_name", "VPN CH"),
            command=lambda: self.open_primary_vpn_slot(1),
        )
        self.btn_primary_vpn_1.pack(fill="x")
        self.register_feature_button("wg", self.btn_primary_vpn_1)
        self.add_tooltip(
            self.btn_primary_vpn_1,
            "Startet die primäre VPN-Verbindung (#1).",
        )

        self.btn_primary_vpn_2 = ttk.Button(
            vpnf,
            text=self.primary_vpn_slots[1].get("custom_name", "VPN DE"),
            command=lambda: self.open_primary_vpn_slot(2),
        )
        self.btn_primary_vpn_2.pack(fill="x", pady=5)
        self.register_feature_button("wg", self.btn_primary_vpn_2)
        self.add_tooltip(
            self.btn_primary_vpn_2,
            "Startet die primäre VPN-Verbindung (#2).",
        )

        self.update_primary_vpn_buttons()

        btn_vpn_off = ttk.Button(
            vpnf,
            text=self.tr("btn_vpn_off", "VPN AUS"),
            command=self.vpn_off,
        )
        btn_vpn_off.pack(fill="x")
        self.register_feature_button("wg", btn_vpn_off)
        self.add_tooltip(btn_vpn_off, "Beendet alle aktiven VPN-Verbindungen.")
        btn_wg_reload = ttk.Button(
            vpnf,
            text=self.tr("btn_reload_wg_profiles", "WireGuard-Profile neu einlesen"),
            command=self.reload_wg_profiles,
        )
        btn_wg_reload.pack(fill="x", pady=(6, 0))
        self.register_feature_button("wg_profiles", btn_wg_reload)
        btn_wg_import = ttk.Button(
            vpnf,
            text=self.tr(
                "btn_import_wg_profiles",
                "WireGuard-Profile (.conf) einlesen",
            ),
            command=self.import_wg_profiles_from_folder,
        )
        btn_wg_import.pack(fill="x", pady=(6, 0))
        self.register_feature_button("wg_profiles", btn_wg_import)

        ksf = ttk.LabelFrame(
            left_controls,
            text=self.tr("group_killswitch", "Killswitch"),
            padding=10,
        )
        ksf.grid(row=2, column=0, sticky="new", pady=(10, 0))

        btn_ks_on = ttk.Button(
            ksf,
            text=self.tr("btn_killswitch_on", "Killswitch AN"),
            command=self.killswitch_on,
        )
        btn_ks_on.pack(fill="x")
        self.register_feature_button("ufw", btn_ks_on)
        self.add_tooltip(
            btn_ks_on,
            "Erlaubt ausgehenden Traffic nur noch über VPN. Schutz bei VPN-Ausfall.",
        )
        btn_ks_off = ttk.Button(
            ksf,
            text=self.tr("btn_killswitch_off", "Killswitch AUS"),
            command=self.killswitch_off,
        )
        btn_ks_off.pack(fill="x", pady=5)
        self.register_feature_button("ufw", btn_ks_off)
        self.add_tooltip(
            btn_ks_off,
            "Deaktiviert den Killswitch und erlaubt wieder normales Routing.",
        )

        ipv6f = ttk.LabelFrame(
            left_controls, text=self.tr("group_ipv6", "IPv6"), padding=10
        )
        ipv6f.grid(row=3, column=0, sticky="new", pady=(10, 0))

        self.btn_ipv6_on = ttk.Button(
            ipv6f,
            text=self.tr("btn_ipv6_on", "IPv6 AN"),
            command=self.ipv6_enable,
        )
        self.btn_ipv6_on.pack(fill="x")
        self.add_tooltip(
            self.btn_ipv6_on,
            "Aktiviert systemweites IPv6. Bei VPN-Profilen ohne IPv6 wird der Button deaktiviert.",
        )
        self.btn_ipv6_off = ttk.Button(
            ipv6f,
            text=self.tr("btn_ipv6_off", "IPv6 AUS"),
            command=self.ipv6_disable,
        )
        self.btn_ipv6_off.pack(fill="x", pady=6)
        self.add_tooltip(
            self.btn_ipv6_off,
            "Deaktiviert systemweites IPv6. Bei VPN-Profilen ohne IPv6 wird der Button deaktiviert.",
        )
        self.ipv6_controls_hint = ttk.Label(
            ipv6f, text="", foreground="#666666", wraplength=220, justify="left"
        )
        self.ipv6_controls_hint.pack(fill="x")

        # Vereinfachte Architektur: keine manuelle IPv4/IPv6-Profilumschaltung mehr im UI.
        ipv6f.grid_remove()

        sshf = ttk.LabelFrame(
            left_controls,
            text=self.tr("group_ssh_rules", "SSH-Regeln"),
            padding=10,
        )
        sshf.grid(row=4, column=0, sticky="new", pady=(10, 0))

        btn_ssh_add = ttk.Button(
            sshf,
            text=self.tr("btn_ssh_rule_add", "SSH Regel setzen (UFW)"),
            command=self.allow_ssh_rule_async,
        )
        btn_ssh_add.pack(fill="x")
        self.register_feature_button("ufw", btn_ssh_add)
        self.add_tooltip(
            btn_ssh_add,
            "Erlaubt eingehende SSH-Verbindungen ueber UFW auf Port 22/TCP.",
        )

        btn_ssh_rm = ttk.Button(
            sshf,
            text=self.tr("btn_ssh_rule_remove", "SSH Regel entfernen (UFW)"),
            command=self.remove_ssh_rule_async,
        )
        btn_ssh_rm.pack(fill="x", pady=(6, 0))
        self.register_feature_button("ufw", btn_ssh_rm)
        self.add_tooltip(
            btn_ssh_rm,
            "Entfernt SSH-Firewallregeln. Achtung: Remote-Zugriff kann dadurch getrennt werden.",
        )

        btn_port_reality = ttk.Button(
            sshf,
            text=self.tr("btn_port_reality_start", "Port Reality Check starten"),
            command=self.port_reality_check_async,
        )
        btn_port_reality.pack(fill="x", pady=(6, 0))
        self.add_tooltip(
            btn_port_reality,
            "Prüft offene Ports auf echte Dienste und erkennt pseudo-open Verhalten.",
        )

        btn_port_reality_last = ttk.Button(
            sshf,
            text=self.tr("btn_show_last_scan", "Letzten Scan anzeigen"),
            command=self.show_last_port_reality_scan,
        )
        btn_port_reality_last.pack(fill="x", pady=(6, 0))
        self.add_tooltip(
            btn_port_reality_last,
            "Zeigt die Ergebnisse des letzten Port Reality Checks im Log.",
        )

        connf = ttk.LabelFrame(
            left_controls,
            text=self.tr("group_connections", "Verbindungen"),
            padding=10,
        )
        connf.grid(row=5, column=0, sticky="new", pady=(10, 0))

        btn_conn = ttk.Button(
            connf,
            text=self.tr(
                "btn_refresh_connections",
                "Verbindungen (PID) aktualisieren",
            ),
            command=self.connections_async,
        )
        btn_conn.pack(fill="x")
        self.register_feature_button("ss", btn_conn)
        self.add_tooltip(
            btn_conn,
            "Lädt aktuelle Netzwerkverbindungen mit zugehörigen Prozessen neu.",
        )

        ttk.Button(
            connf,
            text=self.tr("btn_arp_table", "ARP Tabelle (arp -a)"),
            command=self.arp_async,
        ).pack(fill="x", pady=(6, 0))

        ttk.Button(
            connf,
            text=self.tr("btn_route_table", "Routing Tabelle (ip route)"),
            command=self.route_async,
        ).pack(fill="x", pady=(6, 0))

        iftopf = ttk.LabelFrame(
            left_controls,
            text=self.tr("group_iftop", "iftop"),
            padding=10,
        )
        iftopf.grid(row=6, column=0, sticky="new", pady=(10, 0))

        btn_iftop = ttk.Button(
            iftopf,
            text=self.tr("btn_live_snapshot", "Live Snapshot (10s)"),
            command=self.iftop_async,
        )
        btn_iftop.pack(fill="x")
        self.register_feature_button("iftop", btn_iftop)
        self.add_tooltip(
            btn_iftop,
            "Zeigt einen kurzen Live-Snapshot des aktuellen Netzwerkverkehrs.",
        )

        btn_htop = ttk.Button(
            iftopf,
            text=self.tr("btn_htop_snapshot", "htop Snapshot"),
            command=self.htop_async,
        )
        btn_htop.pack(fill="x", pady=(6, 0))
        self.register_feature_button("htop", btn_htop)
        self.add_tooltip(
            btn_htop,
            "Zeigt eine Momentaufnahme der laufenden Prozesse und Systemlast.",
        )

        secf = ttk.LabelFrame(
            right_core_group,
            text=self.tr("group_security", "Security"),
            padding=10,
        )
        secf.grid(row=0, column=0, sticky="new")
        btn_sec_start = ttk.Button(
            secf,
            text=self.tr("btn_monitoring_start", "Monitoring Start"),
            command=self.security_start_monitoring,
        )
        btn_sec_start.pack(fill="x")
        self.register_feature_button("security", btn_sec_start)
        self.add_tooltip(
            btn_sec_start,
            "Startet die Sicherheitsüberwachung von UFW-Block-Einträgen.",
        )
        btn_sec_stop = ttk.Button(
            secf,
            text=self.tr("btn_monitoring_stop", "Monitoring Stop"),
            command=self.security_stop_monitoring,
        )
        btn_sec_stop.pack(fill="x", pady=(4, 0))
        self.register_feature_button("security", btn_sec_stop)
        self.add_tooltip(
            btn_sec_stop,
            "Stoppt die Sicherheitsüberwachung. Neue Events werden dann nicht mehr eingelesen.",
        )
        btn_sec_log = ttk.Button(
            secf,
            text=self.tr("btn_ufw_logging_high", "UFW Logging HIGH"),
            command=self.security_set_ufw_logging_high,
        )
        btn_sec_log.pack(fill="x", pady=(4, 0))
        self.register_feature_button("security", btn_sec_log)
        self.add_tooltip(
            btn_sec_log,
            "Erhöht die UFW-Logstufe, damit mehr sicherheitsrelevante Ereignisse sichtbar sind.",
        )
        btn_sec_open = ttk.Button(
            secf,
            text=self.tr("btn_open_security", "Security öffnen"),
            command=self.security_open_tab,
        )
        btn_sec_open.pack(fill="x", pady=(4, 0))
        self.add_tooltip(
            btn_sec_open,
            "Öffnet den Security-Tab mit Events, Auswertung und Details.",
        )

        dnsf = ttk.LabelFrame(
            right_core_group,
            text=self.tr("group_dns", "DNS"),
            padding=10,
        )
        dnsf.grid(row=1, column=0, sticky="new", pady=(10, 0))

        self.dns_choice = tk.StringVar(value="System (DHCP/auto)")
        self.dns_combo = ttk.Combobox(
            dnsf,
            textvariable=self.dns_choice,
            values=list(DNS_PROFILES.keys()),
            state="readonly",
        )
        self.dns_combo.pack(fill="x")
        self.add_tooltip(
            self.dns_combo,
            "Bei aktivem VPN wird DNS vom VPN/WireGuard-Profil verwaltet. "
            "Ein manueller DNS-Wechsel würde DNS-Leaks verursachen.",
        )

        self.dot_choice = tk.StringVar(value="Cloudflare")

        self.btn_dns_apply = ttk.Button(
            dnsf,
            text=self.tr("btn_apply_dns", "DNS anwenden"),
            command=self.apply_dns,
        )
        self.btn_dns_apply.pack(fill="x", pady=6)
        self.register_feature_button("nmcli", self.btn_dns_apply)
        self.add_tooltip(
            self.btn_dns_apply,
            "Übernimmt den ausgewählten DNS-Server für die aktive Verbindung. "
            "Bei aktivem VPN ist diese Aktion gesperrt.",
        )
        btn_dns_nm = ttk.Button(
            dnsf,
            text=self.tr("btn_dns_status_nmcli", "DNS Status (nmcli)"),
            command=self.show_dns_status,
        )
        btn_dns_nm.pack(fill="x")
        self.register_feature_button("nmcli", btn_dns_nm)
        self.add_tooltip(
            btn_dns_nm,
            "Zeigt, welche DNS-Server aktuell über NetworkManager gesetzt sind.",
        )
        btn_dns_resolv = ttk.Button(
            dnsf,
            text=self.tr("btn_dns_status_resolv", "DNS Status (resolv*)"),
            command=self.resolved_status_async,
        )
        btn_dns_resolv.pack(fill="x", pady=(6, 0))
        self.register_feature_button("resolvectl", btn_dns_resolv)
        self.add_tooltip(
            btn_dns_resolv,
            "Zeigt den Resolver-Status inkl. /etc/resolv.conf und resolvectl-Ausgabe.",
        )
        ttk.Label(
            dnsf,
            textvariable=self.dns_control_hint_var,
            foreground="#aa7700",
            wraplength=220,
            justify="left",
        ).pack(fill="x", pady=(6, 0))

        tr = ttk.LabelFrame(
            right_diag_group,
            text=self.tr("group_trace", "Trace"),
            padding=10,
        )
        tr.grid(row=0, column=0, sticky="new")
        self.trace_target = tk.StringVar(value=TRACE_TARGETS[0])
        self.trace_combo = ttk.Combobox(
            tr, textvariable=self.trace_target, values=TRACE_TARGETS, state="normal"
        )
        self.trace_combo.pack(fill="x")
        btn_trace = ttk.Button(
            tr,
            text=self.tr("btn_start_trace", "Trace starten"),
            command=self.trace_table_async,
        )
        btn_trace.pack(fill="x", pady=(6, 0))
        self.register_feature_button("traceroute", btn_trace)
        self.add_tooltip(
            btn_trace,
            "Startet eine Routenanalyse und zeigt die Hops im Trace-Tab an.",
        )
        cb_hop_lookup = ttk.Checkbutton(
            tr,
            text=self.tr("check_hop_lookup", "Hop Lookup (Land/Provider)"),
            variable=self.hop_lookup_enabled,
        )
        cb_hop_lookup.pack(fill="x", pady=(6, 0))
        self.register_feature_button("traceroute", cb_hop_lookup)

        mtrf = ttk.LabelFrame(
            right_diag_group,
            text=self.tr("group_mtr", "MTR"),
            padding=10,
        )
        mtrf.grid(row=1, column=0, sticky="new", pady=(10, 0))
        self.mtr_target = tk.StringVar(value=TRACE_TARGETS[0])
        ttk.Entry(mtrf, textvariable=self.mtr_target).pack(fill="x")
        btn_mtr = ttk.Button(
            mtrf,
            text=self.tr("btn_start_mtr", "MTR starten"),
            command=self.mtr_async,
        )
        btn_mtr.pack(fill="x", pady=(6, 0))
        self.register_feature_button("mtr", btn_mtr)
        self.add_tooltip(
            btn_mtr,
            "Startet eine kombinierte Ping- und Traceroute-Analyse zum Ziel.",
        )

        nmapf = ttk.LabelFrame(
            right_diag_group,
            text=self.tr("group_nmap", "Nmap"),
            padding=10,
        )
        nmapf.grid(row=2, column=0, sticky="new", pady=(10, 0))
        self.nmap_target = tk.StringVar(value="127.0.0.1")
        ttk.Entry(nmapf, textvariable=self.nmap_target).pack(fill="x")
        self.nmap_preset = tk.StringVar(value="quick")
        self.nmap_preset_combo = ttk.Combobox(
            nmapf,
            textvariable=self.nmap_preset,
            values=["quick", "service", "aggressive"],
            state="readonly",
        )
        self.nmap_preset_combo.pack(fill="x", pady=(6, 0))
        self.btn_nmap_start = ttk.Button(
            nmapf,
            text=self.tr("btn_start_scan", "Scan starten"),
            command=self.nmap_async,
        )
        self.btn_nmap_start.pack(fill="x", pady=(6, 0))
        self.register_feature_button("nmap", self.btn_nmap_start)
        self.add_tooltip(
            self.btn_nmap_start,
            "Startet einen Portscan für das eingetragene Ziel.",
        )

        # Rechte Reserve: Links/VPN links, HDD/Speedtest rechts
        self.right_button_space.columnconfigure(0, weight=1, minsize=side_panel_width)

        linkf = ttk.LabelFrame(
            self.right_button_space,
            text=self.tr("group_links", "Links"),
            padding=8,
        )
        linkf.grid(row=0, column=0, sticky="new")
        for idx in range(LINK_SLOT_COUNT):
            btn = ttk.Button(
                linkf,
                text=f"Link {idx + 1} frei",
                style="LinkSlot.TButton",
                command=lambda i=idx: self.open_link_slot(i),
            )
            btn.pack(fill="x", pady=(0, 6) if idx < LINK_SLOT_COUNT - 1 else 0)
            self.link_slot_buttons.append(btn)
            self.register_feature_button("links", btn)
            self.add_tooltip(
                btn,
                "Startet den konfigurierten Link-Slot oder zeigt im Log, dass der Slot noch frei ist.",
            )

        vpnf_slots = ttk.LabelFrame(
            self.right_button_space,
            text=self.tr("group_vpn_slots", "VPN-Slots"),
            padding=8,
        )
        vpnf_slots.grid(row=1, column=0, sticky="new", pady=(12, 0))
        vpnf_slots.columnconfigure(0, weight=1)
        for idx in range(VPN_SLOT_COUNT):
            row = ttk.Frame(vpnf_slots)
            row.pack(fill="x", pady=(0, 6) if idx < VPN_SLOT_COUNT - 1 else 0)
            row.columnconfigure(0, weight=1)

            btn = ttk.Button(
                row,
                text=f"VPN {idx + 1} frei",
                style="LinkSlot.TButton",
                command=lambda i=idx: self.open_vpn_slot(i),
            )
            btn.grid(row=0, column=0, sticky="ew")
            self.vpn_slot_buttons.append(btn)
            self.register_feature_button("wg", btn)
            self.add_tooltip(
                btn,
                "Startet das zugeordnete WireGuard-Profil oder zeigt im Log, dass der Slot frei ist.",
            )

            light = tk.Canvas(row, width=14, height=14, highlightthickness=0)
            light.grid(row=0, column=1, padx=(6, 0))
            dot = light.create_oval(2, 2, 12, 12, fill="gray", outline="")
            self.vpn_slot_lights.append((light, dot))

        hddf = ttk.LabelFrame(
            self.right_reserve_space,
            text=self.tr("group_hdd_slots", "HDD Slots"),
            padding=8,
        )
        hddf.grid(row=0, column=0, sticky="new")
        hddf.columnconfigure(0, weight=1)
        for idx, slot in enumerate(self.hdd_slots):
            row = ttk.Frame(hddf)
            row.pack(fill="x", pady=(0, 6))
            row.columnconfigure(0, weight=1)

            label = slot.get("custom_name") or f"Slot {idx + 1} frei"
            name_btn = ttk.Button(
                row, text=label, command=lambda i=idx: self.hdd_info_async(i)
            )
            name_btn.grid(row=0, column=0, sticky="ew")
            self.hdd_slot_name_buttons.append(name_btn)
            self.register_feature_button("hdd_manage", name_btn)

            light = tk.Canvas(row, width=14, height=14, highlightthickness=0)
            light.grid(row=0, column=1, padx=(6, 0))
            dot = light.create_oval(2, 2, 12, 12, fill="gray", outline="")
            self.hdd_slot_lights.append((light, dot))

            btn_mount = ttk.Button(
                row, text="Mount", command=lambda i=idx: self.hdd_mount_async(i)
            )
            btn_mount.grid(row=1, column=0, sticky="ew", pady=(4, 0))
            self.register_feature_button("hdd_manage", btn_mount)
            self.add_tooltip(
                btn_mount,
                "Mountet den zugeordneten Datenträger auf den konfigurierten Ordner.",
            )

            btn_unmount = ttk.Button(
                row, text="Unmount", command=lambda i=idx: self.hdd_unmount_async(i)
            )
            btn_unmount.grid(row=1, column=1, sticky="ew", padx=(6, 0), pady=(4, 0))
            self.register_feature_button("hdd_manage", btn_unmount)
            self.add_tooltip(
                btn_unmount,
                "Hängt den zugeordneten Datenträger sauber aus.",
            )

            btn_info = ttk.Button(
                row, text="Info", command=lambda i=idx: self.hdd_info_async(i)
            )
            btn_info.grid(row=2, column=0, columnspan=2, sticky="ew", pady=(4, 0))
            self.register_feature_button("hdd_manage", btn_info)
            self.add_tooltip(
                btn_info,
                "Zeigt Details und Speicherbelegung des zugeordneten Datenträgers.",
            )

        speedtestf = ttk.LabelFrame(
            self.right_reserve_space, text="Speedtest", padding=8
        )
        speedtestf.grid(row=1, column=0, sticky="new", pady=(12, 0))
        speedtestf.columnconfigure(0, weight=1)

        btn_speedtest_start = ttk.Button(
            speedtestf,
            text="Speedtest starten",
            command=self.speedtest_async,
        )
        btn_speedtest_start.pack(fill="x")
        self.register_feature_button("speedtest", btn_speedtest_start)
        self.add_tooltip(
            btn_speedtest_start,
            "Startet einen nativen Speedtest (Ping, Jitter, Download, Upload).",
        )

        btn_speedtest_history = ttk.Button(
            speedtestf,
            text="Speedtest-Verlauf",
            command=self.show_speedtest_history,
        )
        btn_speedtest_history.pack(fill="x", pady=(6, 0))
        self.register_feature_button("speedtest", btn_speedtest_history)
        self.add_tooltip(
            btn_speedtest_history,
            "Zeigt die letzten gespeicherten Speedtest-Ergebnisse.",
        )

        speedtest_lastf = ttk.LabelFrame(
            self.right_reserve_space, text="Speedtest (letztes Ergebnis)", padding=8
        )
        speedtest_lastf.grid(row=2, column=0, sticky="new", pady=(12, 0))
        speedtest_lastf.columnconfigure(1, weight=1)

        ttk.Label(speedtest_lastf, text="Ping:").grid(row=0, column=0, sticky="w")
        ttk.Label(speedtest_lastf, textvariable=self.speedtest_ping_var).grid(
            row=0, column=1, sticky="w", padx=(8, 0)
        )
        ttk.Label(speedtest_lastf, text="Jitter:").grid(
            row=1, column=0, sticky="w", pady=(2, 0)
        )
        ttk.Label(speedtest_lastf, textvariable=self.speedtest_jitter_var).grid(
            row=1, column=1, sticky="w", padx=(8, 0), pady=(2, 0)
        )
        ttk.Label(speedtest_lastf, text="Download:").grid(
            row=2, column=0, sticky="w", pady=(2, 0)
        )
        ttk.Label(speedtest_lastf, textvariable=self.speedtest_download_var).grid(
            row=2, column=1, sticky="w", padx=(8, 0), pady=(2, 0)
        )
        ttk.Label(speedtest_lastf, text="Upload:").grid(
            row=3, column=0, sticky="w", pady=(2, 0)
        )
        ttk.Label(speedtest_lastf, textvariable=self.speedtest_upload_var).grid(
            row=3, column=1, sticky="w", padx=(8, 0), pady=(2, 0)
        )
        ttk.Label(speedtest_lastf, text="Zeit:").grid(
            row=4, column=0, sticky="w", pady=(2, 0)
        )
        ttk.Label(speedtest_lastf, textvariable=self.speedtest_timestamp_var).grid(
            row=4, column=1, sticky="w", padx=(8, 0), pady=(2, 0)
        )

        self.update_speedtest_panel(self.speedtest_last_result)

        self.nb = ttk.Notebook(self.main_work_area, width=main_work_width)
        self.nb.grid(row=1, column=0, sticky="nsew", pady=(10, 0))

        self.tab_system_main = ttk.Frame(self.nb, padding=10)
        self.system_nb = ttk.Notebook(self.tab_system_main)
        self.system_nb.pack(fill="both", expand=True)

        self.tab_tools = ttk.Frame(self.nb, padding=10)
        self.tools_nb = ttk.Notebook(self.tab_tools)
        self.tools_nb.pack(fill="both", expand=True)

        self.tab_security_main = ttk.Frame(self.nb, padding=10)
        self.security_nb = ttk.Notebook(self.tab_security_main)
        self.security_nb.pack(fill="both", expand=True)

        build_log_tab(self)
        build_trace_tab(self)
        build_connections_tab(self)

        build_dns_tab(self)

        build_chad_tab(self)

        self.nb.add(self.tab_system_main, text=self.tr("tab_system", "System"))
        self.nb.add(self.tab_tools, text=self.tr("tab_tools", "Tools"))
        self.nb.add(self.tab_security_main, text=self.tr("tab_security", "Security"))
        build_settings_tab_ui(self)

        build_basic_tools_tabs(self)

        build_linux_help_tab(self)

        build_system_tab(self)
        build_hdd_tab(self)
        build_links_tab(self)
        build_vpn_slots_tab(self)

        build_security_tabs(self)

        bottom = ttk.Frame(self.page, padding=10)
        bottom.grid(row=2, column=0, columnspan=2, sticky="ew")

        self._install_theme_toggle_button()
        initial_theme = (
            self._theme_start_override or self._preferred_theme_name or "light"
        )
        self.apply_theme(initial_theme, log_change=False, persist=False)
        self.check_keyring_availability()
        self.load_ai_config()

        self.log("Start.")
        self.setup_sudo_auth()
        self.refresh_wg_profiles()
        self.update_primary_vpn_form_controls()
        self.sync_primary_vpn_form_selection()
        self.preflight_dependencies()
        self.update_ipv6_controls_state()
        self.refresh()
        self.after(50, self._activate_startup_log_view)
        if self.settings.get("startup_guard_autorun", True):
            self.after(1800, self.startup_guard_startup_check)
        self.after(1000, self.attach_all_chad_context_menus)
        self.after(1100, self.attach_all_selection_copy_shortcuts)
        self.log_debug(
            f"_security_query_log_level vorhanden = {callable(getattr(self, '_security_query_log_level', None))}"
        )
        self.scan_hdds_async()
        self.hdd_refresh_async()
        self._security_update_status("gray")
        if not self._load_security_baseline():
            self.security_build_baseline()
        else:
            # update after load, run immediate check
            self.security_scan_now()
        self.isp_stealth_apply_startup()
        self.clock_apply_alarm_settings(log_change=False)
        self.log(
            f"LCD-Uhr aktiv (Systemzeit, TZ: {datetime.now().astimezone().tzname() or '-'})"
        )
        self.after(120, self.clock_tick)
        self.after(250, self.system_monitor_refresh)
        self.after(900, self.start_dns_startup_sequence)
        if self.settings.get("path_guard_autorun", False):
            self.after(2500, self.run_path_guard_check)
        self.after(3600, self.auto_refresh)

    # scroll helpers (Linux wheel)
    def _scroll_up(self, _e):
        self.page_canvas.yview_scroll(-1, "units")

    def _scroll_down(self, _e):
        self.page_canvas.yview_scroll(1, "units")

    def on_close(self):
        self.system_monitor_active = False
        clock_after_id = self.__dict__.get("_clock_after_id")
        if clock_after_id:
            try:
                self.after_cancel(clock_after_id)
            except Exception:
                pass
            self._clock_after_id = None
        self.destroy()

    def report_callback_exception(self, exc, val, tb):
        _append_crash_dump("tk.report_callback_exception", exc, val, tb)
        try:
            self.log(f"Unerwarteter GUI-Fehler. Crash-Dump: {CRASH_LOG_PATH}")
        except Exception:
            pass
        try:
            traceback.print_exception(exc, val, tb, file=sys.stderr)
        except Exception:
            pass

    def _normalize_theme_name(self, theme_name):
        name = str(theme_name or "").strip().lower()
        return name if name in {"light", "dark"} else "light"

    def _install_theme_toggle_button(self):
        if self._theme_toggle_btn is not None:
            return
        btn = ttk.Button(
            self, textvariable=self._theme_name_var, command=self.toggle_theme
        )
        btn.place(x=260, y=10, anchor="nw")
        self._theme_toggle_btn = btn
        self.after(250, self._position_theme_button)
        self.bind("<Configure>", self._on_window_configure, add="+")

    def _on_window_configure(self, _event=None):
        self.after(50, self._position_theme_button)

    def _find_terminal_button(self):
        stack = [self]
        while stack:
            node = stack.pop()
            for child in node.winfo_children():
                stack.append(child)
                try:
                    if (
                        isinstance(child, ttk.Button)
                        and str(child.cget("text")) == "Terminal"
                    ):
                        return child
                except Exception:
                    pass
        return None

    def _position_theme_button(self):
        if self._theme_toggle_btn is None:
            return
        try:
            term_btn = self._find_terminal_button()
            if term_btn is None or not term_btn.winfo_ismapped():
                self._theme_toggle_btn.place(x=260, y=10, anchor="nw")
                return

            self.update_idletasks()
            x = term_btn.winfo_rootx() - self.winfo_rootx()
            y = term_btn.winfo_rooty() - self.winfo_rooty()
            btn_h = max(1, self._theme_toggle_btn.winfo_height())
            y = max(8, y - btn_h - 8)
            self._theme_toggle_btn.place(x=x, y=y, anchor="nw")
        except Exception:
            self._theme_toggle_btn.place(x=260, y=10, anchor="nw")

    def toggle_theme(self):
        current = (self._app_theme or LIGHT_THEME).get("name", "light")
        next_theme = "dark" if current == "light" else "light"
        self.apply_theme(next_theme)

    def apply_theme(self, theme_name, log_change=True, persist=True):
        theme = (
            DARK_THEME
            if self._normalize_theme_name(theme_name) == "dark"
            else LIGHT_THEME
        )
        self._app_theme = theme
        self._preferred_theme_name = theme["name"]

        self._apply_ttk_theme(theme)
        self._apply_tk_widget_theme(self, theme)

        if self._theme_toggle_btn is not None:
            self._theme_name_var.set("Light" if theme["name"] == "dark" else "Dark")
            self.after(0, self._position_theme_button)

        if persist and hasattr(self, "privacy_config_path"):
            try:
                self.save_privacy_config()
            except Exception:
                pass

        if log_change:
            self.log(f"Theme gewechselt: {theme['name']}")

    def _apply_ttk_theme(self, theme):
        style = self.style if hasattr(self, "style") else ttk.Style(self)

        style.configure(".", background=theme["panel_bg"], foreground=theme["fg"])
        style.configure("TFrame", background=theme["panel_bg"])
        style.configure(
            "TLabelframe", background=theme["panel_bg"], foreground=theme["fg"]
        )
        style.configure(
            "TLabelframe.Label", background=theme["panel_bg"], foreground=theme["fg"]
        )
        style.configure("TLabel", background=theme["panel_bg"], foreground=theme["fg"])
        style.configure(
            "TButton",
            background=theme["panel_bg"],
            foreground=theme["fg"],
            padding=(6, 3),
        )
        style.configure("LinkSlot.TButton", padding=(12, 10), width=22)
        style.map(
            "TButton",
            background=[("active", theme["tab_selected_bg"])],
            foreground=[("disabled", "#7a8694"), ("active", theme["fg"])],
        )

        style.configure(
            "TCheckbutton",
            background=theme["panel_bg"],
            foreground=theme["fg"],
            padding=(4, 2),
        )
        style.configure(
            "TRadiobutton", background=theme["panel_bg"], foreground=theme["fg"]
        )

        style.configure(
            "TEntry", fieldbackground=theme["entry_bg"], foreground=theme["entry_fg"]
        )
        style.configure(
            "TCombobox", fieldbackground=theme["entry_bg"], foreground=theme["entry_fg"]
        )

        style.configure("TNotebook", background=theme["panel_bg"], borderwidth=0)
        style.configure(
            "TNotebook.Tab",
            background=theme["tab_bg"],
            foreground=theme["tab_fg"],
            padding=(8, 4),
        )
        style.map(
            "TNotebook.Tab",
            background=[
                ("selected", theme["tab_selected_bg"]),
                ("active", theme["tab_bg"]),
            ],
            foreground=[
                ("selected", theme["tab_selected_fg"]),
                ("!selected", theme["tab_fg"]),
            ],
        )

        style.configure(
            "Treeview",
            background=theme["tree_bg"],
            foreground=theme["tree_fg"],
            fieldbackground=theme["tree_bg"],
        )
        style.map(
            "Treeview",
            background=[("selected", theme["select_bg"])],
            foreground=[("selected", theme["select_fg"])],
        )
        style.configure(
            "Treeview.Heading", background=theme["tab_bg"], foreground=theme["fg"]
        )

        try:
            self.configure(bg=theme["bg"])
        except Exception:
            pass

    def _apply_tk_widget_theme(self, widget, theme):
        klass = str(widget.winfo_class())
        preserve_colors = bool(getattr(widget, "_preserve_theme_colors", False))

        try:
            if klass in {"Frame", "Labelframe"}:
                widget.configure(bg=theme["panel_bg"])
            elif klass in {"Label"}:
                if not preserve_colors:
                    widget.configure(bg=theme["panel_bg"], fg=theme["fg"])
            elif klass in {"Button", "Checkbutton", "Radiobutton"}:
                widget.configure(
                    bg=theme["tab_bg"],
                    fg=theme["fg"],
                    activebackground=theme["tab_selected_bg"],
                    activeforeground=theme["fg"],
                    selectcolor=theme["panel_bg"],
                )
            elif klass in {"Entry", "Spinbox"}:
                widget.configure(
                    bg=theme["entry_bg"],
                    fg=theme["entry_fg"],
                    insertbackground=theme["entry_fg"],
                )
            elif klass == "Text":
                widget.configure(
                    bg=theme["text_bg"],
                    fg=theme["text_fg"],
                    insertbackground=theme["text_fg"],
                    selectbackground=theme["select_bg"],
                    selectforeground=theme["select_fg"],
                )
            elif klass == "Listbox":
                widget.configure(
                    bg=theme["text_bg"],
                    fg=theme["text_fg"],
                    selectbackground=theme["select_bg"],
                    selectforeground=theme["select_fg"],
                )
            elif klass == "Canvas":
                widget.configure(bg=theme["panel_bg"], highlightthickness=0)
        except Exception:
            pass

        for child in widget.winfo_children():
            self._apply_tk_widget_theme(child, theme)

        if widget is self:
            self._apply_chad_chat_theme(theme)

    def _apply_chad_chat_theme(self, theme):
        widget = self.__dict__.get("chad_analysis_text")
        if widget is None:
            return

        if theme.get("name") == "dark":
            user_header_fg = "#7ee787"
            assistant_header_fg = "#79c0ff"
        else:
            user_header_fg = "#1a7f37"
            assistant_header_fg = "#0969da"

        try:
            widget.tag_configure(
                "chat_user_header",
                foreground=user_header_fg,
            )
            widget.tag_configure(
                "chat_assistant_header",
                foreground=assistant_header_fg,
            )
            widget.tag_configure(
                "chat_system_header",
                foreground=theme["muted_fg"],
            )
            widget.tag_configure(
                "chat_user_body",
                foreground=theme["text_fg"],
            )
            widget.tag_configure(
                "chat_assistant_body",
                foreground=theme["text_fg"],
            )
            widget.tag_configure(
                "chat_system_body",
                foreground=theme["muted_fg"],
            )
        except Exception:
            pass

    # ---------- UI helpers ----------
    def make_light(self, parent, label):
        return build_light_widget(parent, label, tk_module=tk, ttk_module=ttk)

    def set_light(self, light, color):
        update_light_widget(light, color)

    def _validate_alarm_time(self, value):
        txt = str(value or "").strip()
        if not re.fullmatch(r"([01]\d|2[0-3]):([0-5]\d)", txt):
            return None
        return txt

    def _set_alarm_status_text(self, text, color="#444444"):
        status_var = self.__dict__.get("alarm_status_var")
        if status_var is not None:
            status_var.set(str(text or ""))
        label = self.__dict__.get("clock_alarm_status_label")
        if label is not None:
            try:
                label.configure(foreground=color)
            except Exception:
                pass

    def clock_toggle_alarm(self):
        self.clock_apply_alarm_settings(log_change=True)

    def clock_apply_alarm_settings(self, log_change=True):
        alarm_txt = self._validate_alarm_time(self.alarm_time_var.get())
        enabled = bool(self.alarm_enabled_var.get())

        if enabled and not alarm_txt:
            self.alarm_enabled_var.set(False)
            self._set_alarm_status_text(
                "Ungültige Weckerzeit (HH:MM).", color="#b04545"
            )
            if log_change:
                self.log("Wecker nicht aktiviert: Uhrzeit muss im Format HH:MM sein.")
            try:
                self.save_privacy_config()
            except Exception:
                pass
            return

        if alarm_txt:
            self.alarm_time_var.set(alarm_txt)

        # Re-arm trigger whenever settings are applied.
        self._alarm_last_fired_key = ""

        if enabled:
            self._set_alarm_status_text(
                f"Aktiv: {self.alarm_time_var.get()}", color="#2f6f2f"
            )
            if log_change:
                self.log(f"Wecker aktiviert: {self.alarm_time_var.get()}")
        else:
            self._set_alarm_status_text("Wecker aus", color="#444444")
            if log_change:
                self.log("Wecker deaktiviert.")

        try:
            self.save_privacy_config()
        except Exception:
            pass

    def clock_tick(self):
        now = datetime.now().astimezone()
        self.lcd_time_var.set(now.strftime("%H:%M:%S"))
        self.lcd_date_var.set(now.strftime("%d.%m.%Y"))
        tz_name = now.tzname() or "-"
        self.lcd_tz_var.set(f"TZ: {tz_name}")

        alarm_txt = self._validate_alarm_time(self.alarm_time_var.get())
        enabled = bool(self.alarm_enabled_var.get())
        if enabled and not alarm_txt:
            if not self._alarm_invalid_logged:
                self._set_alarm_status_text(
                    "Ungültige Weckerzeit (HH:MM).", color="#b04545"
                )
                self.log("Wecker: ungültige Zeit erkannt, bitte HH:MM verwenden.")
                self._alarm_invalid_logged = True
        else:
            self._alarm_invalid_logged = False

        if enabled and alarm_txt:
            minute_key = now.strftime("%Y-%m-%d") + " " + alarm_txt
            if (
                now.strftime("%H:%M") == alarm_txt
                and self._alarm_last_fired_key != minute_key
            ):
                self._alarm_last_fired_key = minute_key
                self._set_alarm_status_text(
                    f"ALARM {alarm_txt} ausgelöst ({now.strftime('%H:%M:%S')})",
                    color="#c05000",
                )
                self.log(f"Wecker ausgelöst: {alarm_txt}")
                try:
                    self.bell()
                except Exception:
                    pass

        ms_until_next_second = 1000 - (int(time.time() * 1000) % 1000)
        self._clock_after_id = self.after(
            max(120, ms_until_next_second), self.clock_tick
        )

    def log(self, msg):
        text_widget = self.__dict__.get("text")
        if text_widget is not None:
            try:
                text_widget.insert("end", f"[{ts()}] {msg}\n")
                text_widget.see("end")
                return
            except Exception:
                pass
        print(f"[{ts()}] {msg}")

    def log_debug(self, msg):
        if self.__dict__.get("debug_mode", False):
            self.log(f"[DEBUG] {msg}")

    def log_debug_safe(self, msg):
        if self.__dict__.get("debug_mode", False):
            self.log_safe(f"[DEBUG] {msg}")

    def log_safe(self, msg):
        if threading.current_thread() is threading.main_thread():
            self.log(msg)
            return
        self.after(0, lambda m=msg: self.log(m))

    def log_install_hint(self, title: str, lines):
        parts = [
            "=" * 60,
            (title or "Hinweis").strip(),
            "=" * 60,
        ]
        parts.extend([str(line) for line in (lines or [])])
        parts.append("")
        self.log_safe("\n".join(parts))

    def show_wireshark_install_hint(self):
        self.log_install_hint(
            "Wireshark-/Packet-Capture-Funktionen sind nicht vollständig installiert.",
            [
                "Bitte folgende Befehle ausführen:",
                "",
                "sudo apt update",
                "sudo apt install wireshark tshark",
                "sudo dpkg-reconfigure wireshark-common",
                "sudo usermod -aG wireshark $USER",
                "",
                "Danach bitte ab- und wieder anmelden oder das System neu starten.",
            ],
        )

    def show_dumpcap_permission_hint(self):
        self.log_install_hint(
            "tshark ist installiert, aber dumpcap hat keine ausreichenden Berechtigungen.",
            [
                "Bitte folgende Befehle ausführen:",
                "",
                "sudo dpkg-reconfigure wireshark-common",
                "sudo usermod -aG wireshark $USER",
                "",
                "Wichtig:",
                "Danach bitte ab- und wieder anmelden oder neu starten.",
            ],
        )

    def show_lynis_install_hint(self):
        self.log_install_hint(
            "Lynis ist nicht installiert.",
            [
                "Bitte folgende Befehle ausführen:",
                "",
                "sudo apt update",
                "sudo apt install lynis",
            ],
        )

    def show_rkhunter_install_hint(self):
        self.log_install_hint(
            "rkhunter ist nicht installiert.",
            [
                "Bitte folgende Befehle ausführen:",
                "",
                "sudo apt update",
                "sudo apt install rkhunter",
            ],
        )

    def show_rkhunter_web_cmd_hint(self):
        self.log_install_hint(
            "rkhunter-Konfiguration fehlerhaft: WEB_CMD ist ungültig.",
            [
                "Die Meldung 'Invalid WEB_CMD configuration option: Relative pathname: \"/bin/false\"' kommt aus der rkhunter-Konfiguration.",
                "",
                "Bitte die Konfiguration prüfen, zum Beispiel:",
                "",
                "sudo grep -n '^WEB_CMD' /etc/rkhunter.conf /etc/rkhunter.conf.d/*",
                "sudo editor /etc/rkhunter.conf",
                "",
                "Empfehlung:",
                "WEB_CMD=''",
                "oder einen gültigen absoluten Downloader setzen, z. B.:",
                "WEB_CMD='/usr/bin/wget'",
                "",
                "Danach den Update-Vorgang erneut starten.",
            ],
        )

    def _looks_like_rkhunter_web_cmd_issue(self, text):
        lowered = str(text or "").strip().lower()
        return (
            "invalid web_cmd configuration option" in lowered
            or "web_cmd" in lowered
            and "/bin/false" in lowered
        )

    def _looks_like_dumpcap_permission_issue(self, text):
        lowered = str(text or "").strip().lower()
        if not lowered:
            return False
        if "dumpcap" not in lowered:
            return False
        markers = [
            "permission denied",
            "you do not have permission",
            "operation not permitted",
            "cap_set_proc",
            "couldn't run",
            "could not run",
        ]
        return any(marker in lowered for marker in markers)

    def refresh_safe(self):
        if threading.current_thread() is threading.main_thread():
            self.refresh()
            if "dns_combo" in self.__dict__:
                self.update_dns_controls_state_safe()
            return
        self.after(0, self.refresh)
        if "dns_combo" in self.__dict__:
            self.update_dns_controls_state_safe()

    def update_ipv6_controls_state_safe(self):
        if threading.current_thread() is threading.main_thread():
            self.update_ipv6_controls_state()
            return
        self.after(0, self.update_ipv6_controls_state)

    def _set_stealth_enabled_state(self, enabled):
        self.isp_stealth_enabled.set(enabled)
        self.save_privacy_config()

    def _set_dns_choice_safe(self, choice):
        if threading.current_thread() is threading.main_thread():
            self.dns_choice.set(choice)
            return
        self.after(0, lambda c=choice: self.dns_choice.set(c))

    def is_vpn_active_for_dns_ui(self):
        service = self._dns_control_service()
        if service is not None:
            code, out = service.read_resolvectl_status(timeout=6)
            if code != 0:
                return False
            return service.detect_vpn_active(out)
        code, out = run_cmd(["resolvectl", "status"], timeout=6)
        if code != 0:
            return False
        text = (out or "").lower()
        return any(
            token in text for token in ("proton", "wireguard", "wg", "tun", "vpn")
        )

    def _dns_controls_locked(self):
        if self.is_vpn_active_for_dns_ui():
            return True, "vpn"
        if self._dns_runtime_transition_active():
            return True, "transition"
        return False, ""

    def update_dns_controls_state(self):
        controls_locked, lock_reason = self._dns_controls_locked()

        if controls_locked:
            combo_state = "disabled"
            if lock_reason == "transition":
                hint = "VPN/DNS-Wechsel läuft – DNS wird vorübergehend gesperrt."
            else:
                hint = "VPN aktiv – DNS wird vom VPN-Profil verwaltet."
        else:
            combo_state = "readonly"
            hint = "DNS-Auswahl aktiv."

        combo = self.__dict__.get("dns_combo")
        if combo is not None:
            try:
                combo.configure(state=combo_state)
            except Exception:
                pass

        for attr in ("btn_dns_apply", "btn_dot_on"):
            btn = self.__dict__.get(attr)
            if btn is not None:
                try:
                    btn.configure(state="disabled" if controls_locked else "normal")
                except Exception:
                    pass

        btn_dot_off = self.__dict__.get("btn_dot_off")
        if btn_dot_off is not None:
            try:
                btn_dot_off.configure(state="normal")
            except Exception:
                pass

        hint_var = self.__dict__.get("dns_control_hint_var")
        if hint_var is not None:
            try:
                hint_var.set(hint)
            except Exception:
                pass

    def update_dns_controls_state_safe(self):
        if threading.current_thread() is threading.main_thread():
            self.update_dns_controls_state()
            return
        self.after(0, self.update_dns_controls_state)

    def get_vpn_state(self):
        return self.is_vpn_active()

    def get_active_vpn_dns(self):
        vpn_iface, vpn_dns, _raw = self.find_vpn_dns_servers()
        return vpn_dns if vpn_iface or vpn_dns else []

    def get_selected_user_dns(self):
        choice = self._manual_dns_choice()
        servers = DNS_PROFILES.get(choice)
        if not servers:
            return []
        return [str(server).strip() for server in servers if str(server).strip()]

    def get_selected_dot_dns(self):
        choice_var = self.__dict__.get("dot_choice")
        profile_name = (
            (choice_var.get() if choice_var is not None else "") or ""
        ).strip()
        profile = DOT_PROFILES.get(profile_name, {})
        dns_values = [
            str(ip).strip() for ip in (profile.get("dns") or []) if str(ip).strip()
        ]
        hostname = (profile.get("hostname") or "").strip()
        if hostname:
            dns_values.extend([f"{ip}#{hostname}" for ip in dns_values])
        return dns_values

    def is_dot_active(self):
        if bool(self.__dict__.get("dot_enabled", False)):
            return True
        service = self._dns_control_service()
        if service is not None:
            code, out = service.read_resolvectl_status(timeout=6)
        else:
            code, out = run_cmd(["resolvectl", "status"], timeout=6)
        if code != 0:
            return False
        return "+DNSOverTLS" in out or "DNSOverTLS=yes" in out

    def normalize_dns_value(self, value: str) -> str:
        normalized = str(value or "").strip()
        if not normalized:
            return ""
        if "#" in normalized:
            normalized = normalized.split("#", 1)[0].strip()
        return normalized

    def rebuild_dns_guard_whitelist(self):
        if self.is_vpn_active():
            allowed = self.get_active_vpn_dns()
        elif self.is_dot_active():
            allowed = self.get_selected_dot_dns()
        else:
            allowed = self.get_selected_user_dns()
            if not allowed:
                allowed = self.get_current_dns()
        self.dns_guard_allowed_dns = [
            str(dns).strip() for dns in (allowed or []) if str(dns).strip()
        ]
        return list(self.dns_guard_allowed_dns)

    def _normalized_dns_values(self, values):
        return [
            self.normalize_dns_value(dns)
            for dns in (values or [])
            if self.normalize_dns_value(dns)
        ]

    def get_allowed_dns(self):
        cached = [
            str(dns).strip()
            for dns in (self.__dict__.get("dns_guard_allowed_dns") or [])
            if str(dns).strip()
        ]
        if cached:
            return cached
        return self.rebuild_dns_guard_whitelist()

    def get_current_dns(self):
        return self.get_active_dns_servers()

    def dns_is_clean(self, current, allowed):
        allowed_set = set(self._normalized_dns_values(allowed))
        current_set = set(self._normalized_dns_values(current))
        if not current_set:
            return False
        if not allowed_set:
            return False
        return current_set.issubset(allowed_set)

    def get_dns_guard_status(self):
        guard_var = self.__dict__.get("dns_guard_enabled_var")
        guard_enabled = bool(guard_var.get()) if guard_var is not None else False
        if not guard_enabled:
            return "off"

        if time.time() < float(self.__dict__.get("dns_guard_transition_until", 0.0)):
            return "transition"

        explicit_state = str(self.__dict__.get("dns_guard_state") or "").strip().lower()
        if explicit_state in {"ok", "problem", "off", "unknown", "transition", "error"}:
            return explicit_state

        current_dns = self.get_current_dns()
        allowed_dns = self.get_allowed_dns()
        normalized_allowed = self._normalized_dns_values(allowed_dns)

        if not current_dns:
            return "unknown"

        if not normalized_allowed:
            return "unknown"

        if not self.dns_is_clean(current_dns, allowed_dns):
            return "problem"

        return "ok"

    def update_dns_guard_indicator(self):
        status = self.get_dns_guard_status()
        if status == "ok":
            color = "green"
            text = "DNS Guard: aktiv (OK)"
        elif status == "problem":
            color = "red"
            text = "DNS Guard: Problem erkannt"
        elif status == "transition":
            color = "yellow"
            text = "DNS Guard: Wechsel wird geprüft"
        elif status == "off":
            color = "gray"
            text = "DNS Guard: deaktiviert"
        else:
            color = "orange"
            text = "DNS Guard: unbekannt"

        if "dns_guard_label" in self.__dict__:
            try:
                self.dns_guard_label.configure(text=text)
            except Exception:
                pass
        indicator = self.__dict__.get("dns_guard_indicator")
        if indicator is not None:
            self.set_light(indicator, color)
        return status

    def show_dns_guard_details(self):
        current = self.get_current_dns()
        allowed = self.get_allowed_dns()
        messagebox.showinfo(
            "DNS Guard Details",
            "Aktuelle DNS:\n"
            + ("\n".join(current) if current else "–")
            + "\n\nErlaubte DNS:\n"
            + ("\n".join(allowed) if allowed else "–"),
            parent=self,
        )

    def dns_guard_ui_loop(self):
        try:
            self.update_dns_guard_indicator()
        finally:
            self.after(3000, self.dns_guard_ui_loop)

    def _set_dns_guard_status(self, status: str, last_action: str | None = None):
        self.dns_guard_status_var.set(f"DNS Guard: {status}")
        if last_action is not None:
            self.dns_guard_last_action_var.set(last_action)

    def _start_dns_guard_transition(self, delay_ms=2000):
        self.dns_guard_transition_until = time.time() + max(1.0, delay_ms / 1000.0)
        self.set_dns_guard_state("transition")
        self.after(delay_ms, self.run_dns_guard_check_once)

    def cleanup_old_vpn_dns_links(self):
        actions = []
        code, out = run_cmd(["resolvectl", "status"], timeout=8)
        if code != 0:
            return actions

        candidates = set()
        get_ifaces = getattr(self, "get_wg_interfaces", None)
        if callable(get_ifaces):
            try:
                candidates.update(
                    str(iface).strip() for iface in (get_ifaces() or []) if iface
                )
            except Exception:
                pass

        for line in (out or "").splitlines():
            stripped = line.strip()
            if not stripped.startswith("Link "):
                continue
            low = stripped.lower()
            iface = ""
            try:
                iface = stripped.split("(", 1)[1].split(")", 1)[0].strip()
            except Exception:
                iface = ""
            iface_low = iface.lower()
            if any(
                token in low or token in iface_low
                for token in ("proton", "wireguard", "wg", "tun", "vpn")
            ):
                if iface:
                    candidates.add(iface)

        for iface in sorted(candidate for candidate in candidates if candidate):
            rc, _cmd_out = run_cmd(
                ["sudo", "-n", "resolvectl", "revert", iface], timeout=10
            )
            actions.append(f"resolvectl revert {iface}: rc={rc}")

        return actions

    def prepare_dns_for_vpn_profile_switch(self):
        actions = []
        cleanup_failed = False
        sudo_failed = False

        self.dns_guard_transition_until = time.time() + 4.0
        try:
            self.set_dns_guard_state("transition")
        except Exception:
            pass

        try:
            actions.extend(self.cleanup_old_vpn_dns_links())
        except Exception:
            cleanup_failed = True

        for cmd, label in (
            (
                ["sudo", "-n", "rm", "-f", DOT_DROPIN_PATH],
                f"remove {os.path.basename(DOT_DROPIN_PATH)}",
            ),
            (
                ["sudo", "-n", "resolvectl", "flush-caches"],
                "resolvectl flush-caches",
            ),
            (
                ["sudo", "-n", "systemctl", "restart", "systemd-resolved"],
                "restart systemd-resolved",
            ),
        ):
            rc, out = run_cmd(cmd, timeout=15)
            actions.append(f"{label}: rc={rc}")
            if rc != 0:
                cleanup_failed = True
                low_out = (out or "").lower()
                if "sudo" in low_out or "password" in low_out:
                    sudo_failed = True

        self.__dict__["last_vpn_dns_cleanup_actions"] = list(actions)
        try:
            self.log_safe("VPN-Profilwechsel: DNS bereinigt, Recheck geplant.")
            if sudo_failed:
                self.log_safe("DNS Cleanup benötigt sudo-Rechte.")
        except Exception:
            pass

        if cleanup_failed:
            try:
                self.__dict__["dns_guard_last_cleanup_warning"] = True
            except Exception:
                pass

        return actions

    def finalize_dns_after_vpn_profile_switch(self):
        self.dns_guard_transition_until = 0.0
        self.__dict__.pop("_dns_finalize_scheduled_token", None)

        try:
            status = self._collect_dot_status(run_test=False)
            self._set_dot_state(
                status.get("ok"),
                profile=status.get("profile", ""),
                servers=status.get("servers") or [],
                mode=status.get("mode", ""),
                enabled=status.get("enabled"),
                vpn_conflict=status.get("vpn_conflict"),
                vpn_active=status.get("vpn_active"),
            )
        except Exception:
            pass

        try:
            self.rebuild_dns_guard_whitelist()
        except Exception:
            pass

        try:
            self.run_dns_guard_check_once()
        except Exception:
            pass

        try:
            self.update_dns_controls_state()
        except Exception:
            pass

        try:
            self.run_dns_leak_check(quiet=True)
        except Exception:
            pass

    def _schedule_dns_finalize_after_vpn_profile_switch(self, delay_ms=2500):
        callback = self.finalize_dns_after_vpn_profile_switch
        token = int(self.__dict__.get("_dns_finalize_token", 0)) + 1
        self.__dict__["_dns_finalize_token"] = token
        self.__dict__["_dns_finalize_scheduled_token"] = token

        def guarded_callback():
            if token != int(self.__dict__.get("_dns_finalize_token", 0)):
                return
            try:
                callback()
            except Exception:
                pass

        after = getattr(self, "after", None)
        if callable(after):
            try:
                after(delay_ms, guarded_callback)
                return
            except Exception:
                pass
        try:
            guarded_callback()
        except Exception:
            pass

    def set_dns_guard_state(
        self,
        state: str,
        *,
        status_text: str | None = None,
        last_action: str | None = None,
        log_message: str | None = None,
    ):
        state_value = str(state or "unknown").strip().lower() or "unknown"
        previous = self.__dict__.get("last_dns_guard_state")
        self.dns_guard_state = state_value
        self.last_dns_guard_state = state_value

        status_map = {
            "off": "deaktiviert",
            "ok": "aktiv / sauber",
            "problem": "Problem erkannt",
            "transition": "Wechsel wird geprüft",
            "unknown": "aktiv / unbekannt",
            "error": "aktiv / Fehler",
        }
        text = status_text or status_map.get(state_value, state_value)
        self._set_dns_guard_status(text, last_action)

        if log_message and previous != state_value:
            self.log_safe(log_message)

        self.update_dns_guard_indicator()

    def run_dns_guard_check_once(self, auto_repair=False):
        if self.__dict__.get("dns_guard_check_running", False):
            return

        now = time.time()
        if now < float(self.__dict__.get("dns_guard_transition_until", 0.0)):
            self.set_dns_guard_state("transition")
            return

        self.dns_guard_check_running = True
        try:
            self.last_dns_guard_check_ts = now
            guard_var = self.__dict__.get("dns_guard_enabled_var")
            guard_enabled = bool(guard_var.get()) if guard_var is not None else False
            if not guard_enabled:
                self.set_dns_guard_state("off")
                return

            allowed = self.rebuild_dns_guard_whitelist()
            current = self.get_current_dns()
            normalized_allowed = self._normalized_dns_values(allowed)

            if not current:
                self.set_dns_guard_state("unknown")
                return

            if not normalized_allowed:
                if self.is_vpn_active():
                    status_text = "aktiv / VPN-DNS unklar"
                elif self.is_dot_active():
                    status_text = "aktiv / DoT-Ziel unklar"
                else:
                    status_text = "aktiv / DNS-Ziel unklar"
                self.set_dns_guard_state(
                    "unknown",
                    status_text=status_text,
                    log_message="DNS Guard: DNS-Ziel unklar",
                )
                return

            if self.dns_is_clean(current, allowed):
                if self.is_vpn_active():
                    status_text = "aktiv / VPN-DNS OK"
                elif self.is_dot_active():
                    status_text = "aktiv (DoT-Resolver OK)"
                else:
                    status_text = "aktiv / sauber"
                self.set_dns_guard_state(
                    "ok",
                    status_text=status_text,
                    log_message="DNS Guard: OK",
                )
                return

            if self.is_dot_active() and not self.is_vpn_active():
                problem_text = "Problem (DoT-DNS abweichend)"
            else:
                problem_text = "Problem erkannt"
            self.set_dns_guard_state(
                "problem",
                status_text=problem_text,
                log_message="DNS Guard: Abweichung erkannt",
            )
            if not auto_repair:
                return

            repair_last = float(self.__dict__.get("last_dns_guard_repair_ts", 0.0))
            if (now - repair_last) < 30:
                return

            success, text = self.enforce_dns_clean()
            self.last_dns_guard_repair_ts = time.time()
            self.dns_guard_last_fix_ts = self.last_dns_guard_repair_ts
            ts_text = datetime.fromtimestamp(self.last_dns_guard_repair_ts).strftime(
                "%Y-%m-%d %H:%M:%S"
            )
            self.update_dns_output(text, append=True)
            self.set_dns_guard_state(
                "ok" if success else "problem",
                status_text="aktiv / korrigiert" if success else "Problem erkannt",
                last_action=f"Letzte Korrektur: {ts_text}",
                log_message=(
                    "DNS Guard: Inkonsistenz erkannt → Auto-Korrektur"
                    if success
                    else "DNS Guard: Auto-Korrektur fehlgeschlagen"
                ),
            )
            self.rebuild_dns_guard_whitelist()
        except Exception as exc:
            self.set_dns_guard_state(
                "error",
                last_action=self.dns_guard_last_action_var.get(),
                log_message="DNS Guard: Fehler erkannt",
            )
            self.log_debug_safe(f"DNS Guard Fehler: {exc}")
        finally:
            self.dns_guard_check_running = False

    def enforce_dns_clean(self):
        if not self._auto_dns_control_available():
            return (
                False,
                "=== DNS Guard ===\n"
                "Status: Korrektur übersprungen\n"
                "Grund: sudo -n ist nicht verfügbar.",
            )

        if self.is_vpn_active():
            return self.use_vpn_dns()

        selected = self.get_selected_user_dns()
        if selected:
            choice = self._manual_dns_choice()
            ok = self.apply_dns(
                choice=choice,
                refresh_ui=False,
                update_choice_var=False,
            )
            return (
                ok,
                "=== DNS Guard ===\n"
                f"Aktion: User DNS erzwingen\nProfil: {choice}\n"
                f"Ergebnis: {'erfolgreich' if ok else 'fehlgeschlagen'}",
            )

        return (
            True,
            "=== DNS Guard ===\n"
            "Status: Keine Korrektur nötig\n"
            "Grund: Kein manuelles DNS-Profil aktiv.",
        )

    def _auto_dns_control_available(self):
        return has_cmd("sudo") and check_sudo_nopasswd()

    def _show_auto_dns_policy_result(self, title: str, text: str):
        self.update_dns_output(text, append=True)
        self.refresh_dns_policy_status()
        self.update_dns_controls_state()
        self.log_safe(title)

    def initial_dns_cleanup(self, defer_followup=False):
        if not self._auto_dns_control_available():
            self.log_safe("DNS-Startbereinigung übersprungen: sudo -n nicht verfügbar")
            if not defer_followup:
                self.rebuild_dns_guard_whitelist()
                self.run_dns_guard_check_once()
                self.update_dns_controls_state_safe()
            return
        rc, out = run_cmd(
            ["sudo", "-n", "systemctl", "restart", "systemd-resolved"],
            timeout=15,
        )
        if rc == 0:
            self.log_safe("DNS-Startbereinigung: systemd-resolved neu gestartet")
        else:
            self.log_safe("DNS-Startbereinigung fehlgeschlagen")
            if out:
                self.log_debug_safe(out)
        if defer_followup:
            return
        self.rebuild_dns_guard_whitelist()
        self.run_dns_guard_check_once()
        self.update_dns_controls_state_safe()

    def _dns_startup_transition_active(self):
        return bool(self.__dict__.get("dns_startup_sequence_running", False))

    def _dns_runtime_transition_active(self):
        if self._dns_startup_transition_active():
            return True
        if time.time() < float(self.__dict__.get("dns_guard_transition_until", 0.0)):
            return True
        if bool(self.__dict__.get("dns_mode_running", False)):
            return True
        if bool(self.__dict__.get("dot_running", False)):
            return True
        active_operation = str(self.__dict__.get("network_active_operation") or "")
        active_operation = active_operation.strip().lower()
        return bool(
            active_operation
            and any(token in active_operation for token in ("vpn", "dns"))
        )

    def _schedule_dns_startup_step(self, delay_ms, callback):
        after = getattr(self, "after", None)
        if callable(after):
            try:
                after(delay_ms, callback)
                return
            except Exception:
                pass
        try:
            callback()
        except Exception:
            pass

    def _start_dns_startup_transition(self, duration_ms=4200):
        self.__dict__["dns_startup_sequence_running"] = True
        self.dns_guard_transition_until = time.time() + max(1.0, duration_ms / 1000.0)
        self.set_dns_guard_state(
            "transition",
            log_message="DNS-Startsequenz: Initialisierung läuft",
        )
        try:
            self.update_dns_controls_state()
        except Exception:
            pass

    def _run_dns_startup_cleanup_step(self):
        try:
            self.initial_dns_cleanup(defer_followup=True)
        except Exception as exc:
            self.log_debug_safe(f"DNS-Startsequenz Cleanup Fehler: {exc}")

    def _run_dns_startup_dot_step(self):
        try:
            self.startup_dot_status_check()
        except Exception as exc:
            self.log_debug_safe(f"DNS-Startsequenz DoT Fehler: {exc}")

    def _run_dns_startup_policy_step(self):
        try:
            self.startup_dns_policy_check()
        except Exception as exc:
            self.log_debug_safe(f"DNS-Startsequenz Policy Fehler: {exc}")

    def _finalize_dns_startup_sequence(self):
        self.__dict__["dns_startup_sequence_running"] = False
        self.dns_guard_transition_until = 0.0
        try:
            self.refresh_dns_policy_status()
        except Exception:
            pass
        try:
            self.update_dns_controls_state()
        except Exception:
            pass
        try:
            self.rebuild_dns_guard_whitelist()
        except Exception:
            pass
        try:
            self.run_dns_guard_check_once()
        except Exception:
            pass
        try:
            self.update_dns_guard_indicator()
        except Exception:
            pass

    def start_dns_startup_sequence(self):
        if self.__dict__.get("dns_startup_sequence_running", False):
            return
        self._start_dns_startup_transition(duration_ms=4200)
        self._schedule_dns_startup_step(0, self._run_dns_startup_cleanup_step)
        self._schedule_dns_startup_step(700, self._run_dns_startup_dot_step)
        self._schedule_dns_startup_step(1500, self._run_dns_startup_policy_step)
        self._schedule_dns_startup_step(2400, self._finalize_dns_startup_sequence)
        self._schedule_dns_startup_step(2600, self.dns_guard_ui_loop)
        self._schedule_dns_startup_step(3200, self.monitor_vpn_changes)
        self._schedule_dns_startup_step(3800, self.dns_guard_loop)

    def on_vpn_state_changed(self, vpn_active: bool):
        self.dns_guard_transition_until = time.time() + 5
        self.set_dns_guard_state(
            "transition",
            log_message=f"VPN-Wechsel erkannt: {'aktiv' if vpn_active else 'inaktiv'}",
        )
        try:
            self.update_dns_controls_state()
        except Exception:
            pass
        self._schedule_dns_finalize_after_vpn_profile_switch(delay_ms=2500)

    def on_vpn_connected(self):
        self.log_safe("VPN erkannt → DNS wird bereinigt")
        if not self._auto_dns_control_available():
            result = (
                "=== DNS Control ===\n"
                "Aktion: VPN verbunden\n"
                "Status: automatische DNS-Bereinigung übersprungen\n"
                "Grund: sudo -n ist nicht verfügbar."
            )
            self._show_auto_dns_policy_result(
                "VPN erkannt → DNS-Bereinigung übersprungen",
                result,
            )
            return

        def worker():
            success, text = self.use_vpn_dns()

            def finish():
                self._show_auto_dns_policy_result(
                    (
                        "DNS auf VPN-DNS umgestellt"
                        if success
                        else "DNS-Bereinigung bei VPN fehlgeschlagen"
                    ),
                    text,
                )

            self.after(0, finish)

        threading.Thread(target=worker, daemon=True).start()

    def on_vpn_disconnected(self):
        self.log_safe("VPN getrennt → DNS wird zurückgesetzt")
        if not self._auto_dns_control_available():
            result = (
                "=== DNS Control ===\n"
                "Aktion: VPN getrennt\n"
                "Status: DNS-Reset übersprungen\n"
                "Grund: sudo -n ist nicht verfügbar."
            )
            self._show_auto_dns_policy_result(
                "VPN getrennt → DNS-Reset übersprungen",
                result,
            )
            return

        service = self._dns_control_service()

        def worker():
            if service is None:
                success = False
                text = (
                    "=== DNS Control ===\n"
                    "Aktion: VPN getrennt\n"
                    "Status: fehlgeschlagen\n"
                    "Grund: DNS-Control-Service nicht verfügbar."
                )
            else:
                nm_ok, nm_details = service.apply_nm_dns_policy_to_targets(
                    ignore_auto_dns=False,
                    targets=self._active_nm_dns_targets(),
                )
                rm_ok, rm_detail = service.remove_dropin()
                restart_ok, restart_detail = service.restart_resolved()
                success = nm_ok and rm_ok and restart_ok
                lines = [
                    "=== DNS Control ===",
                    "Aktion: VPN getrennt / DNS zurück auf Standard",
                    f"Ergebnis: {'erfolgreich' if success else 'fehlgeschlagen'}",
                    "",
                    "NetMon DoT Drop-in:",
                    "entfernt" if rm_ok else f"Fehler: {rm_detail or 'unknown'}",
                    "",
                    "NetworkManager:",
                    *nm_details,
                    "systemd-resolved Restart:",
                    "OK" if restart_ok else (restart_detail or "fehlgeschlagen"),
                ]
                text = "\n".join(lines).strip()

            def finish():
                self._show_auto_dns_policy_result(
                    (
                        "DNS zurück auf Standard"
                        if success
                        else "DNS-Reset nach VPN fehlgeschlagen"
                    ),
                    text,
                )

            self.after(0, finish)

        threading.Thread(target=worker, daemon=True).start()

    def monitor_vpn_changes(self):
        try:
            if self._dns_runtime_transition_active():
                self.update_dns_controls_state()
                return
            current = bool(self.get_vpn_state())
            if current != bool(self.last_vpn_state):
                self.on_vpn_state_changed(current)
                self.last_vpn_state = current
            self.update_dns_controls_state()
        except Exception as exc:
            self.log_debug_safe(f"VPN-State-Monitor Fehler: {exc}")
        finally:
            self.after(3000, self.monitor_vpn_changes)

    def dns_guard_loop(self):
        try:
            if self._dns_runtime_transition_active():
                self.set_dns_guard_state("transition")
                return
            self.run_dns_guard_check_once(auto_repair=True)
        except Exception as exc:
            self.log_debug_safe(f"DNS Guard Fehler: {exc}")
            self.set_dns_guard_state("error")
        finally:
            self.after(5000, self.dns_guard_loop)

    def _async_job_label(self, func, job_name=""):
        explicit = str(job_name or "").strip()
        if explicit:
            return explicit
        inferred = str(getattr(func, "__name__", "") or "").strip()
        if inferred and inferred != "<lambda>":
            return inferred
        return "unbenannter Async-Job"

    def _log_async_job_error(self, func, exc, job_name=""):
        debug_logger = self.__dict__.get("log_debug_safe")
        plain_logger = self.__dict__.get("log")
        label = self._async_job_label(func, job_name=job_name)
        message = f"Async-Job Fehler [{label}]: {exc}"
        try:
            if callable(debug_logger):
                debug_logger(message)
            elif callable(plain_logger):
                plain_logger(message)
        except Exception:
            pass
        if callable(debug_logger):
            try:
                debug_logger(traceback.format_exc())
            except Exception:
                pass

    def run_async(self, func, job_name=""):
        def runner():
            try:
                func()
            except Exception as exc:
                self._log_async_job_error(func, exc, job_name=job_name)

        threading.Thread(target=runner, daemon=True).start()

    def _sync_network_operation_state_from_coordinator(self):
        coordinator = self.__dict__.get("network_ops_coordinator")
        if coordinator is None:
            return
        self.network_ops_lock = coordinator.lock
        self.network_active_operation = coordinator.active_operation
        self._network_op_owner = coordinator.owner_thread_id
        self._network_op_depth = coordinator.depth

    def _network_service(self):
        service = self.__dict__.get("network_service")
        if service is not None:
            return service
        if NetmonNetworkService is None:
            return None
        service = NetmonNetworkService(run_cmd)
        self.network_service = service
        return service

    def _state_coordinator(self):
        coordinator = self.__dict__.get("state_coordinator")
        if coordinator is not None:
            return coordinator
        if NetmonStateCoordinator is None:
            return None
        coordinator = NetmonStateCoordinator()
        self.state_coordinator = coordinator
        return coordinator

    def _dns_control_service(self):
        service = self.__dict__.get("dns_control_service")
        if service is not None:
            return service
        if DNSControlService is None:
            return None
        service = DNSControlService(
            run_cmd,
            has_cmd,
            dot_dropin_path=DOT_DROPIN_PATH,
            dot_dropin_dir=DOT_DROPIN_DIR,
        )
        self.dns_control_service = service
        return service

    def _new_rollback_transaction(self, name):
        if RollbackTransaction is None:
            return None
        return RollbackTransaction(
            name,
            log_fn=self.log_safe,
            log_debug_fn=self.log_debug_safe,
        )

    def _security_parser(self):
        parser = self.__dict__.get("security_parser")
        if parser is not None:
            return parser

        parser_cls = NetmonSecurityParser
        if parser_cls is None:
            if self.__dict__.get("_security_parser_load_failed"):
                return None
            module_path = os.path.join(
                os.path.dirname(__file__), "netmon_security_parser.py"
            )
            try:
                spec = importlib.util.spec_from_file_location(
                    "netmon_security_parser_dynamic", module_path
                )
                if spec and spec.loader:
                    mod = importlib.util.module_from_spec(spec)
                    spec.loader.exec_module(mod)
                    parser_cls = getattr(mod, "NetmonSecurityParser", None)
            except Exception as exc:
                parser_cls = None
                self.log_debug(f"Security parser dynamisch laden fehlgeschlagen: {exc}")
            if parser_cls is None:
                self._security_parser_load_failed = True
                return None

        try:
            parser = parser_cls()
        except Exception as exc:
            self.log_debug(f"Security parser Initialisierung fehlgeschlagen: {exc}")
            self._security_parser_load_failed = True
            return None
        self.security_parser = parser
        return parser

    def _security_store(self):
        store = self.__dict__.get("security_store")
        if store is not None:
            return store
        if NetmonSecurityStore is None:
            return None
        store = NetmonSecurityStore()
        self.security_store = store
        return store

    def _try_begin_network_operation(self, operation, log_busy=True):
        coordinator = self.__dict__.get("network_ops_coordinator")
        if coordinator is not None:
            started = coordinator.try_begin(operation, log_busy=log_busy)
            self._sync_network_operation_state_from_coordinator()
            return started

        lock = self.__dict__.get("network_ops_lock")
        if lock is None:
            lock = threading.RLock()
            self.network_ops_lock = lock
        current_tid = threading.get_ident()
        owner = self.__dict__.get("_network_op_owner")
        depth = self.__dict__.get("_network_op_depth", 0)

        if owner == current_tid:
            lock.acquire()
            self._network_op_depth = depth + 1
            return True

        if not lock.acquire(blocking=False):
            if log_busy:
                active = (
                    self.__dict__.get("network_active_operation")
                    or "andere Netzwerkaktion"
                ).strip()
                self.log_safe(
                    f"Netzwerkaktion blockiert: {operation} wird nicht gestartet, da bereits '{active}' läuft."
                )
            return False

        self._network_op_owner = current_tid
        self._network_op_depth = 1
        self.network_active_operation = operation
        return True

    def _end_network_operation(self):
        coordinator = self.__dict__.get("network_ops_coordinator")
        if coordinator is not None:
            coordinator.end()
            self._sync_network_operation_state_from_coordinator()
            return

        lock = self.__dict__.get("network_ops_lock")
        if lock is None:
            return
        current_tid = threading.get_ident()
        owner = self.__dict__.get("_network_op_owner")
        if owner != current_tid:
            return

        depth = max(0, int(self.__dict__.get("_network_op_depth", 1)) - 1)
        self._network_op_depth = depth
        if depth == 0:
            self._network_op_owner = None
            self.network_active_operation = ""
        lock.release()

    def _refresh_security_table_safe(self):
        if threading.current_thread() is threading.main_thread():
            self._refresh_security_table()
            return
        self.after(0, self._refresh_security_table)

    def clear_log(self):
        self.text.delete("1.0", "end")

    def copy_log_selection(self, _event=None):
        try:
            selected = self.text.get("sel.first", "sel.last")
        except tk.TclError:
            return "break"
        self.clipboard_clear()
        self.clipboard_append(selected)
        return "break"

    def prepare_log_selection_chad_analysis(self):
        try:
            selected = self.text.get("sel.first", "sel.last").strip()
        except tk.TclError:
            self.chad_analysis_status_var.set("Bitte zuerst Text im Log markieren.")
            return
        if not selected:
            self.chad_analysis_status_var.set("Bitte zuerst Text im Log markieren.")
            return
        self.prepare_chad_analysis("Log Analyse", selected)

    def _persist_log_overlay_preferences(self):
        if "privacy_config" not in self.__dict__:
            return
        try:
            self.save_privacy_config()
        except Exception:
            pass

    def _default_log_overlay_dir(self):
        pictures_dir = get_user_pictures_dir() or os.path.expanduser("~")
        log_pics_dir = os.path.join(pictures_dir, "Log-Pics")
        try:
            os.makedirs(log_pics_dir, exist_ok=True)
        except Exception:
            pass
        if os.path.isdir(log_pics_dir):
            return log_pics_dir
        if os.path.isdir(pictures_dir):
            return pictures_dir
        return os.path.expanduser("~")

    def _preferred_log_overlay_dir(self):
        picker_dir = str(self.__dict__.get("log_overlay_picker_dir") or "").strip()
        if picker_dir and os.path.isdir(picker_dir):
            return picker_dir

        config = getattr(self, "privacy_config", {})
        if isinstance(config, dict):
            saved = str(config.get("log_overlay_last_dir") or "").strip()
            if saved and os.path.isdir(saved):
                self.log_overlay_picker_dir = saved
                return saved

        default_dir = self._default_log_overlay_dir()
        self.log_overlay_picker_dir = default_dir
        return default_dir

    def _list_overlay_images(self, folder):
        folder_path = os.path.abspath(os.path.expanduser(folder or "")).strip()
        if not folder_path or not os.path.isdir(folder_path):
            return []
        exts = {".png", ".jpg", ".jpeg", ".PNG", ".JPG", ".JPEG"}
        files = []
        try:
            for name in sorted(os.listdir(folder_path)):
                full = os.path.join(folder_path, name)
                if not os.path.isfile(full):
                    continue
                _, ext = os.path.splitext(name)
                if ext in exts:
                    files.append(full)
        except Exception:
            return []
        return files

    def open_log_overlay_picker(self, folder=""):
        picker = self.__dict__.get("log_overlay_picker_frame")
        if picker is None:
            return False

        target_folder = (folder or self.log_overlay_picker_dir or "").strip()
        if not target_folder:
            current_path = str(self.log_overlay_image_path or "").strip()
            if current_path:
                target_folder = os.path.dirname(current_path)

        if not target_folder:
            target_folder = self._preferred_log_overlay_dir()

        target_folder = os.path.abspath(os.path.expanduser(target_folder))
        files = self._list_overlay_images(target_folder)
        if not files:
            self.log_safe("Keine Bilddateien (PNG/JPG) im gewählten Ordner gefunden.")
            return False

        self.log_overlay_picker_open = True
        self.log_overlay_picker_dir = target_folder
        self.log_overlay_picker_files = files

        self.log_overlay_picker_list.delete(0, "end")
        for img_path in files:
            self.log_overlay_picker_list.insert("end", os.path.basename(img_path))

        selected_path = str(self.log_overlay_image_path or "").strip()
        idx = 0
        if selected_path and selected_path in files:
            idx = files.index(selected_path)

        self.log_overlay_picker_index = idx
        self.log_overlay_picker_list.selection_clear(0, "end")
        self.log_overlay_picker_list.selection_set(idx)
        self.log_overlay_picker_list.activate(idx)
        self.log_overlay_picker_list.see(idx)

        picker.place(relx=1.0, rely=1.0, x=-10, y=-52, anchor="se")
        picker.lift()
        self._apply_log_overlay_picker_index(idx)
        return True

    def close_log_overlay_picker(self):
        picker = self.__dict__.get("log_overlay_picker_frame")
        if picker is not None:
            picker.place_forget()
        self.log_overlay_picker_open = False

    def _choose_log_overlay_folder(self):
        initial_dir = self._preferred_log_overlay_dir()
        folder = filedialog.askdirectory(
            title="Ordner mit Overlay-Bildern wählen",
            initialdir=initial_dir,
        )
        if not folder:
            return
        if self.open_log_overlay_picker(folder):
            self.log_safe(f"Overlay-Bilder geladen: {folder}")

    def _apply_log_overlay_picker_index(self, idx):
        files = list(self.log_overlay_picker_files or [])
        if not files:
            return False
        if idx < 0 or idx >= len(files):
            return False

        target = files[idx]
        ok = self._load_log_overlay_image(target)
        if not ok:
            return False

        self.log_overlay_picker_index = idx
        self._render_log_overlay_image()
        self._set_log_overlay_visible(True)
        return True

    def _step_log_overlay_selection(self, step):
        files = list(self.log_overlay_picker_files or [])
        if not files:
            return
        current = int(self.log_overlay_picker_index)
        if current < 0:
            current = 0
        nxt = (current + int(step)) % len(files)
        self.log_overlay_picker_list.selection_clear(0, "end")
        self.log_overlay_picker_list.selection_set(nxt)
        self.log_overlay_picker_list.activate(nxt)
        self.log_overlay_picker_list.see(nxt)
        self._apply_log_overlay_picker_index(nxt)

    def _on_log_overlay_picker_select(self, _event=None):
        selection = self.log_overlay_picker_list.curselection()
        if not selection:
            return
        idx = int(selection[0])
        self._apply_log_overlay_picker_index(idx)

    def _format_log_overlay_image_name(self, image_path):
        base = os.path.basename(str(image_path or "").strip())
        if not base:
            return ""
        stem, _ext = os.path.splitext(base)
        title = (stem or base).replace("_", " ").strip()
        return title or base

    def _update_log_overlay_title(self):
        badge = self.__dict__.get("log_overlay_name_badge")
        text_var = self.__dict__.get("log_overlay_name_var")
        if badge is None or text_var is None:
            return

        if not self.log_overlay_active:
            badge.place_forget()
            text_var.set("")
            return

        title = self._format_log_overlay_image_name(self.log_overlay_image_path)
        if not title:
            badge.place_forget()
            text_var.set("")
            return

        text_var.set(title)
        badge.place(relx=0.5, rely=0.0, y=12, anchor="n")
        badge.lift()

    def _restore_log_overlay_from_config(self):
        self.close_log_overlay_picker()
        config = getattr(self, "privacy_config", {})
        if not isinstance(config, dict):
            self._set_log_overlay_visible(False, persist=False)
            return

        image_path = str(config.get("log_overlay_image_path") or "").strip()
        active = bool(config.get("log_overlay_active", False))
        last_dir = str(config.get("log_overlay_last_dir") or "").strip()

        if last_dir and os.path.isdir(last_dir):
            self.log_overlay_picker_dir = last_dir
        elif image_path:
            guessed = os.path.dirname(image_path)
            if guessed and os.path.isdir(guessed):
                self.log_overlay_picker_dir = guessed
            else:
                self.log_overlay_picker_dir = self._default_log_overlay_dir()
        else:
            self.log_overlay_picker_dir = self._default_log_overlay_dir()

        loaded = False
        if image_path:
            loaded = self._load_log_overlay_image(image_path)
            if loaded:
                self._render_log_overlay_image()

        self._set_log_overlay_visible(active and loaded, persist=False)

    def _set_log_overlay_visible(self, visible, persist=True):
        self.log_overlay_active = bool(visible)
        if not hasattr(self, "log_overlay_frame"):
            return

        if self.log_overlay_active:
            self.log_overlay_frame.place(
                relx=0.0, rely=0.0, relwidth=1.0, relheight=1.0
            )
            self.log_overlay_frame.lift()
            if hasattr(self, "log_overlay_controls"):
                self.log_overlay_controls.lift()
            if self.log_overlay_picker_open and hasattr(
                self, "log_overlay_picker_frame"
            ):
                self.log_overlay_picker_frame.lift()
            if hasattr(self, "log_overlay_toggle_text"):
                self.log_overlay_toggle_text.set("Log")
        else:
            self.log_overlay_frame.place_forget()
            if hasattr(self, "log_overlay_toggle_text"):
                self.log_overlay_toggle_text.set("Bild")

        self._update_log_overlay_title()

        if persist:
            self._persist_log_overlay_preferences()

    def _on_log_overlay_resize(self, _event=None):
        if self.log_overlay_active:
            self._render_log_overlay_image()

    def _log_overlay_max_source_size(self):
        fallback_w, fallback_h = 1920, 1080
        frame = self.__dict__.get("log_view_frame")
        if frame is None:
            return fallback_w, fallback_h

        width = int(frame.winfo_width() or 0)
        height = int(frame.winfo_height() or 0)
        if width <= 1 or height <= 1:
            return fallback_w, fallback_h

        max_w = max(640, min(width * 2, 3840))
        max_h = max(360, min(height * 2, 2160))
        return max_w, max_h

    def _load_log_overlay_image(self, path):
        image_path = os.path.abspath(os.path.expanduser(path or "")).strip()
        if not image_path:
            return False
        if not os.path.isfile(image_path):
            self.log_safe("Log-Bild nicht gefunden.")
            return False

        self.log_overlay_source_image = None
        self.log_overlay_source_photo = None
        self.log_overlay_photo = None

        if Image is not None and ImageTk is not None:
            try:
                max_w, max_h = self._log_overlay_max_source_size()
                with warnings.catch_warnings():
                    bomb_warning = getattr(Image, "DecompressionBombWarning", None)
                    if bomb_warning is not None:
                        warnings.simplefilter("error", bomb_warning)
                    img = Image.open(image_path)
                    img.load()

                if img.mode != "RGB":
                    img = img.convert("RGB")

                if img.width > max_w or img.height > max_h:
                    if hasattr(Image, "Resampling"):
                        img.thumbnail((max_w, max_h), Image.Resampling.LANCZOS)
                    else:
                        img.thumbnail((max_w, max_h), Image.LANCZOS)

                self.log_overlay_source_image = img
                self.log_overlay_image_path = image_path
                self.log_overlay_picker_dir = os.path.dirname(image_path)
                self._persist_log_overlay_preferences()
                return True
            except Exception as exc:
                bomb_warning = getattr(Image, "DecompressionBombWarning", None)
                if bomb_warning is not None and isinstance(exc, bomb_warning):
                    self.log_safe(
                        "Log-Bild abgelehnt: extrem große Bilddatei (Pillow-Schutz aktiv)."
                    )
                else:
                    self.log_safe(f"Log-Bild konnte nicht geladen werden: {exc}")
                return False

        try:
            self.log_overlay_source_photo = tk.PhotoImage(file=image_path)
            self.log_overlay_image_path = image_path
            self.log_overlay_picker_dir = os.path.dirname(image_path)
            self._persist_log_overlay_preferences()
            return True
        except Exception:
            detail = (PIL_IMPORT_ERROR or "").strip()
            if detail:
                self.log_safe(
                    "Bildformat nicht unterstützt. Pillow ist nicht nutzbar: " + detail
                )
            else:
                self.log_safe(
                    "Bildformat nicht unterstützt. Für JPG-Skalierung bitte Pillow installieren (python3-pil)."
                )
            return False

    def _render_log_overlay_image(self):
        if not hasattr(self, "log_overlay_label"):
            return

        if (
            self.log_overlay_source_image is None
            and self.log_overlay_source_photo is None
        ):
            self.log_overlay_label.configure(
                image="", text="Kein Bild geladen.", fg="#f0f0f0"
            )
            self.log_overlay_photo = None
            return

        if (
            self.log_overlay_source_image is not None
            and Image is not None
            and ImageTk is not None
        ):
            width = max(1, int(self.log_view_frame.winfo_width()))
            height = max(1, int(self.log_view_frame.winfo_height()))
            if width <= 1 or height <= 1:
                return

            src = self.log_overlay_source_image
            src_w, src_h = src.size
            if src_w <= 0 or src_h <= 0:
                self.log_overlay_label.configure(
                    image="", text="Bild ungültig.", fg="#f0f0f0"
                )
                self.log_overlay_photo = None
                return

            # cover: Bereich vollständig füllen, Seitenverhältnis erhalten,
            # bei Bedarf zentriert beschneiden.
            scale = max(width / src_w, height / src_h)
            new_w = max(1, int(round(src_w * scale)))
            new_h = max(1, int(round(src_h * scale)))

            if hasattr(Image, "Resampling"):
                resized = src.resize((new_w, new_h), Image.Resampling.LANCZOS)
            else:
                resized = src.resize((new_w, new_h), Image.LANCZOS)

            if resized.mode != "RGB":
                resized = resized.convert("RGB")

            left = max(0, (new_w - width) // 2)
            top = max(0, (new_h - height) // 2)
            right = left + width
            bottom = top + height
            cover_img = resized.crop((left, top, right, bottom))

            self.log_overlay_photo = ImageTk.PhotoImage(cover_img)
            self.log_overlay_label.configure(image=self.log_overlay_photo, text="")
            return

        if self.log_overlay_source_photo is not None:
            self.log_overlay_photo = self.log_overlay_source_photo
            self.log_overlay_label.configure(image=self.log_overlay_photo, text="")
            return

    def choose_log_overlay_image(self):
        folder = self.log_overlay_picker_dir
        if not folder:
            current_path = str(self.log_overlay_image_path or "").strip()
            if current_path:
                folder = os.path.dirname(current_path)
        if not folder or not os.path.isdir(folder):
            folder = self._preferred_log_overlay_dir()

        if self.open_log_overlay_picker(folder):
            self.log_safe("Bildauswahl geöffnet (mit 'Auswahl schließen' beenden).")

    def toggle_log_overlay(self):
        if not self.log_overlay_active:
            if not (
                self.log_overlay_source_image is not None
                or self.log_overlay_source_photo is not None
            ):
                self.choose_log_overlay_image()
                return
            self._render_log_overlay_image()
            self._set_log_overlay_visible(True)
            if self.log_overlay_picker_open and hasattr(
                self, "log_overlay_picker_frame"
            ):
                self.log_overlay_picker_frame.lift()
            self.log_safe("Log-Overlay aktiviert.")
            return

        self._set_log_overlay_visible(False)
        self.log_safe("Log-Overlay deaktiviert.")

    def tr(self, key: str, default: str | None = None) -> str:
        lang = self.__dict__.get("lang")
        if isinstance(lang, dict):
            value = lang.get(key)
            if isinstance(value, str) and value:
                return value
        if default is not None:
            return default
        return key

    def tooltip_lang_key(self, text: str) -> str:
        raw = str(text or "").strip()
        digest = hashlib.sha1(raw.encode("utf-8")).hexdigest()[:12]
        return f"tooltip_auto_{digest}"

    def _translate_tooltip_text(self, text_or_key: str) -> str:
        raw = str(text_or_key or "").strip()
        if not raw:
            return ""

        lang = self.__dict__.get("lang")
        if isinstance(lang, dict):
            value = lang.get(raw)
            if isinstance(value, str) and value:
                return value

            tooltips = lang.get("tooltips")
            if isinstance(tooltips, dict):
                value = tooltips.get(self.tooltip_lang_key(raw))
                if isinstance(value, str) and value:
                    return value

            tooltip_map = lang.get("tooltips_by_source")
            if isinstance(tooltip_map, dict):
                value = tooltip_map.get(raw)
                if isinstance(value, str) and value:
                    return value

        return raw

    def add_tooltip(self, widget, text):
        translated = self._translate_tooltip_text(text)
        tip = SimpleTooltip(widget, translated, self.tooltips_enabled)
        tip._netmon_tooltip_source = text
        self._tooltips.append(tip)

    def open_external_terminal(self):
        """Öffnet ein externes System-Terminal."""
        candidates = [
            ["x-terminal-emulator"],
            ["gnome-terminal"],
            ["konsole"],
            ["xfce4-terminal"],
            ["mate-terminal"],
            ["lxterminal"],
            ["xterm"],
        ]
        for cmd in candidates:
            if shutil.which(cmd[0]):
                try:
                    subprocess.Popen(cmd, start_new_session=True)
                    self.log(f"Externes Terminal gestartet: {cmd[0]}")
                    return
                except Exception as e:
                    self.log(f"Start fehlgeschlagen ({cmd[0]}): {e}")
        messagebox.showerror(
            "Terminal",
            "Kein unterstütztes Terminal-Programm gefunden.",
            parent=self,
        )

    def on_tooltip_toggle(self):
        if self.tooltips_enabled.get():
            return
        for tip in self._tooltips:
            try:
                tip._hide()
            except Exception:
                pass

    # ---------- System Monitor ----------
    def _init_system_monitor_styles(self):
        self.style.configure(
            "SystemGreen.Horizontal.TProgressbar",
            troughcolor="#d9d9d9",
            background="#2e7d32",
        )
        self.style.configure(
            "SystemYellow.Horizontal.TProgressbar",
            troughcolor="#d9d9d9",
            background="#f9a825",
        )
        self.style.configure(
            "SystemRed.Horizontal.TProgressbar",
            troughcolor="#d9d9d9",
            background="#c62828",
        )

    def _build_system_tab(self):
        top = ttk.Frame(self.tab_system)
        top.pack(fill="x", pady=(0, 8))
        ttk.Label(
            top, text="Integrierter Systemmonitor (aktualisiert alle 1-2 Sekunden)."
        ).pack(side="left")
        ttk.Button(
            top, text="Jetzt aktualisieren", command=self.system_monitor_refresh
        ).pack(side="right")

        body = ttk.Frame(self.tab_system)
        body.pack(fill="both", expand=True)
        body.columnconfigure(0, weight=1)
        body.columnconfigure(1, weight=0)

        self.system_cpu_var = tk.StringVar(value="CPU: –")
        self.system_ram_var = tk.StringVar(value="RAM: –")
        self.system_swap_var = tk.StringVar(value="Swap: –")
        self.system_disk_var = tk.StringVar(value="Disk /: –")
        self.system_temp_var = tk.StringVar(value="CPU-Temp: –")
        self.system_uptime_var = tk.StringVar(value="Uptime: –")

        self.system_cpu_bar = self._system_row(
            body, 0, "CPU-Auslastung", self.system_cpu_var
        )
        self.system_ram_bar = self._system_row(body, 1, "RAM", self.system_ram_var)
        self.system_swap_bar = self._system_row(body, 2, "Swap", self.system_swap_var)
        self.system_disk_bar = self._system_row(body, 3, "Disk /", self.system_disk_var)
        self.system_temp_bar = self._system_row(
            body, 4, "CPU-Temperatur", self.system_temp_var
        )

        up = ttk.LabelFrame(body, text="Uptime", padding=8)
        up.grid(row=5, column=0, columnspan=2, sticky="ew", pady=(8, 0))
        ttk.Label(up, textvariable=self.system_uptime_var).pack(anchor="w")

        sysf = ttk.LabelFrame(body, text="Sysinfo", padding=8)
        sysf.grid(row=6, column=0, columnspan=2, sticky="nsew", pady=(10, 0))
        sysf.columnconfigure(0, weight=1)
        sysf.columnconfigure(1, weight=1)
        sysf.rowconfigure(0, weight=1)
        sysf.rowconfigure(1, weight=1)
        body.rowconfigure(6, weight=1)

        self.sysinfo_vars = {}
        self._build_sysinfo_group(
            sysf,
            row=0,
            col=0,
            title="Betriebssystem",
            entries=[
                ("os_distribution", "Distribution"),
                ("os_version", "Version"),
                ("os_kernel", "Kernel"),
                ("os_arch", "Architektur"),
                ("os_hostname", "Hostname"),
                ("os_user", "Benutzer"),
            ],
        )
        self._build_sysinfo_group(
            sysf,
            row=0,
            col=1,
            title="Hardware",
            entries=[
                ("hw_cpu", "CPU"),
                ("hw_cores_threads", "Kerne/Threads"),
                ("hw_gpu", "GPU"),
                ("hw_board", "Mainboard"),
                ("hw_ram_total", "RAM gesamt"),
                ("hw_swap_total", "Swap gesamt"),
            ],
        )
        self._build_sysinfo_group(
            sysf,
            row=1,
            col=0,
            title="System",
            entries=[
                ("sys_root", "Root /"),
                ("sys_swap", "Swap"),
            ],
        )
        self._build_sysinfo_group(
            sysf,
            row=1,
            col=1,
            title="Software",
            entries=[
                ("sw_python", "Python"),
                ("sw_desktop", "Desktop-Session"),
            ],
        )

        self._load_sysinfo_static()
        self._update_sysinfo_dynamic()

    def _system_row(self, parent, row, title, value_var):
        box = ttk.LabelFrame(parent, text=title, padding=8)
        box.grid(row=row, column=0, sticky="ew", pady=(0 if row == 0 else 8, 0))
        box.columnconfigure(0, weight=1)
        ttk.Label(box, textvariable=value_var).grid(row=0, column=0, sticky="w")
        status = tk.Label(box, text="–", fg="#666666")
        status.grid(row=0, column=1, sticky="e", padx=(8, 0))
        bar = ttk.Progressbar(
            box,
            mode="determinate",
            maximum=100,
            style="SystemGreen.Horizontal.TProgressbar",
        )
        bar.grid(row=1, column=0, sticky="ew", pady=(6, 0))
        return {"bar": bar, "status": status}

    def _system_color_band(self, percent):
        p = max(0.0, min(100.0, float(percent)))
        if p < 70:
            return (
                "OK",
                "#2e7d32",
                "SystemGreen.Horizontal.TProgressbar",
            )
        if p < 90:
            return (
                "Warnung",
                "#b26a00",
                "SystemYellow.Horizontal.TProgressbar",
            )
        return (
            "Kritisch",
            "#b71c1c",
            "SystemRed.Horizontal.TProgressbar",
        )

    def _set_system_row_state(self, row_widgets, percent):
        label, color, style = self._system_color_band(percent)
        row_widgets["status"].config(text=label, fg=color)
        row_widgets["bar"]["value"] = percent
        row_widgets["bar"]["style"] = style

    def _system_temp_band(self, temp_c):
        # Gleiche Ampellogik, aber mit praxisnahen CPU-Temperatur-Schwellen.
        t = float(temp_c)
        if t < 70:
            return (
                "OK",
                "#2e7d32",
                "SystemGreen.Horizontal.TProgressbar",
            )
        if t < 85:
            return (
                "Warnung",
                "#b26a00",
                "SystemYellow.Horizontal.TProgressbar",
            )
        return (
            "Kritisch",
            "#b71c1c",
            "SystemRed.Horizontal.TProgressbar",
        )

    def _set_system_temp_state(self, row_widgets, temp_c):
        label, color, style = self._system_temp_band(temp_c)
        row_widgets["status"].config(text=label, fg=color)
        row_widgets["bar"]["value"] = max(0.0, min(100.0, float(temp_c)))
        row_widgets["bar"]["style"] = style

    def _format_uptime(self, seconds):
        total = max(0, int(seconds))
        days, rest = divmod(total, 86400)
        hours, rest = divmod(rest, 3600)
        minutes, _ = divmod(rest, 60)
        if days > 0:
            return f"{days}d {hours}h {minutes}m"
        if hours > 0:
            return f"{hours}h {minutes}m"
        return f"{minutes}m"

    def _build_sysinfo_group(self, parent, row, col, title, entries):
        box = ttk.LabelFrame(parent, text=title, padding=8)
        box.grid(
            row=row,
            column=col,
            sticky="nsew",
            padx=(0 if col == 0 else 6, 0),
            pady=(0 if row == 0 else 6, 0),
        )
        box.columnconfigure(1, weight=1)
        for idx, (key, label_text) in enumerate(entries):
            ttk.Label(box, text=f"{label_text}:").grid(
                row=idx, column=0, sticky="w", padx=(0, 8), pady=(0, 2)
            )
            var = tk.StringVar(value="-")
            ttk.Label(box, textvariable=var).grid(
                row=idx, column=1, sticky="w", pady=(0, 2)
            )
            self.sysinfo_vars[key] = var

    def _set_sysinfo_value(self, key, value):
        var = self.sysinfo_vars.get(key)
        if var is None:
            return
        txt = "-" if value is None else str(value).strip()
        var.set(txt if txt else "-")

    def _safe_read_file(self, path):
        try:
            with open(path, "r", encoding="utf-8") as fh:
                return fh.read().strip()
        except Exception:
            return ""

    def _read_os_release(self):
        info = {}
        path = "/etc/os-release"
        try:
            with open(path, "r", encoding="utf-8") as fh:
                for line in fh:
                    line = line.strip()
                    if not line or "=" not in line:
                        continue
                    key, val = line.split("=", 1)
                    info[key.strip()] = val.strip().strip('"')
        except Exception:
            return {}
        return info

    def _read_cpu_model(self):
        try:
            with open("/proc/cpuinfo", "r", encoding="utf-8") as fh:
                for line in fh:
                    if ":" not in line:
                        continue
                    key, val = line.split(":", 1)
                    if key.strip().lower() in {"model name", "hardware", "processor"}:
                        model = val.strip()
                        if model:
                            return model
        except Exception:
            pass
        rc, out = run_cmd(["lscpu"], timeout=5)
        if rc == 0 and out:
            for line in out.splitlines():
                if ":" not in line:
                    continue
                key, val = line.split(":", 1)
                if key.strip().lower() == "model name":
                    model = val.strip()
                    if model:
                        return model
        return "-"

    def _read_cpu_cores_threads(self):
        cores = None
        threads = None
        if psutil is not None:
            cores = psutil.cpu_count(logical=False)
            threads = psutil.cpu_count(logical=True)
        if cores is None or threads is None:
            rc, out = run_cmd(["lscpu"], timeout=5)
            if rc == 0 and out:
                data = {}
                for line in out.splitlines():
                    if ":" not in line:
                        continue
                    key, val = line.split(":", 1)
                    data[key.strip().lower()] = val.strip()
                try:
                    sockets = int(data.get("socket(s)", "1"))
                    core_per_socket = int(data.get("core(s) per socket", "0"))
                    if core_per_socket > 0:
                        cores = core_per_socket * sockets
                except Exception:
                    pass
                try:
                    threads = int(data.get("cpu(s)", "0")) or threads
                except Exception:
                    pass
        if threads is None:
            threads = os.cpu_count()
        return cores or "-", threads or "-"

    def _read_gpu_model(self):
        rc, out = run_cmd(["lspci"], timeout=6)
        if rc != 0 or not out:
            return "-"
        candidates = []
        for line in out.splitlines():
            low = line.lower()
            if (
                "vga compatible controller" in low
                or "3d controller" in low
                or "display controller" in low
            ):
                candidates.append(line)
        if not candidates:
            return "-"
        first = candidates[0]
        if ":" in first:
            return first.split(":", 2)[-1].strip() or "-"
        return first.strip() or "-"

    def _read_board_name(self):
        vendor = self._safe_read_file("/sys/devices/virtual/dmi/id/board_vendor")
        name = self._safe_read_file("/sys/devices/virtual/dmi/id/board_name")
        version = self._safe_read_file("/sys/devices/virtual/dmi/id/board_version")
        parts = [p for p in [vendor, name, version] if p and p.lower() != "none"]
        if parts:
            return " ".join(parts)
        return "-"

    def _read_mem_totals(self):
        if psutil is not None:
            vm = psutil.virtual_memory()
            sm = psutil.swap_memory()
            return vm.total, sm.total
        mem_total = None
        swap_total = None
        try:
            with open("/proc/meminfo", "r", encoding="utf-8") as fh:
                for line in fh:
                    if line.startswith("MemTotal:"):
                        parts = line.split()
                        if len(parts) >= 2:
                            mem_total = int(parts[1]) * 1024
                    elif line.startswith("SwapTotal:"):
                        parts = line.split()
                        if len(parts) >= 2:
                            swap_total = int(parts[1]) * 1024
        except Exception:
            pass
        return mem_total, swap_total

    def _read_root_usage(self):
        if psutil is not None:
            du = psutil.disk_usage("/")
            return du.total, du.used, du.free
        rc, out = run_cmd(["df", "-B1", "/"], timeout=5)
        if rc != 0 or not out:
            return None, None, None
        lines = out.splitlines()
        if len(lines) < 2:
            return None, None, None
        parts = lines[1].split()
        if len(parts) < 6:
            return None, None, None
        try:
            total = int(parts[1])
            used = int(parts[2])
            free = int(parts[3])
            return total, used, free
        except Exception:
            return None, None, None

    def _read_swap_usage(self):
        if psutil is not None:
            sm = psutil.swap_memory()
            return sm.total, sm.used, sm.free
        total = None
        free = None
        try:
            with open("/proc/meminfo", "r", encoding="utf-8") as fh:
                for line in fh:
                    if line.startswith("SwapTotal:"):
                        parts = line.split()
                        if len(parts) >= 2:
                            total = int(parts[1]) * 1024
                    elif line.startswith("SwapFree:"):
                        parts = line.split()
                        if len(parts) >= 2:
                            free = int(parts[1]) * 1024
        except Exception:
            pass
        if total is None or free is None:
            return total, None, free
        return total, max(0, total - free), free

    def _load_sysinfo_static(self):
        osr = self._read_os_release()
        distro = osr.get("NAME") or osr.get("PRETTY_NAME") or "-"
        version = osr.get("VERSION_ID") or osr.get("VERSION") or "-"
        kernel = platform.release() or "-"
        arch = platform.machine() or "-"
        hostname = socket.gethostname() or "-"
        cpu_model = self._read_cpu_model()
        cores, threads = self._read_cpu_cores_threads()
        gpu_model = self._read_gpu_model()
        board_name = self._read_board_name()
        mem_total, swap_total = self._read_mem_totals()

        self.sysinfo_static = {
            "distribution": distro,
            "version": version,
            "kernel": kernel,
            "arch": arch,
            "hostname": hostname,
            "cpu_model": cpu_model or "-",
            "cores": cores,
            "threads": threads,
            "gpu_model": gpu_model or "-",
            "board_name": board_name or "-",
            "ram_total": human_size(mem_total),
            "swap_total": human_size(swap_total),
        }

        self._set_sysinfo_value("os_distribution", self.sysinfo_static["distribution"])
        self._set_sysinfo_value("os_version", self.sysinfo_static["version"])
        self._set_sysinfo_value("os_kernel", self.sysinfo_static["kernel"])
        self._set_sysinfo_value("os_arch", self.sysinfo_static["arch"])
        self._set_sysinfo_value("os_hostname", self.sysinfo_static["hostname"])
        self._set_sysinfo_value("hw_cpu", self.sysinfo_static["cpu_model"])
        self._set_sysinfo_value(
            "hw_cores_threads",
            f"{self.sysinfo_static['cores']} / {self.sysinfo_static['threads']}",
        )
        self._set_sysinfo_value("hw_gpu", self.sysinfo_static["gpu_model"])
        self._set_sysinfo_value("hw_board", self.sysinfo_static["board_name"])
        self._set_sysinfo_value("hw_ram_total", self.sysinfo_static["ram_total"])
        self._set_sysinfo_value("hw_swap_total", self.sysinfo_static["swap_total"])

    def _update_sysinfo_dynamic(self):
        if not hasattr(self, "sysinfo_vars"):
            return
        total, used, free = self._read_root_usage()
        swap_total, swap_used, _swap_free = self._read_swap_usage()
        desktop = (
            os.environ.get("XDG_CURRENT_DESKTOP")
            or os.environ.get("DESKTOP_SESSION")
            or "-"
        )
        user = getpass.getuser() or "-"
        pyver = platform.python_version() or sys.version.split()[0]

        self._set_sysinfo_value("os_user", user)
        self._set_sysinfo_value(
            "sys_root",
            f"{human_size(free)} frei von {human_size(total)} ({human_size(used)} belegt)",
        )
        self._set_sysinfo_value(
            "sys_swap",
            f"{human_size(swap_used)} / {human_size(swap_total)} genutzt",
        )
        self._set_sysinfo_value("sw_python", pyver)
        self._set_sysinfo_value("sw_desktop", desktop)

    def _read_cpu_temp_c(self):
        temps = {}
        try:
            temps = psutil.sensors_temperatures(fahrenheit=False) or {}
        except Exception:
            temps = {}

        preferred = [
            "coretemp",
            "k10temp",
            "cpu_thermal",
            "x86_pkg_temp",
            "zenpower",
        ]
        values = []
        for key in preferred:
            for entry in temps.get(key, []):
                cur = getattr(entry, "current", None)
                if cur is not None:
                    values.append(float(cur))
            if values:
                return max(values)

        for entries in temps.values():
            for entry in entries:
                cur = getattr(entry, "current", None)
                if cur is not None:
                    values.append(float(cur))
        if values:
            return max(values)

        zone_root = "/sys/class/thermal"
        try:
            entries = os.listdir(zone_root)
        except Exception:
            entries = []
        best = None
        for name in entries:
            if not name.startswith("thermal_zone"):
                continue
            base = os.path.join(zone_root, name)
            temp_path = os.path.join(base, "temp")
            type_path = os.path.join(base, "type")
            try:
                raw = open(temp_path, "r", encoding="utf-8").read().strip()
                milli = float(raw)
                c = milli / 1000.0 if milli > 1000 else milli
            except Exception:
                continue
            try:
                ztype = open(type_path, "r", encoding="utf-8").read().strip().lower()
            except Exception:
                ztype = ""
            if "cpu" in ztype or "x86" in ztype or "pkg" in ztype:
                return c
            if best is None or c > best:
                best = c
        return best

    def _system_monitor_missing_psutil(self):
        self.system_cpu_var.set("CPU: –")
        self.system_ram_var.set("RAM: –")
        self.system_swap_var.set("Swap: –")
        self.system_disk_var.set("Disk /: –")
        self.system_temp_var.set("CPU-Temp: –")
        self.system_uptime_var.set("Uptime: –")
        for bar in [
            self.system_cpu_bar,
            self.system_ram_bar,
            self.system_swap_bar,
            self.system_disk_bar,
            self.system_temp_bar,
        ]:
            self._set_system_row_state(bar, 0)
            bar["status"].config(text="–", fg="#666666")
        if not self.system_monitor_psutil_hint_logged:
            self.log("Systemmonitor benötigt das Python-Modul psutil.")
            self.log("Installationsbefehl:")
            self.log("pip install psutil")
            self.system_monitor_psutil_hint_logged = True
        self._update_sysinfo_dynamic()

    def system_monitor_refresh(self):
        if not self.system_monitor_active:
            return

        if psutil is None:
            self._system_monitor_missing_psutil()
        else:
            cpu = max(0.0, min(100.0, float(psutil.cpu_percent(interval=None))))
            vm = psutil.virtual_memory()
            sm = psutil.swap_memory()
            du = psutil.disk_usage("/")
            temp_c = self._read_cpu_temp_c()
            up_sec = time.time() - float(psutil.boot_time())

            self.system_cpu_var.set(f"CPU: {cpu:.0f} %")
            self.system_ram_var.set(
                f"RAM: {human_size(vm.used)} / {human_size(vm.total)} ({vm.percent:.0f} %)"
            )
            self.system_swap_var.set(
                f"Swap: {human_size(sm.used)} / {human_size(sm.total)} ({sm.percent:.0f} %)"
            )
            self.system_disk_var.set(
                f"Disk /: {human_size(du.used)} / {human_size(du.total)} ({du.percent:.0f} %)"
            )
            if temp_c is None:
                self.system_temp_var.set("CPU-Temp: nicht verfügbar")
                self.system_temp_bar["bar"]["value"] = 0
                self.system_temp_bar["bar"][
                    "style"
                ] = "SystemGreen.Horizontal.TProgressbar"
                self.system_temp_bar["status"].config(text="–", fg="#666666")
            else:
                self.system_temp_var.set(f"CPU-Temp: {temp_c:.1f} °C")
                self._set_system_temp_state(self.system_temp_bar, temp_c)
            self.system_uptime_var.set(f"Uptime: {self._format_uptime(up_sec)}")

            self._set_system_row_state(self.system_cpu_bar, cpu)
            self._set_system_row_state(self.system_ram_bar, vm.percent)
            self._set_system_row_state(self.system_swap_bar, sm.percent)
            self._set_system_row_state(self.system_disk_bar, du.percent)
            self._update_sysinfo_dynamic()

        self.after(self.system_monitor_interval_ms, self.system_monitor_refresh)

    def _select_system_tab(self, tab):
        nb = self.__dict__.get("nb")
        system_main = self.__dict__.get("tab_system_main")
        system_nb = self.__dict__.get("system_nb")
        if nb is None:
            return
        if system_main is not None and system_nb is not None:
            nb.select(system_main)
            system_nb.select(tab)
            return
        nb.select(tab)

    def _select_tools_tab(self, tab):
        nb = self.__dict__.get("nb")
        tools_tab = self.__dict__.get("tab_tools")
        tools_nb = self.__dict__.get("tools_nb")
        if nb is None:
            return
        if tools_tab is not None and tools_nb is not None:
            nb.select(tools_tab)
            tools_nb.select(tab)
            return
        nb.select(tab)

    def _select_security_tab(self, tab):
        nb = self.__dict__.get("nb")
        security_main = self.__dict__.get("tab_security_main")
        security_nb = self.__dict__.get("security_nb")
        if nb is None:
            return
        if security_main is not None and security_nb is not None:
            nb.select(security_main)
            security_nb.select(tab)
            return
        nb.select(tab)

    def _activate_startup_log_view(self):
        nb = self.__dict__.get("nb")
        tab_log = self.__dict__.get("tab_log")
        if nb is not None and tab_log is not None:
            try:
                nb.select(tab_log)
            except Exception:
                pass

        has_overlay_image = bool(
            self.__dict__.get("log_overlay_source_image") is not None
            or self.__dict__.get("log_overlay_source_photo") is not None
            or str(self.__dict__.get("log_overlay_image_path") or "").strip()
        )
        if not has_overlay_image:
            return

        try:
            self._render_log_overlay_image()
            self._set_log_overlay_visible(True, persist=False)
        except Exception:
            pass

    def update_dns_output(self, text: str, append: bool = False):
        if not hasattr(self, "dns_text") or self.dns_text is None:
            return
        self.dns_text.configure(state="normal")
        if not append:
            self.dns_text.delete("1.0", "end")
        if text:
            if append and self.dns_text.index("end-1c") != "1.0":
                self.dns_text.insert("end", "\n\n")
        self.dns_text.insert("end", text)
        self.dns_text.see("end")
        self.dns_text.configure(state="normal")

    def clear_dns_output(self):
        self.update_dns_output("")
        if "chad_analysis_status_var" in self.__dict__:
            self.chad_analysis_status_var.set("DNS-Ausgabe geleert.")

    def _dns_mode_from_label(self, value):
        return dns_controller.dns_mode_from_label(value)

    def _dns_mode_label(self, mode):
        return dns_controller.dns_mode_label(mode)

    def _current_dns_mode(self):
        return dns_controller.current_dns_mode(self)

    def _set_dns_mode(self, mode, persist=True, update_var=True):
        return dns_controller.set_dns_mode(
            self,
            mode,
            persist=persist,
            update_var=update_var,
        )

    def on_dns_mode_selected(self):
        dns_controller.on_dns_mode_selected(self)

    def is_vpn_active(self):
        service = self._dns_control_service()
        if service is None:
            code, out = run_cmd(["resolvectl", "status"], timeout=6)
            if code != 0:
                return False
            return self.detect_vpn_active(out)
        code, out = service.read_resolvectl_status(timeout=6)
        if code != 0:
            return False
        return service.detect_vpn_active(out)

    def _dns_mode_display_name(self, mode):
        return dns_controller.dns_mode_display_name(mode)

    def _manual_dns_choice(self):
        choice_var = self.__dict__.get("dns_choice")
        choice = (
            (choice_var.get() if choice_var is not None else "") or "System (DHCP/auto)"
        ).strip()
        if choice in DNS_PROFILES:
            return choice
        return "System (DHCP/auto)"

    def _dot_profile_for_dns_choice(self, dns_choice):
        choice = str(dns_choice or "").strip()
        mapping = {
            "Cloudflare (1.1.1.1)": "Cloudflare",
            "Google (8.8.8.8)": "Google",
            "Quad9 (malware-blocking)": "Quad9",
        }
        return mapping.get(choice, "Cloudflare")

    def is_physical_dns_candidate(self, name: str, conn_type: str, device: str) -> bool:
        service = self._dns_control_service()
        if service is None:
            return False
        return service.is_physical_dns_candidate(name, conn_type, device)

    def nm_active_connections_detailed(self):
        service = self._dns_control_service()
        if service is None:
            return []
        return service.nm_active_connections_detailed()

    def _active_nm_dns_targets(self):
        service = self._dns_control_service()
        if service is None:
            return []
        iface = self.get_default_iface()
        con = self.nm_active_conn_for_dev(iface)
        return service.filter_physical_nm_targets(
            self.nm_active_connections_detailed(),
            default_iface=iface or "",
            default_connection=con or "",
        )

    def _apply_nm_dns_policy_to_targets(self, *, ignore_auto_dns):
        service = self._dns_control_service()
        if service is None:
            return False, ["DNS-Control-Service nicht verfügbar."]
        return service.apply_nm_dns_policy_to_targets(
            ignore_auto_dns=ignore_auto_dns,
            targets=self._active_nm_dns_targets(),
        )

    def _write_manual_dns_dropin(self, dns_choice):
        profile_name = self._dot_profile_for_dns_choice(dns_choice)
        self._set_dot_profile_choice_safe(profile_name)
        _name, _profile, content = self._build_dot_dropin_content(
            profile_name=profile_name
        )
        service = self._dns_control_service()
        if service is None:
            tmp_path = ""
            try:
                with tempfile.NamedTemporaryFile(
                    mode="w", encoding="utf-8", delete=False
                ) as temp_file:
                    temp_file.write(content)
                    tmp_path = temp_file.name
                rc_dir, out_dir = run_cmd(
                    ["sudo", "-n", "mkdir", "-p", DOT_DROPIN_DIR],
                    timeout=10,
                )
                if rc_dir != 0:
                    return (
                        False,
                        (out_dir or "").strip() or "Drop-in-Verzeichnis fehlt.",
                    )
                rc_install, out_install = run_cmd(
                    ["sudo", "-n", "install", "-m", "644", tmp_path, DOT_DROPIN_PATH],
                    timeout=12,
                )
                if rc_install != 0:
                    return (
                        False,
                        (out_install or "").strip()
                        or "DNS-Drop-in konnte nicht geschrieben werden.",
                    )
                return True, content
            finally:
                if tmp_path:
                    try:
                        os.remove(tmp_path)
                    except Exception:
                        pass

        ok, detail = service.write_dropin(content)
        return ok, content if ok else detail

    def _remove_manual_dns_dropin(self):
        service = self._dns_control_service()
        if service is None:
            rc_rm, out_rm = run_cmd(
                ["sudo", "-n", "rm", "-f", DOT_DROPIN_PATH],
                timeout=10,
            )
            return rc_rm == 0, (out_rm or "").strip()
        return service.remove_dropin()

    def _restart_resolved_for_dns_policy(self):
        service = self._dns_control_service()
        if service is None:
            rc, out = run_cmd(
                ["sudo", "-n", "systemctl", "restart", "systemd-resolved"],
                timeout=15,
            )
            return rc == 0, (out or "").strip()
        return service.restart_resolved()

    def refresh_dns_policy_status(self):
        return dns_controller.refresh_dns_policy_status(self)

    def evaluate_dns_policy(self, resolvectl_output: str) -> dict:
        return dns_controller.evaluate_dns_policy(self, resolvectl_output)

    def _show_dns_policy_result(self, text, *, startup=False):
        dns_controller.show_dns_policy_result(self, text, startup=startup)

    def warn_no_vpn(self):
        result = (
            "=== DNS Control ===\n"
            "Modus: VPN-PRIORITY\n"
            "Status: kein aktives VPN erkannt\n"
            "Aktion: keine Systemänderung durchgeführt\n\n"
            "Empfehlung: VPN verbinden oder auf AUTO/MANUAL wechseln."
        )
        self.refresh_dns_policy_status()
        return False, result

    def warn_manual_dns_with_vpn(self):
        result = (
            "=== DNS Control ===\n"
            "Modus: MANUAL\n"
            "Status: blockiert\n"
            "Grund: VPN ist aktiv.\n\n"
            "NetMon setzt kein globales DNS-Drop-in, solange VPN-DNS aktiv ist,\n"
            "damit nicht gleichzeitig VPN-DNS und manueller Resolver aktiv sind."
        )
        self.refresh_dns_policy_status()
        return False, result

    def use_vpn_dns(self):
        return dns_controller.use_vpn_dns(self)

    def _normalize_dns_profile_choice(self, choice=None):
        return dns_controller.normalize_dns_profile_choice(self, choice)

    def _apply_dns_profile_choice_via_control(self, choice):
        return dns_controller.apply_dns_profile_choice_via_control(self, choice)

    def _format_dns_profile_control_result(self, result, *, heading, action):
        return dns_controller.format_dns_profile_control_result(
            result,
            heading=heading,
            action=action,
        )

    def use_user_dns(self):
        return dns_controller.use_user_dns(self)

    def enforce_dns_policy(self):
        return dns_controller.enforce_dns_policy(self)

    def _run_dns_policy_action(self, *, startup=False, repair_only=False):
        try:
            if repair_only:
                success, text = self.fix_dns_conflicts(startup=startup)
            else:
                success, text = self.enforce_dns_policy()
            self.after(
                0,
                lambda t=text, s=startup: self._show_dns_policy_result(
                    t,
                    startup=s,
                ),
            )
            if not startup:
                if repair_only:
                    self.log_safe(
                        "DNS-Status bewertet"
                        if success
                        else "DNS-Status mit Hinweis bewertet"
                    )
                else:
                    self.log_safe(
                        "DNS-Policy angewendet"
                        if success
                        else "DNS-Policy mit Hinweis beendet"
                    )
        except Exception as exc:
            error_text = f"=== DNS Control ===\nFehler: {exc}"
            self.after(
                0,
                lambda t=error_text, s=startup: self._show_dns_policy_result(
                    t,
                    startup=s,
                ),
            )
            if not startup:
                self.log_safe(
                    "DNS-Statusprüfung fehlgeschlagen"
                    if repair_only
                    else "DNS-Policy fehlgeschlagen"
                )
        finally:
            self.dns_mode_running = False

    def enforce_dns_policy_async(self):
        if self.dns_mode_running:
            self.log_safe("DNS-Policy läuft bereits.")
            return
        mode = self._set_dns_mode(self.dns_mode_var.get())
        if not self.ensure_sudo_session(
            f"DNS Policy anwenden ({self._dns_mode_display_name(mode)})"
        ):
            return
        self.dns_mode_running = True
        self.nb.select(self.tab_dns)
        self.log_safe("DNS-Policy wird angewendet")
        threading.Thread(
            target=self._run_dns_policy_action,
            kwargs={"startup": False, "repair_only": False},
            daemon=True,
        ).start()

    def fix_dns_conflicts(self, startup=False):
        code, out = run_cmd(["resolvectl", "status"], timeout=8)
        if code != 0:
            return (
                False,
                "=== DNS Control ===\n"
                "Status: Prüfung fehlgeschlagen\n"
                "Grund: resolvectl status konnte nicht gelesen werden.\n\n"
                f"Details:\n{out}",
            )

        evaluation = self.evaluate_dns_policy(out)
        lines = [
            "=== DNS Control ===",
            "Aktion: Statusbewertung",
            f"DNS-Modus: {self._dns_mode_display_name(self._current_dns_mode())}",
            f"VPN aktiv: {'JA' if evaluation.get('vpn_active') else 'NEIN'}",
            f"Aktiver Resolver: {evaluation.get('resolver_label', 'unbekannt')}",
            f"NetMon DoT Drop-in: {'vorhanden' if evaluation.get('dot_dropin_exists') else 'nicht vorhanden'}",
            f"Globale DNS-Erzwingung: {'JA' if evaluation.get('dot_dropin_conflict') else 'NEIN'}",
            f"Risiko: {evaluation.get('risk', 'unbekannt')}",
            "",
            "Empfehlungen:",
        ]
        for recommendation in evaluation.get("recommendations") or [
            "Keine Empfehlung verfügbar."
        ]:
            lines.append(f"- {recommendation}")
        if evaluation.get("warnings"):
            lines.append("")
            lines.append("Auffälligkeiten:")
            for warning in evaluation["warnings"]:
                lines.append(f"- {warning}")
        if startup:
            lines.append("")
            lines.append("Hinweis: Beim Start wurden keine Änderungen durchgeführt.")
        return True, "\n".join(lines)

    def fix_dns_conflicts_async(self):
        if self.dns_mode_running:
            self.log_safe("DNS-Konfliktprüfung läuft bereits.")
            return
        self.dns_mode_running = True
        self.nb.select(self.tab_dns)
        self.log_safe("DNS-Status wird bewertet")
        threading.Thread(
            target=self._run_dns_policy_action,
            kwargs={"startup": False, "repair_only": True},
            daemon=True,
        ).start()

    def startup_dns_policy_check(self):
        try:
            self.refresh_dns_policy_status()
            if self.dns_mode_running:
                return
            self.dns_mode_running = True
            threading.Thread(
                target=self._run_dns_policy_action,
                kwargs={"startup": True, "repair_only": True},
                daemon=True,
            ).start()
        except Exception as exc:
            try:
                self.log(f"[WARN] DNS-Policy-Startprüfung fehlgeschlagen: {exc}")
            except Exception:
                pass

    def apply_vpn_dns_policy(self):
        code, out = run_cmd(["resolvectl", "status"], timeout=8)
        if code != 0:
            return (
                False,
                "=== DNS Control ===\n"
                "Aktion: VPN-DNS erzwingen\n"
                "Status: abgebrochen\n"
                "Grund: resolvectl status konnte nicht gelesen werden.\n\n"
                f"Details:\n{out}",
            )
        evaluation = self.evaluate_dns_policy(out)
        if not evaluation.get("vpn_active"):
            return self.warn_no_vpn()
        if not evaluation.get("vpn_dns"):
            return (
                False,
                "=== DNS Control ===\n"
                "Aktion: VPN-DNS erzwingen\n"
                "Status: abgebrochen\n"
                "Grund: Kein VPN-DNS erkannt.",
            )
        return self.use_vpn_dns()

    def apply_vpn_dns_policy_async(self):
        if self.dns_mode_running:
            self.log_safe("DNS-Control läuft bereits.")
            return
        code, out = run_cmd(["resolvectl", "status"], timeout=8)
        if code != 0:
            self.nb.select(self.tab_dns)
            self.update_dns_output(
                "=== DNS Control ===\n"
                "Aktion: VPN-DNS erzwingen\n"
                "Status: Prüfung fehlgeschlagen\n\n"
                f"Details:\n{out}",
                append=True,
            )
            self.log_safe("VPN-DNS-Erzwingung fehlgeschlagen")
            return

        evaluation = self.evaluate_dns_policy(out)
        if not evaluation.get("vpn_active"):
            success, text = self.warn_no_vpn()
            self.nb.select(self.tab_dns)
            self.update_dns_output(text, append=True)
            self.refresh_dns_policy_status()
            self.log_safe("VPN-DNS-Erzwingung nicht möglich")
            return

        ok = messagebox.askyesno(
            self.tr("vpn_dns_enforce_confirm_title", "VPN-DNS erzwingen"),
            self.tr(
                "vpn_dns_enforce_confirm_message",
                "NetMon entfernt das NetMon-DoT-Drop-in, leert manuelle DNS-Werte auf "
                "echten LAN/WLAN-Verbindungen und setzt ignore-auto-dns auf diesen "
                "physischen Connections.\n\n"
                "VPN-Verbindungen, lo, docker0 und veth-Interfaces bleiben unberührt.\n\n"
                "Fortfahren?",
            ),
            parent=self,
        )
        if not ok:
            self.chad_analysis_status_var.set(
                self.tr(
                    "vpn_dns_enforce_status_cancelled",
                    "VPN-DNS-Erzwingung abgebrochen.",
                )
            )
            return
        if not self.ensure_sudo_session("VPN-DNS erzwingen"):
            result = (
                "=== DNS Control ===\n"
                "Aktion: VPN-DNS erzwingen\n"
                "Status: abgebrochen\n"
                "Grund: sudo-Freigabe fehlt."
            )
            self.nb.select(self.tab_dns)
            self.update_dns_output(result, append=True)
            return

        self.dns_mode_running = True
        self.nb.select(self.tab_dns)
        self.log_safe("VPN-DNS-Erzwingung gestartet")

        def worker():
            try:
                success, text = self.apply_vpn_dns_policy()
                self.after(
                    0,
                    lambda t=text: self._show_dns_policy_result(t, startup=False),
                )
                self.after(0, self.run_dns_leak_check)
                self.log_safe(
                    "VPN-DNS-Erzwingung abgeschlossen"
                    if success
                    else "VPN-DNS-Erzwingung mit Hinweis beendet"
                )
            finally:
                self.dns_mode_running = False

        threading.Thread(target=worker, daemon=True).start()

    def load_ai_config(self):
        try:
            if self.ai_online_service is not None:
                env_key = (self.ai_online_service.api_key or "").strip()
                if env_key:
                    self.openai_api_key_var.set(env_key)
                    self.chad_online_status_var.set("Online-Analyse: bereit")

            if not os.path.isfile(self.ai_config_path):
                return

            with open(self.ai_config_path, "r", encoding="utf-8") as f:
                data = json.load(f)

            use_keyring = bool(data.get("use_keyring"))
            if not use_keyring:
                return

            if not self.keyring_available:
                self.chad_analysis_status_var.set(
                    "Keyring nicht verfügbar. API-Key nicht geladen."
                )
                return

            kr = self._get_keyring_module()
            if kr is None:
                self.chad_analysis_status_var.set(
                    "Keyring-Modul nicht verfügbar. API-Key nicht geladen."
                )
                return

            key = kr.get_password(KEYRING_SERVICE, KEYRING_USERNAME)
            if key:
                self.openai_api_key_var.set(key)
                if self.ai_online_service is not None:
                    self.ai_online_service.api_key = key
                self.openai_api_key_saved = True
                self.chad_analysis_status_var.set("API-Key aus System-Keyring geladen.")
                self.chad_online_status_var.set("Online-Analyse: bereit")
            else:
                self.chad_analysis_status_var.set(
                    "Kein API-Key im System-Keyring gefunden."
                )

        except Exception as e:
            try:
                self.chad_analysis_status_var.set(
                    f"Fehler beim Laden des API-Keys: {e}"
                )
            except Exception:
                pass

    def save_ai_config(self):
        try:
            key = (self.openai_api_key_var.get() or "").strip()

            if not key:
                self.chad_analysis_status_var.set("Kein API-Key eingegeben.")
                return

            if not self.keyring_available:
                self.chad_analysis_status_var.set(
                    "Keyring nicht verfügbar. API-Key kann nicht sicher gespeichert werden."
                )
                return

            kr = self._get_keyring_module()
            if kr is None:
                self.chad_analysis_status_var.set(
                    "Keyring-Modul nicht verfügbar. API-Key wurde nicht gespeichert."
                )
                return

            kr.set_password(KEYRING_SERVICE, KEYRING_USERNAME, key)

            ok, err = safe_write_json(self.ai_config_path, {"use_keyring": True})

            if ok:
                self.openai_api_key_saved = True
                if self.ai_online_service is not None:
                    self.ai_online_service.api_key = key
                self.chad_analysis_status_var.set(
                    "API-Key sicher im System-Keyring gespeichert."
                )
                self.chad_online_status_var.set("Online-Analyse: bereit")
            else:
                self.chad_analysis_status_var.set(
                    f"Key gespeichert, aber Config-Fehler: {err}"
                )

        except Exception as e:
            self.chad_analysis_status_var.set(
                f"Fehler beim Speichern des API-Keys: {e}"
            )

    def apply_api_key(self):
        key = (self.openai_api_key_var.get() or "").strip()

        if not key:
            if self.ai_online_service is not None:
                self.ai_online_service.api_key = ""
            self.openai_api_key_saved = False
            self.chad_analysis_status_var.set("Kein API-Key eingegeben.")
            self.chad_online_status_var.set("Online-Analyse: inaktiv")
            return

        if self.ai_online_service is None:
            self.chad_analysis_status_var.set("AI-Online-Service nicht verfügbar.")
            self.chad_online_status_var.set("Online-Analyse: inaktiv")
            return

        self.ai_online_service.api_key = key
        self.chad_analysis_status_var.set("API-Key übernommen.")
        self.chad_online_status_var.set("Online-Analyse: bereit")

    def clear_api_key(self):
        self.openai_api_key_var.set("")
        if self.ai_online_service is not None:
            self.ai_online_service.api_key = ""
        self.openai_api_key_saved = False
        self.chad_analysis_status_var.set("API-Key im Feld gelöscht.")
        self.chad_online_status_var.set("Online-Analyse: inaktiv")

    # Hinweis fuer Paketierung:
    # Abhaengigkeiten:
    # - python3-keyring
    # - gnome-keyring
    def check_keyring_availability(self):
        self.keyring_available = False
        self.keyring_error = ""
        self.keyring_backend_name = ""

        try:
            kr_mod = self._get_keyring_module()
            if kr_mod is None:
                self.keyring_error = "Python-Modul 'keyring' nicht installiert."
                self.log(
                    "[WARN] Keyring nicht verfügbar – API-Key Speicherung deaktiviert."
                )
                return False

            try:
                kr = kr_mod.get_keyring()
                self.keyring_backend_name = kr.__class__.__name__
                try:
                    kr_mod.get_password("netmon_test", "probe")
                    self.keyring_available = True
                except Exception as e:
                    self.keyring_error = f"Keyring Backend Fehler: {e}"
            except Exception as e:
                self.keyring_error = f"Keyring nicht nutzbar: {e}"

        except ImportError:
            self.keyring_error = "Python-Modul 'keyring' nicht installiert."

        if not self.keyring_available:
            self.log(
                "[WARN] Keyring nicht verfügbar – API-Key Speicherung deaktiviert."
            )

        return self.keyring_available

    def _get_keyring_module(self):
        global keyring
        if keyring is not None:
            return keyring
        try:
            keyring = importlib.import_module("keyring")
            return keyring
        except Exception as exc:
            try:
                self.log_debug_safe(f"Keyring Import fehlgeschlagen: {exc}")
            except Exception:
                pass
            return None

    def delete_saved_api_key(self):
        try:
            kr = self._get_keyring_module()
            if kr is not None:
                try:
                    kr.delete_password(KEYRING_SERVICE, KEYRING_USERNAME)
                except Exception:
                    pass

            self.openai_api_key_var.set("")
            if self.ai_online_service is not None:
                self.ai_online_service.api_key = ""
            self.openai_api_key_saved = False

            ok, err = safe_write_json(self.ai_config_path, {"use_keyring": False})

            if ok:
                self.chad_analysis_status_var.set("Gespeicherter API-Key gelöscht.")
            else:
                self.chad_analysis_status_var.set(
                    f"Key gelöscht, aber Config-Fehler: {err}"
                )
            self.chad_online_status_var.set("Online-Analyse: inaktiv")

        except Exception as e:
            self.chad_analysis_status_var.set(f"Fehler beim Löschen des API-Keys: {e}")

    def paste_api_key_from_clipboard(self, _event=None):
        try:
            clip = self.clipboard_get()
        except Exception:
            self.chad_analysis_status_var.set(
                "Zwischenablage leer oder nicht verfügbar."
            )
            return "break"

        key = (clip or "").strip()
        if not key:
            self.chad_analysis_status_var.set("Zwischenablage enthält keinen API-Key.")
            return "break"

        self.openai_api_key_var.set(key)
        self.chad_analysis_status_var.set("API-Key aus Zwischenablage eingefügt.")
        try:
            self.openai_api_key_entry.focus_set()
            self.openai_api_key_entry.icursor("end")
        except Exception:
            pass
        return "break"

    def _log_chad_error(self, source: str, error_text: str):
        chad_controller.log_chad_error(self, source, error_text)

    def append_to_chad_output(self, title: str, text: str):
        chad_controller.append_to_chad_output(self, title, text)

    def scroll_chad_output_to_top(self):
        chad_controller.scroll_chad_output_to_top(self)

    def scroll_chad_output_to_bottom(self):
        chad_controller.scroll_chad_output_to_bottom(self)

    def confirm_chad_send(self, title: str, intro_text: str) -> bool:
        if tk is None or ttk is None:
            return messagebox.askyesno(title, intro_text, parent=self)

        result = {"ok": False}
        win = tk.Toplevel(self)
        win.title(title)
        win.geometry("560x220")
        win.resizable(False, False)
        win.transient(self)

        outer = ttk.Frame(win, padding=12)
        outer.pack(fill="both", expand=True)

        ttk.Label(
            outer,
            text=intro_text,
            justify="left",
            wraplength=520,
        ).pack(anchor="w")

        ttk.Label(
            outer,
            text=self.tr("chad_security_notice_title", "Sicherheitshinweis"),
            justify="left",
            font=("TkDefaultFont", 10, "bold"),
        ).pack(anchor="w", pady=(12, 4))

        ttk.Label(
            outer,
            text=self.tr(
                "chad_security_notice_body",
                "Bitte keine sensiblen Daten absenden.\n"
                "Dazu gehören z. B. Passwörter, API-Keys, Tokens, private Schlüssel, "
                "persönliche Daten oder interne Inhalte.\n\n"
                "Bitte vor dem Senden sorgfältig prüfen und nur notwendige Ausschnitte übertragen.",
            ),
            justify="left",
            wraplength=520,
        ).pack(anchor="w")

        btn_row = ttk.Frame(outer)
        btn_row.pack(fill="x", pady=(14, 0))

        def _close(ok: bool):
            result["ok"] = bool(ok)
            try:
                win.grab_release()
            except Exception:
                pass
            win.destroy()

        ttk.Button(
            btn_row,
            text=self.tr("btn_cancel", "Abbrechen"),
            command=lambda: _close(False),
        ).pack(side="right")
        ttk.Button(
            btn_row,
            text=self.tr("btn_continue", "Fortfahren"),
            command=lambda: _close(True),
        ).pack(side="right", padx=(0, 8))

        win.protocol("WM_DELETE_WINDOW", lambda: _close(False))
        try:
            win.grab_set()
            win.focus_force()
        except Exception:
            pass
        self.wait_window(win)
        return result["ok"]

    def send_chad_chat_message(self):
        return chad_controller.send_chad_chat_message(self)

    def clear_chad_chat(self):
        chad_controller.clear_chad_chat(self)

    def add_current_analysis_to_chad_chat(self):
        chad_controller.add_current_analysis_to_chad_chat(self)

    def extract_pdf_text_for_chad(self, path: str) -> str:
        return chad_controller.extract_pdf_text_for_chad(self, path)

    def extract_docx_text_for_chad(self, path: str) -> str:
        return chad_controller.extract_docx_text_for_chad(self, path)

    def import_file_to_chad_chat(self):
        chad_controller.import_file_to_chad_chat(self)

    def _default_settings_payload(self) -> dict:
        return {
            "language": "de",
            "startup_guard_autorun": True,
            "path_guard_autorun": False,
            "dns_guard_enabled": True,
            "auto_switch_to_chad": True,
            "chad_preview_popup": False,
        }

    def _sync_settings_vars_from_settings(self):
        mapping = {
            "var_language": self.settings.get("language", "de"),
            "var_startup_guard_autorun": bool(
                self.settings.get("startup_guard_autorun", True)
            ),
            "var_path_guard_autorun": bool(
                self.settings.get("path_guard_autorun", False)
            ),
            "var_dns_guard_enabled": bool(self.settings.get("dns_guard_enabled", True)),
            "var_auto_switch_to_chad": bool(
                self.settings.get("auto_switch_to_chad", True)
            ),
            "var_chad_preview_popup": bool(
                self.settings.get("chad_preview_popup", False)
            ),
        }
        for attr_name, value in mapping.items():
            var = self.__dict__.get(attr_name)
            if var is None:
                continue
            try:
                var.set(value)
            except Exception:
                pass

    def _apply_settings_runtime_state(self):
        dns_guard_enabled = bool(self.settings.get("dns_guard_enabled", True))
        self.dns_guard_enabled = dns_guard_enabled

        guard_var = self.__dict__.get("dns_guard_enabled_var")
        if guard_var is not None:
            try:
                guard_var.set(dns_guard_enabled)
            except Exception:
                pass

        self._sync_settings_vars_from_settings()

    def load_settings(self):
        try:
            if not os.path.isfile(self.netmon_config_path):
                return

            with open(self.netmon_config_path, "r", encoding="utf-8") as handle:
                data = json.load(handle)

            if isinstance(data, dict):
                self.settings.update(data)

        except Exception as e:
            try:
                self.log(f"[WARN] Einstellungen konnten nicht geladen werden: {e}")
            except Exception:
                pass

    def save_settings(self, log_success: bool = True):
        try:
            os.makedirs(self.netmon_config_dir, exist_ok=True)

            ok, err = safe_write_json(
                self.netmon_config_path,
                self.settings,
                indent=2,
                ensure_ascii=False,
            )

            if ok:
                if log_success:
                    self.log("Einstellungen gespeichert.")
            else:
                self.log(
                    f"[WARN] Einstellungen konnten nicht gespeichert werden: {err}"
                )

        except Exception as e:
            self.log(f"[WARN] Einstellungen konnten nicht gespeichert werden: {e}")

    def apply_settings_from_ui(self):
        previous_language = str(self.settings.get("language", "de") or "de")
        new_language = str(self.var_language.get() or "de").strip().lower() or "de"
        language_changed = new_language != previous_language

        self.settings["language"] = new_language
        self.settings["startup_guard_autorun"] = bool(
            self.var_startup_guard_autorun.get()
        )
        self.settings["path_guard_autorun"] = bool(self.var_path_guard_autorun.get())
        self.settings["dns_guard_enabled"] = bool(self.var_dns_guard_enabled.get())
        self.settings["auto_switch_to_chad"] = bool(self.var_auto_switch_to_chad.get())
        self.settings["chad_preview_popup"] = bool(self.var_chad_preview_popup.get())

        if self.lang_service is not None:
            self.lang = self.lang_service.load(new_language)

        self._apply_settings_runtime_state()
        self.save_settings()

        if "privacy_config" in self.__dict__:
            self.save_privacy_config()

        if language_changed:
            self.log(
                self.tr(
                    "log_language_changed_restart",
                    "Sprache geändert – Neustart empfohlen.",
                )
            )
        self.log(self.tr("log_settings_applied", "Einstellungen übernommen."))

    def reset_settings_to_defaults(self):
        self.settings.update(self._default_settings_payload())
        self._apply_settings_runtime_state()
        self.save_settings()

        if "privacy_config" in self.__dict__:
            self.save_privacy_config()

        if self.lang_service is not None:
            self.lang = self.lang_service.load(
                str(self.settings.get("language", "de") or "de")
            )
        self.log(self.tr("log_settings_reset", "Einstellungen zurückgesetzt."))

    def build_settings_tab(self):
        frame = ttk.Frame(self.tab_settings, padding=12)
        frame.pack(fill="both", expand=True)

        general = ttk.LabelFrame(
            frame,
            text=self.tr("settings_group_general", "Allgemein"),
            padding=10,
        )
        general.pack(fill="x", pady=(0, 10))
        general.columnconfigure(1, weight=1)

        ttk.Label(
            general,
            text=self.tr("settings_language_label", "Sprache:"),
        ).grid(row=0, column=0, sticky="w")
        lang_combo = ttk.Combobox(
            general,
            textvariable=self.var_language,
            values=["de", "en"],
            state="readonly",
            width=12,
        )
        lang_combo.grid(row=0, column=1, sticky="w", padx=(8, 0))

        ttk.Label(
            general,
            text=self.tr(
                "settings_language_hint",
                "Hinweis: Sprachpakete werden später vollständig umgesetzt.",
            ),
        ).grid(row=1, column=0, columnspan=2, sticky="w", pady=(6, 0))

        security = ttk.LabelFrame(
            frame,
            text=self.tr("settings_group_security", "Security"),
            padding=10,
        )
        security.pack(fill="x", pady=(0, 10))

        ttk.Checkbutton(
            security,
            text=self.tr(
                "settings_startup_guard_autorun",
                "Startup Guard beim Programmstart prüfen",
            ),
            variable=self.var_startup_guard_autorun,
        ).pack(anchor="w")

        ttk.Checkbutton(
            security,
            text=self.tr(
                "settings_path_guard_autorun",
                "PATH Guard beim Programmstart prüfen",
            ),
            variable=self.var_path_guard_autorun,
        ).pack(anchor="w")

        ttk.Checkbutton(
            security,
            text=self.tr(
                "settings_dns_guard_enabled",
                "DNS Guard aktivieren",
            ),
            variable=self.var_dns_guard_enabled,
        ).pack(anchor="w")

        chad = ttk.LabelFrame(
            frame,
            text=self.tr("settings_group_chad", "Chad Analyse"),
            padding=10,
        )
        chad.pack(fill="x", pady=(0, 10))

        ttk.Checkbutton(
            chad,
            text=self.tr(
                "settings_auto_switch_to_chad",
                "Beim Senden an Chad automatisch zum Chad-Tab wechseln",
            ),
            variable=self.var_auto_switch_to_chad,
        ).pack(anchor="w")

        ttk.Checkbutton(
            chad,
            text=self.tr(
                "settings_chad_preview_popup",
                "Vorschau-Popup bei Chad-Analyse anzeigen",
            ),
            variable=self.var_chad_preview_popup,
        ).pack(anchor="w")

        buttons = ttk.Frame(frame)
        buttons.pack(fill="x", pady=(10, 0))

        ttk.Button(
            buttons,
            text=self.tr("settings_save_button", "Einstellungen speichern"),
            command=self.apply_settings_from_ui,
        ).pack(side="left")

        ttk.Button(
            buttons,
            text=self.tr("settings_reset_button", "Zurücksetzen"),
            command=self.reset_settings_to_defaults,
        ).pack(side="left", padx=(8, 0))

    def build_chad_analysis_tab(self):
        top = ttk.Frame(self.tab_chad_analysis)
        top.pack(fill="x", pady=(0, 8))

        ttk.Label(
            top,
            text=self.tr("chad_title", "Chad Analyse"),
            font=("TkDefaultFont", 12, "bold"),
        ).pack(side="left")

        ttk.Label(
            top,
            textvariable=self.chad_analysis_status_var,
        ).pack(side="left", padx=(12, 0))

        toggle_row = ttk.Frame(self.tab_chad_analysis)
        toggle_row.pack(fill="x", pady=(0, 8))

        self.chad_enabled_check = ttk.Checkbutton(
            toggle_row,
            text=self.tr("chad_check_enabled", "Chad-Anbindung aktiv"),
            variable=self.chad_enabled_var,
            onvalue=True,
            offvalue=False,
            command=self._update_chad_controls_state,
        )
        self.chad_enabled_check.pack(anchor="w")

        api_frame = ttk.LabelFrame(
            self.tab_chad_analysis,
            text=self.tr("chad_group_api", "OpenAI API"),
            padding=8,
        )
        api_frame.pack(fill="x", pady=(0, 8))

        ttk.Label(
            api_frame,
            text=self.tr("chad_label_api_key", "API Key:"),
        ).pack(side="left")

        self.openai_api_key_entry = ttk.Entry(
            api_frame,
            textvariable=self.openai_api_key_var,
            show="*",
            width=40,
        )
        self.openai_api_key_entry.pack(side="left", padx=(6, 6))
        self.openai_api_key_entry.bind("<Control-v>", self.paste_api_key_from_clipboard)
        self.openai_api_key_entry.bind("<Control-V>", self.paste_api_key_from_clipboard)
        self.openai_api_key_entry.bind(
            "<Shift-Insert>", self.paste_api_key_from_clipboard
        )

        self.btn_chad_api_paste = ttk.Button(
            api_frame,
            text=self.tr("btn_paste", "Einfügen"),
            command=self.paste_api_key_from_clipboard,
            width=12,
        )
        self.btn_chad_api_paste.pack(side="left")

        self.btn_chad_api_apply = ttk.Button(
            api_frame,
            text=self.tr("btn_apply", "Übernehmen"),
            command=self.apply_api_key,
            width=12,
        )
        self.btn_chad_api_apply.pack(side="left", padx=(6, 0))

        self.btn_chad_api_save = ttk.Button(
            api_frame,
            text=self.tr("btn_save", "Speichern"),
            command=self.save_ai_config,
            width=12,
        )
        if not self.keyring_available:
            self.btn_chad_api_save.configure(state="disabled")
        self.btn_chad_api_save.pack(side="left", padx=(6, 0))

        self.btn_chad_api_clear = ttk.Button(
            api_frame,
            text=self.tr("chad_btn_clear_field", "Feld leeren"),
            command=self.clear_api_key,
            width=12,
        )
        self.btn_chad_api_clear.pack(side="left", padx=(6, 0))

        self.btn_chad_api_delete = ttk.Button(
            api_frame,
            text=self.tr("chad_btn_delete_key", "Key löschen"),
            command=self.delete_saved_api_key,
            width=12,
        )
        self.btn_chad_api_delete.pack(side="left", padx=(6, 0))

        ttk.Label(
            self.tab_chad_analysis,
            text=self.tr(
                "chad_api_hint",
                "API-Key wird optional im System-Keyring gespeichert, nicht im Projektordner.",
            ),
        ).pack(anchor="w", pady=(0, 8))

        keyring_info_text = (
            f"Keyring Backend: {self.keyring_backend_name}"
            if self.keyring_available and self.keyring_backend_name
            else f"Keyring Status: {self.keyring_error or 'unbekannt'}"
        )

        ttk.Label(
            api_frame,
            text=keyring_info_text,
            justify="left",
        ).pack(anchor="w", pady=(6, 0))

        if not getattr(self, "keyring_available", False):
            warn_text = (
                "API-Key kann nicht sicher gespeichert werden.\n"
                "Benötigt: python3-keyring + gnome-keyring\n"
                "Installation:\n"
                "sudo apt install python3-keyring gnome-keyring"
            )

            ttk.Label(
                api_frame,
                text=warn_text,
                foreground="orange",
                justify="left",
            ).pack(anchor="w", pady=(6, 0))

        input_frame = ttk.LabelFrame(
            self.tab_chad_analysis,
            text=self.tr("chad_group_input", "Chat / Analyse Eingabe"),
            padding=8,
        )
        input_frame.pack(fill="x", pady=(8, 8))

        input_row = ttk.Frame(input_frame)
        input_row.pack(fill="x")

        self.chad_chat_input = tk.Text(input_row, wrap="word", height=3)
        self.chad_chat_input.pack(side="left", fill="x", expand=True)
        self.chad_chat_input.bind(
            "<Control-Return>",
            lambda _event: self.send_chad_chat_message(),
        )

        input_btns = ttk.Frame(input_row)
        input_btns.pack(side="right", padx=(8, 0), fill="y")

        self.btn_chad_send = ttk.Button(
            input_btns,
            text=self.tr("chad_btn_send", "Senden"),
            command=self.send_chad_chat_message,
        )
        self.btn_chad_send.pack(fill="x")

        self.btn_chad_take_analysis = ttk.Button(
            input_btns,
            text=self.tr("chad_btn_take_analysis", "Analyse übernehmen"),
            command=self.add_current_analysis_to_chad_chat,
        )
        self.btn_chad_take_analysis.pack(fill="x", pady=(6, 0))

        control_btns = ttk.Frame(input_frame)
        control_btns.pack(anchor="center", pady=(8, 0))

        self.btn_chad_import_file = ttk.Button(
            control_btns,
            text=self.tr("chad_btn_import_file", "Datei einfügen"),
            command=self.import_file_to_chad_chat,
        )
        self.btn_chad_import_file.pack(side="left", padx=4)

        self.btn_chad_online = ttk.Button(
            control_btns,
            text=self.tr("chad_btn_online_analysis", "Online analysieren"),
            command=self.run_chad_online_analysis,
        )
        self.btn_chad_online.pack(side="left", padx=4)

        self.btn_chad_copy = ttk.Button(
            control_btns,
            text=self.tr("chad_btn_copy_clipboard", "In Zwischenablage kopieren"),
            command=self.copy_chad_analysis_to_clipboard,
        )
        self.btn_chad_copy.pack(side="left", padx=4)

        self.btn_chad_clear = ttk.Button(
            control_btns,
            text=self.tr("chad_btn_clear_analysis", "Analyse leeren"),
            command=self.clear_chad_analysis,
        )
        self.btn_chad_clear.pack(side="left", padx=4)

        self.btn_chad_top = ttk.Button(
            control_btns,
            text=self.tr("chad_btn_top", "Zum Anfang"),
            command=self.scroll_chad_output_to_top,
        )
        self.btn_chad_top.pack(side="left", padx=4)

        self.btn_chad_bottom = ttk.Button(
            control_btns,
            text=self.tr("chad_btn_bottom", "Zum Ende"),
            command=self.scroll_chad_output_to_bottom,
        )
        self.btn_chad_bottom.pack(side="left", padx=4)

        self.chad_auto_scroll_check = ttk.Checkbutton(
            control_btns,
            text=self.tr("chad_check_auto_scroll", "Auto-Scroll"),
            variable=self.chad_auto_scroll_var,
            onvalue=True,
            offvalue=False,
        )
        self.chad_auto_scroll_check.pack(side="left", padx=8)

        status_row = ttk.Frame(input_frame)
        status_row.pack(fill="x", pady=(6, 0))

        ttk.Label(
            status_row,
            textvariable=self.chad_online_status_var,
        ).pack(side="left")

        ttk.Label(
            status_row,
            textvariable=self.chad_chat_status_var,
        ).pack(side="left", padx=(16, 0))

        ttk.Label(
            status_row,
            textvariable=self.chad_analysis_status_var,
        ).pack(side="left", padx=(16, 0))

        ttk.Label(
            input_frame,
            textvariable=self.chad_analysis_source_var,
        ).pack(anchor="w", pady=(4, 0))

        output_frame = ttk.LabelFrame(
            self.tab_chad_analysis,
            text=self.tr("chad_group_output", "Ausgabe / Analyse / Chatverlauf"),
            padding=8,
        )
        output_frame.pack(fill="both", expand=True, pady=(0, 0))

        analysis_text_frame = ttk.Frame(output_frame)
        analysis_text_frame.pack(fill="both", expand=True)

        analysis_scroll = ttk.Scrollbar(analysis_text_frame, orient="vertical")
        analysis_scroll.pack(side="right", fill="y")

        self.chad_analysis_text = tk.Text(
            analysis_text_frame,
            wrap="word",
            yscrollcommand=analysis_scroll.set,
        )
        self.chad_analysis_text.pack(side="left", fill="both", expand=True)
        analysis_scroll.configure(command=self.chad_analysis_text.yview)

        def _scroll_chad_analysis_text(event):
            try:
                self.chad_analysis_text.yview_scroll(
                    int(-1 * (event.delta / 120)), "units"
                )
            except Exception:
                pass
            return "break"

        def _scroll_chad_analysis_text_linux_up(_event):
            self.chad_analysis_text.yview_scroll(-1, "units")
            return "break"

        def _scroll_chad_analysis_text_linux_down(_event):
            self.chad_analysis_text.yview_scroll(1, "units")
            return "break"

        self.chad_analysis_text.bind(
            "<MouseWheel>", _scroll_chad_analysis_text, add="+"
        )
        self.chad_analysis_text.bind(
            "<Button-4>", _scroll_chad_analysis_text_linux_up, add="+"
        )
        self.chad_analysis_text.bind(
            "<Button-5>", _scroll_chad_analysis_text_linux_down, add="+"
        )
        self.chad_analysis_text.tag_configure(
            "chat_user_header",
            font=("TkDefaultFont", 10, "bold"),
            spacing1=8,
            spacing3=2,
        )
        self.chad_analysis_text.tag_configure(
            "chat_assistant_header",
            font=("TkDefaultFont", 10, "bold"),
            spacing1=8,
            spacing3=2,
        )
        self.chad_analysis_text.tag_configure(
            "chat_system_header",
            font=("TkDefaultFont", 10, "bold"),
            spacing1=8,
            spacing3=2,
        )
        self.chad_analysis_text.tag_configure(
            "chat_user_body",
            lmargin1=14,
            lmargin2=14,
            spacing3=6,
        )
        self.chad_analysis_text.tag_configure(
            "chat_assistant_body",
            lmargin1=14,
            lmargin2=14,
            spacing3=6,
        )
        self.chad_analysis_text.tag_configure(
            "chat_system_body",
            lmargin1=14,
            lmargin2=14,
            spacing3=6,
        )
        self.chad_analysis_text.tag_configure("chat_spacer", spacing3=4)
        self._apply_chad_chat_theme(self._app_theme or LIGHT_THEME)
        self.chad_analysis_text.configure(state="disabled")
        self._update_chad_controls_state()

    def is_chad_enabled(self) -> bool:
        return chad_controller.is_chad_enabled(self)

    def _ensure_chad_enabled(self, message: str = "Chad-Anbindung ist deaktiviert."):
        return chad_controller.ensure_chad_enabled(self, message=message)

    def _show_chad_activation_required_popup(self):
        chad_controller.show_chad_activation_required_popup(self)

    def _update_chad_controls_state(self):
        chad_controller.update_chad_controls_state(self)

    def _build_chad_analysis_prompt(self, title: str, raw_text: str) -> str:
        return chad_controller.build_chad_analysis_prompt(self, title, raw_text)

    def _cancel_active_context_menu_timeout(self):
        after_id = self.__dict__.get("_active_context_menu_after_id")
        if after_id is None:
            return
        self._active_context_menu_after_id = None
        try:
            self.after_cancel(after_id)
        except Exception:
            pass

    def _close_active_context_menu(self):
        self._cancel_active_context_menu_timeout()
        menu = self.__dict__.get("_active_context_menu")
        self._active_context_menu = None
        if menu is None:
            return
        try:
            menu.unpost()
        except Exception:
            pass
        try:
            menu.grab_release()
        except Exception:
            pass

    def _schedule_active_context_menu_close(self):
        self._cancel_active_context_menu_timeout()
        self._active_context_menu_after_id = self.after(
            10000, self._close_active_context_menu
        )

    def _handle_app_click_for_context_menu(self, _event=None):
        self._close_active_context_menu()

    def _show_context_menu(self, menu, event):
        self._close_active_context_menu()
        self._active_context_menu = menu
        self._schedule_active_context_menu_close()
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            try:
                menu.grab_release()
            except Exception:
                pass
        return "break"

    def _copy_text_widget_selection(self, text_widget, _event=None):
        try:
            selected = text_widget.get("sel.first", "sel.last")
        except Exception:
            return "break"
        if not selected:
            return "break"
        self.clipboard_clear()
        self.clipboard_append(selected)
        return "break"

    def _copy_listbox_selection(self, listbox, _event=None):
        try:
            selection = listbox.curselection()
        except Exception:
            selection = ()
        if not selection:
            return "break"

        lines = []
        for index in selection:
            try:
                value = listbox.get(index)
            except Exception:
                value = ""
            value = str(value or "").strip()
            if value:
                lines.append(value)

        if not lines:
            return "break"

        self.clipboard_clear()
        self.clipboard_append("\n".join(lines))
        return "break"

    def get_treeview_selection_text(self, treeview) -> str:
        try:
            selection = tuple(treeview.selection())
        except Exception:
            selection = ()
        if not selection:
            return ""

        try:
            columns = tuple(treeview["columns"])
        except Exception:
            columns = ()

        headings = []
        for column in columns:
            try:
                heading = treeview.heading(column, "text")
            except Exception:
                heading = ""
            headings.append(str(heading or column))

        lines = []
        if headings:
            lines.append("\t".join(headings))

        for item_id in selection:
            try:
                values = treeview.item(item_id, "values")
            except Exception:
                values = ()
            row = [str(value or "").strip() for value in values]
            if row:
                lines.append("\t".join(row))

        return "\n".join(lines).strip()

    def _copy_treeview_selection(self, treeview, _event=None):
        text = self.get_treeview_selection_text(treeview)
        if not text:
            return "break"
        self.clipboard_clear()
        self.clipboard_append(text)
        return "break"

    def _paste_text_widget_clipboard(self, text_widget, _event=None):
        try:
            if str(text_widget.cget("state")) == "disabled":
                return "break"
        except Exception:
            pass

        try:
            text_widget.focus_force()
        except Exception:
            pass

        try:
            text_widget.event_generate("<<Paste>>")
            return "break"
        except Exception:
            pass

        try:
            pasted = self.clipboard_get()
        except Exception:
            pasted = ""

        if pasted:
            try:
                text_widget.insert("insert", pasted)
            except Exception:
                pass
        return "break"

    def _paste_entry_widget_clipboard(self, entry_widget, _event=None):
        try:
            state = str(entry_widget.cget("state"))
            if state in {"disabled", "readonly"}:
                return "break"
        except Exception:
            pass

        try:
            entry_widget.focus_force()
        except Exception:
            pass

        try:
            entry_widget.event_generate("<<Paste>>")
            return "break"
        except Exception:
            pass

        try:
            pasted = self.clipboard_get()
        except Exception:
            pasted = ""

        if pasted:
            try:
                entry_widget.insert("insert", pasted)
            except Exception:
                pass
        return "break"

    def attach_selection_copy_shortcuts(self, widget, kind: str):
        if widget is None:
            return

        attached = self.__dict__.setdefault("_selection_copy_shortcuts_attached", set())
        widget_id = str(widget)
        if widget_id in attached:
            return
        attached.add(widget_id)

        if kind == "text":

            def handler(event):
                return self._copy_text_widget_selection(widget, event)

        elif kind == "listbox":

            def handler(event):
                return self._copy_listbox_selection(widget, event)

        elif kind == "treeview":

            def handler(event):
                return self._copy_treeview_selection(widget, event)

        else:
            return

        widget.bind("<Control-c>", handler, add="+")
        widget.bind("<Control-C>", handler, add="+")
        widget.bind("<Control-Insert>", handler, add="+")
        if kind == "text":
            widget.bind(
                "<Control-v>",
                lambda event: self._paste_text_widget_clipboard(widget, event),
                add="+",
            )
            widget.bind(
                "<Control-V>",
                lambda event: self._paste_text_widget_clipboard(widget, event),
                add="+",
            )
            widget.bind(
                "<Shift-Insert>",
                lambda event: self._paste_text_widget_clipboard(widget, event),
                add="+",
            )

    def attach_entry_edit_shortcuts(self, entry_widget):
        if entry_widget is None:
            return

        attached = self.__dict__.setdefault("_selection_copy_shortcuts_attached", set())
        widget_id = str(entry_widget)
        if widget_id in attached:
            return
        attached.add(widget_id)

        def copy_handler(_event=None):
            try:
                selected = entry_widget.selection_get()
            except Exception:
                return "break"
            if not selected:
                return "break"
            self.clipboard_clear()
            self.clipboard_append(selected)
            return "break"

        def paste_handler(event):
            return self._paste_entry_widget_clipboard(entry_widget, event)

        entry_widget.bind("<Control-c>", copy_handler, add="+")
        entry_widget.bind("<Control-C>", copy_handler, add="+")
        entry_widget.bind("<Control-Insert>", copy_handler, add="+")
        entry_widget.bind("<Control-v>", paste_handler, add="+")
        entry_widget.bind("<Control-V>", paste_handler, add="+")
        entry_widget.bind("<Shift-Insert>", paste_handler, add="+")

    def safe_attach_entry_edit_shortcuts(self, widget_name: str):
        widget = self.__dict__.get(widget_name)
        if widget is None:
            return

        try:
            self.attach_entry_edit_shortcuts(widget)
        except Exception as exc:
            try:
                self.log(
                    f"[WARN] Edit-Shortcuts für {widget_name} konnten nicht gesetzt werden: {exc}"
                )
            except Exception:
                pass

    def safe_attach_selection_copy_shortcuts(self, widget_name: str, kind: str):
        widget = self.__dict__.get(widget_name)
        if widget is None:
            return

        try:
            self.attach_selection_copy_shortcuts(widget, kind)
        except Exception as exc:
            try:
                self.log(
                    f"[WARN] Kopier-Shortcuts für {widget_name} konnten nicht gesetzt werden: {exc}"
                )
            except Exception:
                pass

    def attach_all_selection_copy_shortcuts(self):
        text_targets = [
            "text",
            "dns_text",
            "nmap_text",
            "mtr_text",
            "iftop_text",
            "htop_text",
            "linux_help_text",
            "hdd_info_text",
            "security_detail",
            "clamav_output_text",
            "packet_capture_text",
            "lynis_text",
            "rkhunter_text",
            "startup_guard_details_text",
            "chad_chat_input",
            "chad_analysis_text",
        ]
        listbox_targets = [
            "log_overlay_picker_list",
            "linux_help_listbox",
        ]
        treeview_targets = [
            "trace_table",
            "conn_table",
            "hdd_table",
            "link_table",
            "vpn_table",
            "security_table",
            "startup_guard_tree",
        ]
        entry_targets = [
            "conn_filter_entry",
            "hdd_name_entry",
            "link_target_entry",
            "vpn_custom_name_entry",
            "manual_wg_profile_name_entry",
            "manual_wg_private_key_entry",
            "manual_wg_public_key_entry",
            "manual_wg_endpoint_entry",
            "manual_wg_allowed_ips_entry",
            "manual_wg_dns_entry",
            "manual_wg_keepalive_entry",
            "openai_api_key_entry",
        ]

        for widget_name in text_targets:
            self.safe_attach_selection_copy_shortcuts(widget_name, "text")
        for widget_name in listbox_targets:
            self.safe_attach_selection_copy_shortcuts(widget_name, "listbox")
        for widget_name in treeview_targets:
            self.safe_attach_selection_copy_shortcuts(widget_name, "treeview")
        for widget_name in entry_targets:
            self.safe_attach_entry_edit_shortcuts(widget_name)

    def attach_chad_context_menu(self, text_widget, title: str):
        attached = self.__dict__.setdefault("_chad_context_menu_attached", set())
        widget_id = str(text_widget)
        if widget_id in attached:
            return
        attached.add(widget_id)
        self.attach_selection_copy_shortcuts(text_widget, "text")

        menu = tk.Menu(text_widget, tearoff=0)

        def get_selected_text():
            try:
                return text_widget.get("sel.first", "sel.last").strip()
            except Exception:
                return ""

        def copy_selection():
            selected = get_selected_text()
            if not selected:
                return
            self.clipboard_clear()
            self.clipboard_append(selected)

        def paste_from_clipboard():
            try:
                if str(text_widget.cget("state")) == "disabled":
                    return
            except Exception:
                pass

            try:
                text_widget.focus_force()
            except Exception:
                pass

            try:
                text_widget.event_generate("<<Paste>>")
                return
            except Exception:
                pass

            try:
                pasted = self.clipboard_get()
                if pasted:
                    text_widget.insert("insert", pasted)
            except Exception:
                pass

        def select_all():
            try:
                text_widget.tag_add("sel", "1.0", "end")
                text_widget.mark_set("insert", "1.0")
                text_widget.see("insert")
            except Exception:
                pass

        def send_selection_to_chad():
            if not self.is_chad_enabled():
                self._show_chad_activation_required_popup()
                return

            selected = get_selected_text()
            if not selected:
                try:
                    self.chad_analysis_status_var.set(
                        "Keine Auswahl für Chad Analyse markiert."
                    )
                except Exception:
                    pass
                return

            self.prepare_chad_analysis(title, selected)

        def copy_selection_as_prompt():
            selected = get_selected_text()
            if not selected:
                try:
                    self.chad_analysis_status_var.set(
                        "Keine Auswahl zum Kopieren markiert."
                    )
                except Exception:
                    pass
                return

            prompt = self._build_chad_analysis_prompt(title, selected)
            self.clipboard_clear()
            self.clipboard_append(prompt)

            try:
                self.chad_analysis_status_var.set(
                    "Chad-Prompt in Zwischenablage kopiert."
                )
            except Exception:
                pass

        menu.add_command(label="Kopieren", command=copy_selection)
        menu.add_command(label="Einfügen", command=paste_from_clipboard)
        menu.add_command(label="Alles auswählen", command=select_all)
        menu.add_separator()
        menu.add_command(
            label="Auswahl an Chad Analyse senden",
            command=send_selection_to_chad,
        )
        menu.add_command(
            label="Auswahl als Chad-Prompt kopieren",
            command=copy_selection_as_prompt,
        )

        def show_menu(event):
            return self._show_context_menu(menu, event)

        text_widget.bind("<Button-3>", show_menu, add="+")
        text_widget._chad_context_menu = menu

    def safe_attach_chad_context_menu(self, widget_name: str, title: str):
        widget = self.__dict__.get(widget_name)
        if widget is None:
            return

        try:
            self.attach_chad_context_menu(widget, title)
        except Exception as exc:
            try:
                self.log(
                    f"[WARN] Kontextmenü für {widget_name} konnte nicht gesetzt werden: {exc}"
                )
            except Exception:
                pass

    def attach_all_chad_context_menus(self):
        targets = [
            ("text", "NetMon Log Analyse"),
            ("dns_text", "DNS Analyse"),
            ("nmap_text", "Nmap Analyse"),
            ("mtr_text", "MTR Analyse"),
            ("iftop_text", "iftop Analyse"),
            ("htop_text", "htop Analyse"),
            ("linux_help_text", "Linux Hilfe Analyse"),
            ("hdd_info_text", "HDD Analyse"),
            ("security_detail", "Security Analyse"),
            ("clamav_output_text", "ClamAV Analyse"),
            ("packet_capture_text", "Packet Capture Analyse"),
            ("lynis_text", "Lynis Analyse"),
            ("rkhunter_text", "rkhunter Analyse"),
            (
                "startup_guard_details_text",
                "Startup Guard Eintrag Analyse",
            ),
            ("chad_chat_input", "Chad Chat Eingabe"),
            ("chad_analysis_text", "Chad Verlauf Analyse"),
        ]

        for widget_name, title in targets:
            self.safe_attach_chad_context_menu(widget_name, title)

    def set_chad_analysis(self, title: str, raw_text: str):
        chad_controller.set_chad_analysis(self, title, raw_text)

    def copy_chad_analysis_to_clipboard(self):
        chad_controller.copy_chad_analysis_to_clipboard(self)

    def clear_chad_analysis(self):
        chad_controller.clear_chad_analysis(self)

    def run_chad_online_analysis(self):
        chad_controller.run_chad_online_analysis(self)

    def prepare_chad_analysis(
        self, title: str, raw_text: str, show_preview: bool | None = None
    ):
        chad_controller.prepare_chad_analysis(
            self,
            title,
            raw_text,
            show_preview=show_preview,
        )

    def open_ai_analysis_preview(self, title: str, raw_text: str):
        chad_controller.open_ai_analysis_preview(self, title, raw_text)

    def linux_help_placeholder_action(self):
        self.linux_help_status_var.set(
            self.tr(
                "linux_help_status_display_only",
                "Linux Hilfe: Anzeige aktiv. Ausführung folgt erst im nächsten Integrationsschritt.",
            )
        )
        if hasattr(self, "linux_help_text") and self.linux_help_text is not None:
            self.linux_help_text.delete("1.0", "end")
            self.linux_help_text.insert(
                "1.0",
                self.tr(
                    "linux_help_placeholder_body",
                    "Aktuell ist 'Linux Hilfe' an die CSV-Daten angebunden.\n\n"
                    "Die spätere sichere Ausführung und Terminal-Integration "
                    "folgen in einem separaten Schritt.",
                ),
            )

    def _validate_linux_help_command_for_terminal(self, cmd: str) -> tuple[bool, str]:
        if not isinstance(cmd, str):
            return False, "Befehl ist kein gültiger Textwert."
        text = cmd.strip()
        if not text:
            return False, "Befehl ist leer."
        if "\x00" in text:
            return False, "Befehl enthält ungültige Nullbytes."
        if "\n" in text or "\r" in text:
            return False, "Mehrzeilige Befehle sind hier nicht erlaubt."
        return True, ""

    def _build_external_terminal_command(self, script: str):
        candidates = [
            (
                "x-terminal-emulator",
                ["x-terminal-emulator", "-e", "bash", "-lc", script],
            ),
            ("gnome-terminal", ["gnome-terminal", "--", "bash", "-lc", script]),
            ("konsole", ["konsole", "-e", "bash", "-lc", script]),
            (
                "xfce4-terminal",
                ["xfce4-terminal", "--command", f"bash -lc {shlex.quote(script)}"],
            ),
            ("mate-terminal", ["mate-terminal", "--", "bash", "-lc", script]),
            ("lxterminal", ["lxterminal", "-e", f"bash -lc {shlex.quote(script)}"]),
            ("xterm", ["xterm", "-e", "bash", "-lc", script]),
        ]
        for binary, command in candidates:
            if shutil.which(binary):
                return command
        return None

    def _open_command_in_external_terminal(self, cmd: str):
        script = f"{cmd}; exec bash"
        launch_cmd = self._build_external_terminal_command(script)
        if not launch_cmd:
            raise RuntimeError("Kein unterstütztes Terminal-Programm gefunden.")
        subprocess.Popen(launch_cmd, start_new_session=True)
        return launch_cmd[0]

    def _set_linux_help_dirty(self, dirty=True):
        linux_help_controller.set_linux_help_dirty(self, dirty)

    def _linux_help_csv_candidates(self):
        return linux_help_controller.linux_help_csv_candidates()

    def _linux_help_csv_path(self):
        return linux_help_controller.linux_help_csv_path()

    def _linux_help_slugify(self, text: str) -> str:
        return linux_help_controller.linux_help_slugify(text)

    def _linux_help_make_id(self, raw, existing_ids=None):
        return linux_help_controller.linux_help_make_id(raw, existing_ids=existing_ids)

    def linux_help_get_risk_level(self, cmd: str) -> str:
        return linux_help_controller.linux_help_get_risk_level(cmd)

    def normalize_linux_help_entry(self, raw, existing_ids=None):
        return linux_help_controller.normalize_linux_help_entry(
            self,
            raw,
            existing_ids=existing_ids,
        )

    def _rebuild_linux_help_indexes(self):
        linux_help_controller.rebuild_linux_help_indexes(self)

    def import_linux_help_csv(self, path):
        return linux_help_controller.import_linux_help_csv(self, path)

    def export_linux_help_csv(self, path):
        linux_help_controller.export_linux_help_csv(self, path)

    def load_linux_help_data(self):
        return linux_help_controller.load_linux_help_data(self)

    def save_linux_help_data(self, update_status=True):
        return linux_help_controller.save_linux_help_data(
            self,
            update_status=update_status,
        )

    def get_linux_help_categories(self):
        return linux_help_controller.get_linux_help_categories(self)

    def get_linux_help_commands_for_category(self, category):
        return linux_help_controller.get_linux_help_commands_for_category(
            self,
            category,
        )

    def get_linux_help_command_by_id(self, command_id):
        return linux_help_controller.get_linux_help_command_by_id(self, command_id)

    def _sync_linux_help_form_categories(self):
        linux_help_controller.sync_linux_help_form_categories(self)

    def _set_linux_help_form_from_entry(self, item=None):
        linux_help_controller.set_linux_help_form_from_entry(self, item)

    def _set_linux_help_risk_light(self, risk):
        linux_help_controller.set_linux_help_risk_light(self, risk)

    def _update_linux_help_notes(self, item=None):
        linux_help_controller.update_linux_help_notes(self, item)

    def refresh_linux_help_list(self, category=None):
        linux_help_controller.refresh_linux_help_list(self, category=category)

    def on_linux_help_select(self, _event=None):
        linux_help_controller.on_linux_help_select(self, _event)

    def linux_help_apply_form_to_entry(self, mark_dirty=True):
        return linux_help_controller.linux_help_apply_form_to_entry(
            self,
            mark_dirty=mark_dirty,
        )

    def linux_help_new_entry(self):
        linux_help_controller.linux_help_new_entry(self)

    def linux_help_edit_selected(self):
        linux_help_controller.linux_help_edit_selected(self)

    def linux_help_delete_selected(self):
        linux_help_controller.linux_help_delete_selected(self)

    def linux_help_import_csv_dialog(self):
        linux_help_controller.linux_help_import_csv_dialog(self)

    def linux_help_export_csv_dialog(self):
        linux_help_controller.linux_help_export_csv_dialog(self)

    def send_linux_help_command_to_terminal(self):
        return linux_help_controller.send_linux_help_command_to_terminal(self)

    def reload_linux_help_data(self):
        linux_help_controller.reload_linux_help_data(self)

    def _set_dot_profile_choice_safe(self, choice):
        dot_controller.set_dot_profile_choice_safe(self, choice)

    def _update_dot_light(self):
        dot_controller.update_dot_light(self)

    def _update_dot_light_safe(self):
        dot_controller.update_dot_light_safe(self)

    def _set_dot_state(
        self,
        ok,
        *,
        profile="",
        servers=None,
        mode="",
        enabled=None,
        timestamp=None,
        vpn_conflict=None,
        vpn_active=None,
    ):
        dot_controller.set_dot_state(
            self,
            ok,
            profile=profile,
            servers=servers,
            mode=mode,
            enabled=enabled,
            timestamp=timestamp,
            vpn_conflict=vpn_conflict,
            vpn_active=vpn_active,
        )

    def _set_dot_state_safe(self, ok, **kwargs):
        dot_controller.set_dot_state_safe(self, ok, **kwargs)

    def _dot_profile_config(self, profile_name=None):
        return dot_controller.dot_profile_config(self, profile_name)

    def _build_dot_dropin_content(self, profile_name=None):
        return dot_controller.build_dot_dropin_content(self, profile_name)

    def _parse_dot_dropin(self, text):
        return dot_controller.parse_dot_dropin(text)

    def _read_dot_dropin(self):
        return dot_controller.read_dot_dropin(self)

    def is_vpn_dns_active(self):
        return dot_controller.is_vpn_dns_active(self)

    def detect_netmon_dot_dropin_conflict(self, path=None):
        return dot_controller.detect_netmon_dot_dropin_conflict(self, path)

    def _dot_leak_conflict_warning_text(self):
        return dot_controller.dot_leak_conflict_warning_text()

    def _infer_dot_profile_name(self, servers):
        return dot_controller.infer_dot_profile_name(servers)

    def _parse_dot_resolvectl_status(self, status_text):
        return dot_controller.parse_dot_resolvectl_status(status_text)

    def parse_dot_status_from_resolvectl(self, output: str):
        return dot_controller.parse_dot_status_from_resolvectl(output)

    def _dot_preflight(self):
        return dot_controller.dot_preflight()

    def _dot_tls_probe(self, server, hostname, timeout=4):
        return dot_controller.dot_tls_probe(server, hostname, timeout=timeout)

    def _collect_dot_status(self, run_test=False):
        return dot_controller.collect_dot_status(self, run_test=run_test)

    def _show_dot_status_result(
        self,
        status,
        *,
        append=False,
        startup=False,
        show_output=None,
        emit_chad_warning=None,
    ):
        dot_controller.show_dot_status_result(
            self,
            status,
            append=append,
            startup=startup,
            show_output=show_output,
            emit_chad_warning=emit_chad_warning,
        )

    def startup_dot_status_check(self):
        dot_controller.startup_dot_status_check(self)

    def show_dot_status_async(self, startup=False):
        dot_controller.show_dot_status_async(self, startup=startup)

    def enable_dot_async(self):
        dot_controller.enable_dot_async(self)

    def disable_dot_async(self):
        dot_controller.disable_dot_async(self)

    def test_dot_async(self):
        dot_controller.test_dot_async(self)

    def build_packet_capture_tab(self):
        config = ttk.LabelFrame(
            self.tab_packet_capture, text="Konfiguration", padding=6
        )
        config.pack(fill="x")
        config.columnconfigure(1, weight=1)

        ttk.Label(config, text="Interface:").grid(row=0, column=0, sticky="w")
        self.packet_capture_interface_combo = ttk.Combobox(
            config,
            textvariable=self.packet_capture_interface_var,
            state="readonly",
            values=[],
        )
        self.packet_capture_interface_combo.grid(
            row=0, column=1, sticky="ew", padx=(8, 8)
        )
        self.btn_packet_capture_refresh = ttk.Button(
            config,
            text="Interfaces aktualisieren",
            command=self.refresh_packet_capture_interfaces,
        )
        self.btn_packet_capture_refresh.grid(row=0, column=2, sticky="ew")
        self.register_feature_button(
            "packet_capture_ui", self.btn_packet_capture_refresh
        )
        self.add_tooltip(
            self.btn_packet_capture_refresh,
            "Aktualisiert die verfügbare Liste der Capture-Interfaces.",
        )

        ttk.Label(config, text="Capture-Filter:").grid(
            row=1, column=0, sticky="w", pady=(6, 0)
        )
        ttk.Entry(config, textvariable=self.packet_capture_filter_var).grid(
            row=1,
            column=1,
            columnspan=2,
            sticky="ew",
            padx=(8, 0),
            pady=(6, 0),
        )

        ttk.Label(config, text="Datei:").grid(row=2, column=0, sticky="w", pady=(6, 0))
        ttk.Entry(config, textvariable=self.packet_capture_file_var).grid(
            row=2, column=1, sticky="ew", padx=(8, 8), pady=(6, 0)
        )
        self.btn_packet_capture_choose = ttk.Button(
            config,
            text="Datei wählen",
            command=self.choose_packet_capture_file,
        )
        self.btn_packet_capture_choose.grid(row=2, column=2, sticky="ew", pady=(6, 0))
        self.register_feature_button(
            "packet_capture_ui", self.btn_packet_capture_choose
        )
        self.add_tooltip(
            self.btn_packet_capture_choose,
            "Legt Speicherort und Dateinamen für die Capture-Datei fest.",
        )

        actions = ttk.LabelFrame(self.tab_packet_capture, text="Aktionen", padding=6)
        actions.pack(fill="x", pady=(8, 0))

        self.btn_packet_capture_start = ttk.Button(
            actions,
            text="Capture starten",
            command=self.packet_capture_start_async,
        )
        self.btn_packet_capture_start.pack(side="left")
        self.register_feature_button("packet_capture", self.btn_packet_capture_start)
        self.add_tooltip(
            self.btn_packet_capture_start,
            "Startet einen Mitschnitt mit tshark oder dumpcap.",
        )

        self.btn_packet_capture_stop = ttk.Button(
            actions,
            text="Capture stoppen",
            command=self.packet_capture_stop,
        )
        self.btn_packet_capture_stop.pack(side="left", padx=(8, 0))
        self.register_feature_button("packet_capture", self.btn_packet_capture_stop)
        self.add_tooltip(
            self.btn_packet_capture_stop,
            "Beendet einen laufenden Mitschnitt sauber.",
        )

        self.btn_packet_capture_open_wireshark = ttk.Button(
            actions,
            text="In Wireshark öffnen",
            command=self.open_packet_capture_in_wireshark,
        )
        self.btn_packet_capture_open_wireshark.pack(side="left", padx=(8, 0))
        self.register_feature_button(
            "wireshark_open", self.btn_packet_capture_open_wireshark
        )
        self.add_tooltip(
            self.btn_packet_capture_open_wireshark,
            "Öffnet die aktuelle Capture-Datei in Wireshark, falls installiert.",
        )

        self.btn_packet_capture_open_file = ttk.Button(
            actions,
            text="Capture-Datei öffnen",
            command=self.open_packet_capture_file_location,
        )
        self.btn_packet_capture_open_file.pack(side="left", padx=(8, 0))
        self.register_feature_button(
            "packet_capture", self.btn_packet_capture_open_file
        )
        self.add_tooltip(
            self.btn_packet_capture_open_file,
            "Öffnet die Capture-Datei oder ihren Ordner im Dateimanager.",
        )

        self.btn_packet_capture_info = ttk.Button(
            actions,
            text="Datei-Info",
            command=self.packet_capture_show_file_info,
        )
        self.btn_packet_capture_info.pack(side="left", padx=(8, 0))
        self.register_feature_button("packet_capture_ui", self.btn_packet_capture_info)
        self.add_tooltip(
            self.btn_packet_capture_info,
            "Zeigt Metadaten und Größe der aktuellen Capture-Datei an.",
        )

        status = ttk.LabelFrame(self.tab_packet_capture, text="Status", padding=6)
        status.pack(fill="x", pady=(8, 0))
        for col in range(4):
            status.columnconfigure(col, weight=1)

        ttk.Label(status, text="Status:").grid(row=0, column=0, sticky="w")
        ttk.Label(status, textvariable=self.packet_capture_status_var).grid(
            row=0, column=1, sticky="w", padx=(8, 12)
        )
        ttk.Label(status, text="Laufzeit:").grid(row=0, column=2, sticky="w")
        ttk.Label(status, textvariable=self.packet_capture_runtime_var).grid(
            row=0, column=3, sticky="w", padx=(8, 0)
        )
        ttk.Label(status, text="Pakete:").grid(row=1, column=0, sticky="w", pady=(4, 0))
        ttk.Label(status, textvariable=self.packet_capture_packets_var).grid(
            row=1, column=1, sticky="w", padx=(8, 12), pady=(4, 0)
        )
        ttk.Label(status, text="Dateigröße:").grid(
            row=1, column=2, sticky="w", pady=(4, 0)
        )
        ttk.Label(status, textvariable=self.packet_capture_size_var).grid(
            row=1, column=3, sticky="w", padx=(8, 0), pady=(4, 0)
        )

        output_frame = ttk.LabelFrame(
            self.tab_packet_capture, text="Capture-Status", padding=6
        )
        output_frame.pack(fill="both", expand=True, pady=(8, 0))
        self.packet_capture_text = tk.Text(output_frame, wrap="word", height=14)
        self.packet_capture_text.pack(fill="both", expand=True)
        self.packet_capture_text.insert(
            "1.0",
            "Packet-Capture-Status und Dateiinformationen erscheinen hier.\n",
        )
        self.packet_capture_text.configure(state="disabled")
        self.packet_capture_set_running_state(False)

    def packet_capture_set_running_state(self, running):
        packet_capture_controller.packet_capture_set_running_state(self, running)

    def append_packet_capture_output(self, text: str, replace: bool = False):
        packet_capture_controller.append_packet_capture_output(
            self,
            text,
            replace=replace,
        )

    def clear_packet_capture_output(self):
        packet_capture_controller.clear_packet_capture_output(self)

    def packet_capture_generate_default_path(self):
        return packet_capture_controller.packet_capture_generate_default_path(self)

    def choose_packet_capture_file(self):
        packet_capture_controller.choose_packet_capture_file(self)

    def refresh_packet_capture_interfaces(self):
        packet_capture_controller.refresh_packet_capture_interfaces(self)

    def packet_capture_update_runtime(self):
        packet_capture_controller.packet_capture_update_runtime(self)

    def packet_capture_show_file_info(self):
        packet_capture_controller.packet_capture_show_file_info(self)

    def packet_capture_start_async(self):
        packet_capture_controller.packet_capture_start_async(self)

    def packet_capture_stop(self):
        packet_capture_controller.packet_capture_stop(self)

    def open_packet_capture_in_wireshark(self):
        packet_capture_controller.open_packet_capture_in_wireshark(self)

    def open_packet_capture_file_location(self):
        packet_capture_controller.open_packet_capture_file_location(self)

    def build_lynis_tab(self):
        lynis_top = ttk.Frame(self.tab_lynis)
        lynis_top.pack(fill="x", pady=(0, 8))
        lynis_row1 = ttk.Frame(lynis_top)
        lynis_row1.pack(fill="x")
        lynis_row2 = ttk.Frame(lynis_top)
        lynis_row2.pack(fill="x", pady=(6, 0))

        self.btn_lynis_quick = ttk.Button(
            lynis_row1,
            text="Quick Scan",
            command=self.run_lynis_quick,
        )
        self.btn_lynis_quick.pack(side="left")
        self.register_feature_button("lynis", self.btn_lynis_quick)
        self.add_tooltip(
            self.btn_lynis_quick,
            "Startet einen schnellen Lynis-Scan ohne root-Rechte.",
        )

        self.btn_lynis_full = ttk.Button(
            lynis_row1,
            text="Full Security Scan (sudo)",
            command=self.run_lynis_full,
        )
        self.btn_lynis_full.pack(side="left", padx=(8, 0))
        self.register_feature_button("lynis", self.btn_lynis_full)
        self.add_tooltip(
            self.btn_lynis_full,
            "Startet eine vollständige Lynis-Sicherheitsprüfung mit sudo.",
        )

        self.btn_lynis_clear = ttk.Button(
            lynis_row1,
            text="Lynis-Ausgabe leeren",
            command=self.clear_lynis_output,
        )
        self.btn_lynis_clear.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_lynis_clear,
            "Leert die aktuelle Lynis-Ausgabe im Tab.",
        )

        self.btn_lynis_issues = ttk.Button(
            lynis_row1,
            text="Nur Probleme & Empfehlungen",
            command=self.show_lynis_issues_only,
        )
        self.btn_lynis_issues.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_lynis_issues,
            "Zeigt nur Warnungen, Empfehlungen und auffällige Lynis-Zeilen an.",
        )

        self.btn_lynis_full_output = ttk.Button(
            lynis_row2,
            text="Vollständige Ausgabe",
            command=self.show_lynis_full_output,
        )
        self.btn_lynis_full_output.pack(side="left")
        self.add_tooltip(
            self.btn_lynis_full_output,
            "Stellt die vollständige bereinigte Lynis-Ausgabe wieder her.",
        )

        self.btn_lynis_chad_issues = ttk.Button(
            lynis_row2,
            text="Probleme mit Chad analysieren",
            command=self.analyze_lynis_issues_with_chad,
        )
        self.btn_lynis_chad_issues.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_lynis_chad_issues,
            "Bereitet nur Lynis-Probleme und Empfehlungen für Chad vor.",
        )

        self.btn_lynis_hint = ttk.Button(
            lynis_row2,
            text="Installationshinweis",
            command=self.show_lynis_install_hint,
        )
        self.btn_lynis_hint.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_lynis_hint,
            "Zeigt kopierbare Installationshinweise für Lynis im Log an.",
        )

        ttk.Label(lynis_row2, text="Status:").pack(side="left", padx=(16, 4))
        ttk.Label(lynis_row2, textvariable=self.lynis_status_var).pack(side="left")

        lynis_summary = ttk.LabelFrame(
            self.tab_lynis, text="Zusammenfassung", padding=6
        )
        lynis_summary.pack(fill="x", pady=(0, 8))
        ttk.Label(
            lynis_summary,
            textvariable=self.lynis_summary_var,
            justify="left",
        ).pack(anchor="w")

        self.lynis_text = tk.Text(self.tab_lynis, wrap="word", height=20)
        self.lynis_text.pack(fill="both", expand=True)
        self.lynis_text.insert("1.0", "Lynis-Ausgabe erscheint hier.\n")
        self.lynis_text.configure(state="disabled")

    def open_lynis_tab(self):
        security_audit_controller.open_lynis_tab(self)

    def clear_lynis_output(self):
        security_audit_controller.clear_lynis_output(self)

    def update_lynis_output(self, text: str, append: bool = False):
        security_audit_controller.update_lynis_output(
            self,
            text,
            append=append,
            normalize_output_fn=normalize_lynis_output,
        )

    def extract_lynis_issues_and_recommendations(self, text: str) -> str:
        return security_audit_controller.extract_lynis_issues_and_recommendations(
            text,
            normalize_output_fn=normalize_lynis_output,
        )

    def show_lynis_issues_only(self):
        security_audit_controller.show_lynis_issues_only(
            self,
            normalize_output_fn=normalize_lynis_output,
        )

    def show_lynis_full_output(self):
        security_audit_controller.show_lynis_full_output(self)

    def analyze_lynis_issues_with_chad(self):
        security_audit_controller.analyze_lynis_issues_with_chad(
            self,
            normalize_output_fn=normalize_lynis_output,
        )

    def _update_lynis_summary_from_output(self, output: str, rc: int):
        security_audit_controller.update_lynis_summary_from_output(
            self,
            output,
            rc,
            normalize_output_fn=normalize_lynis_output,
        )

    def _set_lynis_buttons_state(self, enabled: bool):
        security_audit_controller.set_lynis_buttons_state(self, enabled)

    def _finish_lynis_audit(self, rc, out, *, mode: str = "quick"):
        security_audit_controller.finish_lynis_audit(
            self,
            rc,
            out,
            mode=mode,
            normalize_output_fn=normalize_lynis_output,
        )

    def _lynis_scan_worker(self, cmd, timeout: int, mode: str, header: str):
        security_audit_controller.lynis_scan_worker(
            self,
            cmd,
            timeout,
            mode,
            header,
            normalize_output_fn=normalize_lynis_output,
        )

    def _start_lynis_scan(self, *, mode: str, cmd, timeout: int, header: str):
        security_audit_controller.start_lynis_scan(
            self,
            mode=mode,
            cmd=cmd,
            timeout=timeout,
            header=header,
        )

    def run_lynis_quick(self):
        security_audit_controller.run_lynis_quick(self)

    def run_lynis_full(self):
        security_audit_controller.run_lynis_full(self)

    def lynis_audit_async(self):
        security_audit_controller.lynis_audit_async(self)

    def build_rkhunter_tab(self):
        rkh_top = ttk.Frame(self.tab_rkhunter)
        rkh_top.pack(fill="x", pady=(0, 8))

        self.btn_rkhunter_update = ttk.Button(
            rkh_top,
            text="rkhunter Update",
            command=self.rkhunter_update_async,
        )
        self.btn_rkhunter_update.pack(side="left")
        self.register_feature_button("rkhunter", self.btn_rkhunter_update)
        self.add_tooltip(
            self.btn_rkhunter_update,
            "Aktualisiert die rkhunter-Datenbasis ohne weitere Systemeingriffe.",
        )

        self.btn_rkhunter_scan = ttk.Button(
            rkh_top,
            text="rkhunter Scan starten",
            command=self.rkhunter_scan_async,
        )
        self.btn_rkhunter_scan.pack(side="left", padx=(8, 0))
        self.register_feature_button("rkhunter", self.btn_rkhunter_scan)
        self.add_tooltip(
            self.btn_rkhunter_scan,
            "Startet einen nicht-interaktiven rkhunter-Check ohne Reparaturaktionen.",
        )

        self.btn_rkhunter_repair = ttk.Button(
            rkh_top,
            text="Repair RKHunter update",
            command=self.rkhunter_repair_update_async,
        )
        self.btn_rkhunter_repair.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_rkhunter_repair,
            "Sichert /etc/rkhunter.conf, repariert typische Update-Werte und startet das Update erneut.",
        )
        self.btn_rkhunter_repair.configure(state="disabled")

        self.btn_rkhunter_clear = ttk.Button(
            rkh_top,
            text="Ausgabe leeren",
            command=self.clear_rkhunter_output,
        )
        self.btn_rkhunter_clear.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_rkhunter_clear,
            "Leert die aktuelle rkhunter-Ausgabe im Tab.",
        )

        ttk.Label(rkh_top, text="Status:").pack(side="left", padx=(16, 4))
        ttk.Label(rkh_top, textvariable=self.rkhunter_status_var).pack(side="left")
        ttk.Label(rkh_top, text=" | Update-Stand:").pack(side="left", padx=(12, 4))
        ttk.Label(rkh_top, textvariable=self.rkhunter_last_update_var).pack(side="left")

        rkh_summary = ttk.LabelFrame(
            self.tab_rkhunter, text="Zusammenfassung", padding=6
        )
        rkh_summary.pack(fill="x", pady=(0, 8))
        ttk.Label(
            rkh_summary,
            textvariable=self.rkhunter_summary_var,
            justify="left",
        ).pack(anchor="w")

        self.rkhunter_text = tk.Text(self.tab_rkhunter, wrap="word", height=20)
        self.rkhunter_text.pack(fill="both", expand=True)
        self.rkhunter_text.insert("1.0", "rkhunter-Ausgabe erscheint hier.\n")
        self.rkhunter_text.configure(state="disabled")

    def _set_rkhunter_repair_available(self, enabled):
        security_audit_controller.set_rkhunter_repair_available(self, enabled)

    def open_rkhunter_tab(self):
        security_audit_controller.open_rkhunter_tab(self)

    def update_rkhunter_output(self, text: str, append: bool = False):
        security_audit_controller.update_rkhunter_output(
            self,
            text,
            append=append,
        )

    def clear_rkhunter_output(self):
        security_audit_controller.clear_rkhunter_output(self)

    def _update_rkhunter_summary_from_output(self, output: str, rc: int):
        security_audit_controller.update_rkhunter_summary_from_output(self, output, rc)

    def _finish_rkhunter_update(self, rc, out):
        security_audit_controller.finish_rkhunter_update(
            self,
            rc,
            out,
            detect_update_failure_fn=detect_update_failure,
        )

    def _rkhunter_update_worker(self):
        security_audit_controller.rkhunter_update_worker(
            self,
            run_rkhunter_update_fn=run_rkhunter_update,
        )

    def rkhunter_update_async(self):
        security_audit_controller.rkhunter_update_async(
            self,
            is_rkhunter_installed_fn=is_rkhunter_installed,
        )

    def _finish_rkhunter_repair(self, backup_rc, backup_out, repair_rc, repair_out):
        security_audit_controller.finish_rkhunter_repair(
            self,
            backup_rc,
            backup_out,
            repair_rc,
            repair_out,
            backup_path=RKHUNTER_BACKUP_PATH,
        )

    def _rkhunter_repair_worker(self):
        security_audit_controller.rkhunter_repair_worker(
            self,
            rkhunter_backup_config_fn=rkhunter_backup_config,
            rkhunter_repair_config_fn=rkhunter_repair_config,
        )

    def rkhunter_repair_update_async(self):
        security_audit_controller.rkhunter_repair_update_async(
            self,
            is_rkhunter_installed_fn=is_rkhunter_installed,
        )

    def _finish_rkhunter_scan(self, rc, out):
        security_audit_controller.finish_rkhunter_scan(self, rc, out)

    def _rkhunter_scan_worker(self):
        security_audit_controller.rkhunter_scan_worker(self)

    def rkhunter_scan_async(self):
        security_audit_controller.rkhunter_scan_async(
            self,
            is_rkhunter_installed_fn=is_rkhunter_installed,
        )

    def build_startup_guard_tab(self):
        startup_top = ttk.Frame(self.tab_startup_guard)
        startup_top.pack(fill="x", pady=(0, 8))

        # UI Cleanup inventory (Startup Guard):
        # Basis: Baseline erstellen -> startup_guard_create_baseline
        # Basis: Startup prüfen -> startup_guard_check
        # Analyse: Eintrag/Änderungen mit Chad -> analyze_selected... / analyze_changes...
        # Vertrauen: Als bekannt markieren / entfernen / Vertrauensliste anzeigen
        startup_row1 = ttk.LabelFrame(
            startup_top,
            text=self.tr("startup_guard_group_basic", "Basis"),
            padding=6,
        )
        startup_row1.pack(fill="x")
        startup_row2 = ttk.LabelFrame(
            startup_top,
            text=self.tr("startup_guard_group_analysis", "Analyse"),
            padding=6,
        )
        startup_row2.pack(fill="x", pady=(6, 0))
        startup_row3 = ttk.LabelFrame(
            startup_top,
            text=self.tr("startup_guard_group_trust", "Vertrauen"),
            padding=6,
        )
        startup_row3.pack(fill="x", pady=(6, 0))
        startup_status_row = ttk.Frame(startup_top)
        startup_status_row.pack(fill="x", pady=(6, 0))

        self.btn_startup_guard_baseline = ttk.Button(
            startup_row1,
            text=self.tr("startup_guard_btn_baseline", "Baseline erstellen"),
            command=self.startup_guard_create_baseline,
        )
        self.btn_startup_guard_baseline.pack(side="left")
        self.add_tooltip(
            self.btn_startup_guard_baseline,
            "Sammelt aktuelle Startup-/Autostart-Einträge und speichert sie als Baseline.",
        )

        self.btn_startup_guard_check = ttk.Button(
            startup_row1,
            text=self.tr("startup_guard_btn_check", "Startup prüfen"),
            command=self.startup_guard_check,
        )
        self.btn_startup_guard_check.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_startup_guard_check,
            "Vergleicht den aktuellen Startup-Zustand mit der gespeicherten Baseline.",
        )

        self.btn_startup_guard_clear = ttk.Button(
            startup_row1,
            text=self.tr("startup_guard_btn_clear", "Ergebnis leeren"),
            command=self.clear_startup_guard_output,
        )
        self.btn_startup_guard_clear.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_startup_guard_clear,
            "Leert die aktuelle Startup-Guard-Ausgabe im Tab.",
        )

        self.btn_startup_guard_chad_selected = ttk.Button(
            startup_row2,
            text=self.tr(
                "startup_guard_btn_chad_selected",
                "Eintrag mit Chad analysieren",
            ),
            command=self.analyze_selected_startup_entry_with_chad,
        )
        self.btn_startup_guard_chad_selected.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_startup_guard_chad_selected,
            "Übergibt den aktuell ausgewählten Startup-Eintrag an den Chad Analyse-Bereich.",
        )

        self.btn_startup_guard_chad_changes = ttk.Button(
            startup_row2,
            text=self.tr(
                "startup_guard_btn_chad_changes",
                "Alle Änderungen mit Chad analysieren",
            ),
            command=self.analyze_startup_guard_changes_with_chad,
        )
        self.btn_startup_guard_chad_changes.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_startup_guard_chad_changes,
            "Übergibt neue, entfernte und geänderte Startup-Einträge gesammelt an den Chad Analyse-Bereich.",
        )

        self.btn_startup_guard_mark_trusted = ttk.Button(
            startup_row3,
            text=self.tr("startup_guard_btn_mark_trusted", "Als bekannt markieren"),
            command=self.startup_guard_mark_selected_trusted,
        )
        self.btn_startup_guard_mark_trusted.pack(side="left")
        self.add_tooltip(
            self.btn_startup_guard_mark_trusted,
            "Markiert den ausgewählten Startup-Eintrag als bekannt, solange Pfad und Hash unverändert bleiben.",
        )

        self.btn_startup_guard_remove_trusted = ttk.Button(
            startup_row3,
            text=self.tr(
                "startup_guard_btn_remove_trusted",
                "Aus Vertrauensliste entfernen",
            ),
            command=self.startup_guard_remove_selected_trusted,
        )
        self.btn_startup_guard_remove_trusted.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_startup_guard_remove_trusted,
            "Entfernt den ausgewählten Startup-Eintrag aus der Vertrauensliste.",
        )

        self.btn_startup_guard_show_trustlist = ttk.Button(
            startup_row3,
            text=self.tr(
                "startup_guard_btn_show_trustlist",
                "Vertrauensliste anzeigen",
            ),
            command=self.startup_guard_show_trustlist,
        )
        self.btn_startup_guard_show_trustlist.pack(side="left", padx=(8, 0))
        self.add_tooltip(
            self.btn_startup_guard_show_trustlist,
            "Zeigt alle aktuell als bekannt markierten Startup-Einträge an.",
        )

        ttk.Label(
            startup_status_row,
            text=self.tr("startup_guard_label_status", "Status:"),
        ).pack(side="left")
        ttk.Label(
            startup_status_row,
            textvariable=self.startup_guard_status_var,
        ).pack(side="left")

        content = ttk.Panedwindow(self.tab_startup_guard, orient="vertical")
        content.pack(fill="both", expand=True)

        tree_frame = ttk.Frame(content)
        details_frame = ttk.LabelFrame(content, text="Details", padding=6)
        content.add(tree_frame, weight=3)
        content.add(details_frame, weight=2)

        tree_scroll_y = ttk.Scrollbar(tree_frame, orient="vertical")
        tree_scroll_y.pack(side="right", fill="y")
        tree_scroll_x = ttk.Scrollbar(tree_frame, orient="horizontal")
        tree_scroll_x.pack(side="bottom", fill="x")

        self.startup_guard_tree = ttk.Treeview(
            tree_frame,
            columns=("status", "source", "name", "path"),
            show="headings",
            height=12,
            yscrollcommand=tree_scroll_y.set,
            xscrollcommand=tree_scroll_x.set,
        )
        self.startup_guard_tree.heading("status", text="Status")
        self.startup_guard_tree.heading("source", text="Quelle")
        self.startup_guard_tree.heading("name", text="Name")
        self.startup_guard_tree.heading("path", text="Pfad")
        self.startup_guard_tree.column("status", width=110, anchor="center")
        self.startup_guard_tree.column("source", width=150, anchor="w")
        self.startup_guard_tree.column("name", width=220, anchor="w")
        self.startup_guard_tree.column("path", width=520, anchor="w")
        self.startup_guard_tree.pack(fill="both", expand=True)
        tree_scroll_y.configure(command=self.startup_guard_tree.yview)
        tree_scroll_x.configure(command=self.startup_guard_tree.xview)
        self.startup_guard_tree.bind(
            "<<TreeviewSelect>>",
            self.on_startup_guard_select,
        )

        self.startup_guard_details_text = tk.Text(
            details_frame,
            wrap="word",
            height=14,
        )
        self.startup_guard_details_text.pack(fill="both", expand=True)
        self.startup_guard_details_text.insert(
            "1.0",
            "Startup Guard Details erscheinen hier.\n",
        )
        self.startup_guard_details_text.configure(state="disabled")

    def open_startup_guard_tab(self):
        security_audit_controller.open_startup_guard_tab(self)

    def update_startup_guard_output(self, text: str, append: bool = False):
        security_audit_controller.update_startup_guard_output(
            self,
            text,
            append=append,
        )

    def clear_startup_guard_output(self):
        security_audit_controller.clear_startup_guard_output(self)

    def _format_startup_guard_entry(self, entry: dict) -> str:
        return security_audit_controller.format_startup_guard_entry(entry)

    def _render_startup_guard_result(self, result: dict, current_count: int) -> str:
        return security_audit_controller.render_startup_guard_result(
            self,
            result,
            current_count,
        )

    def _startup_guard_status_entry(self, status: str, entry) -> dict:
        return security_audit_controller.startup_guard_status_entry(status, entry)

    def show_startup_guard_result(self, result):
        security_audit_controller.show_startup_guard_result(self, result)

    def format_startup_guard_entry_details(self, entry: dict) -> str:
        return security_audit_controller.format_startup_guard_entry_details(entry)

    def on_startup_guard_select(self, _event=None):
        security_audit_controller.on_startup_guard_select(self, _event=_event)

    def analyze_selected_startup_entry_with_chad(self):
        security_audit_controller.analyze_selected_startup_entry_with_chad(self)

    def analyze_startup_guard_changes_with_chad(self):
        security_audit_controller.analyze_startup_guard_changes_with_chad(self)

    def format_path_guard_result(self, result: dict) -> str:
        return security_audit_controller.format_path_guard_result(result)

    def set_path_guard_light(self, state: str, message: str = ""):
        security_audit_controller.set_path_guard_light(
            self,
            state,
            message=message,
        )

    def run_path_guard_check(self):
        security_audit_controller.run_path_guard_check(self)

    def analyze_path_guard_with_chad(self):
        security_audit_controller.analyze_path_guard_with_chad(self)

    def startup_guard_mark_selected_trusted(self):
        security_audit_controller.startup_guard_mark_selected_trusted(self)

    def startup_guard_remove_selected_trusted(self):
        security_audit_controller.startup_guard_remove_selected_trusted(self)

    def startup_guard_show_trustlist(self):
        security_audit_controller.startup_guard_show_trustlist(self)

    def set_startup_guard_light(self, state: str, message: str = ""):
        security_audit_controller.set_startup_guard_light(
            self,
            state,
            message=message,
        )

    def startup_guard_startup_check(self):
        security_audit_controller.startup_guard_startup_check(self)

    def startup_guard_create_baseline(self):
        security_audit_controller.startup_guard_create_baseline(self)

    def startup_guard_check(self):
        security_audit_controller.startup_guard_check(self)

    # ---------- Security ----------
    def security_open_tab(self):
        self._select_security_tab(self.tab_security)

    def prepare_security_chad_analysis(self):
        summary = self.build_security_chad_summary()
        self.prepare_chad_analysis("Security Analyse", summary)

    def build_security_chad_summary(self) -> str:
        events = list(getattr(self, "security_events", []) or [])
        visible_events = list(getattr(self, "security_visible_events", []) or [])

        total = len(events)
        visible = len(visible_events)

        top_ip = "–"
        top_port = "–"
        block_count = 0
        portscan_count = 0

        if getattr(self, "security_store", None) is not None:
            try:
                total_stats, block_count, portscan_count, top_ip, top_port = (
                    self.security_store.compute_stats(events)
                )
                total = total_stats
            except Exception:
                pass

        relevant = []
        for event in events:
            level = str(event.get("level", "")).upper()
            if (
                "WARN" in level
                or "ALERT" in level
                or "CRITICAL" in level
                or event.get("type") == "Portscan-Verdacht"
            ):
                relevant.append(event)

        latest_relevant = relevant[-20:]
        latest_external = [e for e in events if e.get("external")][-20:]

        lines = []
        lines.append("NetMon Security-Zusammenfassung")
        lines.append("=" * 32)
        lines.append("")
        lines.append(f"Gesamtzahl Events: {total}")
        lines.append(f"Sichtbare Events im aktuellen Filter: {visible}")
        lines.append(f"Externe UFW BLOCK Events: {block_count}")
        lines.append(f"Portscan-Verdacht Events: {portscan_count}")
        lines.append(f"Top externe IP: {top_ip}")
        lines.append(f"Top Port: {top_port}")
        lines.append("")

        lines.append("Letzte relevante Events:")
        if latest_relevant:
            for event in latest_relevant:
                lines.append(
                    "- "
                    f"Zeit={event.get('time', '-')}; "
                    f"Typ={event.get('type', '-')}; "
                    f"Level={event.get('level', '-')}; "
                    f"Quelle={event.get('src', '-')}; "
                    f"Port={event.get('port', '-')}; "
                    f"Proto={event.get('proto', '-')}; "
                    f"Aktion={event.get('action', '-')}"
                )
        else:
            lines.append(
                "- Keine relevanten WARN/ALERT/CRITICAL/Portscan Events gefunden."
            )

        lines.append("")
        lines.append("Letzte externe Events:")
        if latest_external:
            for event in latest_external:
                lines.append(
                    "- "
                    f"Zeit={event.get('time', '-')}; "
                    f"Typ={event.get('type', '-')}; "
                    f"Level={event.get('level', '-')}; "
                    f"Quelle={event.get('src', '-')}; "
                    f"Port={event.get('port', '-')}; "
                    f"Proto={event.get('proto', '-')}; "
                    f"Aktion={event.get('action', '-')}"
                )
        else:
            lines.append("- Keine externen Events gefunden.")

        aggregated = getattr(self, "security_aggregated", {}) or {}
        if aggregated:
            lines.append("")
            lines.append("Aggregierte interne Events:")
            for key, value in sorted(aggregated.items()):
                lines.append(f"- {key}: {value}")

        return "\n".join(lines)

    def hash_file_sha256(self, path):
        try:
            import hashlib

            h = hashlib.sha256()
            with open(path, "rb") as f:
                for chunk in iter(lambda: f.read(8192), b""):
                    h.update(chunk)
            return h.hexdigest()
        except Exception:
            return ""

    def _integrity_watch_files(self):
        files = [
            self.vpn_config_path,
            self.primary_vpn_config_path,
            self.link_config_path,
            self.privacy_config_path,
            self.hdd_config_path,
            self.security_baseline_path,
            self.clamav_findings_path,
            os.path.abspath(__file__),
        ]

        wgdir = "/etc/wireguard"
        if os.path.isdir(wgdir):
            try:
                for entry in os.listdir(wgdir):
                    full = os.path.join(wgdir, entry)
                    if os.path.isfile(full):
                        files.append(full)
            except Exception:
                pass

        return files

    def build_integrity_baseline(self):
        baseline = {}
        for path in self._integrity_watch_files():
            if not os.path.isfile(path):
                continue
            try:
                st = os.stat(path)
                baseline[path] = {
                    "sha": self.hash_file_sha256(path),
                    "size": st.st_size,
                    "mtime": st.st_mtime,
                }
            except Exception:
                continue

        self.integrity_baseline = baseline
        return baseline

    def scan_integrity_changes(self):
        current = {}
        events = []

        for path in self._integrity_watch_files():
            if not os.path.isfile(path):
                continue
            try:
                st = os.stat(path)
            except Exception:
                continue
            current[path] = {
                "sha": self.hash_file_sha256(path),
                "size": st.st_size,
                "mtime": st.st_mtime,
            }

        old_paths = set(self.integrity_baseline.keys())
        new_paths = set(current.keys())

        for path in sorted(new_paths - old_paths):
            events.append(self._security_create_file_event("FILE NEW", path))

        for path in sorted(old_paths - new_paths):
            events.append(self._security_create_file_event("FILE DELETED", path))

        for path in sorted(old_paths & new_paths):
            old = self.integrity_baseline.get(path, {})
            cur = current.get(path, {})
            if not old or not cur:
                continue
            if (
                old.get("sha") != cur.get("sha")
                or old.get("size") != cur.get("size")
                or abs(old.get("mtime", 0) - cur.get("mtime", 0)) > 1e-6
            ):
                events.append(self._security_create_file_event("FILE MODIFIED", path))

        self.integrity_baseline = current
        return events

    def _device_watch_paths(self):
        devices = []
        input_dir = "/dev/input"
        if os.path.isdir(input_dir):
            try:
                for name in os.listdir(input_dir):
                    full = os.path.join(input_dir, name)
                    if os.path.exists(full):
                        devices.append(full)
            except Exception:
                pass

        sys_input = "/sys/class/input"
        if os.path.isdir(sys_input):
            try:
                for name in os.listdir(sys_input):
                    full = os.path.join(sys_input, name)
                    if os.path.islink(full):
                        devices.append(full)
            except Exception:
                pass

        return sorted(set(devices))

    def build_device_baseline(self):
        baseline = {}
        for path in self._device_watch_paths():
            baseline[path] = self._device_type_from_path(path)
        self.device_baseline = baseline
        return baseline

    def scan_device_changes(self):
        current = {
            path: self._device_type_from_path(path)
            for path in self._device_watch_paths()
        }
        old_paths = set(self.device_baseline.keys())
        new_paths = set(current.keys())
        events = []

        for path in sorted(new_paths - old_paths):
            events.append(
                self._security_create_device_event(
                    "DEVICE NEW", path, current.get(path, "unknown")
                )
            )

        for path in sorted(old_paths - new_paths):
            events.append(
                self._security_create_device_event(
                    "DEVICE REMOVED", path, self.device_baseline.get(path, "unknown")
                )
            )

        self.device_baseline = current
        return events

    def _device_type_from_path(self, path):
        name = os.path.basename(path).lower()
        if name.startswith("event"):
            return "input-event"
        if "mouse" in name:
            return "mouse"
        if "kbd" in name or "keyboard" in name:
            return "keyboard"
        return "hid"

    def _security_input_access_whitelisted(self, proc_name, pid):
        if not proc_name:
            return False
        if pid == os.getpid():
            return True
        if proc_name in self.security_input_access_whitelist:
            return True
        return False

    def _scan_input_access_state(self):
        state = {}
        input_prefix = "/dev/input"

        # Prefer psutil for process names and open files when available
        if psutil is not None:
            try:
                for proc in psutil.process_iter(["pid", "name"]):
                    try:
                        pid = proc.pid
                        pname = (proc.name() or "").strip()
                        for open_file in proc.open_files():
                            path = open_file.path
                            if path.startswith(input_prefix):
                                state.setdefault(path, {})[str(pid)] = pname
                    except (
                        psutil.NoSuchProcess,
                        psutil.AccessDenied,
                        psutil.ZombieProcess,
                    ):
                        continue
            except Exception as e:
                self.log(f"Security: psutil Input-Access-Scan fehlgeschlagen: {e}")
            else:
                if state:
                    return state

        # Fallback: /proc filesystem scan
        if os.path.isdir("/proc"):
            for entry in os.listdir("/proc"):
                if not entry.isdigit():
                    continue
                pid = entry
                fd_dir = os.path.join("/proc", pid, "fd")

                try:
                    pname = ""
                    with open(
                        os.path.join("/proc", pid, "comm"),
                        "r",
                        encoding="utf-8",
                        errors="ignore",
                    ) as f:
                        pname = f.read().strip()
                except Exception:
                    pname = ""

                if not os.path.isdir(fd_dir):
                    continue

                try:
                    fd_entries = os.listdir(fd_dir)
                except (PermissionError, OSError):
                    continue

                for fd in fd_entries:
                    try:
                        link = os.readlink(os.path.join(fd_dir, fd))
                    except OSError:
                        continue
                    if link.startswith(input_prefix):
                        state.setdefault(link, {})[pid] = pname
        return state

    def build_input_access_baseline(self):
        try:
            self.input_access_baseline = self._scan_input_access_state()
        except Exception as e:
            self.input_access_baseline = {}
            self.log(f"Security: Input-Access-Baseline fehlgeschlagen: {e}")
        return self.input_access_baseline

    def scan_input_access_changes(self):
        try:
            current = self._scan_input_access_state()
        except Exception as e:
            self.log(f"Security: Input-Access-Scan fehlgeschlagen: {e}")
            return []
        events = []

        old_paths = set(self.input_access_baseline.keys())
        new_paths = set(current.keys())

        for path in sorted(new_paths):
            old_procs = self.input_access_baseline.get(path, {})
            new_procs = current.get(path, {})
            for pid, pname in sorted(
                new_procs.items(), key=lambda kv: int(kv[0]) if kv[0].isdigit() else 0
            ):
                if pid in old_procs:
                    continue
                if self._security_input_access_whitelisted(
                    pname, int(pid) if pid.isdigit() else None
                ):
                    continue
                events.append(
                    self._security_create_input_access_event(
                        "INPUT ACCESS START", path, pid, pname
                    )
                )

        for path in sorted(old_paths - new_paths):
            old_procs = self.input_access_baseline.get(path, {})
            for pid, pname in sorted(
                old_procs.items(), key=lambda kv: int(kv[0]) if kv[0].isdigit() else 0
            ):
                if self._security_input_access_whitelisted(
                    pname, int(pid) if pid.isdigit() else None
                ):
                    continue
                events.append(
                    self._security_create_input_access_event(
                        "INPUT ACCESS STOP", path, pid, pname
                    )
                )

        # existing process stops for paths still present but pid removed
        for path in sorted(old_paths & new_paths):
            old_procs = self.input_access_baseline.get(path, {})
            new_procs = current.get(path, {})
            for pid, pname in sorted(
                old_procs.items(), key=lambda kv: int(kv[0]) if kv[0].isdigit() else 0
            ):
                if pid in new_procs:
                    continue
                if self._security_input_access_whitelisted(
                    pname, int(pid) if pid.isdigit() else None
                ):
                    continue
                events.append(
                    self._security_create_input_access_event(
                        "INPUT ACCESS STOP", path, pid, pname
                    )
                )

        self.input_access_baseline = current
        return events

    def _security_create_input_access_event(self, kind, path, pid, pname):
        level = "INFO"
        raw = f"Security: {kind} {path} pid={pid} name={pname}"
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        key = f"{kind}|{path}|{pid}"
        return {
            "time": now,
            "type": kind,
            "src": pname or "-",
            "port": pid,
            "proto": "-",
            "level": level,
            "action": "DETECT",
            "raw": raw,
            "key": key,
            "epoch": time.time(),
            "external": False,
        }

    def _security_create_file_event(self, kind, path):
        level = "INFO"
        if "/etc/wireguard" in path:
            level = "WARN"
        elif path.endswith("netmon.py") or os.path.basename(path) in {
            VPN_SLOT_CONFIG_FILE,
            PRIMARY_VPN_SLOT_CONFIG_FILE,
            LINK_SLOT_CONFIG_FILE,
            PRIVACY_CONFIG_FILE,
            HDD_SLOT_CONFIG_FILE,
        }:
            level = "WARN"

        raw = f"Security: {kind} {path}"
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        key = f"{kind}|{path}"
        return {
            "time": now,
            "type": kind,
            "src": "-",
            "port": "-",
            "proto": "-",
            "level": level,
            "action": "DETECT",
            "raw": raw,
            "key": key,
            "epoch": time.time(),
            "external": False,
        }

    def _security_create_device_event(self, kind, path, devtype):
        level = "INFO"
        if os.path.basename(path).startswith("event"):
            level = "WARN"

        raw = f"Security: {kind} {path} ({devtype})"
        now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        key = f"{kind}|{path}"
        return {
            "time": now,
            "type": kind,
            "src": "-",
            "port": "-",
            "proto": "-",
            "level": level,
            "action": "DETECT",
            "raw": raw,
            "key": key,
            "epoch": time.time(),
            "external": False,
        }
        rc, out = run_cmd(["sudo", "-n", "ufw", "status", "verbose"], timeout=10)
        if rc != 0:
            rc, out = run_cmd(["ufw", "status", "verbose"], timeout=10)
        if rc != 0:
            return "–"
        m = re.search(r"Logging:\s*([A-Za-z0-9_/-]+)", out or "")
        if not m:
            return "–"
        return m.group(1).upper()

    def _security_update_status(self, color):
        security_monitoring_controller.security_update_status(self, color)

    def _read_ufw_log_tail(self, lines=1200):
        return security_monitoring_controller.read_ufw_log_tail(lines=lines)

    def _parse_ufw_block_line(self, line):
        return security_monitoring_controller.parse_ufw_block_line(
            self,
            line,
            is_local_or_lan_ip_fn=is_local_or_lan_ip,
        )

    def _security_detect_portscan(self):
        store = self._security_store()
        if store is not None:
            return store.detect_portscan(
                self.security_events, self.security_last_portscan_alert
            )

        now = time.time()
        recent = [
            e
            for e in self.security_events
            if now - e.get("epoch", now) <= 90
            and e.get("external")
            and e.get("type") == "UFW BLOCK"
        ]
        by_ip = {}
        for e in recent:
            ip = e.get("src") or "-"
            entry = by_ip.setdefault(
                ip, {"count": 0, "ports": set(), "latest": e.get("time", "-")}
            )
            entry["count"] += 1
            p = e.get("port")
            if p and p != "-":
                entry["ports"].add(p)
            entry["latest"] = e.get("time", entry["latest"])

        def severity_for(count, unique_ports):
            if count >= 40 or unique_ports >= 25:
                return 5
            if count >= 25 or unique_ports >= 15:
                return 4
            if count >= 12 or unique_ports >= 8:
                return 3
            if count >= 6 or unique_ports >= 4:
                return 2
            if count >= 1:
                return 1
            return 0

        def severity_label(level):
            return {
                1: "INFO",
                2: "WARN",
                3: "ALERT",
                4: "CRITICAL",
                5: "CRITICAL",
            }.get(level, "INFO")

        findings = []
        for ip, info in by_ip.items():
            sev = severity_for(info["count"], len(info["ports"]))
            if sev < 2:
                continue
            last_alert = self.security_last_portscan_alert.get(ip, 0.0)
            if now - last_alert < 120:
                continue
            self.security_last_portscan_alert[ip] = now
            sev_label = severity_label(sev)
            findings.append(
                {
                    "time": info["latest"],
                    "type": "Portscan-Verdacht",
                    "src": ip,
                    "port": str(len(info["ports"])),
                    "proto": "-",
                    "level": f"L{sev} {sev_label}",
                    "action": "SUSPECT",
                    "raw": (
                        f"Portscan-Verdacht: SRC={ip}, Level=L{sev} {sev_label}, "
                        f"Events={info['count']}, verschiedene Ports={len(info['ports'])}"
                    ),
                    "key": f"portscan|{ip}|L{sev}|{int(now // 120)}",
                    "epoch": now,
                    "external": True,
                    "severity": sev,
                }
            )
        return findings

    def _security_filtered_events(self):
        return security_monitoring_controller.security_filtered_events(self)

    def _security_compute_stats(self):
        return security_monitoring_controller.security_compute_stats(self)

    def _security_update_stats_view(self):
        security_monitoring_controller.security_update_stats_view(self)

    def _refresh_security_table(self):
        security_monitoring_controller.refresh_security_table(self)

    def on_security_event_select(self, _event=None):
        security_monitoring_controller.on_security_event_select(self, _event=_event)

    def security_refresh_table_action(self):
        security_monitoring_controller.security_refresh_table_action(self)

    def security_filter_portscan(self):
        security_monitoring_controller.security_filter_portscan(self)

    def security_filter_relevant(self):
        security_monitoring_controller.security_filter_relevant(self)

    def security_filter_external(self):
        security_monitoring_controller.security_filter_external(self)

    def security_filter_all(self):
        security_monitoring_controller.security_filter_all(self)

    def security_clear_events(self, confirm=False):
        security_monitoring_controller.security_clear_events(self, confirm=confirm)

    def _security_export_dir(self):
        return security_monitoring_controller.security_export_dir()

    def export_security_logs(self):
        security_monitoring_controller.export_security_logs(self)

    def _security_poll_async(self):
        security_monitoring_controller.security_poll_async(self)

    def _security_query_log_level(self):
        # Ermittelt das aktuelle UFW-Logging-Level (wenn verfügbar), sonst fallback.
        # Rückgabe: KANN HIGH/MEDIUM/LOW/NONE/– sein.
        cmd = ["sudo", "-n", "ufw", "status", "verbose"]
        for attempt in (cmd, ["ufw", "status", "verbose"]):
            try:
                rc, out = run_cmd(attempt, timeout=12)
            except Exception as e:
                self.log(
                    f"Security: UFW-Level-Abfrage fehlgeschlagen (cmd={' '.join(attempt)}): {e}"
                )
                continue
            if rc != 0 or not out:
                continue
            # präzises Zeilenmatching
            for line in str(out).splitlines():
                m = re.search(r"Logging:\s*([A-Za-z0-9_-]+)", line, re.IGNORECASE)
                if m:
                    return m.group(1).upper()
            # alternative: UFW kann in anderen Formaten liefern
            m2 = re.search(r"UFW logging:\s*([A-Za-z0-9_-]+)", str(out), re.IGNORECASE)
            if m2:
                return m2.group(1).upper()
        return "–"

    def security_start_monitoring(self):
        security_monitoring_controller.security_start_monitoring(self)

    def security_stop_monitoring(self):
        security_monitoring_controller.security_stop_monitoring(self)

    def security_build_baseline(self):
        self.build_integrity_baseline()
        self.build_device_baseline()
        self.build_input_access_baseline()
        self._save_security_baseline()
        self.log(
            self.tr(
                "security_log_baseline_rebuilt",
                "Security: Baseline (File+Device+Input) neu erstellt.",
            )
        )
        self._refresh_security_table()

    def _load_security_baseline(self):
        try:
            if os.path.isfile(self.security_baseline_path):
                with open(self.security_baseline_path, "r", encoding="utf-8") as f:
                    data = json.load(f)
                self.integrity_baseline = data.get("integrity", {})
                self.device_baseline = data.get("device", {})
                self.input_access_baseline = data.get("input_access", {})
                return True
        except Exception:
            pass
        return False

    def _save_security_baseline(self):
        payload = {
            "integrity": self.integrity_baseline,
            "device": self.device_baseline,
            "input_access": self.input_access_baseline,
            "ts": datetime.now().isoformat(),
        }
        ok, err = safe_write_json(
            self.security_baseline_path,
            payload,
            indent=2,
            ensure_ascii=False,
        )
        if not ok:
            self.log_debug(f"Security-Baseline speichern fehlgeschlagen: {err}")
        return ok

    def security_scan_now(self):
        security_monitoring_controller.security_scan_now(self)

    def security_input_access_scan(self):
        security_monitoring_controller.security_input_access_scan(self)

    def update_clamav_panel(
        self,
        status=None,
        signatures=None,
        last_scan=None,
        last_target=None,
        findings=None,
        result=None,
    ):
        clamav_controller.update_clamav_panel(
            self,
            status=status,
            signatures=signatures,
            last_scan=last_scan,
            last_target=last_target,
            findings=findings,
            result=result,
        )

    def update_clamav_panel_safe(
        self,
        status=None,
        signatures=None,
        last_scan=None,
        last_target=None,
        findings=None,
        result=None,
    ):
        clamav_controller.update_clamav_panel_safe(
            self,
            status=status,
            signatures=signatures,
            last_scan=last_scan,
            last_target=last_target,
            findings=findings,
            result=result,
        )

    def _clamav_findings_open(self):
        return clamav_controller.clamav_findings_open(self)

    def _clamav_findings_quarantined(self):
        return clamav_controller.clamav_findings_quarantined(self)

    def _update_clamav_findings_buttons_state(self):
        clamav_controller.update_clamav_findings_buttons_state(self)

    def _update_clamav_findings_buttons_state_safe(self):
        clamav_controller.update_clamav_findings_buttons_state_safe(self)

    def _write_clamav_output(self, lines, replace=True):
        clamav_controller.write_clamav_output(self, lines, replace=replace)

    def _format_runtime_hms(self, seconds):
        return clamav_controller.format_runtime_hms(seconds)

    def _clamav_progress_apply(self):
        clamav_controller.clamav_progress_apply(self)

    def _clamav_progress_tick(self):
        clamav_controller.clamav_progress_tick(self)

    def _clamav_progress_start(self, expected_files=None):
        clamav_controller.clamav_progress_start(self, expected_files=expected_files)

    def _clamav_progress_update_safe(self, scanned_files):
        clamav_controller.clamav_progress_update_safe(self, scanned_files)

    def _clamav_progress_finish_safe(self):
        clamav_controller.clamav_progress_finish_safe(self)

    def _clamav_is_excluded_root(self, path, excluded_roots):
        real = os.path.realpath(path)
        for root in excluded_roots or []:
            rr = os.path.realpath(root)
            if real == rr or real.startswith(rr + os.sep):
                return True
        return False

    def _clamav_system_targets(self):
        home = os.path.expanduser("~")
        candidates = [home, "/etc", "/opt", "/usr/local"]
        targets = []
        seen = set()
        for raw in candidates:
            p = os.path.realpath(os.path.normpath(raw))
            if p in seen:
                continue
            seen.add(p)
            if os.path.isdir(p):
                targets.append(p)
        return targets

    def _clamav_system_excluded_roots(self):
        return [
            "/proc",
            "/sys",
            "/dev",
            "/run",
            "/tmp",
            "/var/tmp",
        ]

    def _clamav_count_files(self, targets, excluded_roots=None, limit=600000):
        total = 0
        excluded_roots = list(excluded_roots or [])
        for target in targets or []:
            if not os.path.isdir(target):
                continue
            for root, dirs, files in os.walk(target, topdown=True):
                if self._clamav_is_excluded_root(root, excluded_roots):
                    dirs[:] = []
                    continue
                dirs[:] = [
                    d
                    for d in dirs
                    if not self._clamav_is_excluded_root(
                        os.path.join(root, d), excluded_roots
                    )
                ]
                total += len(files)
                if total >= limit:
                    return total
        return total

    def _clamav_parse_scanned_files(self, output):
        return clamav_controller.clamav_parse_scanned_files(output)

    def _run_clamscan_with_progress(self, cmd, timeout=3600, progress_cb=None):
        proc = None
        cancelled = False
        timed_out = False
        out_lines = []
        scanned = 0

        self.clamav_cancel_requested = False
        self._set_clamav_abort_button_state(True)
        try:
            proc = subprocess.Popen(
                cmd,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                bufsize=1,
                start_new_session=True,
            )
            self.clamav_active_process = proc
            started = time.time()

            while True:
                if self.clamav_cancel_requested and proc.poll() is None:
                    cancelled = True
                    try:
                        os.killpg(proc.pid, signal.SIGTERM)
                    except Exception:
                        try:
                            proc.terminate()
                        except Exception:
                            pass

                if (time.time() - started) > float(timeout) and proc.poll() is None:
                    timed_out = True
                    try:
                        os.killpg(proc.pid, signal.SIGTERM)
                    except Exception:
                        try:
                            proc.terminate()
                        except Exception:
                            pass

                line = proc.stdout.readline() if proc.stdout is not None else ""
                if line:
                    out_lines.append(line)
                    t = line.strip()
                    if t.endswith(" OK") or t.endswith(" FOUND"):
                        scanned += 1
                        if callable(progress_cb):
                            progress_cb(scanned)

                if proc.poll() is not None:
                    # rest output
                    try:
                        rest = proc.stdout.read() if proc.stdout is not None else ""
                    except Exception:
                        rest = ""
                    if rest:
                        out_lines.append(rest)
                    break

                if not line:
                    time.sleep(0.05)

            rc = proc.returncode if proc.returncode is not None else 1
            if timed_out:
                rc = 124
            elif cancelled and rc == 0:
                rc = -15
            return rc, "".join(out_lines), cancelled, timed_out, scanned
        finally:
            self.clamav_active_process = None
            self._set_clamav_abort_button_state(False)
            self.clamav_cancel_requested = False

    def _clamav_select_items_dialog(self, items, title, mode="found"):
        if not items:
            return None

        win = tk.Toplevel(self)
        win.title(title)
        win.transient(self)
        win.grab_set()
        win.resizable(True, True)
        win.minsize(760, 360)

        container = ttk.Frame(win, padding=10)
        container.pack(fill="both", expand=True)

        hint = (
            "Mehrfachauswahl mit Strg/Shift möglich."
            if mode == "found"
            else "Wähle Quarantäne-Einträge zur Wiederherstellung aus."
        )
        ttk.Label(container, text=hint).pack(anchor="w", pady=(0, 8))

        list_wrap = ttk.Frame(container)
        list_wrap.pack(fill="both", expand=True)
        list_wrap.columnconfigure(0, weight=1)
        list_wrap.rowconfigure(0, weight=1)

        listbox = tk.Listbox(
            list_wrap,
            selectmode=tk.EXTENDED,
            exportselection=False,
            font=("Courier", 10),
        )
        listbox.grid(row=0, column=0, sticky="nsew")
        scroll = ttk.Scrollbar(list_wrap, orient="vertical", command=listbox.yview)
        scroll.grid(row=0, column=1, sticky="ns")
        listbox.configure(yscrollcommand=scroll.set)

        for idx, item in enumerate(items, start=1):
            src = (item.get("path") or "-").strip() or "-"
            sig = (item.get("signature") or "-").strip() or "-"
            marker = ""
            if mode == "quarantine":
                qpath = (item.get("quarantine_path") or "").strip()
                marker = f" | Q: {qpath}" if qpath else ""
            line = f"{idx:>2}. {src} | Sig: {sig}{marker}"
            if len(line) > 240:
                line = line[:237] + "..."
            listbox.insert("end", line)

        result = {"selection": None}

        def on_select_all():
            listbox.selection_set(0, "end")

        def on_clear():
            listbox.selection_clear(0, "end")

        def on_cancel():
            result["selection"] = None
            win.destroy()

        def on_ok():
            selected_idx = listbox.curselection()
            if not selected_idx:
                result["selection"] = []
            else:
                result["selection"] = [items[i] for i in selected_idx if i < len(items)]
            win.destroy()

        btns = ttk.Frame(container)
        btns.pack(fill="x", pady=(10, 0))
        ttk.Button(btns, text="Alle", command=on_select_all).pack(side="left")
        ttk.Button(btns, text="Keine", command=on_clear).pack(side="left", padx=(8, 0))
        ttk.Button(btns, text="Abbrechen", command=on_cancel).pack(side="right")
        ttk.Button(btns, text="Übernehmen", command=on_ok).pack(
            side="right", padx=(0, 8)
        )

        win.protocol("WM_DELETE_WINDOW", on_cancel)
        win.wait_window()
        return result["selection"]

    def _set_clamav_abort_button_state(self, enabled):
        btn = getattr(self, "btn_clamav_abort_scan", None)
        if btn is None:
            return
        state = "normal" if enabled else "disabled"
        if threading.current_thread() is threading.main_thread():
            try:
                btn.configure(state=state)
            except Exception:
                pass
            return
        self.after(0, lambda s=state: btn.configure(state=s))

    def _run_clamscan_with_cancel(self, cmd, timeout=900):
        proc = None
        cancelled = False
        out = ""
        rc = 1
        self.clamav_cancel_requested = False
        self._set_clamav_abort_button_state(True)
        try:
            proc = subprocess.Popen(
                cmd,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                start_new_session=True,
            )
            self.clamav_active_process = proc

            while True:
                if self.clamav_cancel_requested and proc.poll() is None:
                    cancelled = True
                    try:
                        os.killpg(proc.pid, signal.SIGTERM)
                    except Exception:
                        try:
                            proc.terminate()
                        except Exception:
                            pass
                try:
                    out, _ = proc.communicate(timeout=0.25)
                    rc = proc.returncode if proc.returncode is not None else 1
                    break
                except subprocess.TimeoutExpired:
                    if self.clamav_cancel_requested:
                        cancelled = True
                        continue

            if cancelled and rc not in (0,):
                rc = -15
            return rc, out or "", cancelled
        finally:
            self.clamav_active_process = None
            self._set_clamav_abort_button_state(False)
            self.clamav_cancel_requested = False

    def clamav_abort_scan(self):
        clamav_controller.clamav_abort_scan(self)

    def export_clamav_findings(self):
        clamav_controller.export_clamav_findings(self)

    def _load_clamav_findings(self):
        clamav_controller.load_clamav_findings(self)

    def _save_clamav_findings(self):
        clamav_controller.save_clamav_findings(self)

    def _clamav_extract_findings(self, output, scan_type, timestamp):
        return clamav_controller.clamav_extract_findings(
            output,
            scan_type,
            timestamp,
        )

    def _clamav_replace_open_findings(self, findings):
        clamav_controller.clamav_replace_open_findings(self, findings)

    def _clamav_next_quarantine_path(self, source_path):
        base = os.path.basename(source_path.rstrip(os.sep)) or "file"
        stem, ext = os.path.splitext(base)
        stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
        candidate = os.path.join(
            self.clamav_quarantine_dir,
            f"{stem}__quarantine__{stamp}{ext}",
        )
        idx = 1
        while os.path.exists(candidate):
            candidate = os.path.join(
                self.clamav_quarantine_dir,
                f"{stem}__quarantine__{stamp}_{idx}{ext}",
            )
            idx += 1
        return candidate

    def _clamav_next_restore_path(self, original_path):
        target = os.path.normpath(original_path)
        if not os.path.exists(target):
            return target

        base = os.path.basename(target.rstrip(os.sep)) or "file"
        stem, ext = os.path.splitext(base)
        stamp = datetime.now().strftime("%Y-%m-%d_%H%M%S")
        parent = os.path.dirname(target) or os.path.expanduser("~")
        candidate = os.path.join(parent, f"{stem}__restored__{stamp}{ext}")
        idx = 1
        while os.path.exists(candidate):
            candidate = os.path.join(
                parent,
                f"{stem}__restored__{stamp}_{idx}{ext}",
            )
            idx += 1
        return candidate

    def _clamav_path_is_within(self, candidate_path, root_path):
        try:
            candidate_real = os.path.realpath(candidate_path)
            root_real = os.path.realpath(root_path)
            return os.path.commonpath([candidate_real, root_real]) == root_real
        except Exception:
            return False

    def _clamav_validate_source_for_quarantine(self, source_path):
        src = os.path.realpath(os.path.normpath((source_path or "").strip()))
        if not src:
            return None, "leer"
        if not os.path.isabs(src):
            return None, "nicht absolut"
        if not os.path.exists(src):
            return None, "nicht vorhanden"
        if os.path.isdir(src):
            return None, "ist ein Ordner"
        home_root = os.path.realpath(os.path.expanduser("~"))
        if not self._clamav_path_is_within(src, home_root):
            return None, "außerhalb des Home-Verzeichnisses"
        return src, ""

    def _clamav_validate_quarantine_path(self, quarantine_path):
        qpath = os.path.realpath(os.path.normpath((quarantine_path or "").strip()))
        if not qpath:
            return None, "leer"
        if not os.path.isabs(qpath):
            return None, "nicht absolut"
        if not self._clamav_path_is_within(qpath, self.clamav_quarantine_dir):
            return None, "außerhalb des Quarantäne-Ordners"
        return qpath, ""

    def _clamav_safe_restore_source(self, original_path):
        original = os.path.realpath(os.path.normpath((original_path or "").strip()))
        home_root = os.path.realpath(os.path.expanduser("~"))
        if original and self._clamav_path_is_within(original, home_root):
            return original
        base = os.path.basename((original_path or "").strip()) or "restored_file"
        return os.path.join(home_root, base)

    def _clamav_quick_targets(self):
        home = os.path.expanduser("~")
        candidates = [
            os.path.join(home, "Downloads"),
            os.path.join(home, "Desktop"),
            os.path.join(home, "Dokumente"),
            os.path.join(home, "Documents"),
        ]
        seen = set()
        targets = []
        for path in candidates:
            norm = os.path.normpath(path)
            if norm in seen:
                continue
            seen.add(norm)
            if os.path.isdir(norm):
                targets.append(norm)
        if targets:
            return targets, False
        return [home], True

    def _clamav_extract_error_message(self, output):
        return clamav_controller.clamav_extract_error_message(output)

    def _clamav_is_nonfatal_scan_issue(self, output):
        return clamav_controller.clamav_is_nonfatal_scan_issue(output)

    def _clamav_is_nonfatal_update_issue(self, output):
        return clamav_controller.clamav_is_nonfatal_update_issue(output)

    def _clamav_parse_infected_count(self, output):
        return clamav_controller.clamav_parse_infected_count(output)

    def security_clamav_update_async(self):
        clamav_controller.security_clamav_update_async(self)

    def security_clamav_quickscan_async(self):
        clamav_controller.security_clamav_quickscan_async(self)

    def security_clamav_folder_scan_async(self):
        clamav_controller.security_clamav_folder_scan_async(self)

    def security_clamav_systemscan_async(self):
        clamav_controller.security_clamav_systemscan_async(self)

    def _log_clamav_quarantine_preview(self, items, label="Trockenlauf"):
        clamav_controller.log_clamav_quarantine_preview(self, items, label=label)

    def _log_clamav_restore_preview(self, items, label="Trockenlauf"):
        clamav_controller.log_clamav_restore_preview(self, items, label=label)

    def show_clamav_findings(self):
        clamav_controller.show_clamav_findings(self)

    def show_clamav_quarantine_history(self):
        clamav_controller.show_clamav_quarantine_history(self)

    def quarantine_clamav_findings_async(self):
        clamav_controller.quarantine_clamav_findings_async(self)

    def quarantine_selected_clamav_findings_async(self):
        clamav_controller.quarantine_selected_clamav_findings_async(self)

    def restore_clamav_quarantine_async(self):
        clamav_controller.restore_clamav_quarantine_async(self)

    def restore_selected_clamav_quarantine_async(self):
        clamav_controller.restore_selected_clamav_quarantine_async(self)

    def security_clamav_scan_async(self):
        clamav_controller.security_clamav_scan_async(self)

    def security_set_ufw_logging_high(self):
        firewall_controller.security_set_ufw_logging_high(self)

    def ensure_sudo_session(self, reason=""):
        if not has_cmd("sudo"):
            self.log("Sudo-Auth: 'sudo' nicht gefunden.")
            return False
        if check_sudo_nopasswd():
            return True
        ok = renew_sudo_timestamp(self)
        if not ok:
            if reason:
                self.log(f"Sudo-Auth abgebrochen: {reason}.")
            else:
                self.log("Sudo-Auth abgebrochen.")
            return False
        if reason:
            self.log(f"Sudo-Auth erneuert für: {reason}.")
        else:
            self.log("Sudo-Auth erneuert.")
        return True

    def setup_sudo_auth(self):
        if not has_cmd("sudo"):
            self.log("Sudo-Auth: 'sudo' nicht gefunden.")
            return
        if check_sudo_nopasswd():
            self.log("Sudo-Auth: sudo-Session bereits gültig (sudo -n true).")
            return
        if renew_sudo_timestamp(self):
            self.log("Sudo-Auth: Session erfolgreich freigeschaltet (Timestamp aktiv).")
            return
        self.log(
            "Sudo-Auth fehlgeschlagen. Rechtepflichtige Funktionen können fehlschlagen."
        )

    def _confirm_danger_action(
        self,
        action_key,
        *,
        title,
        prompt,
        cooldown_s=None,
        skip_guard=False,
        require_phrase="",
        force_guard=False,
    ):
        if skip_guard:
            return True
        guard_enabled = self.__dict__.get("danger_guard_enabled", True)
        if not bool(guard_enabled) and not bool(force_guard):
            return True

        now = time.time()
        cooldown = (
            self.__dict__.get("danger_confirm_cooldown_s", 15)
            if cooldown_s is None
            else max(0, int(cooldown_s))
        )
        danger_history = self.__dict__.get("_danger_last_confirm")
        if not isinstance(danger_history, dict):
            danger_history = {}
            self._danger_last_confirm = danger_history
        last = float(danger_history.get(action_key, 0.0))
        if cooldown > 0 and now - last <= cooldown:
            return True

        parent_window = self if self.__dict__.get("tk") is not None else None
        if parent_window is None:
            return True

        ok = messagebox.askyesno(
            title,
            prompt,
            parent=self,
        )
        if not ok:
            return False

        phrase = (require_phrase or "").strip()
        if phrase:
            entered = simpledialog.askstring(
                title,
                f'Bitte zur Bestätigung exakt "{phrase}" eingeben:',
                parent=self,
            )
            if entered is None:
                return False
            if (entered or "").strip() != phrase:
                messagebox.showwarning(
                    title,
                    "Bestätigungswort stimmt nicht. Aktion wurde abgebrochen.",
                    parent=self,
                )
                return False

        danger_history[action_key] = now
        return True

    def execute_selected_system_action(self):
        action = (self.system_action_var.get() or "").strip()
        prompts = {
            "Reboot": "Möchtest du den Rechner wirklich neu starten?",
            "Abmelden": "Möchtest du dich wirklich abmelden?",
            "Ausschalten": "Möchtest du den Rechner wirklich ausschalten?",
            "Ausschalten - 15 Min": "Möchtest du den Rechner in 15 Minuten ausschalten?",
            "Ausschalten - 30 Min": "Möchtest du den Rechner in 30 Minuten ausschalten?",
            "Ausschalten - 45 Min": "Möchtest du den Rechner in 45 Minuten ausschalten?",
            "Ausschalten - 60 Min": "Möchtest du den Rechner in 60 Minuten ausschalten?",
            "Shutdown abbrechen": "Möchtest du einen aktiven Shutdown-Timer wirklich abbrechen?",
        }
        if action not in prompts:
            self.log("Bitte zuerst eine Systemaktion auswählen.")
            return
        self.run_system_action_async(action)

    def _system_action_commands(self, action):
        shutdown_timers = {
            "Ausschalten - 15 Min": 15,
            "Ausschalten - 30 Min": 30,
            "Ausschalten - 45 Min": 45,
            "Ausschalten - 60 Min": 60,
        }
        if action == "Reboot":
            return [(["systemctl", "reboot"], True)]
        if action == "Ausschalten":
            return [(["systemctl", "poweroff"], True)]
        if action in shutdown_timers:
            minutes = shutdown_timers[action]
            return [(["shutdown", "-h", f"+{minutes}"], True)]
        if action == "Shutdown abbrechen":
            return [(["shutdown", "-c"], True)]
        if action == "Abmelden":
            commands = []
            if has_cmd("gnome-session-quit"):
                commands.append(
                    (["gnome-session-quit", "--logout", "--no-prompt"], False)
                )
            session_id = (os.environ.get("XDG_SESSION_ID") or "").strip()
            if session_id and has_cmd("loginctl"):
                commands.append((["loginctl", "terminate-session", session_id], False))
            if has_cmd("loginctl"):
                commands.append(
                    (["loginctl", "terminate-user", getpass.getuser()], True)
                )
            return commands
        return []

    def run_system_action_async(self, action, skip_guard=False):
        commands = self._system_action_commands(action)
        if not commands:
            self.log(f"Systemaktion nicht verfügbar: {action}.")
            return

        system_prompts = {
            "Reboot": "Diese Aktion startet das System sofort neu.",
            "Abmelden": "Diese Aktion meldet die aktuelle Sitzung ab.",
            "Ausschalten": "Diese Aktion fährt das System sofort herunter.",
            "Ausschalten - 15 Min": "Diese Aktion plant ein Herunterfahren in 15 Minuten.",
            "Ausschalten - 30 Min": "Diese Aktion plant ein Herunterfahren in 30 Minuten.",
            "Ausschalten - 45 Min": "Diese Aktion plant ein Herunterfahren in 45 Minuten.",
            "Ausschalten - 60 Min": "Diese Aktion plant ein Herunterfahren in 60 Minuten.",
            "Shutdown abbrechen": "Diese Aktion bricht einen aktiven Shutdown-Timer ab.",
        }
        prompt = system_prompts.get(action, "Diese Aktion verändert den Systemzustand.")
        if not self._confirm_danger_action(
            f"system:{action}",
            title="Systemaktion bestätigen",
            prompt=f"{prompt}\n\nWirklich fortfahren?",
            cooldown_s=0,
            skip_guard=skip_guard,
            require_phrase="AUSFUEHREN",
            force_guard=True,
        ):
            self.log(f"Systemaktion abgebrochen: {action}.")
            return

        if any(needs_sudo for _cmd, needs_sudo in commands):
            if not self.ensure_sudo_session(f"Systemaktion {action}"):
                return

        self.log(f"Systemaktion gestartet: {action}.")

        def job():
            last_out = ""
            for cmd, needs_sudo in commands:
                exec_cmd = ["sudo", "-n", *cmd] if needs_sudo else cmd
                rc, out = run_cmd(exec_cmd, timeout=20)
                last_out = out or last_out
                if rc == 0:
                    self.after(
                        0, lambda a=action: self.log(f"Systemaktion ausgelöst: {a}.")
                    )
                    return
            msg = last_out.strip() or "Unbekannter Fehler"
            self.after(
                0,
                lambda a=action, m=msg: self.log(
                    f"Systemaktion fehlgeschlagen ({a}): {m}"
                ),
            )

        threading.Thread(target=job, daemon=True).start()

    def register_feature_button(self, feature, widget):
        self.feature_buttons.setdefault(feature, []).append(widget)

    def set_feature_enabled(self, feature, enabled):
        state = "normal" if enabled else "disabled"
        for widget in self.feature_buttons.get(feature, []):
            try:
                widget.configure(state=state)
            except Exception:
                pass

    def _find_vscode_desktop_file(self):
        env_desktop = os.environ.get("VSCODE_DESKTOP_FILE", "").strip()
        if env_desktop and os.path.isfile(env_desktop):
            return env_desktop

        search_dirs = [
            os.path.expanduser("~/.local/share/applications"),
            "/usr/local/share/applications",
            "/usr/share/applications",
        ]
        for base in search_dirs:
            for name in VSCODE_DESKTOP_FILES:
                p = os.path.join(base, name)
                if os.path.isfile(p):
                    return p
        return None

    def _desktop_exec_command(self, desktop_file):
        if not desktop_file or not os.path.isfile(desktop_file):
            return None
        try:
            with open(desktop_file, "r", encoding="utf-8") as f:
                for line in f:
                    line = line.strip()
                    if line.startswith("Exec="):
                        exec_line = line.split("=", 1)[1].strip()
                        if not exec_line:
                            return None
                        parts = shlex.split(exec_line)
                        cleaned = []
                        for p in parts:
                            if p.startswith("%"):
                                continue
                            p = re.sub(r"%[fFuUdDnNickvm]", "", p)
                            if p:
                                cleaned.append(p)
                        return cleaned or None
        except Exception:
            return None
        return None

    def _desktop_search_dirs(self):
        # User + System + Flatpak Export-Pfade (für bessere Menü-Abdeckung).
        return [
            os.path.expanduser("~/.local/share/applications"),
            "/usr/share/applications",
            os.path.expanduser("~/.local/share/flatpak/exports/share/applications"),
            "/var/lib/flatpak/exports/share/applications",
        ]

    def _desktop_language_keys(self):
        lang = (
            os.environ.get("LC_MESSAGES")
            or os.environ.get("LANG")
            or os.environ.get("LANGUAGE")
            or ""
        ).split(":", 1)[0]
        lang = lang.split(".", 1)[0].strip()
        if not lang:
            return []
        keys = [lang]
        if "_" in lang:
            keys.append(lang.split("_", 1)[0])
        return keys

    def _parse_desktop_entry(self, desktop_file):
        if not desktop_file or not os.path.isfile(desktop_file):
            return None
        data = {}
        in_entry = False
        try:
            with open(desktop_file, "r", encoding="utf-8") as f:
                for raw_line in f:
                    line = raw_line.strip()
                    if not line or line.startswith("#"):
                        continue
                    if line.startswith("["):
                        in_entry = line == "[Desktop Entry]"
                        continue
                    if not in_entry or "=" not in line:
                        continue
                    key, value = line.split("=", 1)
                    data[key.strip()] = value.strip()
        except Exception:
            return None

        if (data.get("Type") or "Application").strip() != "Application":
            return None
        if _desktop_bool(data.get("Hidden")) or _desktop_bool(data.get("NoDisplay")):
            return None

        name = ""
        for key in self._desktop_language_keys():
            localized = data.get(f"Name[{key}]")
            if localized:
                name = localized.strip()
                break
        if not name:
            name = (data.get("Name") or "").strip()
        if not name:
            return None

        exec_line = (data.get("Exec") or "").strip()
        exec_cmd = None
        if exec_line:
            try:
                parts = shlex.split(exec_line)
            except Exception:
                parts = exec_line.split()
            cleaned = []
            for part in parts:
                if part.startswith("%"):
                    continue
                part = re.sub(r"%[fFuUdDnNickvm]", "", part)
                if part:
                    cleaned.append(part)
            exec_cmd = cleaned or None

        desktop_id = os.path.basename(desktop_file)
        return {
            "name": name,
            "desktop_id": desktop_id,
            "desktop_file": desktop_file,
            "exec": exec_cmd,
            "icon": (data.get("Icon") or "").strip(),
        }

    def _load_desktop_app_entries(self):
        entries_by_id = {}
        for base in self._desktop_search_dirs():
            if not os.path.isdir(base):
                continue
            try:
                names = sorted(os.listdir(base), key=str.casefold)
            except Exception:
                continue
            for name in names:
                if not name.endswith(".desktop"):
                    continue
                desktop_file = os.path.join(base, name)
                entry = self._parse_desktop_entry(desktop_file)
                if not entry:
                    continue
                # Dedup über desktop_id: erster Treffer gewinnt (User-Pfad zuerst).
                did = (entry.get("desktop_id") or "").strip().lower()
                if did and did not in entries_by_id:
                    entries_by_id[did] = entry

        entries = list(entries_by_id.values())
        name_counts = {}
        for entry in entries:
            name_counts[entry["name"]] = name_counts.get(entry["name"], 0) + 1
        for entry in entries:
            label = entry["name"]
            if name_counts[label] > 1:
                label = f"{label} ({entry['desktop_id']})"
            entry["label"] = label
        entries.sort(key=lambda item: item["label"].casefold())
        return entries

    def refresh_desktop_app_entries(self):
        previous = (
            (self.link_system_app_var.get() or "").strip()
            if hasattr(self, "link_system_app_var")
            else ""
        )
        entries = self._load_desktop_app_entries()
        self.desktop_app_entries = entries
        self.desktop_app_labels = [entry["label"] for entry in entries]
        self.desktop_app_by_label = {entry["label"]: entry for entry in entries}
        self.desktop_app_by_id = {
            entry["desktop_id"]: entry for entry in entries if entry.get("desktop_id")
        }

        combo = getattr(self, "link_system_combo", None)
        if combo is not None:
            try:
                combo.configure(values=self.desktop_app_labels)
            except Exception:
                pass

        if previous and previous in self.desktop_app_by_label:
            self.link_system_app_var.set(previous)
        elif self.desktop_app_labels:
            self.link_system_app_var.set(self.desktop_app_labels[0])
            self.on_systemlink_app_select()
        else:
            self.link_system_app_var.set("")

        self.log(
            f"Programmliste aktualisiert: {len(self.desktop_app_labels)} Einträge."
        )

    def _default_link_slots(self):
        return [
            {
                "slot": i + 1,
                "custom_name": "",
                "link_type": "",
                "target": "",
                "app_name": "",
                "desktop_id": "",
                "desktop_file": "",
            }
            for i in range(LINK_SLOT_COUNT)
        ]

    def _default_vpn_slots(self):
        return [
            {"slot": i + 1, "custom_name": "", "profile": ""}
            for i in range(VPN_SLOT_COUNT)
        ]

    def _default_primary_vpn_slots(self):
        return [
            {"slot": 1, "custom_name": "Proton CH", "profile": VPN_CH},
            {"slot": 2, "custom_name": "Proton DE", "profile": VPN_DE},
        ]

    def load_primary_vpn_slots_config(self):
        return vpn_controller.load_primary_vpn_slots_config(self)

    def get_privacy_vpn_options(self):
        return vpn_controller.get_privacy_vpn_options(self)

    def _privacy_vpn_label_profile_maps(self):
        return vpn_controller.privacy_vpn_label_profile_maps(self)

    def update_privacy_vpn_combo(self):
        vpn_controller.update_privacy_vpn_combo(self)

    def _privacy_vpn_profile_from_value(self, value):
        return vpn_controller.privacy_vpn_profile_from_value(self, value)

    def _privacy_vpn_label_for_profile(self, profile):
        return vpn_controller.privacy_vpn_label_for_profile(self, profile)

    def save_primary_vpn_slots_config(self):
        vpn_controller.save_primary_vpn_slots_config(self)

    def load_vpn_slots_config(self):
        return vpn_controller.load_vpn_slots_config(self)

    def save_vpn_slots_config(self):
        return vpn_controller.save_vpn_slots_config(self)

    def load_link_slots_config(self):
        slots = self._default_link_slots()
        try:
            if os.path.isfile(self.link_config_path):
                with open(self.link_config_path, "r", encoding="utf-8") as f:
                    raw = json.load(f)
                if isinstance(raw, list):
                    for item in raw:
                        if not isinstance(item, dict):
                            continue
                        slot_no = int(item.get("slot", 0))
                        if 1 <= slot_no <= LINK_SLOT_COUNT:
                            link_type = (item.get("link_type") or "").strip()
                            if link_type not in ("Systemlink", "Externer Link"):
                                link_type = ""
                            slots[slot_no - 1] = {
                                "slot": slot_no,
                                "custom_name": (item.get("custom_name") or "").strip(),
                                "link_type": link_type,
                                "target": (item.get("target") or "").strip(),
                                "app_name": (item.get("app_name") or "").strip(),
                                "desktop_id": (item.get("desktop_id") or "").strip(),
                                "desktop_file": (
                                    item.get("desktop_file") or ""
                                ).strip(),
                            }
        except Exception as e:
            self.log(f"Link-Konfig konnte nicht geladen werden: {e}")
        return slots

    def save_link_slots_config(self):
        ok, err = safe_write_json(
            self.link_config_path,
            self.link_slots,
            indent=2,
            ensure_ascii=False,
        )
        if not ok:
            self.log(f"Link-Konfig konnte nicht gespeichert werden: {err}")

    def link_slot_configured(self, slot):
        return bool(
            (slot.get("link_type") or "").strip() and (slot.get("target") or "").strip()
        )

    def _link_slot_display_name(self, idx):
        slot = self.link_slots[idx]
        custom = (slot.get("custom_name") or "").strip()
        if custom:
            return custom
        if self.link_slot_configured(slot):
            return f"Link {idx + 1}"
        return f"Link {idx + 1} frei"

    def update_link_slot_buttons(self):
        for idx, btn in enumerate(self.link_slot_buttons):
            try:
                btn.configure(text=self._link_slot_display_name(idx))
            except Exception:
                pass

    def _link_slot_summary(self, slot):
        name = (slot.get("custom_name") or "").strip() or f"Link {slot.get('slot')}"
        ltype = (slot.get("link_type") or "-").strip() or "-"
        target = (slot.get("target") or "").strip()
        if ltype == "Systemlink":
            target = (slot.get("app_name") or "").strip() or target
        target = target or "-"
        return name, ltype, target

    def refresh_link_table(self):
        if not hasattr(self, "link_table"):
            return
        for item in self.link_table.get_children():
            self.link_table.delete(item)
        lines = []
        for slot in self.link_slots:
            name, ltype, target = self._link_slot_summary(slot)
            iid = f"link-slot-{slot['slot']}"
            self.link_table.insert(
                "",
                "end",
                iid=iid,
                values=(slot["slot"], name, ltype, target),
            )
            lines.append(f"Slot {slot['slot']}: {name} | {ltype} | {target}")

    def load_link_slot_into_form(self, slot_no):
        if not (1 <= slot_no <= LINK_SLOT_COUNT):
            return
        slot = self.link_slots[slot_no - 1]
        self.link_slot_choice.set(str(slot_no))
        self.link_custom_name_var.set((slot.get("custom_name") or "").strip())
        self.link_type_var.set(
            (slot.get("link_type") or "Systemlink").strip() or "Systemlink"
        )
        self.link_target_var.set((slot.get("target") or "").strip())
        self.link_system_app_var.set("")

        if (slot.get("link_type") or "").strip() == "Systemlink":
            desktop_id = (slot.get("desktop_id") or "").strip()
            app_name = (slot.get("app_name") or "").strip()
            entry = None
            if desktop_id:
                entry = self.desktop_app_by_id.get(desktop_id)
            if entry:
                self.link_system_app_var.set(entry["label"])
            elif app_name:
                self.link_system_app_var.set(app_name)

        self.sync_link_target_widgets()

    def on_link_slot_select(self, _event=None):
        try:
            slot_no = int((self.link_slot_choice.get() or "0").strip())
        except ValueError:
            slot_no = 0
        self.load_link_slot_into_form(slot_no)

    def on_link_table_select(self, _event=None):
        sel = self.link_table.selection() if hasattr(self, "link_table") else ()
        if not sel:
            return
        values = self.link_table.item(sel[0], "values")
        if not values:
            return
        try:
            slot_no = int(values[0])
        except (TypeError, ValueError):
            return
        self.load_link_slot_into_form(slot_no)

    def on_link_type_change(self, _event=None):
        self.sync_link_target_widgets()

    def on_systemlink_app_select(self, _event=None):
        entry = self.desktop_app_by_label.get(
            (self.link_system_app_var.get() or "").strip()
        )
        if entry:
            self.link_target_var.set(f"desktop:{entry['desktop_id']}")

    def sync_link_target_widgets(self):
        link_type = (self.link_type_var.get() or "").strip()
        if link_type == "Systemlink":
            self.link_target_label.configure(text="Programm:")
            self.link_target_entry.grid_remove()
            self.link_system_frame.grid()
        else:
            self.link_target_label.configure(text="Ziel / Pfad / URL:")
            self.link_system_frame.grid_remove()
            self.link_target_entry.grid()

    def _build_systemlink_target(self):
        if (self.link_type_var.get() or "").strip() != "Systemlink":
            return (self.link_target_var.get() or "").strip(), "", "", ""

        entry = self.desktop_app_by_label.get(
            (self.link_system_app_var.get() or "").strip()
        )
        if not entry:
            return "", "", "", ""
        return (
            f"desktop:{entry['desktop_id']}",
            entry["name"],
            entry["desktop_id"],
            entry["desktop_file"],
        )

    def _resolve_manual_system_command(self, target):
        raw = (target or "").strip()
        if not raw:
            return []
        try:
            parts = shlex.split(raw)
        except ValueError:
            # Fallback: einfache Leerzeichen-Trennung, falls Quotes unvollständig sind.
            parts = raw.split()
        if not parts:
            return []
        bin_part = parts[0]

        # PATH command
        if has_cmd(bin_part):
            return parts

        # absolute or relative executable path
        expanded_bin = os.path.expanduser(bin_part)
        if os.path.isfile(expanded_bin) and os.access(expanded_bin, os.X_OK):
            parts[0] = expanded_bin
            return parts

        # last fallback: resolve via which() for unusual PATH behavior
        resolved = shutil.which(bin_part)
        if resolved:
            parts[0] = resolved
            return parts

        return []

    def _find_wireshark_desktop_file(self):
        env_desktop = os.environ.get("WIRESHARK_DESKTOP_FILE", "").strip()
        if env_desktop and os.path.isfile(env_desktop):
            return env_desktop

        search_dirs = [
            os.path.expanduser("~/.local/share/applications"),
            "/usr/local/share/applications",
            "/usr/share/applications",
        ]
        candidates = [
            "wireshark.desktop",
            "org.wireshark.Wireshark.desktop",
        ]
        for base in search_dirs:
            for name in candidates:
                candidate = os.path.join(base, name)
                if os.path.isfile(candidate):
                    return candidate
        return ""

    def _resolve_wireshark_launcher(self):
        for candidate in [
            get_cmd_path("wireshark"),
            "/usr/bin/wireshark",
            "/usr/local/bin/wireshark",
            "/snap/bin/wireshark",
            "/var/lib/flatpak/exports/bin/org.wireshark.Wireshark",
            os.path.expanduser(
                "~/.local/share/flatpak/exports/bin/org.wireshark.Wireshark"
            ),
        ]:
            if (
                candidate
                and os.path.isfile(candidate)
                and os.access(candidate, os.X_OK)
            ):
                return {"kind": "binary", "target": candidate}

        desktop_entry = self._find_desktop_entry_by_query("wireshark")
        if desktop_entry:
            desktop_id = (desktop_entry.get("desktop_id") or "").strip()
            if desktop_id:
                return {"kind": "desktop", "target": desktop_id}

        desktop_file = self._find_wireshark_desktop_file()
        if desktop_file:
            return {"kind": "desktop_file", "target": desktop_file}

        flatpak_id = self._find_flatpak_app_id_by_query("wireshark")
        if flatpak_id:
            return {"kind": "flatpak", "target": flatpak_id}

        return {"kind": "", "target": ""}

    def _find_desktop_entry_by_query(self, query):
        q = (query or "").strip().lower()
        if not q:
            return None

        entries = list(getattr(self, "desktop_app_entries", []) or [])
        best = None

        # 1) exact matches first
        for e in entries:
            name = (e.get("name") or "").strip().lower()
            label = (e.get("label") or "").strip().lower()
            did = (e.get("desktop_id") or "").strip().lower()
            if q in {name, label, did, did.replace(".desktop", "")}:
                return e

        # 2) substring fallback (prefer shorter name for relevance)
        candidates = []
        for e in entries:
            name = (e.get("name") or "").strip().lower()
            label = (e.get("label") or "").strip().lower()
            did = (e.get("desktop_id") or "").strip().lower()
            if q in name or q in label or q in did:
                rank = min(len(name) or 999, len(label) or 999, len(did) or 999)
                candidates.append((rank, e))

        if candidates:
            candidates.sort(key=lambda item: item[0])
            best = candidates[0][1]

        return best

    def _find_flatpak_app_id_by_query(self, query):
        q = (query or "").strip().lower()
        if not q or not has_cmd("flatpak"):
            return ""

        rc, out = run_cmd(
            ["flatpak", "list", "--app", "--columns=application,name"],
            timeout=8,
        )
        if rc != 0 or not out:
            return ""

        rows = []
        for raw in out.splitlines():
            line = (raw or "").strip()
            if not line:
                continue
            parts = [p.strip() for p in line.split("	") if p.strip()]
            if not parts:
                continue
            app_id = parts[0]
            app_name = parts[1] if len(parts) > 1 else ""
            rows.append((app_id, app_name))

        # Exact match first.
        for app_id, app_name in rows:
            app_id_l = app_id.lower()
            app_name_l = app_name.lower()
            if q in {app_id_l, app_name_l}:
                return app_id
            if q == app_id_l.split(".")[-1]:
                return app_id

        # Substring fallback.
        candidates = []
        for app_id, app_name in rows:
            app_id_l = app_id.lower()
            app_name_l = app_name.lower()
            if q in app_id_l or (app_name_l and q in app_name_l):
                rank = min(len(app_id_l), len(app_name_l) if app_name_l else 999)
                candidates.append((rank, app_id))
        if candidates:
            candidates.sort(key=lambda item: item[0])
            return candidates[0][1]

        return ""

    def assign_link_slot(self):
        try:
            slot_no = int((self.link_slot_choice.get() or "0").strip())
        except ValueError:
            slot_no = 0
        if not (1 <= slot_no <= LINK_SLOT_COUNT):
            self.log("Ungültiger Link-Slot ausgewählt.")
            return

        link_type = (self.link_type_var.get() or "").strip()
        custom_name = (self.link_custom_name_var.get() or "").strip()
        if link_type not in ("Systemlink", "Externer Link"):
            self.log("Bitte einen gültigen Link-Typ auswählen.")
            return
        app_name = ""
        desktop_id = ""
        desktop_file = ""
        if link_type == "Systemlink":
            target, app_name, desktop_id, desktop_file = self._build_systemlink_target()
        else:
            target = normalize_external_url_target(self.link_target_var.get())
        if not target:
            if link_type == "Systemlink":
                self.log(
                    "Bitte ein Programm auswählen oder eine manuelle Eingabe hinterlegen."
                )
            else:
                self.log("Bitte eine gültige HTTP/HTTPS-URL eintragen.")
            return

        if link_type == "Systemlink":
            lowered = target.strip().lower()
            expanded = os.path.expanduser(target)
            if not (
                target.startswith("desktop:")
                or os.path.exists(expanded)
                or lowered in ("vscode", "code", "codium", "code-insiders")
            ):
                self.log(
                    "Unsicherer Systemlink blockiert. Bitte Desktop-App oder lokalen Pfad verwenden."
                )
                return

        slot = self.link_slots[slot_no - 1]
        slot["custom_name"] = custom_name
        slot["link_type"] = link_type
        slot["target"] = target
        slot["app_name"] = app_name
        slot["desktop_id"] = desktop_id
        slot["desktop_file"] = desktop_file
        self.save_link_slots_config()
        self.update_link_slot_buttons()
        self.refresh_link_table()
        target_label = custom_name or app_name or target
        self.log(f"Link-Zuordnung übernommen: Slot {slot_no} -> {target_label}.")

    def clear_link_slot(self, confirm=False):
        try:
            slot_no = int((self.link_slot_choice.get() or "0").strip())
        except ValueError:
            slot_no = 0
        if not (1 <= slot_no <= LINK_SLOT_COUNT):
            self.log("Ungültiger Link-Slot ausgewählt.")
            return
        slot = self.link_slots[slot_no - 1]
        has_assignment = bool(
            (slot.get("custom_name") or "").strip()
            or (slot.get("link_type") or "").strip()
            or (slot.get("target") or "").strip()
        )
        if confirm and has_assignment:
            ok = messagebox.askyesno(
                "Link-Zuordnung aufheben",
                f"Slot {slot_no} wird geleert. Fortfahren?",
                parent=self,
            )
            if not ok:
                self.log(f"Link-Zuordnung aufheben abgebrochen: Slot {slot_no}.")
                return
        slot["custom_name"] = ""
        slot["link_type"] = ""
        slot["target"] = ""
        slot["app_name"] = ""
        slot["desktop_id"] = ""
        slot["desktop_file"] = ""
        self.save_link_slots_config()
        self.load_link_slot_into_form(slot_no)
        self.update_link_slot_buttons()
        self.refresh_link_table()
        self.log(f"Link-Zuordnung aufgehoben: Slot {slot_no}.")

    def _build_form_link_slot(self):
        target = (self.link_target_var.get() or "").strip()
        app_name = ""
        desktop_id = ""
        desktop_file = ""
        if (self.link_type_var.get() or "").strip() == "Systemlink":
            target, app_name, desktop_id, desktop_file = self._build_systemlink_target()
        return {
            "slot": int((self.link_slot_choice.get() or "1").strip() or 1),
            "custom_name": (self.link_custom_name_var.get() or "").strip(),
            "link_type": (self.link_type_var.get() or "").strip(),
            "target": target,
            "app_name": app_name,
            "desktop_id": desktop_id,
            "desktop_file": desktop_file,
        }

    def test_link_slot(self):
        slot = self._build_form_link_slot()
        if not self.link_slot_configured(slot):
            self.log("Bitte erst Typ und Ziel für den Link-Test ausfüllen.")
            return
        self._execute_link_slot(slot, test_mode=True)

    def open_link_slot(self, idx):
        if idx < 0 or idx >= len(self.link_slots):
            return
        slot = self.link_slots[idx]
        if not self.link_slot_configured(slot):
            self.log(f"Link-Slot {idx + 1} ist frei.")
            if hasattr(self, "tab_links"):
                self._select_system_tab(self.tab_links)
                self.load_link_slot_into_form(idx + 1)
            return
        self._execute_link_slot(slot)

    def _open_local_target(self, target):
        if target.startswith("desktop:"):
            desktop_id = target.split(":", 1)[1].strip()
            if not desktop_id:
                raise RuntimeError("Ungültiger Desktop-Starter.")
            entry = self.desktop_app_by_id.get(desktop_id)
            if has_cmd("gtk-launch"):
                subprocess.Popen(
                    ["gtk-launch", desktop_id],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                    cwd=os.path.expanduser("~"),
                )
                self.log(f"Systemlink gestartet: {desktop_id}")
                return True
            if entry and entry.get("exec"):
                proc = subprocess.Popen(
                    entry["exec"],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                    cwd=os.path.expanduser("~"),
                )
                self.log(f"Systemlink gestartet: {entry['exec'][0]} (PID {proc.pid}).")
                return True
            raise RuntimeError("Desktop-Starter konnte nicht ausgeführt werden.")

        expanded = os.path.expanduser(target)
        if os.path.exists(expanded):
            if has_cmd("xdg-open"):
                subprocess.Popen(
                    ["xdg-open", expanded],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
                self.log(f"Systemlink geöffnet: {expanded}")
                return True
            raise RuntimeError("xdg-open ist nicht verfügbar.")

        lowered = target.strip().lower()
        if lowered in ("vscode", "code", "codium", "code-insiders"):
            self.vscode_async()
            return True

        parts = self._resolve_manual_system_command(target)
        if parts:
            proc = subprocess.Popen(
                parts,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                cwd=os.path.expanduser("~"),
            )
            self.log(f"Systemlink gestartet: {parts[0]} (PID {proc.pid}).")
            return True

        # Fallback 1: Desktop-ID direkt per gtk-launch versuchen
        launch_candidate = (target or "").strip()
        if has_cmd("gtk-launch") and launch_candidate:
            launch_ids = [launch_candidate]
            if not launch_candidate.endswith(".desktop"):
                launch_ids.append(f"{launch_candidate}.desktop")
            for launch_id in launch_ids:
                try:
                    rc = subprocess.run(
                        ["gtk-launch", launch_id],
                        stdout=subprocess.DEVNULL,
                        stderr=subprocess.DEVNULL,
                        cwd=os.path.expanduser("~"),
                        timeout=5,
                        check=False,
                    ).returncode
                    if rc == 0:
                        self.log(f"Systemlink gestartet: {launch_id}")
                        return True
                except Exception:
                    continue

        # Fallback 2: Desktop-Eintrag anhand Name/Label/Desktop-ID suchen
        matched_entry = self._find_desktop_entry_by_query(launch_candidate)
        if matched_entry:
            desktop_id = (matched_entry.get("desktop_id") or "").strip()
            if desktop_id and has_cmd("gtk-launch"):
                try:
                    rc = subprocess.run(
                        ["gtk-launch", desktop_id],
                        stdout=subprocess.DEVNULL,
                        stderr=subprocess.DEVNULL,
                        cwd=os.path.expanduser("~"),
                        timeout=5,
                        check=False,
                    ).returncode
                    if rc == 0:
                        self.log(f"Systemlink gestartet: {desktop_id}")
                        return True
                except Exception:
                    pass
            exec_cmd = matched_entry.get("exec") or []
            if exec_cmd:
                try:
                    proc = subprocess.Popen(
                        exec_cmd,
                        stdout=subprocess.DEVNULL,
                        stderr=subprocess.DEVNULL,
                        cwd=os.path.expanduser("~"),
                    )
                    self.log(f"Systemlink gestartet: {exec_cmd[0]} (PID {proc.pid}).")
                    return True
                except Exception:
                    pass

        # Fallback 3: xdg-open für Dateipfade/sonstige Targets
        if has_cmd("xdg-open") and launch_candidate:
            try:
                rc = subprocess.run(
                    ["xdg-open", os.path.expanduser(launch_candidate)],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                    timeout=8,
                    check=False,
                ).returncode
                if rc == 0:
                    self.log(f"Systemlink geöffnet: {launch_candidate}")
                    return True
            except Exception:
                pass

        # Fallback 4: Flatpak-App anhand Name/App-ID starten.
        flatpak_app_id = self._find_flatpak_app_id_by_query(launch_candidate)
        if flatpak_app_id and has_cmd("flatpak"):
            try:
                proc = subprocess.Popen(
                    ["flatpak", "run", flatpak_app_id],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                    cwd=os.path.expanduser("~"),
                )
                self.log(
                    f"Systemlink gestartet: flatpak run {flatpak_app_id} (PID {proc.pid})."
                )
                return True
            except Exception:
                pass

        raise RuntimeError(
            "Manueller Befehl/Starter konnte nicht ausgeführt werden. Bitte gültigen Befehl, Desktop-ID oder Dateipfad verwenden."
        )

    def _open_external_target(self, target):
        url = normalize_external_url_target(target)
        if not url:
            raise RuntimeError("Ungültige oder unsichere URL (nur HTTP/HTTPS erlaubt).")
        if has_cmd("xdg-open"):
            subprocess.Popen(
                ["xdg-open", url],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
            )
        else:
            opened = webbrowser.open(url)
            if not opened:
                raise RuntimeError("Standardbrowser konnte nicht geöffnet werden.")
        self.log(f"Externer Link geöffnet: {url}")
        return True

    def _execute_link_slot(self, slot, test_mode=False):
        link_type = (slot.get("link_type") or "").strip()
        target = (slot.get("target") or "").strip()
        if not link_type or not target:
            self.log("Link-Slot ist nicht vollständig konfiguriert.")
            return False
        try:
            if link_type == "Systemlink":
                return self._open_local_target(target)
            if link_type == "Externer Link":
                return self._open_external_target(target)
            self.log(f"Unbekannter Link-Typ: {link_type}")
            return False
        except Exception as e:
            prefix = (
                "Link-Test fehlgeschlagen"
                if test_mode
                else "Link konnte nicht geöffnet werden"
            )
            self.log(f"{prefix}: {e}")
            return False

    def _vpn_slot_display_name(self, idx):
        slot = self.vpn_slots[idx]
        custom = (slot.get("custom_name") or "").strip()
        profile = self._base_wg_profile_name(slot.get("profile"))
        if not profile:
            return f"VPN {idx + 1} frei"
        return custom or profile

    def _vpn_slot_status(self, slot):
        profile = self._base_wg_profile_name(slot.get("profile"))
        if not profile:
            return "frei", "gray"

        active_profile = self._base_wg_profile_name(self.active_wg_profile())
        if active_profile and active_profile == profile:
            return "aktiv", "green"

        available_profiles = set(getattr(self, "wg_profiles", []) or [])
        if available_profiles and profile not in available_profiles:
            return "fehlt", "red"

        return "bereit", "red"

    def open_primary_vpn_slot(self, slot_no):
        return vpn_controller.open_primary_vpn_slot(self, slot_no)

    def save_primary_vpn_slot(self, slot_no):
        vpn_controller.save_primary_vpn_slot(self, slot_no)

    def update_primary_vpn_buttons(self):
        vpn_controller.update_primary_vpn_buttons(self)

    def update_primary_vpn_form_controls(self):
        vpn_controller.update_primary_vpn_form_controls(self)

    def update_vpn_slot_buttons(self):
        if threading.current_thread() is not threading.main_thread():
            self.after(0, self.update_vpn_slot_buttons)
            return
        for idx, btn in enumerate(self.vpn_slot_buttons):
            try:
                btn.configure(text=self._vpn_slot_display_name(idx))
            except Exception:
                pass
            if idx < len(self.vpn_slot_lights):
                _status, color = self._vpn_slot_status(self.vpn_slots[idx])
                self.set_light(self.vpn_slot_lights[idx], color)

    def refresh_vpn_slots_table(self):
        vpn_controller.refresh_vpn_slots_table(self)

    def load_vpn_slot_into_form(self, slot_no):
        vpn_controller.load_vpn_slot_into_form(self, slot_no)

    def sync_vpn_profile_form_selection(self):
        vpn_controller.sync_vpn_profile_form_selection(self)

    def sync_primary_vpn_form_selection(self):
        vpn_controller.sync_primary_vpn_form_selection(self)

    def _on_vpn_form_field_changed(self, *_args):
        if getattr(self, "_vpn_form_loading", False):
            return
        self._vpn_form_dirty = True

    def on_vpn_slot_select(self, _event=None):
        vpn_controller.on_vpn_slot_select(self, _event=_event)

    def on_vpn_table_select(self, _event=None):
        vpn_controller.on_vpn_table_select(self, _event=_event)

    def _resolve_vpn_profile_input(self, profile_value, refresh_profiles=False):
        return vpn_controller.resolve_vpn_profile_input(
            self,
            profile_value,
            refresh_profiles=refresh_profiles,
        )

    def assign_vpn_slot(self):
        vpn_controller.assign_vpn_slot(self)

    def _is_vpn_profile_referenced_in_other_slots(self, profile, skip_slot_no=None):
        return vpn_controller.is_vpn_profile_referenced_in_other_slots(
            self,
            profile,
            skip_slot_no=skip_slot_no,
        )

    def _collect_referenced_vpn_profiles(self, include_active=True):
        return vpn_controller.collect_referenced_vpn_profiles(
            self,
            include_active=include_active,
        )

    def _sync_manual_wg_slot_controls(self):
        if not hasattr(self, "manual_wg_slot_combo"):
            return
        state = "readonly" if self.manual_wg_save_slot_var.get() else "disabled"
        self.manual_wg_slot_combo.configure(state=state)

    def _set_manual_wg_status(self, message, level="info"):
        text_var = self.__dict__.get("manual_wg_status_var")
        label = self.__dict__.get("manual_wg_status_label")
        if text_var is None:
            return

        palette = {
            "ok": "#1f4f1f",
            "info": "#1f3f74",
            "warn": "#7a5a00",
            "error": "#8b1a1a",
        }
        color = palette.get((level or "info").strip().lower(), palette["info"])
        text_var.set((message or "").strip() or "-")
        if label is not None:
            try:
                label.configure(fg=color)
            except Exception:
                pass

    def clear_manual_wireguard_form(self):
        vpn_controller.clear_manual_wireguard_form(self)

    def _is_valid_wireguard_key(self, value):
        return vpn_controller.is_valid_wireguard_key(self, value)

    def _parse_manual_wireguard_endpoint(self, endpoint):
        return vpn_controller.parse_manual_wireguard_endpoint(self, endpoint)

    def _validate_manual_wireguard_form(self):
        return vpn_controller.validate_manual_wireguard_form(self)

    def _build_manual_wireguard_config_text(self, payload):
        return vpn_controller.build_manual_wireguard_config_text(self, payload)

    def _write_manual_wg_profile(self, profile_name, config_text):
        return vpn_controller.write_manual_wg_profile(self, profile_name, config_text)

    def _assign_manual_profile_to_slot(self, profile_name, slot_no):
        return vpn_controller.assign_manual_profile_to_slot(
            self,
            profile_name,
            slot_no,
        )

    def apply_manual_wireguard_profile(self):
        vpn_controller.apply_manual_wireguard_profile(self)

    def _collect_interface_rules_for_profiles(
        self, profiles, include_family_candidates=False
    ):
        return vpn_controller.collect_interface_rules_for_profiles(
            self,
            profiles,
            include_family_candidates=include_family_candidates,
        )

    def _collect_interface_rules_for_profiles_safe(
        self, profiles, include_family_candidates=False
    ):
        return vpn_controller.collect_interface_rules_for_profiles_safe(
            self,
            profiles,
            include_family_candidates=include_family_candidates,
        )

    def _collect_endpoint_rules_for_profiles(self, profiles):
        return vpn_controller.collect_endpoint_rules_for_profiles(self, profiles)

    def ensure_profile_family_ready_for_use(self, profile_name, log_failures=True):
        return vpn_controller.ensure_profile_family_ready_for_use(
            self,
            profile_name,
            log_failures=log_failures,
        )

    def _apply_profile_ufw_rules_non_killswitch(self, interfaces, endpoints):
        return vpn_controller.apply_profile_ufw_rules_non_killswitch(
            self,
            interfaces,
            endpoints,
        )

    def reconcile_vpn_slot_firewall_state(
        self, reason="", touched_profiles=None, refresh_ui=False
    ):
        return vpn_controller.reconcile_vpn_slot_firewall_state(
            self,
            reason=reason,
            touched_profiles=touched_profiles,
            refresh_ui=refresh_ui,
        )

    def _remove_interface_rule(self, iface, direction):
        direction = (direction or "").strip().lower()
        iface = (iface or "").strip()
        if direction not in {"in", "out"} or not iface:
            return "failed", 0, "ungueltige Interface-Regelparameter"

        cmd = [
            "sudo",
            "-n",
            "ufw",
            "delete",
            "allow",
            direction,
            "on",
            iface,
        ]

        removed_count = 0
        last_out = ""
        for _ in range(4):
            rc, out = run_cmd(cmd, timeout=20)
            last_out = out or ""
            if rc == 0:
                removed_count += 1
                continue

            text = (out or "").strip().lower()
            if (
                "could not delete non-existent rule" in text
                or "could not delete non existent rule" in text
                or "could not find rule" in text
            ):
                if removed_count > 0:
                    return "removed", removed_count, out
                return "missing", 0, out

            return "failed", removed_count, out

        return "removed", removed_count, last_out

    def _remove_udp_endpoint_rule(self, endpoint_ip, endpoint_port):
        cmd = [
            "sudo",
            "-n",
            "ufw",
            "delete",
            "allow",
            "out",
            "to",
            endpoint_ip,
            "port",
            endpoint_port,
            "proto",
            "udp",
        ]
        rc, out = run_cmd(cmd, timeout=20)
        if rc == 0:
            return "removed", out

        text = (out or "").strip().lower()
        if (
            "could not delete non-existent rule" in text
            or "could not delete non existent rule" in text
            or "could not find rule" in text
        ):
            return "missing", out
        return "failed", out

    def _reconcile_vpn_firewall_rules(
        self,
        old_interfaces,
        new_interfaces,
        old_endpoints,
        new_endpoints,
        removed_profile=None,
    ):
        old_iface_set = {
            (iface or "").strip()
            for iface in (old_interfaces or set())
            if (iface or "").strip()
        }
        new_iface_set = {
            (iface or "").strip()
            for iface in (new_interfaces or set())
            if (iface or "").strip()
        }
        old_ep_set = {
            (str(ip).strip(), str(port).strip())
            for ip, port in (old_endpoints or set())
            if str(ip).strip() and str(port).strip()
        }
        new_ep_set = {
            (str(ip).strip(), str(port).strip())
            for ip, port in (new_endpoints or set())
            if str(ip).strip() and str(port).strip()
        }

        stale_ifaces = sorted(old_iface_set - new_iface_set)
        kept_ifaces = sorted(old_iface_set & new_iface_set)
        stale_endpoints = sorted(old_ep_set - new_ep_set)
        kept_endpoints = sorted(old_ep_set & new_ep_set)

        if removed_profile:
            self.log(
                f"Firewall-Reconcile: Profil '{removed_profile}' wurde geändert/entfernt – prüfe verwaiste Interface- und UDP-Endpoint-Regeln."
            )

        if old_iface_set:
            self.log_debug(
                "Firewall-Reconcile: Interface-Altzustand = "
                + ", ".join(sorted(old_iface_set))
            )
        else:
            self.log_debug("Firewall-Reconcile: Interface-Altzustand leer.")

        if new_iface_set:
            self.log_debug(
                "Firewall-Reconcile: Interface-Neuzustand = "
                + ", ".join(sorted(new_iface_set))
            )
        else:
            self.log_debug("Firewall-Reconcile: Interface-Neuzustand leer.")

        if kept_ifaces:
            self.log_debug(
                "Firewall-Reconcile: Diese Interface-Regeln bleiben erhalten (weiter referenziert): "
                + ", ".join(kept_ifaces)
            )

        if old_ep_set:
            old_ep_text = ", ".join([f"{ip}:{port}" for ip, port in sorted(old_ep_set)])
            self.log_debug(f"Firewall-Reconcile: Endpoint-Altzustand = {old_ep_text}")
        else:
            self.log_debug("Firewall-Reconcile: Endpoint-Altzustand leer.")

        if new_ep_set:
            new_ep_text = ", ".join([f"{ip}:{port}" for ip, port in sorted(new_ep_set)])
            self.log_debug(f"Firewall-Reconcile: Endpoint-Neuzustand = {new_ep_text}")
        else:
            self.log_debug("Firewall-Reconcile: Endpoint-Neuzustand leer.")

        if kept_endpoints:
            kept_ep_text = ", ".join([f"{ip}:{port}" for ip, port in kept_endpoints])
            self.log_debug(
                f"Firewall-Reconcile: Diese UDP-Endpoint-Regeln bleiben erhalten (weiter referenziert): {kept_ep_text}"
            )

        if not stale_ifaces and not stale_endpoints:
            self.log_debug(
                "Firewall-Reconcile: Keine verwaisten Interface- oder UDP-Endpoint-Regeln zu löschen."
            )
            return True

        if not has_cmd("ufw"):
            self.log(
                "Firewall-Reconcile: UFW nicht verfügbar – Firewall-Cleanup übersprungen."
            )
            return True

        if not self._try_begin_network_operation("VPN Firewall-Regeln bereinigen"):
            return False

        try:
            with self.ufw_lock:
                rc_status, out_status = self._read_ufw_status()

            if rc_status != 0:
                if not self.ensure_sudo_session("VPN Firewall-Regeln bereinigen"):
                    self.log(
                        "Firewall-Reconcile: Keine sudo-Berechtigung für Firewall-Cleanup."
                    )
                    return False
                with self.ufw_lock:
                    rc_status, out_status = self._read_ufw_status()

            if rc_status != 0:
                self.log("Firewall-Reconcile fehlgeschlagen: UFW-Status nicht lesbar.")
                if out_status:
                    self.log_debug_safe(out_status)
                return False

            if not self._ufw_status_is_active(out_status):
                self.log(
                    "Firewall-Reconcile: UFW ist deaktiviert – verwaiste Regeln werden nicht entfernt."
                )
                return True

            removed_interface = []
            missing_interface = []
            failed_interface = []
            removed_endpoints = []
            missing_endpoints = []
            failed_endpoints = []

            with self.ufw_lock:
                for iface in stale_ifaces:
                    for direction in ("out", "in"):
                        result, removed_count, detail = self._remove_interface_rule(
                            iface, direction
                        )
                        rule_label = f"{direction} on {iface}"
                        if result == "removed":
                            removed_interface.append((rule_label, removed_count))
                        elif result == "missing":
                            missing_interface.append(rule_label)
                        else:
                            failed_interface.append((rule_label, detail))

                for endpoint_ip, endpoint_port in stale_endpoints:
                    result, detail = self._remove_udp_endpoint_rule(
                        endpoint_ip, endpoint_port
                    )
                    endpoint_label = f"{endpoint_ip}:{endpoint_port}"
                    if result == "removed":
                        removed_endpoints.append(endpoint_label)
                    elif result == "missing":
                        missing_endpoints.append(endpoint_label)
                    else:
                        failed_endpoints.append((endpoint_label, detail))

            if removed_interface:
                details = ", ".join(
                    [
                        f"{label}{' (x' + str(count) + ')' if count > 1 else ''}"
                        for label, count in removed_interface
                    ]
                )
                self.log(f"Firewall-Reconcile: Interface-Regeln entfernt: {details}")
            if missing_interface:
                self.log_debug(
                    "Firewall-Reconcile: Interface-Regeln waren bereits nicht vorhanden: "
                    + ", ".join(missing_interface)
                )
            if failed_interface:
                self.log(
                    "Firewall-Reconcile: Entfernen folgender Interface-Regeln fehlgeschlagen: "
                    + ", ".join([label for label, _ in failed_interface])
                )
                first_detail = failed_interface[0][1]
                if first_detail:
                    self.log_debug_safe(first_detail)

            if removed_endpoints:
                self.log(
                    "Firewall-Reconcile: UDP-Endpoint-Regeln entfernt: "
                    + ", ".join(removed_endpoints)
                )
            if missing_endpoints:
                self.log_debug(
                    "Firewall-Reconcile: UDP-Endpoint-Regeln waren bereits nicht vorhanden: "
                    + ", ".join(missing_endpoints)
                )
            if failed_endpoints:
                self.log(
                    "Firewall-Reconcile: Entfernen folgender UDP-Endpoint-Regeln fehlgeschlagen: "
                    + ", ".join([label for label, _ in failed_endpoints])
                )
                first_detail = failed_endpoints[0][1]
                if first_detail:
                    self.log_debug_safe(first_detail)

            return not failed_interface and not failed_endpoints
        finally:
            self._end_network_operation()

    def _remove_firewall_rules_for_vpn_profile(self, profile, removed_from_slot=None):
        return vpn_controller.remove_firewall_rules_for_vpn_profile(
            self,
            profile,
            removed_from_slot=removed_from_slot,
        )

    def clear_vpn_slot(self, confirm=False):
        vpn_controller.clear_vpn_slot(self, confirm=confirm)

    def delete_selected_vpn_profile(self):
        vpn_controller.delete_selected_vpn_profile(self)

    def _start_vpn_slot_profile(self, profile, slot_no=None, test_mode=False):
        return vpn_controller.start_vpn_slot_profile(
            self,
            profile,
            slot_no=slot_no,
            test_mode=test_mode,
        )

    def test_vpn_slot(self):
        vpn_controller.test_vpn_slot(self)

    def open_vpn_slot(self, idx):
        vpn_controller.open_vpn_slot(self, idx)

    def _default_hdd_slots(self):
        return hdd_controller.default_hdd_slots()

    def load_hdd_slots_config(self):
        return hdd_controller.load_hdd_slots_config(self)

    def save_hdd_slots_config(self):
        hdd_controller.save_hdd_slots_config(self)

    def save_hdd_config(self):
        hdd_controller.save_hdd_config(self)

    def load_privacy_config(self):
        config = {
            "enabled": False,
            "autostart": False,
            "disable_ipv6": False,
            "enable_killswitch": False,
            "strict_protection": True,
            "vpn_profile": "",
            "dns_profile": "System (DHCP/auto)",
            "dns_mode": "auto",
            "dns_guard_enabled": False,
            "wg_import_dir": "",
            "log_overlay_active": False,
            "log_overlay_image_path": "",
            "log_overlay_last_dir": "",
            "clock_alarm_enabled": False,
            "clock_alarm_time": "07:00",
            "theme": "light",
        }
        try:
            if os.path.isfile(self.privacy_config_path):
                with open(self.privacy_config_path, "r", encoding="utf-8") as f:
                    loaded = json.load(f)
                if isinstance(loaded, dict):
                    config.update(loaded)
        except Exception as e:
            self.log(f"Privacy-Konfig konnte nicht geladen werden: {e}")
        return config

    def _apply_privacy_config_to_vars(self, config):
        dns_profile = (config.get("dns_profile") or "System (DHCP/auto)").strip()
        if dns_profile not in DNS_PROFILES:
            dns_profile = "System (DHCP/auto)"
        vpn_profile = self._privacy_vpn_profile_from_value(config.get("vpn_profile"))
        self.isp_stealth_enabled.set(bool(config.get("enabled", False)))
        self.isp_stealth_autostart.set(bool(config.get("autostart", False)))
        self.isp_stealth_disable_ipv6.set(bool(config.get("disable_ipv6", False)))
        self.isp_stealth_enable_killswitch.set(
            bool(config.get("enable_killswitch", False))
        )
        strict_enabled = bool(config.get("strict_protection", True))
        strict_var = self.__dict__.get("strict_protection_var")
        if strict_var is not None:
            strict_var.set(strict_enabled)
            self.danger_guard_enabled = bool(strict_var.get())
        else:
            self.danger_guard_enabled = strict_enabled
        self.isp_stealth_vpn_profile.set(vpn_profile)
        self.isp_stealth_vpn_choice.set(
            self._privacy_vpn_label_for_profile(vpn_profile)
        )
        self.isp_stealth_dns_profile.set(dns_profile)
        self._set_dns_mode((config.get("dns_mode") or "auto"), persist=False)
        dns_guard_var = self.__dict__.get("dns_guard_enabled_var")
        if dns_guard_var is not None:
            dns_guard_var.set(bool(config.get("dns_guard_enabled", False)))

        alarm_time = (
            self._validate_alarm_time(config.get("clock_alarm_time")) or "07:00"
        )
        alarm_time_var = self.__dict__.get("alarm_time_var")
        if alarm_time_var is not None:
            alarm_time_var.set(alarm_time)
        alarm_enabled_var = self.__dict__.get("alarm_enabled_var")
        if alarm_enabled_var is not None:
            alarm_enabled_var.set(bool(config.get("clock_alarm_enabled", False)))
        theme_name = str(config.get("theme") or "light").strip().lower()
        self._preferred_theme_name = (
            theme_name if theme_name in {"light", "dark"} else "light"
        )

    def save_privacy_config(self):
        vpn_profile = self._privacy_vpn_profile_from_value(
            self.isp_stealth_vpn_choice.get()
        )
        self.isp_stealth_vpn_profile.set(vpn_profile)
        current_config = getattr(self, "privacy_config", {})
        wg_import_dir = ""
        if isinstance(current_config, dict):
            wg_import_dir = (current_config.get("wg_import_dir") or "").strip()
        strict_protection = bool(self.strict_protection_var.get())
        self.danger_guard_enabled = strict_protection
        log_overlay_active = bool(self.__dict__.get("log_overlay_active", False))
        log_overlay_image_path = str(
            self.__dict__.get("log_overlay_image_path") or ""
        ).strip()
        log_overlay_last_dir = str(
            self.__dict__.get("log_overlay_picker_dir") or ""
        ).strip()
        if log_overlay_last_dir and not os.path.isdir(log_overlay_last_dir):
            log_overlay_last_dir = ""
        config = {
            "enabled": self.isp_stealth_enabled.get(),
            "autostart": self.isp_stealth_autostart.get(),
            "disable_ipv6": self.isp_stealth_disable_ipv6.get(),
            "enable_killswitch": self.isp_stealth_enable_killswitch.get(),
            "strict_protection": strict_protection,
            "vpn_profile": vpn_profile,
            "dns_profile": self.isp_stealth_dns_profile.get(),
            "dns_mode": self._current_dns_mode(),
            "dns_guard_enabled": bool(self.dns_guard_enabled_var.get()),
            "wg_import_dir": wg_import_dir,
            "log_overlay_active": log_overlay_active,
            "log_overlay_image_path": log_overlay_image_path,
            "log_overlay_last_dir": log_overlay_last_dir,
            "clock_alarm_enabled": bool(self.alarm_enabled_var.get()),
            "clock_alarm_time": self._validate_alarm_time(self.alarm_time_var.get())
            or "07:00",
            "theme": (self._app_theme or LIGHT_THEME).get("name", "light"),
        }
        ok, err = safe_write_json(
            self.privacy_config_path,
            config,
            indent=2,
            ensure_ascii=False,
        )
        if ok:
            self.privacy_config = config
            if "settings" in self.__dict__:
                dns_guard_enabled = bool(self.dns_guard_enabled_var.get())
                self.settings["dns_guard_enabled"] = dns_guard_enabled
                self.dns_guard_enabled = dns_guard_enabled
                settings_var = self.__dict__.get("var_dns_guard_enabled")
                if settings_var is not None:
                    try:
                        settings_var.set(dns_guard_enabled)
                    except Exception:
                        pass
                self.save_settings(log_success=False)
        else:
            self.log(f"Privacy-Konfig konnte nicht gespeichert werden: {err}")

    def get_wg_import_dir(self):
        config = getattr(self, "privacy_config", {})
        if not isinstance(config, dict):
            return ""
        path = (config.get("wg_import_dir") or "").strip()
        if path and os.path.isdir(path):
            return path
        return ""

    def set_wg_import_dir(self, folder):
        path = (folder or "").strip()
        if not path or not os.path.isdir(path):
            return
        config = getattr(self, "privacy_config", None)
        if not isinstance(config, dict):
            self.privacy_config = {}
        self.privacy_config["wg_import_dir"] = path
        self.save_privacy_config()

    # old methods replaced by dynamic primary-vpn-based versions earlier in this class

    def on_privacy_vpn_selected(self):
        selected_profile = self._privacy_vpn_profile_from_value(
            self.isp_stealth_vpn_choice.get()
        )
        self.isp_stealth_vpn_profile.set(selected_profile)
        self.isp_stealth_vpn_choice.set(
            self._privacy_vpn_label_for_profile(selected_profile)
        )
        self.save_privacy_config()
        self.update_ipv6_controls_state()

    def isp_stealth_apply_startup(self):
        self._apply_privacy_config_to_vars(self.privacy_config)
        if self.privacy_config.get("enabled", False) and self.privacy_config.get(
            "autostart", False
        ):
            self.isp_stealth_on(from_startup=True)

    def isp_stealth_on(self, from_startup=False):
        if self.isp_stealth_running:
            self.log("ISP-Stealthmodus läuft bereits.")
            return
        if not self.ensure_sudo_session("ISP-Stealthmodus aktivieren"):
            return

        disable_ipv6 = self.isp_stealth_disable_ipv6.get()
        enable_killswitch = self.isp_stealth_enable_killswitch.get()
        vpn_profile = self.isp_stealth_vpn_profile.get().strip()
        dns_profile = self.isp_stealth_dns_profile.get().strip()
        previous_state = self._capture_stealth_runtime_state()

        def job():
            with self.isp_stealth_lock:
                self.isp_stealth_running = True
                tx = self._new_rollback_transaction("ISP-Stealthmodus AN")
                if tx is not None:
                    tx.add_rollback(
                        "Vorzustand wiederherstellen",
                        lambda: self._restore_stealth_runtime_state(
                            previous_state, refresh_ui=False
                        ),
                    )

                def rollback_and_abort(msg):
                    rollback_ok = True
                    if tx is not None:
                        rollback_ok = tx.rollback()
                    else:
                        try:
                            self._restore_stealth_runtime_state(
                                previous_state, refresh_ui=False
                            )
                        except Exception as rollback_exc:
                            rollback_ok = False
                            self.log_debug_safe(
                                f"ISP-Stealthmodus AN: Rollback-Fehler: {rollback_exc}"
                            )

                    self.after(0, lambda: self._set_stealth_enabled_state(False))
                    self.log_safe(msg)
                    if not rollback_ok:
                        self.log_safe(
                            "Rollback unvollständig – bitte Zustand manuell prüfen (VPN/Killswitch/DNS/IPv6)."
                        )
                    self.refresh_safe()
                    return

                try:
                    if disable_ipv6 and not previous_state.get("ipv6_was_disabled"):
                        self.ipv6_disable(refresh_ui=False)

                    if vpn_profile:
                        vpn_ok = self.vpn_switch(
                            vpn_profile,
                            ensure_sudo=False,
                            refresh_ui=False,
                            force_reconnect=True,
                        )
                        if not (
                            vpn_ok and self.wait_for_vpn_ready(timeout=12, interval=0.5)
                        ):
                            rollback_and_abort(
                                "VPN-Start fehlgeschlagen – Stealthmodus abgebrochen und Vorzustand wiederhergestellt."
                            )
                            return

                    if enable_killswitch:
                        if not self.killswitch_is_on() and not self.killswitch_on(
                            async_mode=False,
                            ensure_sudo=False,
                            refresh_ui=False,
                        ):
                            rollback_and_abort(
                                "Killswitch konnte nicht aktiviert werden – Stealthmodus abgebrochen und Vorzustand wiederhergestellt."
                            )
                            return

                    if dns_profile and dns_profile in DNS_PROFILES:
                        self._set_dns_choice_safe(dns_profile)
                        if not self.apply_dns(
                            choice=dns_profile,
                            refresh_ui=False,
                            update_choice_var=False,
                        ):
                            rollback_and_abort(
                                "DNS-Profil konnte nicht gesetzt werden – Stealthmodus abgebrochen und Vorzustand wiederhergestellt."
                            )
                            return

                    if tx is not None:
                        tx.commit()
                    self.isp_stealth_runtime_state = previous_state
                    self.after(0, lambda: self._set_stealth_enabled_state(True))
                    self.log_safe(
                        "ISP-Stealthmodus per Autostart aktiviert."
                        if from_startup
                        else "ISP-Stealthmodus aktiviert."
                    )
                    self.update_ipv6_controls_state_safe()
                    self.refresh_safe()
                except Exception as e:
                    rollback_and_abort(f"Fehler beim Aktivieren des Stealthmodus: {e}")
                finally:
                    self.isp_stealth_running = False

        threading.Thread(target=job, daemon=True).start()

    def isp_stealth_off(self):
        if self.isp_stealth_running:
            self.log("ISP-Stealthmodus läuft bereits.")
            return
        if not self.ensure_sudo_session("ISP-Stealthmodus deaktivieren"):
            return

        previous_state = dict(
            self.isp_stealth_runtime_state or self._capture_stealth_runtime_state()
        )
        current_state = self._capture_stealth_runtime_state()

        def job():
            with self.isp_stealth_lock:
                self.isp_stealth_running = True
                tx = self._new_rollback_transaction("ISP-Stealthmodus AUS")
                if tx is not None:
                    tx.add_rollback(
                        "Aktiven Stealth-Zustand wiederherstellen",
                        lambda: self._restore_stealth_runtime_state(
                            current_state, refresh_ui=False
                        ),
                    )
                try:
                    self._restore_stealth_runtime_state(
                        previous_state, refresh_ui=False
                    )
                    if tx is not None:
                        tx.commit()
                    self.isp_stealth_runtime_state = None
                    self.after(0, lambda: self._set_stealth_enabled_state(False))
                    self.log_safe("ISP-Stealthmodus deaktiviert.")
                    self.update_ipv6_controls_state_safe()
                    self.refresh_safe()
                except Exception as e:
                    rollback_ok = True
                    if tx is not None:
                        rollback_ok = tx.rollback()
                    self.log_safe(f"Fehler beim Deaktivieren des Stealthmodus: {e}")
                    if not rollback_ok:
                        self.log_safe(
                            "Rollback unvollständig – bitte Zustand manuell prüfen (VPN/Killswitch/DNS/IPv6)."
                        )
                    self.refresh_safe()
                finally:
                    self.isp_stealth_running = False

        threading.Thread(target=job, daemon=True).start()

    def hdd_slot_configured(self, slot):
        return hdd_controller.hdd_slot_configured(slot)

    def _hdd_find_device_by_uuid(self, device_uuid):
        return hdd_controller.hdd_find_device_by_uuid(self, device_uuid)

    def _hdd_resolve_slot_device(self, slot):
        return hdd_controller.hdd_resolve_slot_device(self, slot)

    def scan_hdds(self):
        if not has_cmd("lsblk"):
            return []
        rc, out = run_cmd(
            [
                "lsblk",
                "-J",
                "-p",
                "-b",
                "-o",
                "NAME,PATH,SIZE,FSTYPE,LABEL,UUID,MOUNTPOINT,MODEL,TYPE",
            ],
            timeout=10,
        )
        if rc != 0 or not out.strip():
            return []
        try:
            data = json.loads(out)
        except Exception:
            return []

        detected = []

        def walk(dev):
            dtype = (dev.get("type") or "").strip()
            path = (dev.get("path") or "").strip()
            fstype = (dev.get("fstype") or "").strip()
            mountpoint = (dev.get("mountpoint") or "").strip()
            label = (dev.get("label") or "").strip()

            def iter_descendants(node):
                for child in node.get("children") or []:
                    yield child
                    yield from iter_descendants(child)

            include_direct = False
            if dtype in ("part", "lvm") and path:
                include_direct = True
            elif dtype == "disk" and path:
                include_direct = bool((fstype or "").strip() or mountpoint or label)

            if include_direct:
                try:
                    size_bytes = (
                        int(dev.get("size")) if dev.get("size") is not None else None
                    )
                except (TypeError, ValueError):
                    size_bytes = None
                if size_bytes is None or size_bytes < 1073741824:
                    for child in dev.get("children") or []:
                        walk(child)
                    return

                is_luks = (fstype or "").lower() == "crypto_luks"
                luks_mapper_path = ""
                luks_mapper_name = ""
                luks_unlocked = False
                effective_label = label
                effective_fstype = fstype
                effective_uuid = (dev.get("uuid") or "").strip()
                effective_mountpoint = mountpoint

                descendants = list(iter_descendants(dev))
                if is_luks:
                    for child in descendants:
                        child_path = (child.get("path") or "").strip()
                        if (
                            child_path.startswith("/dev/mapper/")
                            and not luks_mapper_path
                        ):
                            luks_mapper_path = child_path
                            luks_mapper_name = os.path.basename(child_path)
                            luks_unlocked = True
                        child_mountpoint = (child.get("mountpoint") or "").strip()
                        if child_mountpoint and not effective_mountpoint:
                            effective_mountpoint = child_mountpoint
                        child_label = (child.get("label") or "").strip()
                        if child_label and not effective_label:
                            effective_label = child_label
                        child_fstype = (child.get("fstype") or "").strip()
                        if (
                            child_fstype
                            and child_fstype.lower() != "crypto_luks"
                            and (
                                not effective_fstype
                                or effective_fstype.lower() == "crypto_luks"
                            )
                        ):
                            effective_fstype = child_fstype
                        child_uuid = (child.get("uuid") or "").strip()
                        if child_uuid and not effective_uuid:
                            effective_uuid = child_uuid
                else:
                    for child in descendants:
                        child_mountpoint = (child.get("mountpoint") or "").strip()
                        if child_mountpoint and not effective_mountpoint:
                            effective_mountpoint = child_mountpoint
                        child_label = (child.get("label") or "").strip()
                        if child_label and not effective_label:
                            effective_label = child_label

                is_mounted = bool(effective_mountpoint)
                if is_luks:
                    status = (
                        "gemountet"
                        if is_mounted
                        else (
                            "verschlüsselt/offen"
                            if luks_unlocked
                            else "verschlüsselt/gesperrt"
                        )
                    )
                else:
                    status = "gemountet" if is_mounted else "nicht gemountet"

                detected.append(
                    {
                        "device": path,
                        "label": effective_label,
                        "fstype": effective_fstype,
                        "uuid": effective_uuid,
                        "mountpoint": effective_mountpoint,
                        "model": (dev.get("model") or "").strip(),
                        "size_bytes": size_bytes,
                        "mounted": is_mounted,
                        "status": status,
                        "is_luks": is_luks,
                        "luks_locked": is_luks and not luks_unlocked,
                        "luks_unlocked": luks_unlocked,
                        "luks_mapper_path": luks_mapper_path,
                        "luks_mapper_name": luks_mapper_name,
                        "type": dtype,
                    }
                )
            for child in dev.get("children") or []:
                walk(child)

        for root in data.get("blockdevices", []):
            walk(root)

        return detected

    def _slot_for_device(self, device):
        return hdd_controller.slot_for_device(self, device)

    def _refresh_hdd_table(self, keep_selection=True):
        hdd_controller.refresh_hdd_table(self, keep_selection=keep_selection)

    def _slot_display_name(self, idx):
        return hdd_controller.slot_display_name(self, idx)

    def update_hdd_slot_button_labels(self):
        hdd_controller.update_hdd_slot_button_labels(self)

    def update_hdd_buttons(self):
        hdd_controller.update_hdd_buttons(self)

    def on_hdd_name_focus_in(self, _event=None):
        hdd_controller.on_hdd_name_focus_in(self, _event)

    def on_hdd_name_focus_out(self, _event=None):
        hdd_controller.on_hdd_name_focus_out(self, _event)

    def on_hdd_name_changed(self, *_args):
        hdd_controller.on_hdd_name_changed(self, *_args)

    def set_hdd_name_from_model(self, name):
        hdd_controller.set_hdd_name_from_model(self, name)

    def scan_hdds_async(self):
        def job():
            disks = self.scan_hdds()

            def done():
                self.detected_disks = disks
                self.detected_disks_by_device = {
                    (d.get("device") or ""): d for d in disks if d.get("device")
                }
                self._refresh_hdd_table(keep_selection=True)
                self.hdd_refresh_async()
                self.log(f"Festplatten eingelesen: {len(disks)} Einträge.")

            self.after(0, done)

        threading.Thread(target=job, daemon=True).start()

    def on_hdd_table_select(self, _event=None):
        hdd_controller.on_hdd_table_select(self, _event)

    def assign_selected_disk_to_slot(self):
        hdd_controller.assign_selected_disk_to_slot(self)

    def clear_slot_assignment(self, confirm=False):
        hdd_controller.clear_slot_assignment(self, confirm=confirm)

    def _hdd_sanitize_mapper_name(self, device):
        return hdd_controller.hdd_sanitize_mapper_name(device)

    def _hdd_list_mappers(self):
        return hdd_controller.hdd_list_mappers()

    def _hdd_probe_live_device(self, device):
        return hdd_controller.hdd_probe_live_device(device)

    def _hdd_luks_info_for_device(self, device):
        return hdd_controller.hdd_luks_info_for_device(self, device)

    def _hdd_find_mount_source(self, mapper_path):
        return hdd_controller.hdd_find_mount_source(mapper_path)

    def _hdd_is_mounted(self, device, mountpoint):
        return hdd_controller.hdd_is_mounted(self, device, mountpoint)

    def _hdd_usage_for_mountpoint(self, mountpoint):
        return hdd_controller.hdd_usage_for_mountpoint(mountpoint)

    def _hdd_collect_slot_state(self, idx):
        return hdd_controller.hdd_collect_slot_state(self, idx)

    def _hdd_apply_slot_lights(self):
        hdd_controller.hdd_apply_slot_lights(self)

    def hdd_refresh_async(self):
        hdd_controller.hdd_refresh_async(self)

    def _hdd_render_usage_bar(self, usage):
        hdd_controller.hdd_render_usage_bar(self, usage)

    def _hdd_render_info(self, state):
        hdd_controller.hdd_render_info(self, state)

    def hdd_info_async(self, idx):
        hdd_controller.hdd_info_async(self, idx)

    def hdd_mount_async(self, idx):
        hdd_controller.hdd_mount_async(self, idx)

    def hdd_unmount_async(self, idx):
        hdd_controller.hdd_unmount_async(self, idx)

    def preflight_recheck(self):
        self.log(self.tr("preflight_log_started", "Preflight-Prüfung gestartet..."))
        self.preflight_dependencies(manual_trigger=True)
        self.log(self.tr("preflight_log_finished", "Preflight-Prüfung abgeschlossen."))

    def _log_preflight_tool_state(self, label, ok, path=""):
        if ok and path:
            self.log(
                self.tr("preflight_tool_ok_path", "{label}: OK ({path})").format(
                    label=label, path=path
                )
            )
        elif ok:
            self.log(self.tr("preflight_tool_ok", "{label}: OK").format(label=label))
        else:
            self.log(
                self.tr("preflight_tool_missing", "{label}: FEHLT").format(label=label)
            )

    def _log_strict_protection_state(self, prefix="Strenger Schutz"):
        strict_var = self.__dict__.get("strict_protection_var")
        if strict_var is not None:
            try:
                enabled = bool(strict_var.get())
            except Exception:
                enabled = bool(getattr(self, "danger_guard_enabled", True))
        else:
            enabled = bool(getattr(self, "danger_guard_enabled", True))
        self.log(
            self.tr(
                "preflight_strict_protection_state",
                "{prefix}: {state}",
            ).format(
                prefix=prefix,
                state=(
                    self.tr(
                        "preflight_strict_protection_on",
                        "AKTIV (Bestätigungswort AUSFUEHREN erforderlich)",
                    )
                    if enabled
                    else self.tr(
                        "preflight_strict_protection_off",
                        "AUS (nur Standard-Bestätigung)",
                    )
                ),
            )
        )

    def _log_preflight_manual_summary(self, checks):
        ordered_tools = [
            ("traceroute", "traceroute"),
            ("mtr", "mtr"),
            ("nmap", "nmap"),
            ("nmcli", "nmcli"),
            ("resolvectl", "resolvectl"),
            ("wg", "wg"),
            ("wg-quick", "wg-quick"),
            ("ufw", "ufw"),
            ("ss", "ss"),
            ("iftop", "iftop"),
            ("htop", "htop"),
            ("tshark", "tshark"),
            ("dumpcap", "dumpcap"),
            ("Wireshark GUI", "wireshark"),
            ("capinfos", "capinfos"),
            ("lynis", "lynis"),
            ("rkhunter", "rkhunter"),
            ("whois", "whois"),
        ]
        for label, key in ordered_tools:
            self._log_preflight_tool_state(
                label, checks.get(key, False), get_cmd_path(key)
            )

        speedtest_path = get_cmd_path("speedtest") or get_cmd_path("speedtest-cli")
        self._log_preflight_tool_state(
            "Speedtest CLI", checks.get("speedtest", False), speedtest_path
        )
        self._log_preflight_tool_state(
            "Speedtest Modul", checks.get("speedtest_module", False)
        )

        self._log_preflight_tool_state(
            "ClamAV Scan", checks.get("clamav", False), get_cmd_path("clamscan")
        )
        self._log_preflight_tool_state(
            "ClamAV Update",
            checks.get("freshclam", False),
            get_cmd_path("freshclam"),
        )
        self._log_preflight_tool_state("Pillow (Log-Bild)", checks.get("pillow", False))
        if not checks.get("pillow", False) and PIL_IMPORT_ERROR:
            self.log("Pillow Detail: " + PIL_IMPORT_ERROR)
        self._log_strict_protection_state("Strenger Schutz")

    def preflight_dependencies(self, manual_trigger=False):
        vscode_desktop = self._find_vscode_desktop_file()
        wireshark_launcher = self._resolve_wireshark_launcher()
        wireshark_available = bool(wireshark_launcher.get("target"))
        speedtest_module_ok, speedtest_module_detail = check_speedtest_module_available(
            app_file=__file__,
            repo_root=NETMON_REPO_ROOT,
            search_paths=sys.path,
        )
        checks = {
            "sudo": has_cmd("sudo"),
            "ip": has_cmd("ip"),
            "ping": has_cmd("ping"),
            "curl": has_cmd("curl"),
            "systemctl": has_cmd("systemctl"),
            "wg": has_cmd("wg"),
            "wg-quick": has_cmd("wg-quick"),
            "ufw": has_cmd("ufw"),
            "nmcli": has_cmd("nmcli"),
            "dig": has_cmd("dig"),
            "nslookup": has_cmd("nslookup"),
            "tracepath": has_cmd("tracepath"),
            "traceroute": has_cmd("traceroute"),
            "ss": has_cmd("ss"),
            "resolvectl": has_cmd("resolvectl"),
            "find": has_cmd("find"),
            "install": has_cmd("install"),
            "mtr": has_cmd("mtr"),
            "nmap": has_cmd("nmap"),
            "speedtest": has_cmd("speedtest") or has_cmd("speedtest-cli"),
            "speedtest_module": speedtest_module_ok,
            "iftop": has_cmd("iftop"),
            "htop": has_cmd("htop"),
            "tshark": has_cmd("tshark"),
            "dumpcap": has_cmd("dumpcap"),
            "wireshark": wireshark_available,
            "capinfos": has_cmd("capinfos"),
            "lynis": has_cmd("lynis"),
            "rkhunter": has_cmd("rkhunter"),
            "clamav": has_cmd("clamscan"),
            "freshclam": has_cmd("freshclam"),
            "pillow": Image is not None and ImageTk is not None,
            "pypdf2": PyPDF2 is not None,
            "docx": docx is not None,
            "whois": has_cmd("whois"),
            "lsblk": has_cmd("lsblk"),
            "findmnt": has_cmd("findmnt"),
            "df": has_cmd("df"),
            "mount": has_cmd("mount"),
            "umount": has_cmd("umount"),
            "mkdir": has_cmd("mkdir"),
            "udisksctl": has_cmd("udisksctl"),
            "cryptsetup": has_cmd("cryptsetup"),
            "tail": has_cmd("tail"),
            "gtk-launch": has_cmd("gtk-launch"),
            "code": has_cmd("code"),
            "codium": has_cmd("codium"),
            "code-insiders": has_cmd("code-insiders"),
            "vscode_desktop": bool(vscode_desktop),
        }
        checks["wireshark_launcher"] = wireshark_launcher
        self.dep_status = checks

        if not self.security_clamav_running and not self.security_clamav_update_running:
            if checks.get("clamav"):
                sig_status = "aktuell" if checks.get("freshclam") else "unbekannt"
                self.update_clamav_panel(status="bereit", signatures=sig_status)
            else:
                sig_status = (
                    "nicht aktualisierbar"
                    if not checks.get("freshclam")
                    else "unbekannt"
                )
                self.update_clamav_panel(
                    status="nicht installiert", signatures=sig_status
                )

        self.sudo_nopasswd = check_sudo_nopasswd()
        sudo_ready = self.sudo_nopasswd

        if manual_trigger:
            self._log_preflight_manual_summary(checks)
            if not checks.get("speedtest_module") and speedtest_module_detail:
                self.log(
                    self.tr(
                        "preflight_speedtest_module_detail",
                        "Speedtest Modul Detail: {detail}",
                    ).format(detail=speedtest_module_detail)
                )
            if not (
                checks.get("tshark")
                and checks.get("dumpcap")
                and checks.get("wireshark")
            ):
                self.show_wireshark_install_hint()
            if not checks.get("lynis"):
                self.show_lynis_install_hint()
            if not checks.get("rkhunter"):
                self.show_rkhunter_install_hint()
        else:
            required = ["sudo", "ip", "ping", "curl", "wg", "wg-quick", "ufw"]
            missing_required = [name for name in required if not checks.get(name)]
            if missing_required:
                self.log(
                    self.tr(
                        "preflight_missing_required_tools",
                        "Fehlende Pflicht-Tools: {tools}",
                    ).format(tools=", ".join(missing_required))
                )
                self.log(
                    self.tr(
                        "preflight_install_required_hint",
                        "Install (Debian/RPi): sudo apt install -y wireguard-tools ufw iproute2 iputils-ping curl",
                    )
                )
            else:
                self.log(
                    self.tr(
                        "preflight_required_tools_ok", "Preflight: Pflicht-Tools OK."
                    )
                )

            if not sudo_ready:
                self.log(
                    self.tr(
                        "preflight_warning_no_sudo",
                        "WARNUNG: Keine gültige Sudo-Authentifizierung aktiv – Aktionen können fehlschlagen.",
                    )
                )
            else:
                self.log(
                    self.tr(
                        "preflight_sudo_session_valid",
                        "Preflight: sudo-Session gültig (sudo -n true).",
                    )
                )

            optional_missing = []
            if not checks.get("nmcli"):
                optional_missing.append("nmcli (network-manager)")
            if not (checks.get("dig") or checks.get("nslookup")):
                optional_missing.append("dnsutils (dig/nslookup)")
            if not (checks.get("tracepath") or checks.get("traceroute")):
                optional_missing.append("tracepath/traceroute")
            if not checks.get("traceroute"):
                optional_missing.append("traceroute (Hops)")
            if not checks.get("ss"):
                optional_missing.append("ss (iproute2)")
            if not checks.get("mtr"):
                optional_missing.append("mtr")
            if not checks.get("nmap"):
                optional_missing.append("nmap")
            if not checks.get("speedtest"):
                optional_missing.append("speedtest/speedtest-cli")
            if not checks.get("speedtest_module"):
                optional_missing.append(
                    "Speedtest Python-Modul (netmon/modules/speedtest.py)"
                )
            if not checks.get("iftop"):
                optional_missing.append("iftop")
            if not checks.get("htop"):
                optional_missing.append("htop")
            if not (checks.get("tshark") or checks.get("dumpcap")):
                optional_missing.append("tshark/dumpcap (Packet Capture)")
            if not checks.get("wireshark"):
                optional_missing.append("wireshark (GUI öffnen)")
            if not checks.get("capinfos"):
                optional_missing.append("capinfos (Capture-Metadaten)")
            if not checks.get("lynis"):
                optional_missing.append("lynis")
            if not checks.get("rkhunter"):
                optional_missing.append("rkhunter")
            if not checks.get("clamav"):
                optional_missing.append("clamav (clamscan)")
            if not checks.get("pillow"):
                optional_missing.append("python3-pil (Log-Bild/JPG)")
            if not checks.get("pypdf2"):
                optional_missing.append("python3-pypdf2 (PDF-Import)")
            if not checks.get("docx"):
                optional_missing.append("python3-docx (DOCX-Import)")
            if not checks.get("whois"):
                optional_missing.append("whois")
            if not checks.get("lsblk"):
                optional_missing.append("lsblk (HDD)")
            if not checks.get("findmnt"):
                optional_missing.append("findmnt (HDD)")
            if not checks.get("df"):
                optional_missing.append("df (HDD-Usage)")
            if not checks.get("mount") or not checks.get("umount"):
                optional_missing.append("mount/umount (HDD Aktionen)")
            if not checks.get("udisksctl"):
                optional_missing.append("udisksctl (HDD Fallback)")
            if not checks.get("cryptsetup"):
                optional_missing.append("cryptsetup (LUKS)")
            if not checks.get("tail"):
                optional_missing.append("tail (Security Log)")
            if not (
                checks.get("vscode_desktop")
                or checks.get("code")
                or checks.get("codium")
                or checks.get("code-insiders")
            ):
                optional_missing.append("VS Code Starter (desktop/code/codium)")
            if optional_missing:
                self.log(
                    self.tr(
                        "preflight_optional_tools_missing",
                        "Optionale Tools fehlen: {tools}",
                    ).format(tools=", ".join(optional_missing))
                )
            if not (
                checks.get("tshark")
                and checks.get("dumpcap")
                and checks.get("wireshark")
            ):
                self.show_wireshark_install_hint()
            if not checks.get("lynis"):
                self.show_lynis_install_hint()
            if not checks.get("rkhunter"):
                self.show_rkhunter_install_hint()

            self.log("Speedtest CLI: " + ("OK" if checks.get("speedtest") else "FEHLT"))
            self.log(
                "Speedtest Modul: "
                + ("OK" if checks.get("speedtest_module") else "FEHLT")
            )
            self._log_strict_protection_state("Strenger Schutz")
            if checks.get("speedtest") and not checks.get("speedtest_module"):
                self.log(
                    self.tr(
                        "preflight_speedtest_cli_detected", "Speedtest CLI erkannt."
                    )
                )
                self.log(
                    self.tr(
                        "preflight_speedtest_module_detail",
                        "Speedtest Modul Detail: {detail}",
                    ).format(detail=speedtest_module_detail)
                )

            self.log("ClamAV Scan: " + ("OK" if checks.get("clamav") else "FEHLT"))
            self.log("ClamAV Update: " + ("OK" if checks.get("freshclam") else "FEHLT"))
            self.log(
                "Packet Capture: "
                + ("OK" if (checks.get("tshark") or checks.get("dumpcap")) else "FEHLT")
            )
            self.log("Wireshark GUI: " + ("OK" if checks.get("wireshark") else "FEHLT"))
            self.log("Lynis: " + ("OK" if checks.get("lynis") else "FEHLT"))
            self.log("rkhunter: " + ("OK" if checks.get("rkhunter") else "FEHLT"))
            self.log(
                "Pillow (Log-Bild): " + ("OK" if checks.get("pillow") else "FEHLT")
            )
            if not checks.get("pillow") and PIL_IMPORT_ERROR:
                self.log("Pillow Detail: " + PIL_IMPORT_ERROR)
            self.log("PDF Import: " + ("OK" if checks.get("pypdf2") else "FEHLT"))
            if not checks.get("pypdf2") and PYPDF2_IMPORT_ERROR:
                self.log("PDF Import Detail: " + PYPDF2_IMPORT_ERROR)
            self.log("DOCX Import: " + ("OK" if checks.get("docx") else "FEHLT"))
            if not checks.get("docx") and DOCX_IMPORT_ERROR:
                self.log("DOCX Import Detail: " + DOCX_IMPORT_ERROR)
            guided_install_tools = [
                ("traceroute", "traceroute"),
                ("mtr", "mtr"),
                ("htop", "htop"),
                ("iftop", "iftop"),
                ("nmap", "nmap"),
                ("speedtest", "speedtest"),
                ("wireshark", "tshark"),
                ("lynis", "lynis"),
                ("rkhunter", "rkhunter"),
                ("clamav", "clamav"),
                ("python3-pil", "pillow"),
                ("python3-pypdf2", "pypdf2"),
                ("python3-docx", "docx"),
                ("whois", "whois"),
                ("cryptsetup", "cryptsetup"),
            ]
            guided_missing = [
                package
                for package, check_key in guided_install_tools
                if not checks.get(check_key)
            ]
            if guided_missing:
                self.log(
                    self.tr(
                        "preflight_missing_programs_detected",
                        "Fehlende Programme erkannt:",
                    )
                )
                for name in guided_missing:
                    self.log(name)
                self.log(
                    self.tr(
                        "preflight_install_command_label",
                        "Installationsbefehl (kopieren und im Terminal ausführen):",
                    )
                )
                self.log(f"sudo apt install {' '.join(guided_missing)}")
                self.log(
                    self.tr(
                        "preflight_install_steps_title",
                        "So installieren Sie die Programme:",
                    )
                )
                self.log(
                    self.tr(
                        "preflight_install_step_1",
                        "1. Installationsbefehl im Log markieren",
                    )
                )
                self.log(self.tr("preflight_install_step_2", "2. Strg+C drücken"))
                self.log(self.tr("preflight_install_step_3", "3. Terminal öffnen"))
                self.log(self.tr("preflight_install_step_4", "4. Strg+V drücken"))
                self.log(self.tr("preflight_install_step_5", "5. Enter drücken"))

        self.set_feature_enabled(
            "wg", checks.get("wg", False) and checks.get("wg-quick", False)
        )
        self.set_feature_enabled(
            "wg_profiles",
            checks.get("install", False) and checks.get("sudo", False),
        )
        self.set_feature_enabled(
            "ufw", checks.get("ufw", False) and checks.get("sudo", False)
        )
        self.set_feature_enabled("nmcli", checks.get("nmcli", False))
        self.set_feature_enabled(
            "dot",
            checks.get("sudo", False)
            and checks.get("install", False)
            and checks.get("systemctl", False)
            and checks.get("resolvectl", False),
        )
        self.set_feature_enabled(
            "dns_consistency", checks.get("dig", False) or checks.get("nslookup", False)
        )
        self.set_feature_enabled(
            "trace_log",
            checks.get("tracepath", False) or checks.get("traceroute", False),
        )
        self.set_feature_enabled("traceroute", checks.get("traceroute", False))
        self.set_feature_enabled("ss", checks.get("ss", False))
        self.set_feature_enabled("mtr", checks.get("mtr", False))
        self.set_feature_enabled("nmap", checks.get("nmap", False))
        self.set_feature_enabled(
            "speedtest",
            checks.get("speedtest", False) and checks.get("speedtest_module", False),
        )
        self.set_feature_enabled("iftop", checks.get("iftop", False))
        self.set_feature_enabled("htop", checks.get("htop", False))
        self.set_feature_enabled("packet_capture_ui", checks.get("ip", False))
        self.set_feature_enabled(
            "packet_capture",
            checks.get("tshark", False) or checks.get("dumpcap", False),
        )
        self.set_feature_enabled("wireshark_open", checks.get("wireshark", False))
        self.set_feature_enabled("lynis", checks.get("lynis", False))
        self.set_feature_enabled("rkhunter", checks.get("rkhunter", False))
        self.set_feature_enabled("clamav", checks.get("clamav", False))
        self.set_feature_enabled("clamav_scan", checks.get("clamav", False))
        self.set_feature_enabled("clamav_update", checks.get("freshclam", False))
        self.set_feature_enabled(
            "hdd_manage",
            checks.get("lsblk", False)
            and checks.get("mount", False)
            and checks.get("umount", False),
        )
        self.set_feature_enabled(
            "security",
            checks.get("ufw", False) and checks.get("tail", False),
        )
        self.set_feature_enabled(
            "vscode",
            checks.get("vscode_desktop", False)
            or checks.get("code", False)
            or checks.get("codium", False)
            or checks.get("code-insiders", False),
        )

        self._update_clamav_findings_buttons_state()

    # ---------- IPv6 ----------
    def ipv6_is_disabled(self):
        checks = [
            "net.ipv6.conf.all.disable_ipv6",
            "net.ipv6.conf.default.disable_ipv6",
        ]
        values = []

        for key in checks:
            rc, out = run_cmd(["sysctl", "-n", key], timeout=3)
            if rc == 0:
                value = (out or "").strip()
                if value in {"0", "1"}:
                    values.append(value == "1")
                    continue
            proc_path = f"/proc/sys/{key.replace('.', '/')}"
            try:
                with open(proc_path, "r", encoding="utf-8") as f:
                    value = f.read().strip()
                if value in {"0", "1"}:
                    values.append(value == "1")
            except Exception:
                continue

        if values:
            # Konservativ: IPv6 gilt als deaktiviert, wenn all/default deaktiviert sind.
            return all(values)
        return False

    def ipv6_runtime_available_on_iface(self, iface):
        name = (iface or "").strip()
        if not name or self.ipv6_is_disabled():
            return False
        rc, out = run_cmd(
            ["ip", "-6", "-o", "addr", "show", "dev", name, "scope", "global"],
            timeout=4,
        )
        if rc != 0:
            return False
        for line in (out or "").splitlines():
            if " inet6 " in line:
                return True
        return False

    def update_ipv6_controls_state(self):
        active_profile = (self.active_wg_profile() or "").strip()
        selected_profile = ""
        if hasattr(self, "isp_stealth_vpn_profile"):
            selected_profile = (self.isp_stealth_vpn_profile.get() or "").strip()
        if not selected_profile and hasattr(self, "isp_stealth_vpn_choice"):
            selected_profile = self._privacy_vpn_profile_from_value(
                self.isp_stealth_vpn_choice.get()
            )

        buttons_state = "normal"
        system_ipv6_enabled = not self.ipv6_is_disabled()
        active_iface = (self.get_wg_iface() or "").strip()
        active_supports_ipv6 = bool(
            active_profile and self.profile_supports_ipv6(active_profile)
        )
        active_runtime_ipv6 = bool(
            active_iface and self.ipv6_runtime_available_on_iface(active_iface)
        )

        if not active_profile:
            vpn_status = "kein aktives Profil"
        elif active_runtime_ipv6 or (active_iface and system_ipv6_enabled):
            vpn_status = "nutzt IPv6"
        elif active_supports_ipv6 and not system_ipv6_enabled:
            vpn_status = "erwartet IPv6, System deaktiviert"
        else:
            vpn_status = "nutzt nur IPv4"

        system_status = "aktiv" if system_ipv6_enabled else "deaktiviert"
        ipv6_hint_lines = [
            "IPv6 Status:",
            f"System:        {system_status}",
            f"VPN-Profil:    {vpn_status}",
        ]

        ipv6_hint = "\n".join(ipv6_hint_lines)

        for widget_name in ("btn_ipv6_on", "btn_ipv6_off"):
            widget = getattr(self, widget_name, None)
            if widget is None:
                continue
            try:
                widget.configure(state=buttons_state)
            except Exception:
                try:
                    widget["state"] = buttons_state
                except Exception:
                    pass

        if hasattr(self, "ipv6_controls_hint"):
            if not active_profile:
                hint_color = "#666666"
            elif active_runtime_ipv6:
                hint_color = "#2e7d32"
            else:
                hint_color = "#c62828"
            self.ipv6_controls_hint.config(text=ipv6_hint, foreground=hint_color)

        privacy_profile = active_profile or selected_profile
        privacy_enabled = True
        privacy_hint = ""
        if privacy_profile and not self.profile_supports_ipv6(privacy_profile):
            privacy_enabled = False
            if active_profile:
                privacy_hint = "Aktives VPN-Profil nutzt nur IPv4"
            else:
                privacy_hint = "Gewaehltes Privacy-VPN-Profil nutzt nur IPv4"

        privacy_state = "normal" if privacy_enabled else "disabled"
        widget = getattr(self, "cb_privacy_ipv6", None)
        if widget is not None:
            try:
                widget.configure(state=privacy_state)
            except Exception:
                try:
                    widget["state"] = privacy_state
                except Exception:
                    pass
        if hasattr(self, "privacy_ipv6_hint"):
            self.privacy_ipv6_hint.config(text=privacy_hint)

    def ipv6_disable(self, refresh_ui=True):
        ok = True
        cmds = [
            ["sudo", "-n", "sysctl", "-w", "net.ipv6.conf.all.disable_ipv6=1"],
            ["sudo", "-n", "sysctl", "-w", "net.ipv6.conf.default.disable_ipv6=1"],
            ["sudo", "-n", "sysctl", "-w", "net.ipv6.conf.lo.disable_ipv6=1"],
        ]
        for c in cmds:
            rc, out = run_cmd(c, timeout=6)
            if rc != 0:
                ok = False
                if out:
                    self.log_safe(out)
        self.log_safe("IPv6: deaktiviert (sysctl; nicht persistent nach Reboot).")
        if refresh_ui:
            self.refresh_safe()
        return ok

    def ipv6_enable(self, refresh_ui=True):
        ok = True
        cmds = [
            ["sudo", "-n", "sysctl", "-w", "net.ipv6.conf.all.disable_ipv6=0"],
            ["sudo", "-n", "sysctl", "-w", "net.ipv6.conf.default.disable_ipv6=0"],
            ["sudo", "-n", "sysctl", "-w", "net.ipv6.conf.lo.disable_ipv6=0"],
        ]
        for c in cmds:
            rc, out = run_cmd(c, timeout=6)
            if rc != 0:
                ok = False
                if out:
                    self.log_safe(out)
        self.log_safe("IPv6: aktiviert (sysctl; nicht persistent nach Reboot).")
        if refresh_ui:
            self.refresh_safe()
        return ok

    # ---------- Network detect ----------
    def get_default_iface(self):
        _, out = run_cmd(["ip", "route", "show", "default"], timeout=3)
        parts = out.split()
        for i, p in enumerate(parts):
            if p == "dev" and i + 1 < len(parts):
                return parts[i + 1]
        return "eth0"

    def nm_active_connections(self):
        _, out = run_cmd(
            ["nmcli", "-t", "-f", "NAME,DEVICE", "con", "show", "--active"], timeout=5
        )
        res = []
        for line in (out or "").splitlines():
            if ":" in line:
                n, d = line.split(":", 1)
                if n and d:
                    res.append((n, d))
        return res

    def nm_active_conn_for_dev(self, dev: str):
        for n, d in self.nm_active_connections():
            if d == dev:
                return n
        return None

    def check_internet(self):
        try:
            p = subprocess.run(
                ["ping", "-c", "1", "-W", "1", "1.1.1.1"],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=2,
                check=False,
                text=True,
            )
            return p.returncode == 0
        except Exception:
            return False

    def _internet_transition_active(self):
        if self._dns_runtime_transition_active():
            return True
        active_operation = str(self.__dict__.get("network_active_operation") or "")
        active_operation = active_operation.strip().lower()
        return bool(
            active_operation
            and any(token in active_operation for token in ("vpn", "internet", "dns"))
        )

    def _update_internet_indicator_state(self, inet_ok: bool):
        if self._internet_transition_active():
            self.__dict__["internet_indicator_state"] = "transition"
            return "transition"

        if inet_ok:
            self.__dict__["internet_fail_streak"] = 0
            self.__dict__["internet_indicator_state"] = "ok"
            return "ok"

        fail_streak = int(self.__dict__.get("internet_fail_streak", 0)) + 1
        self.__dict__["internet_fail_streak"] = fail_streak
        state = "warning" if fail_streak == 1 else "error"
        self.__dict__["internet_indicator_state"] = state
        return state

    def internet_disable(self, interactive_guard=True):
        if not self._try_begin_network_operation("Internet AUS"):
            return False
        try:
            if interactive_guard and not self._confirm_danger_action(
                "network:internet_off",
                title=self.tr("internet_off_confirm_title", "Internet AUS bestätigen"),
                prompt=self.tr(
                    "internet_off_confirm_message",
                    "Diese Aktion deaktiviert Netzwerk systemweit über NetworkManager.\n"
                    "Aktive Verbindungen werden getrennt.\n\n"
                    "Wirklich fortfahren?",
                ),
                cooldown_s=0,
                require_phrase="AUSFUEHREN",
            ):
                self.log("Internet AUS abgebrochen.")
                return False

            if not has_cmd("nmcli"):
                self.log("Internet AN/AUS nicht verfügbar: nmcli fehlt.")
                messagebox.showerror(
                    self.tr(
                        "networkmanager_missing_title",
                        "NetworkManager nicht verfügbar",
                    ),
                    self.tr(
                        "networkmanager_missing_message",
                        "nmcli (NetworkManager) ist nicht installiert.\n"
                        "Internet AN/AUS ist nicht möglich.",
                    ),
                    parent=self,
                )
                return False

            if not self.ensure_sudo_session("Internet deaktivieren"):
                return False

            state_coordinator = self._state_coordinator()
            if state_coordinator is not None:
                suspend_state = state_coordinator.capture_internet_suspend_state(
                    ufw_was_active=self.check_ufw_active(),
                    killswitch_was_active=self.killswitch_is_on(),
                    vpn_was_active=bool(self.get_wg_iface()),
                    vpn_profile=self.active_wg_profile() or "",
                    vpn_interface=self.get_wg_iface() or "",
                )
            else:
                suspend_state = {
                    "ufw_was_active": self.check_ufw_active(),
                    "killswitch_was_active": self.killswitch_is_on(),
                    "vpn_was_active": bool(self.get_wg_iface()),
                    "vpn_profile": self.active_wg_profile() or "",
                    "vpn_interface": self.get_wg_iface() or "",
                }
            self.internet_suspend_state = dict(suspend_state)
            self.log("Internet AUS: Vorzustand gespeichert.")

            if suspend_state["killswitch_was_active"]:
                self.log("Killswitch vorübergehend deaktiviert.")
                if not self.killswitch_off(
                    async_mode=False, ensure_sudo=False, refresh_ui=False
                ):
                    self.log(
                        "Internet AUS abgebrochen: Killswitch konnte nicht sicher deaktiviert werden."
                    )
                    self.internet_suspend_state = self._default_internet_suspend_state()
                    return False

            rc, out = run_cmd(["sudo", "-n", "nmcli", "networking", "off"], timeout=10)
            if rc == 0:
                self.log("Internet systemweit deaktiviert.")
                self.after(500, self.refresh)
                return True

            self.log(f"Internet AUS fehlgeschlagen: {out}")
            if suspend_state["killswitch_was_active"]:
                if self.killswitch_on(
                    async_mode=False, ensure_sudo=False, refresh_ui=False
                ):
                    self.log("Rollback: Killswitch wieder aktiviert.")
                else:
                    self.log(
                        "Rollback fehlgeschlagen: Killswitch konnte nach Internet-AUS-Fehler nicht wieder aktiviert werden."
                    )
            messagebox.showerror(
                self.tr("internet_off_failed_title", "Internet AUS fehlgeschlagen"),
                self.tr(
                    "internet_off_failed_message",
                    "Netzwerk konnte nicht deaktiviert werden:\n{details}",
                ).format(details=out),
                parent=self,
            )
            self.internet_suspend_state = self._default_internet_suspend_state()
            return False
        finally:
            self._end_network_operation()

    def internet_enable(self, interactive_guard=True):
        if interactive_guard and not self._confirm_danger_action(
            "network:internet_on",
            title=self.tr("internet_on_confirm_title", "Internet AN bestätigen"),
            prompt=self.tr(
                "internet_on_confirm_message",
                "Diese Aktion aktiviert Netzwerk systemweit über NetworkManager.\n"
                "Anschließend wird ein gespeicherter VPN/Killswitch-Zustand ggf. wiederhergestellt.\n\n"
                "Wirklich fortfahren?",
            ),
            cooldown_s=0,
            require_phrase="AUSFUEHREN",
        ):
            self.log("Internet AN abgebrochen.")
            return False

        if not has_cmd("nmcli"):
            self.log("Internet AN/AUS nicht verfügbar: nmcli fehlt.")
            messagebox.showerror(
                self.tr(
                    "networkmanager_missing_title",
                    "NetworkManager nicht verfügbar",
                ),
                self.tr(
                    "networkmanager_missing_message",
                    "nmcli (NetworkManager) ist nicht installiert.\n"
                    "Internet AN/AUS ist nicht möglich.",
                ),
                parent=self,
            )
            return False

        if not self.ensure_sudo_session("Internet aktivieren"):
            return False

        # Async ausführen, um nicht im UI zu blocken
        def job():
            self._internet_on_sync()

        self.run_async(job)
        return True

    def _default_internet_suspend_state(self):
        coordinator = self._state_coordinator()
        if coordinator is not None:
            return coordinator.default_internet_suspend_state()
        return {
            "ufw_was_active": False,
            "killswitch_was_active": False,
            "vpn_was_active": False,
            "vpn_profile": "",
            "vpn_interface": "",
        }

    def _capture_stealth_runtime_state(self):
        dns_choice_var = self.__dict__.get("dns_choice")
        dns_choice = "System (DHCP/auto)"
        if dns_choice_var is not None:
            dns_choice = (dns_choice_var.get() or "System (DHCP/auto)").strip()

        state_coordinator = self._state_coordinator()
        if state_coordinator is not None:
            return state_coordinator.capture_stealth_runtime_state(
                ipv6_was_disabled=self.ipv6_is_disabled(),
                killswitch_was_on=self.killswitch_is_on(),
                vpn_was_active=bool(self.active_wg_profile() or self.get_wg_iface()),
                vpn_profile=self.active_wg_profile() or "",
                dns_choice=dns_choice,
            )

        return {
            "ipv6_was_disabled": self.ipv6_is_disabled(),
            "killswitch_was_on": self.killswitch_is_on(),
            "vpn_was_active": bool(self.active_wg_profile() or self.get_wg_iface()),
            "vpn_profile": self.active_wg_profile() or "",
            "dns_choice": dns_choice or "System (DHCP/auto)",
        }

    def _restore_stealth_runtime_state(self, state, refresh_ui=False):
        state = state or {}
        previous_dns = (state.get("dns_choice") or "System (DHCP/auto)").strip()
        previous_vpn = (state.get("vpn_profile") or "").strip()
        previous_vpn_active = bool(state.get("vpn_was_active"))
        previous_killswitch = bool(state.get("killswitch_was_on"))
        previous_ipv6_disabled = bool(state.get("ipv6_was_disabled"))

        if not previous_killswitch and self.killswitch_is_on():
            self.killswitch_off(async_mode=False, ensure_sudo=False, refresh_ui=False)

        current_vpn_active = bool(self.active_wg_profile() or self.get_wg_iface())
        if previous_vpn_active and previous_vpn:
            if self.active_wg_profile() != self._base_wg_profile_name(previous_vpn):
                self.vpn_switch(
                    previous_vpn,
                    ensure_sudo=False,
                    refresh_ui=False,
                    force_reconnect=True,
                )
                self.wait_for_vpn_ready(timeout=12, interval=0.5)
        elif current_vpn_active:
            self.vpn_off(ensure_sudo=False, refresh_ui=False)

        if previous_ipv6_disabled:
            self.ipv6_disable(refresh_ui=False)
        else:
            self.ipv6_enable(refresh_ui=False)

        self._set_dns_choice_safe(previous_dns)
        if previous_dns in DNS_PROFILES:
            self.apply_dns(
                choice=previous_dns,
                refresh_ui=False,
                update_choice_var=False,
            )

        if previous_killswitch and not self.killswitch_is_on():
            self.killswitch_on(async_mode=False, ensure_sudo=False, refresh_ui=False)

        if refresh_ui:
            self.refresh_safe()

    def _schedule_refresh_burst(self, delays_ms=None):
        delays = delays_ms or (0, 1500, 3000)
        for delay in delays:
            self.after(delay, self.refresh)

    def _vpn_route_uses_iface(self, iface):
        name = (iface or "").strip()
        if not name:
            return False
        rc, out = run_cmd(["ip", "route", "show", "default"], timeout=3)
        if rc != 0:
            return False
        for line in (out or "").splitlines():
            if f" dev {name}" in f" {line} ":
                return True
        return False

    def _check_public_ip_quick(self):
        if not has_cmd("curl"):
            return False
        rc, out = run_cmd(
            ["curl", "-4", "-fsS", "--max-time", "3", "https://api.ipify.org"],
            timeout=4,
        )
        if rc != 0:
            return False
        ip = (out or "").strip()
        try:
            ipaddress.ip_address(ip)
            return True
        except Exception:
            return False

    def wait_for_vpn_ready(self, timeout=15, interval=0.5):
        deadline = time.time() + max(0, timeout)
        last_iface = ""
        last_internet_ok = False

        while time.time() <= deadline:
            wg_iface = str((self.get_wg_iface() or "")).strip()
            internet_ok = False
            if wg_iface:
                try:
                    internet_ok = bool(self.check_internet())
                except Exception:
                    internet_ok = False

            last_iface = wg_iface
            last_internet_ok = internet_ok

            if wg_iface and internet_ok:
                return True

            time.sleep(max(0.1, interval))

        log_debug_safe = self.__dict__.get("log_debug_safe")
        if callable(log_debug_safe):
            log_debug_safe(
                "VPN-Ready-Timeout: "
                + f"iface={'set' if last_iface else 'missing'}, "
                + f"internet_ok={last_internet_ok}"
            )
        return False

    def _internet_on_sync(self):
        if not self._try_begin_network_operation("Internet AN"):
            return False
        try:
            rc, out = run_cmd(["sudo", "-n", "nmcli", "networking", "on"], timeout=10)
            if rc != 0:
                self.after(0, lambda: self.log(f"Internet AN fehlgeschlagen: {out}"))
                self.after(
                    0,
                    lambda: messagebox.showerror(
                        self.tr(
                            "internet_on_failed_title", "Internet AN fehlgeschlagen"
                        ),
                        self.tr(
                            "internet_on_failed_message",
                            "Netzwerk konnte nicht aktiviert werden:\n{details}",
                        ).format(details=out),
                        parent=self,
                    ),
                )
                return False

            self.after(0, lambda: self.log("Internet systemweit aktiviert."))
            time.sleep(1)

            state_coordinator = self._state_coordinator()
            if state_coordinator is not None:
                suspend_state = state_coordinator.normalize_internet_suspend_state(
                    self.internet_suspend_state
                )
            else:
                suspend_state = dict(self.internet_suspend_state)
            restore_complete = True
            vpn_restore_ok = False
            if suspend_state.get("vpn_was_active"):
                restore_complete = False
                profile_to_restore = suspend_state.get("vpn_profile", "").strip()
                if profile_to_restore:
                    self.after(
                        0,
                        lambda: self.log(
                            f"VPN-Wiederherstellung gestartet (Profil: {profile_to_restore})."
                        ),
                    )
                    vpn_ok = self.vpn_switch(
                        profile_to_restore,
                        ensure_sudo=False,
                        refresh_ui=False,
                        force_reconnect=True,
                    )
                    ready_ok = self.wait_for_vpn_ready() if vpn_ok else False
                    if vpn_ok and ready_ok:
                        vpn_restore_ok = True
                        restore_complete = True
                        self.after(0, lambda: self.log("VPN-Tunnel ist bereit."))
                    else:
                        self.after(
                            0, lambda: self.log("VPN-Wiederherstellung fehlgeschlagen.")
                        )
                        if not vpn_ok:
                            self.after(
                                0,
                                lambda: self.log_debug_safe(
                                    "VPN-Wiederherstellung: Reconnect-Aufruf fehlgeschlagen."
                                ),
                            )
                        elif not ready_ok:
                            self.after(
                                0,
                                lambda: self.log_debug_safe(
                                    "VPN-Wiederherstellung: Tunnel wurde innerhalb des Zeitlimits nicht bereit."
                                ),
                            )

            if suspend_state.get("killswitch_was_active"):
                if vpn_restore_ok or not suspend_state.get("vpn_was_active"):
                    if self.killswitch_on(
                        async_mode=False, ensure_sudo=False, refresh_ui=False
                    ):
                        self.after(0, lambda: self.log("Killswitch wieder aktiviert."))
                    else:
                        restore_complete = False
                        self.after(
                            0,
                            lambda: self.log(
                                "Killswitch konnte nach VPN-Wiederherstellung nicht wieder aktiviert werden."
                            ),
                        )
                else:
                    restore_complete = False
                    self.after(
                        0,
                        lambda: self.log(
                            "Killswitch bleibt aus, damit Internetzugang nicht blockiert wird."
                        ),
                    )

            if state_coordinator is not None:
                self.internet_suspend_state = (
                    state_coordinator.finalize_internet_restore_state(
                        suspend_state,
                        restore_complete,
                    )
                )
            elif restore_complete:
                self.internet_suspend_state = self._default_internet_suspend_state()
            else:
                self.internet_suspend_state = suspend_state

            if not restore_complete:
                self.after(
                    0,
                    lambda: self.log(
                        "Wiederherstellungszustand bleibt gespeichert, damit der Restore erneut versucht werden kann."
                    ),
                )
            self._schedule_refresh_burst()
            return True
        finally:
            self._end_network_operation()

    def _is_internal_ipv4_profile(self, profile_name):
        name = (profile_name or "").strip()
        return bool(name) and name.endswith(WG_IPV4_PROFILE_SUFFIX)

    def _base_wg_profile_name(self, profile_name):
        name = (profile_name or "").strip()
        if self._is_internal_ipv4_profile(name):
            return name[: -len(WG_IPV4_PROFILE_SUFFIX)]
        return name

    def resolve_real_interface(self, profile_name):
        raw = (profile_name or "").strip()
        base = self._base_wg_profile_name(raw)
        real_iface = (base or raw).strip()
        if profile_name:
            self.log_debug(
                f"[DEBUG] Profil {profile_name} -> reales Interface: {real_iface or '-'}"
            )
        return real_iface

    def _wg_profile_path(self, profile_name):
        name = (profile_name or "").strip()
        if not name:
            return ""
        return f"/etc/wireguard/{name}.conf"

    def _wg_profile_exists(self, profile_name):
        path = self._wg_profile_path(profile_name)
        if not path:
            return False

        user_exists = os.path.isfile(path)
        if user_exists:
            self.log_debug(f"os.path.isfile({path}) = True")
            return True
        self.log_debug(f"os.path.isfile({path}) = False")

        rc, out = run_cmd(["sudo", "-n", "test", "-f", path], timeout=5)
        self.log_debug(f"sudo test -f {path} rc={rc}, out={out!r}")
        return rc == 0

    def active_wg_profile(self):
        ifaces = self.get_wg_interfaces()
        if ifaces:
            return self._base_wg_profile_name(ifaces[0])
        return None

    def get_wg_interfaces(self):
        _, out = run_cmd(["wg", "show", "interfaces"], timeout=3)
        return [p for p in out.split() if p]

    def get_wg_iface(self):
        ifaces = self.get_wg_interfaces()
        if not ifaces:
            return None
        return ifaces[0]

    def _list_wg_profile_files(self, include_internal=False):
        profiles = set()
        cmd = [
            "find",
            "/etc/wireguard",
            "-maxdepth",
            "1",
            "-type",
            "f",
            "-name",
            "*.conf",
            "-printf",
            "%f\n",
        ]
        rc, out = run_cmd(["sudo", "-n", *cmd], timeout=5)
        if rc != 0:
            rc, out = run_cmd(cmd, timeout=5)
        if rc != 0:
            return []
        for line in (out or "").splitlines():
            name = line.strip()
            if name.endswith(".conf"):
                name = name[:-5]
            if not name:
                continue
            if self._is_internal_ipv4_profile(name):
                if include_internal:
                    profiles.add(name)
                continue
            profiles.add(name)
        return sorted(profiles)

    def _read_wg_profile_text(self, profile_name):
        path = self._wg_profile_path(profile_name)
        if not path:
            return 1, ""
        rc, out = run_cmd(["sudo", "-n", "cat", path], timeout=5)
        if rc != 0:
            rc, out = run_cmd(["cat", path], timeout=5)
        return rc, out

    def parse_vpn_endpoint(self, profile_name):
        if not profile_name:
            return None, None

        rc, out = self._read_wg_profile_text(profile_name)
        if rc != 0:
            return None, None

        for raw_line in (out or "").splitlines():
            line = raw_line.split("#", 1)[0].split(";", 1)[0].strip()
            if not line or "=" not in line:
                continue
            key, value = [part.strip() for part in line.split("=", 1)]
            if key.lower() != "endpoint" or not value:
                continue

            if value.startswith("[") and "]:" in value:
                host, port = value[1:].split("]:", 1)
            else:
                host, sep, port = value.rpartition(":")
                if not sep:
                    continue

            host = host.strip()
            port = port.strip()
            if not host or not port.isdigit():
                continue
            return host, port

        return None, None

    def profile_supports_ipv6(self, profile_name):
        if not profile_name:
            return False

        rc, out = self._read_wg_profile_text(profile_name)
        if rc != 0:
            return False

        def has_ipv6_token(token):
            value = (token or "").strip()
            if not value:
                return False
            if value.startswith("[") and value.endswith("]"):
                value = value[1:-1].strip()
            try:
                if "/" in value:
                    return ipaddress.ip_network(value, strict=False).version == 6
                return ipaddress.ip_address(value).version == 6
            except ValueError:
                return False

        for raw_line in (out or "").splitlines():
            line = raw_line.split("#", 1)[0].split(";", 1)[0].strip()
            if not line or "=" not in line:
                continue
            key, value = [part.strip() for part in line.split("=", 1)]
            key_l = key.lower()
            if not value:
                continue

            if key_l in {"address", "dns", "allowedips"}:
                tokens = [part.strip() for part in value.split(",") if part.strip()]
                if any(has_ipv6_token(token) for token in tokens):
                    return True
                continue

            if key_l == "endpoint":
                host = value
                if value.startswith("[") and "]:" in value:
                    host = value[1:].split("]:", 1)[0].strip()
                else:
                    host, sep, port = value.rpartition(":")
                    if sep and port.strip().isdigit():
                        host = host.strip()
                if has_ipv6_token(host):
                    return True

        return False

    def _extract_ipv4_tokens(self, value):
        tokens = []
        for raw_token in (value or "").split(","):
            token = raw_token.strip()
            if not token:
                continue
            candidate = token
            if candidate.startswith("[") and candidate.endswith("]"):
                candidate = candidate[1:-1].strip()
            try:
                if "/" in candidate:
                    if ipaddress.ip_network(candidate, strict=False).version == 4:
                        tokens.append(token)
                elif ipaddress.ip_address(candidate).version == 4:
                    tokens.append(token)
            except ValueError:
                continue
        return tokens

    def _endpoint_uses_ipv6_literal(self, value):
        endpoint = (value or "").strip()
        if not endpoint:
            return False
        if endpoint.startswith("[") and "]:" in endpoint:
            return True
        host, sep, port = endpoint.rpartition(":")
        if not sep or not port.strip().isdigit():
            return False
        host = host.strip()
        try:
            return ipaddress.ip_address(host).version == 6
        except ValueError:
            return False

    def build_ipv4_only_wg_profile_content(self, profile_name):
        rc, out = self._read_wg_profile_text(profile_name)
        if rc != 0:
            return "", "WireGuard-Profil konnte nicht gelesen werden."

        filtered_lines = []
        relevant_keys = {"address", "dns", "allowedips"}

        for raw_line in (out or "").splitlines():
            stripped = raw_line.strip()
            if not stripped:
                filtered_lines.append("")
                continue
            if stripped.startswith("[") and stripped.endswith("]"):
                filtered_lines.append(stripped)
                continue
            if "=" not in raw_line:
                filtered_lines.append(raw_line.rstrip())
                continue

            key, value = [part.strip() for part in raw_line.split("=", 1)]
            key_l = key.lower()

            if key_l in relevant_keys:
                ipv4_tokens = self._extract_ipv4_tokens(value)
                if not ipv4_tokens:
                    continue
                filtered_lines.append(f"{key} = {', '.join(ipv4_tokens)}")
                continue

            if key_l == "endpoint" and self._endpoint_uses_ipv6_literal(value):
                return (
                    "",
                    "Endpoint nutzt eine IPv6-Adresse und kann nicht als IPv4-only Kopie übernommen werden.",
                )

            filtered_lines.append(raw_line.rstrip())

        content = "\n".join(filtered_lines).strip() + "\n"
        return content, ""

    def ensure_internal_ipv4_profile(self, profile_name, log_failures=True):
        base_profile = self._base_wg_profile_name(profile_name)
        if not base_profile:
            return False

        ipv4_profile = f"{base_profile}{WG_IPV4_PROFILE_SUFFIX}"
        if self._wg_profile_exists(ipv4_profile):
            return True

        content, reason = self.build_ipv4_only_wg_profile_content(base_profile)
        if not content:
            if log_failures and reason:
                self.log_safe(reason)
            return False

        tmp_path = ""
        try:
            with tempfile.NamedTemporaryFile(
                mode="w", encoding="utf-8", delete=False
            ) as temp_file:
                temp_file.write(content)
                tmp_path = temp_file.name
            target = self._wg_profile_path(ipv4_profile)
            rc, out = run_cmd(
                ["sudo", "-n", "install", "-m", "600", tmp_path, target],
                timeout=12,
            )
            if rc != 0:
                if log_failures:
                    self.log_safe(
                        (out or "").strip()
                        or f"Interne IPv4-Kopie konnte nicht erstellt werden: {ipv4_profile}"
                    )
                return False
            return True
        finally:
            if tmp_path:
                try:
                    os.remove(tmp_path)
                except Exception:
                    pass

    def sync_internal_ipv4_profiles(self, log_failures=True):
        # Architektur vereinfacht: kein __ipv4-Sync mehr notwendig.
        self.log_debug("[DEBUG] __ipv4-Sync deaktiviert (ein Profil = ein Interface).")

    def resolve_profile_family_runtime(self, profile_name, log_missing=True):
        base_profile = self._base_wg_profile_name(profile_name)
        system_ipv6_disabled = self.ipv6_is_disabled()
        ipv4_profile = f"{base_profile}{WG_IPV4_PROFILE_SUFFIX}" if base_profile else ""
        runtime = {
            "base_profile": base_profile,
            "ipv4_profile": ipv4_profile or base_profile,
            "effective_ipv4_interface": "",
            "effective_ipv6_interface": None,
            "supports_ipv6": (
                self.profile_supports_ipv6(base_profile) if base_profile else False
            ),
            "runtime_profile": "",
            "system_ipv6_disabled": system_ipv6_disabled,
            "endpoint_profiles": [],
        }
        if not base_profile:
            return runtime

        base_exists = self._wg_profile_exists(base_profile)
        ipv4_exists = bool(ipv4_profile and self._wg_profile_exists(ipv4_profile))
        if system_ipv6_disabled and not ipv4_exists:
            ipv4_exists = self.ensure_internal_ipv4_profile(
                base_profile,
                log_failures=log_missing,
            )
        if base_exists:
            runtime["effective_ipv4_interface"] = base_profile
            runtime["runtime_profile"] = base_profile
            runtime["endpoint_profiles"] = [base_profile]
        if ipv4_exists:
            runtime["ipv4_profile"] = ipv4_profile
            if system_ipv6_disabled:
                runtime["effective_ipv4_interface"] = ipv4_profile
                runtime["runtime_profile"] = ipv4_profile
            if ipv4_profile not in runtime["endpoint_profiles"]:
                runtime["endpoint_profiles"].append(ipv4_profile)

        self.log_debug(
            f"[DEBUG] Resolver {base_profile} -> IPv4={runtime['effective_ipv4_interface'] or '-'}, "
            f"IPv6={runtime['effective_ipv6_interface'] or 'None'}, supports_ipv6={runtime['supports_ipv6']}"
        )
        return runtime

    def resolve_runtime_vpn_profile(self, profile_name, log_missing=True):
        base_profile = self._base_wg_profile_name(profile_name)
        if not base_profile:
            return ""
        if self.ipv6_is_disabled():
            ipv4_profile = f"{base_profile}{WG_IPV4_PROFILE_SUFFIX}"
            if self._wg_profile_exists(ipv4_profile):
                return ipv4_profile
            if self.ensure_internal_ipv4_profile(
                base_profile,
                log_failures=log_missing,
            ) and self._wg_profile_exists(ipv4_profile):
                return ipv4_profile
            if log_missing:
                self.log_safe(
                    f"VPN Start abgebrochen: Fuer Profil '{base_profile}' ist keine interne IPv4-only Kopie verfuegbar."
                )
                self.log_safe(
                    "Bitte IPv6 aktivieren oder ein Profil mit nutzbaren IPv4-Eintraegen verwenden."
                )
            return ""
        if self._wg_profile_exists(base_profile):
            return base_profile
        if log_missing:
            self.log_safe(
                f"VPN Start abgebrochen: Profil '{base_profile}' nicht gefunden."
            )
        return ""

    def get_killswitch_vpn_profiles(self):
        return firewall_controller.get_killswitch_vpn_profiles(self)

    def get_killswitch_vpn_endpoints(self):
        return firewall_controller.get_killswitch_vpn_endpoints(self)

    def list_wg_profiles(self):
        profiles = set()
        for default_profile in (VPN_CH, VPN_DE):
            if self._wg_profile_exists(default_profile):
                profiles.add(default_profile)
        for name in self._list_wg_profile_files(include_internal=False):
            profiles.add(name)
        for iface in self.get_wg_interfaces():
            profiles.add(self._base_wg_profile_name(iface))
        return sorted(profile for profile in profiles if profile)

    def refresh_wg_profiles(self):
        profiles = self.list_wg_profiles()
        self.wg_profiles = profiles

    def reload_wg_profiles(self):
        self.sync_internal_ipv4_profiles(log_failures=False)
        self.refresh_wg_profiles()
        if hasattr(self, "vpn_profile_combo"):
            self.vpn_profile_combo["values"] = self.wg_profiles
        self.sync_primary_vpn_form_selection()
        self.sync_vpn_profile_form_selection()
        self.update_vpn_slot_buttons()
        self.update_primary_vpn_form_controls()
        self.refresh_vpn_slots_table()
        if hasattr(self, "isp_stealth_vpn_combo"):
            self.update_privacy_vpn_combo()

        self.update_ipv6_controls_state()
        if self.wg_profiles:
            self.log(f"WireGuard-Profile: {', '.join(self.wg_profiles)}")
        else:
            self.log("WireGuard-Profile: keine gefunden")

    def import_wg_profiles_from_folder(self):
        initialdir = self.get_wg_import_dir() or os.path.expanduser("~")
        files = filedialog.askopenfilenames(
            title="WireGuard-Profile auswählen",
            initialdir=initialdir,
            filetypes=[("WireGuard config", "*.conf"), ("Alle Dateien", "*.*")],
        )
        if not files:
            self.log("Import abgebrochen.")
            return

        first_dir = os.path.dirname(str(files[0]).strip()) if files else ""
        if first_dir:
            self.set_wg_import_dir(first_dir)

        def job():
            selected_files = sorted([str(p).strip() for p in files if str(p).strip()])

            imported = 0
            skipped = 0
            failed = 0
            imported_profiles = []

            for src in selected_files:
                base = os.path.basename(src)
                name = base[:-5] if base.endswith(".conf") else base
                if not re.match(r"^[A-Za-z0-9_.-]+$", name):
                    skipped += 1
                    self.after(
                        0,
                        lambda n=name: self.log(
                            f"Übersprungen (ungültiger Profilname): {n}"
                        ),
                    )
                    continue

                target = f"/etc/wireguard/{name}.conf"
                rc_cp, out_cp = run_cmd(
                    ["sudo", "-n", "install", "-m", "600", src, target], timeout=10
                )
                if rc_cp == 0:
                    imported += 1
                    imported_profiles.append(name)
                    self.after(0, lambda n=name: self.log(f"Profil importiert: {n}"))
                else:
                    failed += 1
                    self.after(
                        0,
                        lambda n=name, o=out_cp: self.log(
                            f"Import fehlgeschlagen ({n}): {o}"
                        ),
                    )

            def finish():
                if imported_profiles:
                    self.reconcile_vpn_slot_firewall_state(
                        reason="Import-Reconcile",
                        touched_profiles=imported_profiles,
                        refresh_ui=False,
                    )
                self.sync_internal_ipv4_profiles(log_failures=False)
                self.refresh_wg_profiles()
                if hasattr(self, "vpn_profile_combo"):
                    self.vpn_profile_combo["values"] = self.wg_profiles
                self.update_primary_vpn_form_controls()
                self.sync_primary_vpn_form_selection()
                self.sync_vpn_profile_form_selection()
                self.update_vpn_slot_buttons()
                self.refresh_vpn_slots_table()
                self.log(
                    f"WireGuard-Import beendet: {imported} importiert, {skipped} übersprungen, {failed} fehlgeschlagen."
                )

            self.after(0, finish)

        threading.Thread(target=job, daemon=True).start()

    # ---------- UFW status robust + LOCK ----------
    def _read_ufw_status(self, verbose=False):
        return firewall_controller.read_ufw_status(verbose=verbose)

    def _ufw_status_is_active(self, status_text):
        return firewall_controller.ufw_status_is_active(status_text)

    def _ufw_rule_lines(self, status_text):
        return firewall_controller.ufw_rule_lines(status_text)

    def _ufw_default_policy_matches(self, status_text, incoming=None, outgoing=None):
        return firewall_controller.ufw_default_policy_matches(
            status_text,
            incoming=incoming,
            outgoing=outgoing,
        )

    def _ufw_has_interface_rule(self, status_text, iface, direction):
        return firewall_controller.ufw_has_interface_rule(
            status_text,
            iface,
            direction,
        )

    def _ufw_has_endpoint_rule(self, status_text, endpoint_ip, endpoint_port):
        return firewall_controller.ufw_has_endpoint_rule(
            status_text,
            endpoint_ip,
            endpoint_port,
        )

    def _ufw_existing_rule_keys(self, status_text):
        return firewall_controller.ufw_existing_rule_keys(status_text)

    def ensure_ufw_rule_exists(
        self,
        rule,
        *,
        network_service=None,
        known_keys=None,
    ):
        return firewall_controller.ensure_ufw_rule_exists(
            rule,
            network_service=network_service,
            known_keys=known_keys,
        )

    def _parse_ufw_status_numbered(self, status_numbered_text):
        return firewall_controller.parse_ufw_status_numbered(status_numbered_text)

    def _dedupe_vpn_ufw_rules(self):
        return firewall_controller.dedupe_vpn_ufw_rules()

    def _ufw_has_outgoing_port_rule(self, status_text, port, proto):
        return firewall_controller.ufw_has_outgoing_port_rule(
            status_text,
            port,
            proto,
        )

    def _verify_killswitch_rules(
        self, status_text, enabled, wg_ifaces=None, handshake_rules=None
    ):
        return firewall_controller.verify_killswitch_rules(
            self,
            status_text,
            enabled,
            wg_ifaces=wg_ifaces,
            handshake_rules=handshake_rules,
        )

    def check_ufw_active(self):
        return firewall_controller.check_ufw_active(self)

    def check_ssh_rule_active(self):
        return firewall_controller.check_ssh_rule_active(self)

    def killswitch_is_on(self):
        return firewall_controller.killswitch_is_on(self)

    # ---------- NM DNS header ----------
    def nm_get_dns_list(self, iface: str):
        _, out = run_cmd(["nmcli", "-g", "IP4.DNS", "dev", "show", iface], timeout=5)
        if not out.strip():
            return []
        return [ln.strip() for ln in out.splitlines() if ln.strip()]

    def update_dns_header(self):
        iface = self.get_default_iface()
        dns = self.nm_get_dns_list(iface) if has_cmd("nmcli") else []
        v6 = "IPv6 AUS" if self.ipv6_is_disabled() else "IPv6 AN"
        if dns:
            self.nm_dns_header_text = (
                f"NM DNS ({iface}): " + ", ".join(dns) + f" | {v6}"
            )
        else:
            self.nm_dns_header_text = f"NM DNS ({iface}): – | {v6}"
        self.dns_header.config(text=self.nm_dns_header_text)

    # ---------- WG RX/TX ----------
    def read_wg_transfer(self):
        res = {}
        rc, out = run_cmd(["sudo", "-n", "wg", "show", "all", "transfer"], timeout=6)
        if out.strip() and rc == 0:
            for line in out.splitlines():
                parts = line.strip().split()
                if len(parts) >= 3:
                    iface = parts[0]
                    try:
                        rx = int(parts[1])
                        tx = int(parts[2])
                        res[iface] = (rx, tx)
                    except ValueError:
                        pass
        if res:
            return res

        wg_if = self.get_wg_iface()
        if not wg_if:
            return {}

        rc, out = run_cmd(["sudo", "-n", "wg", "show", wg_if, "transfer"], timeout=6)
        total_rx = 0
        total_tx = 0
        if out.strip() and rc == 0:
            for line in out.splitlines():
                parts = line.strip().split()
                if len(parts) >= 3:
                    try:
                        rx = int(parts[1])
                        tx = int(parts[2])
                        total_rx += rx
                        total_tx += tx
                    except ValueError:
                        pass
            res[wg_if] = (total_rx, total_tx)
        return res

    def update_wg_rates(self):
        now = time.time()
        transfer = self.read_wg_transfer()
        wg_if = self.get_wg_iface()

        if not wg_if or wg_if not in transfer:
            self.wg_rx_rate = 0.0
            self.wg_tx_rate = 0.0
            self.last_wg_transfer = transfer
            self.last_wg_ts = now
            return wg_if, None

        cur_rx, cur_tx = transfer[wg_if]

        if (
            self.last_wg_transfer is None
            or self.last_wg_ts is None
            or wg_if not in self.last_wg_transfer
        ):
            self.last_wg_transfer = transfer
            self.last_wg_ts = now
            self.wg_rx_rate = 0.0
            self.wg_tx_rate = 0.0
            return wg_if, (cur_rx, cur_tx)

        dt = max(0.5, now - self.last_wg_ts)
        prev_rx, prev_tx = self.last_wg_transfer[wg_if]

        self.wg_rx_rate = (cur_rx - prev_rx) / dt
        self.wg_tx_rate = (cur_tx - prev_tx) / dt

        self.last_wg_transfer = transfer
        self.last_wg_ts = now
        return wg_if, (cur_rx, cur_tx)

    # ---------- DNS apply ----------
    def apply_dns(self, choice=None, refresh_ui=True, update_choice_var=True):
        return dns_controller.apply_dns(
            self,
            choice=choice,
            refresh_ui=refresh_ui,
            update_choice_var=update_choice_var,
        )

    def show_dns_status(self):
        self.nb.select(self.tab_dns)
        self.log(self.tr("dns_log_status_updated", "DNS-Status aktualisiert"))
        if not has_cmd("nmcli"):
            self.update_dns_output(
                "=== DNS Status (nmcli) ===\n"
                "Ergebnis: fehlgeschlagen\nDetails: nmcli fehlt – NetworkManager ist nicht verfügbar."
            )
            return
        iface = self.get_default_iface()
        _, out = run_cmd(["nmcli", "dev", "show", iface], timeout=8)
        if not out:
            self.update_dns_output(
                "=== DNS Status (nmcli) ===\n"
                f"Interface: {iface}\nErgebnis: keine Ausgabe"
            )
            return
        lines = [ln for ln in out.splitlines() if "IP4.DNS" in ln or "IP4.DOMAIN" in ln]
        body = "\n".join(lines) if lines else out
        self.update_dns_output(
            "=== DNS Status (nmcli) ===\n" f"Interface: {iface}\n\n{body}"
        )

    # ---------- DNS Consistency check ----------
    def enable_dns_consistency_test(self):
        was_enabled = self.dns_consistency_enabled
        self.dns_consistency_enabled = True
        self.dns_consistency_one_shot = False
        self.last_dns_consistency_ok = None
        self.last_dns_consistency_ts = 0.0
        self.nb.select(self.tab_dns)
        self.update_dns_output(
            "=== DNS-Konsistenztest ===\nStatus: gestartet\n"
            f"Zeit: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        )
        self.log(
            self.tr(
                "dns_log_consistency_started",
                "DNS-Konsistenztest gestartet",
            )
        )
        self.refresh()
        if not was_enabled or not self.dns_consistency_running:
            self.dns_consistency_check_async(force=True)

    def disable_dns_consistency_test(self):
        self.dns_consistency_enabled = False
        self.dns_consistency_one_shot = False
        self.last_dns_consistency_ok = None
        self.last_dns_consistency_ts = 0.0
        self.nb.select(self.tab_dns)
        self.update_dns_output(
            "=== DNS-Konsistenztest ===\nStatus: gestoppt\n"
            f"Zeit: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        )
        self.log(
            self.tr(
                "dns_log_consistency_stopped",
                "DNS-Konsistenztest gestoppt",
            )
        )
        self.refresh()

    def dns_consistency_check_async(self, force=False):
        if not self.dns_consistency_enabled and not force:
            return
        if self.dns_consistency_running:
            return
        if not force and (time.time() - self.last_dns_consistency_ts) < 30:
            return

        if not (has_cmd("dig") or has_cmd("nslookup")):
            self.nb.select(self.tab_dns)
            self.update_dns_output(
                "=== DNS-Konsistenztest ===\n"
                "Ergebnis: nicht ausführbar\n"
                "Details: weder 'dig' noch 'nslookup' gefunden. Install: sudo apt install dnsutils"
            )
            self.log(
                self.tr(
                    "dns_log_consistency_start_failed",
                    "DNS-Konsistenztest konnte nicht gestartet werden",
                )
            )
            self.last_dns_consistency_ok = None
            self.last_dns_consistency_ts = time.time()
            return

        self.dns_consistency_running = True
        choice = self.dns_choice.get()

        def job():
            try:
                selected = DNS_PROFILES.get(choice)
                test_domain = "example.com"
                killswitch_active = self.killswitch_is_on()
                vpn_active = bool(self.active_wg_profile())

                def sys_query():
                    if has_cmd("dig"):
                        _, out = run_cmd(
                            ["dig", "+time=2", "+tries=1", "+short", "A", test_domain],
                            timeout=5,
                        )
                        return out.strip()
                    _, out = run_cmd(["nslookup", "-timeout=2", test_domain], timeout=5)
                    return out.strip()

                def explicit_query(server: str):
                    if has_cmd("dig"):
                        _, out = run_cmd(
                            [
                                "dig",
                                "+time=2",
                                "+tries=1",
                                "+short",
                                "A",
                                test_domain,
                                f"@{server}",
                            ],
                            timeout=5,
                        )
                        return out.strip()
                    _, out = run_cmd(
                        ["nslookup", "-timeout=2", test_domain, server], timeout=5
                    )
                    return out.strip()

                def http_query():
                    if not has_cmd("curl"):
                        return None
                    rc, _ = run_cmd(
                        ["curl", "-fsS", "--max-time", "4", "https://example.com"],
                        timeout=6,
                    )
                    return rc == 0

                def build_network_context_lines(dns_ok, http_ok):
                    lines = [f"DNS-Auflösung: {'OK' if dns_ok else '—'}"]
                    if http_ok is None:
                        lines.append("HTTP-Test: übersprungen (curl fehlt)")
                    else:
                        lines.append(f"HTTP-Test: {'OK' if http_ok else 'blockiert'}")
                    if killswitch_active and not vpn_active:
                        lines.append(
                            "Hinweis: Killswitch ist AN – ohne VPN evtl. kein Internet (gewollt)."
                        )
                    return lines

                sys_ok = bool(sys_query())
                http_ok = http_query()
                timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

                if selected is None:
                    self.last_dns_consistency_ok = None
                    self.last_dns_consistency_ts = time.time()
                    detail_text = "\n".join(
                        [
                            "=== DNS-Konsistenztest ===",
                            f"Zeit: {timestamp}",
                            f"Profil: {choice}",
                            "Ergebnis: System(auto)",
                            f"System DNS answer: {'OK' if sys_ok else '—'}",
                            "",
                            *build_network_context_lines(sys_ok, http_ok),
                        ]
                    )

                    def show_auto_result():
                        self.nb.select(self.tab_dns)
                        self.update_dns_output(detail_text, append=True)
                        self.log("DNS-Konsistenztest aktualisiert")

                    self.after(0, show_auto_result)
                    return

                explicit_results = [(s, bool(explicit_query(s))) for s in selected]
                any_explicit_ok = any(ok for _, ok in explicit_results)
                dns_ok = sys_ok or any_explicit_ok

                if (not sys_ok) and (not any_explicit_ok):
                    self.last_dns_consistency_ok = None
                    self.last_dns_consistency_ts = time.time()
                    detail_text = "\n".join(
                        [
                            "=== DNS-Konsistenztest ===",
                            f"Zeit: {timestamp}",
                            f"Profil: {choice}",
                            "Ergebnis: nicht testbar (keine DNS-Antworten)",
                            "",
                            *build_network_context_lines(dns_ok, http_ok),
                        ]
                    )

                    def show_no_dns_result():
                        self.nb.select(self.tab_dns)
                        self.update_dns_output(detail_text, append=True)
                        self.log("DNS-Konsistenztest aktualisiert")

                    self.after(0, show_no_dns_result)
                    return

                ok = sys_ok and any_explicit_ok
                self.last_dns_consistency_ok = ok
                self.last_dns_consistency_ts = time.time()
                explicit_line = ", ".join(
                    [f"{s}->{('OK' if ok2 else '—')}" for s, ok2 in explicit_results]
                )
                detail_text = "\n".join(
                    [
                        "=== DNS-Konsistenztest ===",
                        f"Zeit: {timestamp}",
                        f"Profil: {choice}",
                        f"System DNS answer: {'OK' if sys_ok else '—'}",
                        f"Explizit gewählt: {explicit_line}",
                        "Ergebnis: "
                        + ("OK" if ok else "MÖGLICHE Abweichung (System/Resolver)"),
                        "",
                        *build_network_context_lines(dns_ok, http_ok),
                    ]
                )

                def show_result():
                    self.nb.select(self.tab_dns)
                    self.update_dns_output(detail_text, append=True)
                    self.log("DNS-Konsistenztest aktualisiert")

                self.after(0, show_result)

            finally:
                self.dns_consistency_running = False
                if self.__dict__.get("dns_consistency_one_shot", False):
                    self.dns_consistency_one_shot = False
                    self.dns_consistency_enabled = False
                    self.after(0, lambda: self.log("DNS-Konsistenztest gestoppt"))
                self._update_dns_consistency_light_safe()

        threading.Thread(target=job, daemon=True).start()

    # ---------- Actions ----------
    def fw_enable(self, interactive_guard=True):
        firewall_controller.fw_enable(self, interactive_guard=interactive_guard)

    def _parse_ufw_action(self, action_text):
        return firewall_controller.parse_ufw_action(action_text)

    def _extract_ip_port_proto_from_ufw_to(self, to_text):
        return firewall_controller.extract_ip_port_proto_from_ufw_to(to_text)

    def _is_ufw_non_rule_line(self, line):
        return firewall_controller.is_ufw_non_rule_line(line)

    def _parse_ufw_structured_rule(self, raw_line):
        return firewall_controller.parse_ufw_structured_rule(raw_line)

    def _ip_sort_key(self, value):
        return firewall_controller.ip_sort_key(value)

    def _ufw_rule_sort_key(self, rule):
        return firewall_controller.ufw_rule_sort_key(rule)

    def _format_ufw_rule_lines(self, rules):
        return firewall_controller.format_ufw_rule_lines(rules)

    def _format_ufw_status_grouped(
        self, status_label, default_in, default_out, rules, metadata=None
    ):
        return firewall_controller.format_ufw_status_grouped(
            status_label=status_label,
            default_in=default_in,
            default_out=default_out,
            rules=rules,
            metadata=metadata,
        )

    def fw_status(self):
        firewall_controller.fw_status(self)

    def fw_disable(self, interactive_guard=True):
        firewall_controller.fw_disable(self, interactive_guard=interactive_guard)

    def vpn_switch(
        self, target, ensure_sudo=True, refresh_ui=True, force_reconnect=False
    ):
        if not self._try_begin_network_operation(f"VPN wechseln auf {target}"):
            if refresh_ui:
                self.refresh_safe()
            return False
        try:
            target = self._base_wg_profile_name(target)
            self.log_debug(f"Basisprofil = {target!r}")
            if not target:
                self.log_safe(
                    self.tr(
                        "vpn_log_no_profile_selected", "Kein WireGuard-Profil gewählt."
                    )
                )
                self._recover_firewall_after_killswitch_failure()
                if refresh_ui:
                    self.refresh_safe()
                return False

            if ensure_sudo and not self.ensure_sudo_session(
                f"VPN umschalten auf {target}"
            ):
                return False

            self.sync_internal_ipv4_profiles(log_failures=False)
            resolved = self.resolve_profile_family_runtime(target, log_missing=True)
            runtime_target = (resolved.get("runtime_profile") or "").strip()
            preferred_runtime = (
                self.resolve_runtime_vpn_profile(target, log_missing=not runtime_target)
                or ""
            ).strip()
            if preferred_runtime:
                runtime_target = preferred_runtime
            self.log_debug(f"Runtimeprofil = {runtime_target!r}")
            self.log_debug(f"IPv6 deaktiviert = {resolved.get('system_ipv6_disabled')}")
            self.log_debug(
                f"[DEBUG] Resolver {target} -> IPv4={resolved.get('effective_ipv4_interface') or '-'}, "
                f"IPv6={resolved.get('effective_ipv6_interface') or 'None'}, supports_ipv6={resolved.get('supports_ipv6')}"
            )

            if not runtime_target:
                runtime_target = self.resolve_runtime_vpn_profile(
                    target, log_missing=True
                )
            if not runtime_target:
                self._recover_firewall_after_killswitch_failure()
                if refresh_ui:
                    self.refresh_safe()
                return False

            active_iface = (self.get_wg_iface() or "").strip()
            reconnecting_same_profile = active_iface == runtime_target
            if reconnecting_same_profile:
                if not force_reconnect:
                    self.log_safe(
                        self.tr(
                            "vpn_log_profile_already_active",
                            "VPN-Profil bereits aktiv.",
                        )
                    )
                    if refresh_ui:
                        self.refresh_safe()
                    return True
                self.log_safe(
                    self.tr(
                        "vpn_log_force_reconnect",
                        "Vorhandenes Runtimeprofil erkannt - Reconnect wird erzwungen.",
                    )
                )

            if len(runtime_target) > 15:
                self.log_safe(
                    self.tr(
                        "vpn_log_runtime_name_too_long",
                        "VPN Start abgebrochen: Runtime-Profilname '{name}' ist zu lang fuer ein WireGuard-Interface (max 15 Zeichen).",
                    ).format(name=runtime_target)
                )
                self.log_safe(
                    self.tr(
                        "vpn_log_shorten_runtime_name",
                        "Bitte interne IPv4-Runtime-Namen kuerzer bilden.",
                    )
                )
                if refresh_ui:
                    self.refresh_safe()
                return False

            network_service = self._network_service()
            runtime_path = self._wg_profile_path(runtime_target)
            self._check_wireguard_profile_for_multiple_endpoints(
                runtime_target, runtime_path
            )
            if network_service is not None:
                rc_check, out_check = network_service.file_exists(runtime_path)
            else:
                rc_check, out_check = run_cmd(
                    ["sudo", "-n", "test", "-f", runtime_path], timeout=5
                )
            self.log_debug_safe(f"Runtimepfad = {runtime_path}")
            self.log_debug_safe(f"sudo test -f Runtimepfad rc = {rc_check}")
            if rc_check != 0:
                self.log_safe(
                    self.tr(
                        "vpn_log_runtime_file_missing",
                        "VPN Start abgebrochen: Runtime-Datei nicht gefunden: {path}",
                    ).format(path=runtime_path)
                )
                if out_check:
                    self.log_safe(out_check)
                if refresh_ui:
                    self.refresh_safe()
                return False

            self.prepare_dns_for_vpn_profile_switch()

            for iface in self.get_wg_interfaces():
                if iface != runtime_target or reconnecting_same_profile:
                    if network_service is not None:
                        network_service.wg_quick_down(iface, timeout=12)
                    else:
                        run_cmd(["sudo", "-n", "wg-quick", "down", iface], timeout=12)

            if self.ipv6_is_disabled() and runtime_target == target:
                self.log_safe(
                    "IPv6 ist global deaktiviert - aktiviere IPv6 vor VPN-Start."
                )
                self.ipv6_enable(refresh_ui=False)

            if network_service is not None:
                rc, out = network_service.wg_quick_up(runtime_target, timeout=20)
            else:
                rc, out = run_cmd(
                    ["sudo", "-n", "wg-quick", "up", runtime_target], timeout=20
                )
            if rc == 0:
                if force_reconnect:
                    self.log_safe(
                        self.tr(
                            "vpn_log_reconnected",
                            "VPN Verbindung neu aufgebaut.",
                        )
                    )
                else:
                    self.log_safe(
                        self.tr("vpn_log_connected", "VPN Verbindung aufgebaut.")
                    )
                self.log_safe(
                    self.tr("vpn_log_started", "VPN gestartet: {target}").format(
                        target=target
                    )
                )
                self.log_debug_safe(out)
            else:
                self.log_safe(
                    self.tr(
                        "vpn_log_start_failed_rc", "VPN Start fehlgeschlagen: rc={rc}"
                    ).format(rc=rc)
                )
                if out:
                    self.log_safe((out or "").strip())
                self.log_debug_safe(out)

            self.last_wg_transfer = None
            self.last_wg_ts = None
            self.refresh_wg_profiles()
            self.update_primary_vpn_form_controls()
            self.sync_primary_vpn_form_selection()
            if rc == 0 and self.killswitch_is_on():
                self._killswitch_on_sync(refresh_ui=False)
            if rc == 0:
                self._schedule_dns_finalize_after_vpn_profile_switch(delay_ms=2500)
                self.run_async(lambda t=target: self.verify_vpn_exit_country(t))
            else:
                self._schedule_dns_finalize_after_vpn_profile_switch(delay_ms=750)

            self.update_ipv6_controls_state_safe()
            if refresh_ui:
                self.refresh_safe()
            return rc == 0
        finally:
            self._end_network_operation()

    def vpn_off(self, ensure_sudo=True, refresh_ui=True):
        if not self._try_begin_network_operation("VPN deaktivieren"):
            if refresh_ui:
                self.refresh_safe()
            return False
        try:
            if ensure_sudo and not self.ensure_sudo_session("VPN deaktivieren"):
                return False
            network_service = self._network_service()
            targets = sorted(
                set(self.wg_profiles + self.get_wg_interfaces() + [VPN_CH, VPN_DE])
            )
            for target in targets:
                if network_service is not None:
                    network_service.wg_quick_down(target, timeout=12)
                else:
                    run_cmd(["sudo", "-n", "wg-quick", "down", target], timeout=12)
            self.log_safe(self.tr("vpn_log_disabled", "VPN deaktiviert."))
            self.prepare_dns_for_vpn_profile_switch()
            self._schedule_dns_finalize_after_vpn_profile_switch(delay_ms=2500)
            if self.killswitch_is_on():
                self.log_safe(
                    self.tr(
                        "killswitch_log_notice_enabled_without_vpn",
                        "HINWEIS: Killswitch ist AN – ohne VPN evtl. kein Internet (gewollt).",
                    )
                )
            self.refresh_wg_profiles()
            self.update_ipv6_controls_state_safe()
            if refresh_ui:
                self.refresh_safe()
            return True
        finally:
            self._end_network_operation()

    def _recover_firewall_after_killswitch_failure(self):
        firewall_controller.recover_firewall_after_killswitch_failure(self)

    def _killswitch_on_sync(self, refresh_ui=True):
        return firewall_controller.killswitch_on_sync(self, refresh_ui=refresh_ui)

    def killswitch_on(
        self,
        async_mode=True,
        ensure_sudo=True,
        refresh_ui=True,
        interactive_guard=True,
    ):
        return firewall_controller.killswitch_on(
            self,
            async_mode=async_mode,
            ensure_sudo=ensure_sudo,
            refresh_ui=refresh_ui,
            interactive_guard=interactive_guard,
        )

    def _killswitch_off_sync(self, refresh_ui=True):
        return firewall_controller.killswitch_off_sync(self, refresh_ui=refresh_ui)

    def killswitch_off(
        self,
        async_mode=True,
        ensure_sudo=True,
        refresh_ui=True,
        interactive_guard=True,
    ):
        return firewall_controller.killswitch_off(
            self,
            async_mode=async_mode,
            ensure_sudo=ensure_sudo,
            refresh_ui=refresh_ui,
            interactive_guard=interactive_guard,
        )

    def allow_ssh_rule_async(self):
        firewall_controller.allow_ssh_rule_async(self)

    def remove_ssh_rule_async(self, auto=False):
        firewall_controller.remove_ssh_rule_async(self, auto=auto)

    def _current_port_reality_target(self):
        candidates = []
        for attr in ("nmap_target", "trace_target", "mtr_target"):
            var = self.__dict__.get(attr)
            if var is None:
                continue
            try:
                value = (var.get() or "").strip()
            except Exception:
                value = ""
            if value:
                candidates.append(value)

        for candidate in candidates:
            if is_safe_trace_target(candidate):
                return candidate
        return "1.1.1.1"

    def _append_port_check_debug_log(self, payload):
        try:
            log_dir = os.path.join(self.netmon_state_dir, "logs")
            os.makedirs(log_dir, exist_ok=True)
            log_path = os.path.join(log_dir, "port_check_debug.log")
            record = dict(payload or {})
            record["ts"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            with open(log_path, "a", encoding="utf-8") as fh:
                fh.write(json.dumps(record, ensure_ascii=False) + "\n")
        except Exception as exc:
            self.log_debug_safe(
                f"Port-Reality-Debuglog konnte nicht geschrieben werden: {exc}"
            )

    def _log_port_reality_result(self, result):
        payload = {
            "event": "port_check",
            "host": result.get("host", "-"),
            "port": int(result.get("port", 0) or 0),
            "state": result.get("state", "unknown"),
            "reason": result.get("reason", ""),
        }
        self._append_port_check_debug_log(payload)

        state = payload["state"]
        host = payload["host"]
        port = payload["port"]
        shown = PORT_REALITY_DISPLAY_MAP.get(state, state)
        self.log_safe(f"Port {port} on {host} → {shown}")

        if state == "pseudo-open":
            self.log_safe(
                "[!] Suspicious port behavior detected (possible NAT/proxy interference)"
            )

    def _store_port_reality_scan(self, host, results):
        self.port_reality_last_scan = {
            "host": host,
            "results": list(results or []),
            "ts": time.time(),
        }

    def _run_port_reality_checks(self, host, ports, source_label="manual"):
        checked_ports = []
        for port in ports:
            try:
                checked_ports.append(int(port))
            except (TypeError, ValueError):
                continue

        if not checked_ports:
            self.log_safe("Port Reality Check: keine gueltigen Ports uebergeben.")
            return []

        self.log_safe(
            f"Port Reality Check gestartet ({source_label}): {host} | Ports: {', '.join(str(p) for p in checked_ports)}"
        )
        results = [
            check_port_real_service(host, port, timeout=3) for port in checked_ports
        ]
        self._store_port_reality_scan(host, results)

        pseudo_count = 0
        for item in results:
            if item.get("state") == "pseudo-open":
                pseudo_count += 1
            self._log_port_reality_result(item)

        if pseudo_count:
            penalty = 5 * pseudo_count
            self.log_safe(
                f"Stability-Hinweis: {pseudo_count} pseudo-open Port(s) erkannt (empfohlener Score-Abzug: -{penalty})."
            )

        return results

    def _run_port_reality_checks_async(self, host, ports, source_label="manual"):
        def job():
            self._run_port_reality_checks(host, ports, source_label=source_label)

        threading.Thread(target=job, daemon=True).start()

    def port_reality_check_async(self):
        host = self._current_port_reality_target()
        if not is_safe_trace_target(host):
            self.log("Port Reality Check: Ziel ungültig/unsicher.")
            return

        ports = tuple(self.port_reality_default_ports)
        self._run_port_reality_checks_async(host, ports, source_label="manual")

    def show_last_port_reality_scan(self):
        scan = dict(self.port_reality_last_scan or {})
        results = list(scan.get("results") or [])
        host = (scan.get("host") or "-").strip() or "-"
        ts = float(scan.get("ts") or 0.0)
        if not results:
            self.log("Port Reality Check: noch kein Scan vorhanden.")
            return

        stamp = datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S") if ts else "-"
        self.log(f"Letzter Port Reality Check ({stamp}) für {host}:")
        for item in results:
            try:
                port = int(item.get("port", 0) or 0)
            except (TypeError, ValueError):
                port = 0
            state = (item.get("state") or "unknown").strip() or "unknown"
            shown = PORT_REALITY_DISPLAY_MAP.get(state, state)
            self.log(f"Port {port} → {shown}")

    def _speedtest_logs_dir(self):
        return speedtest_controller.speedtest_logs_dir(self)

    def _speedtest_history_path(self):
        return speedtest_controller.speedtest_history_path(self)

    def _load_speedtest_history(self):
        speedtest_controller.load_speedtest_history(self)

    def _persist_speedtest_history(self):
        speedtest_controller.persist_speedtest_history(self)

    def _append_speedtest_history(self, result):
        speedtest_controller.append_speedtest_history(self, result)

    def _speedtest_warn_messages(self, result):
        return speedtest_controller.speedtest_warn_messages(result)

    def speedtest_async(self):
        return speedtest_controller.speedtest_async(
            self,
            run_speedtest_fn=run_speedtest,
            format_summary_fn=format_speedtest_summary,
            format_table_fn=format_speedtest_table,
            run_cmd_fn=run_cmd,
            has_cmd_fn=has_cmd,
        )

    def update_speedtest_panel(self, result):
        speedtest_controller.update_speedtest_panel(self, result)

    def show_speedtest_history(self):
        speedtest_controller.show_speedtest_history(self)

    # ---------- Trace (Table of hops) ----------
    def _hop_lookup_log_once(self, ip: str, message: str):
        key = f"{ip}:{message}"
        if key in self.hop_lookup_warned_ips:
            return
        self.hop_lookup_warned_ips.add(key)
        try:
            self.after(0, lambda m=message: self.log(m))
        except RuntimeError:
            pass

    def hop_lookup(self, ip: str, enabled=None):
        ip = (ip or "").strip()
        if not ip or ip == "*" or ":" in ip:
            return "-", "unbekannt"

        if ip in self.hop_cache:
            return self.hop_cache[ip]

        lookup_enabled = self.hop_lookup_enabled.get() if enabled is None else enabled
        if not lookup_enabled:
            return "-", "-"

        rc, out = run_cmd(
            ["curl", "-s", "--max-time", "3", f"https://ipwho.is/{ip}"], timeout=5
        )
        if rc != 0 or not out:
            self.hop_cache[ip] = ("-", "unbekannt")
            self._hop_lookup_log_once(
                ip,
                f"Hop Lookup: keine Daten für {ip} (Netzwerk/Timeout).",
            )
            return self.hop_cache[ip]

        try:
            data = json.loads(out)
            if not data.get("success", True):
                self.hop_cache[ip] = ("-", "unbekannt")
                msg = (data.get("message") or "").strip()
                if msg:
                    self._hop_lookup_log_once(ip, f"Hop Lookup: {ip} -> {msg}")
                return self.hop_cache[ip]

            cc = (data.get("country_code") or data.get("country") or "").strip()
            org = (
                data.get("org")
                or data.get("connection", {}).get("org")
                or data.get("connection", {}).get("isp")
                or ""
            ).strip()
            if not cc:
                cc = "-"
            if not org:
                org = "unbekannt"
            self.hop_cache[ip] = (cc, org)
            return cc, org
        except Exception:
            self.hop_cache[ip] = ("-", "unbekannt")
            self._hop_lookup_log_once(ip, f"Hop Lookup: ungültige Antwort für {ip}.")
            return self.hop_cache[ip]

    def trace_table_async(self):
        target = (self.trace_target.get() or "").strip()
        lookup_enabled = bool(self.hop_lookup_enabled.get())
        if not target:
            self.log(
                self.tr("trace_log_enter_target", "Bitte Ziel für Trace eingeben.")
            )
            return
        if not is_safe_trace_target(target):
            self.log(
                self.tr(
                    "trace_log_invalid_target",
                    "Trace-Ziel ungültig/unsicher. Erlaubt: IP oder Hostname (ohne Sonderzeichen/Leerzeichen).",
                )
            )
            return
        if not has_cmd("traceroute"):
            self.log(
                self.tr(
                    "trace_log_traceroute_missing",
                    "Für Hops-Tabelle wird 'traceroute' benötigt (sudo apt install traceroute).",
                )
            )
            return

        def clear_table():
            for item in self.trace_table.get_children():
                self.trace_table.delete(item)

        def job():
            self.after(
                0,
                lambda: (
                    self.nb.select(self.tab_trace),
                    clear_table(),
                    self.log(
                        self.tr(
                            "trace_log_started",
                            "Traceroute zu {target} gestartet",
                        ).format(target=target)
                    ),
                ),
            )
            rc, out = run_cmd(
                ["traceroute", "-n", "-q", "1", "-m", "20", target], timeout=35
            )
            hops = parse_traceroute_hops(out)
            rows = []
            for h in hops:
                rtt = "" if h["rtt_ms"] is None else f"{h['rtt_ms']:.2f}"
                cc, org = self.hop_lookup(h["ip"], enabled=lookup_enabled)
                rows.append((h["hop"], h["ip"], rtt, cc or "-", org or "unbekannt"))

            def fill():
                self.nb.select(self.tab_trace)
                if not out.strip():
                    self.log(
                        self.tr(
                            "trace_log_finished_no_output",
                            "Traceroute zu {target} ohne Ausgabe beendet",
                        ).format(target=target)
                    )
                    return
                if rc != 0:
                    self.log(
                        self.tr(
                            "trace_log_finished_rc",
                            "Traceroute zu {target} abgeschlossen mit rc={rc}",
                        ).format(target=target, rc=rc)
                    )
                else:
                    self.log(
                        self.tr(
                            "trace_log_finished",
                            "Traceroute zu {target} abgeschlossen",
                        ).format(target=target)
                    )

                for row in rows:
                    self.trace_table.insert("", "end", values=row)

                if rc != 0 and not rows:
                    self.log(
                        self.tr(
                            "trace_log_no_hops",
                            "Traceroute lieferte keine auswertbaren Hops.",
                        )
                    )

            self.after(0, fill)

        threading.Thread(target=job, daemon=True).start()

    # ---------- MTR ----------
    def mtr_async(self):
        if not has_cmd("mtr"):
            self.log(self.tr("mtr_log_missing", "MTR fehlt (sudo apt install -y mtr)."))
            return

        target = (self.mtr_target.get() or "").strip()
        if not target:
            self.log(self.tr("mtr_log_enter_target", "Bitte Ziel für MTR eingeben."))
            return
        if not is_safe_trace_target(target):
            self.log(self.tr("mtr_log_invalid_target", "MTR-Ziel ungültig/unsicher."))
            return

        self._select_tools_tab(self.tab_mtr)
        self.mtr_text.delete("1.0", "end")
        self.mtr_text.insert(
            "end",
            self.tr("mtr_output_starting", "MTR startet für {target}...\n").format(
                target=target
            ),
        )
        self.mtr_text.see("1.0")

        def job():
            rc, out = run_cmd(
                ["mtr", "--report", "--report-cycles", "6", "--no-dns", target],
                timeout=45,
            )

            def show():
                self._select_tools_tab(self.tab_mtr)
                self.mtr_text.delete("1.0", "end")
                if out.strip():
                    self.mtr_text.insert("end", out)
                else:
                    self.mtr_text.insert(
                        "end",
                        self.tr(
                            "mtr_output_none",
                            "Keine MTR-Ausgabe erhalten.",
                        ),
                    )
                self.mtr_text.see("1.0")
                self.log(
                    self.tr(
                        "mtr_log_finished_rc",
                        "MTR abgeschlossen für {target} (rc={rc}).",
                    ).format(target=target, rc=rc)
                )

            self.after(0, show)

        threading.Thread(target=job, daemon=True).start()

    # ---------- Nmap ----------
    def _set_nmap_running_state(self, running, status_text=None, error_text=None):
        self.nmap_running = bool(running)
        if status_text is not None:
            self.nmap_status_var.set(status_text)
        if error_text is not None:
            self.nmap_error_var.set(error_text)

        btn = self.__dict__.get("btn_nmap_start")
        if btn is not None:
            try:
                btn.configure(state="disabled" if self.nmap_running else "normal")
            except Exception:
                pass

    def _abort_nmap_process(self):
        proc = self.nmap_active_process
        if proc is None:
            return
        try:
            if proc.poll() is None:
                proc.terminate()
                try:
                    proc.wait(timeout=3)
                except Exception:
                    proc.kill()
        except Exception:
            pass

    def clear_nmap_output(self):
        self.nmap_cancel_requested = True
        self._abort_nmap_process()
        self.nmap_active_process = None
        self._set_nmap_running_state(False, status_text="bereit", error_text="")
        if hasattr(self, "nmap_text"):
            self.nmap_text.delete("1.0", "end")
        self.log(
            self.tr(
                "nmap_log_cleared",
                "Nmap-Ausgabe und Status zurückgesetzt.",
            )
        )

    def update_nmap_output(self, text):
        if not hasattr(self, "nmap_text"):
            return
        self.nmap_text.delete("1.0", "end")
        self.nmap_text.insert("end", text or "")
        self.nmap_text.see("1.0")

    def _run_nmap_with_timeout(self, cmd, timeout):
        self.nmap_cancel_requested = False
        self.nmap_active_process = None
        try:
            proc = subprocess.Popen(
                cmd,
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
            )
            self.nmap_active_process = proc
            try:
                out, _ = proc.communicate(timeout=timeout)
                rc = proc.returncode
                if self.nmap_cancel_requested:
                    return (
                        130,
                        out
                        or self.tr("nmap_output_scan_cancelled", "Scan abgebrochen."),
                        "cancelled",
                    )
                return rc, out or "", "ok"
            except subprocess.TimeoutExpired:
                try:
                    proc.terminate()
                    out, _ = proc.communicate(timeout=3)
                except Exception:
                    try:
                        proc.kill()
                    except Exception:
                        pass
                    out = ""
                return (
                    124,
                    out
                    or self.tr("nmap_output_command_timed_out", "Command timed out"),
                    "timeout",
                )
        except FileNotFoundError:
            return (
                127,
                self.tr("nmap_output_not_found", "nmap nicht gefunden."),
                "error",
            )
        except Exception as exc:
            return (
                1,
                self.tr(
                    "nmap_output_start_failed",
                    "Nmap-Start fehlgeschlagen: {error}",
                ).format(error=exc),
                "error",
            )
        finally:
            self.nmap_active_process = None

    def nmap_async(self):
        if self.nmap_running:
            self.log(self.tr("nmap_log_already_running", "Nmap-Scan läuft bereits."))
            return

        if not has_cmd("nmap"):
            self.log(
                self.tr("nmap_log_missing", "nmap fehlt (sudo apt install -y nmap).")
            )
            return

        target = (self.nmap_target.get() or "").strip()
        preset = (self.nmap_preset.get() or "quick").strip().lower()
        if not target:
            self.log(
                self.tr(
                    "nmap_log_enter_target",
                    "Bitte Zielnetz/Zielhost für nmap eingeben.",
                )
            )
            return
        if not is_safe_nmap_target(target):
            self.log(self.tr("nmap_log_invalid_target", "nmap-Ziel ungültig/unsicher."))
            return

        preset_map = {
            "quick": (["-Pn", "-sT", "--top-ports", "100"], 90),
            "service": (["-Pn", "-sT", "-sV", "--top-ports", "500"], 420),
            "aggressive": (["-Pn", "-sT", "-A", "--top-ports", "500"], 720),
        }
        args, timeout = preset_map.get(preset, preset_map["quick"])

        self._set_nmap_running_state(
            True,
            status_text=f"läuft ({preset})",
            error_text="",
        )

        def job():
            try:
                cmd = ["nmap", *args, target]
                rc, out, state = self._run_nmap_with_timeout(cmd, timeout=timeout)
                open_ports = parse_nmap_open_tcp_ports(out) if rc == 0 else []
            except Exception as exc:
                rc, out, state = (
                    1,
                    self.tr(
                        "nmap_output_worker_error",
                        "Nmap-Worker-Fehler: {error}",
                    ).format(error=exc),
                    "error",
                )
                open_ports = []

            def show():
                self._select_tools_tab(self.tab_nmap)
                if out.strip():
                    display_text = out
                elif state == "timeout":
                    display_text = self.tr(
                        "nmap_output_timeout",
                        "Nmap-Scan abgebrochen: Timeout.",
                    )
                elif state == "cancelled":
                    display_text = self.tr(
                        "nmap_output_cancelled",
                        "Nmap-Scan wurde abgebrochen.",
                    )
                else:
                    display_text = self.tr(
                        "nmap_output_none",
                        "Keine nmap-Ausgabe erhalten.",
                    )
                self.update_nmap_output(display_text)

                if state == "timeout":
                    self.nmap_error_var.set("Timeout")
                    self.nmap_status_var.set("bereit")
                    self.log(
                        self.tr(
                            "nmap_log_timeout",
                            "Nmap-Timeout ({preset}) für {target} nach {timeout}s.",
                        ).format(preset=preset, target=target, timeout=timeout)
                    )
                elif state == "cancelled":
                    self.nmap_error_var.set("Abgebrochen")
                    self.nmap_status_var.set("bereit")
                    self.log(
                        self.tr(
                            "nmap_log_cancelled",
                            "Nmap-Scan abgebrochen ({preset}) für {target}.",
                        ).format(preset=preset, target=target)
                    )
                else:
                    self.nmap_error_var.set("" if rc == 0 else f"Fehler rc={rc}")
                    self.nmap_status_var.set("bereit")
                    if rc != 0:
                        self.log(
                            self.tr(
                                "nmap_log_error_rc",
                                "Nmap-Fehler ({preset}) für {target}: rc={rc}.",
                            ).format(preset=preset, target=target, rc=rc)
                        )

                if rc == 0 and open_ports:
                    self._run_port_reality_checks_async(
                        target,
                        open_ports,
                        source_label="nmap-auto",
                    )

                self._set_nmap_running_state(False)
                self.nmap_cancel_requested = False
                self.nmap_active_process = None

            self.after(0, show)

        threading.Thread(target=job, daemon=True).start()

    # ---------- iftop ----------
    def iftop_async(self):
        if not has_cmd("iftop"):
            self.log(
                self.tr("iftop_log_missing", "iftop fehlt (sudo apt install -y iftop).")
            )
            return

        iface = self.get_default_iface()

        def job():
            rc, out = run_cmd(
                ["sudo", "-n", "iftop", "-t", "-s", "10", "-L", "25", "-i", iface],
                timeout=25,
            )
            if rc != 0 or not out.strip():
                rc, out = run_cmd(
                    ["iftop", "-t", "-s", "10", "-L", "25", "-i", iface], timeout=25
                )

            def show():
                self._select_tools_tab(self.tab_iftop)
                self.iftop_text.delete("1.0", "end")
                if out.strip():
                    self.iftop_text.insert("end", out)
                else:
                    self.iftop_text.insert(
                        "end",
                        self.tr(
                            "iftop_output_none",
                            "Keine iftop-Ausgabe erhalten (Root/NOPASSWD ggf. erforderlich).",
                        ),
                    )
                self.iftop_text.see("1.0")
                self.log(
                    self.tr(
                        "iftop_log_finished_rc",
                        "iftop Snapshot für {iface} abgeschlossen (rc={rc}).",
                    ).format(iface=iface, rc=rc)
                )

            self.after(0, show)

        threading.Thread(target=job, daemon=True).start()

    # ---------- htop ----------
    def htop_async(self):
        if not has_cmd("htop"):
            self.log(
                self.tr("htop_log_missing", "htop fehlt (sudo apt install -y htop).")
            )
            return

        def job():
            # htop 3.x kann je nach Build kein Batch-Flag (-b) haben.
            rc, out = run_cmd(["htop", "-b", "-n", "1"], timeout=20)
            source = "htop"

            if rc != 0 or "invalid option" in (out or "").lower():
                if has_cmd("top"):
                    rc, out = run_cmd(["top", "-b", "-n", "1"], timeout=20)
                    source = "top"
                elif has_cmd("ps"):
                    rc, out = run_cmd(
                        ["ps", "aux", "--sort=-%cpu", "--sort=-%mem"], timeout=20
                    )
                    source = "ps"

            def show():
                self._select_tools_tab(self.tab_htop)
                self.htop_text.delete("1.0", "end")
                if out.strip():
                    self.htop_text.insert("end", out)
                else:
                    self.htop_text.insert(
                        "end",
                        self.tr(
                            "htop_output_none",
                            "Keine Prozessausgabe erhalten (htop/top/ps).",
                        ),
                    )
                self.htop_text.see("1.0")
                self.log(
                    self.tr(
                        "htop_log_finished_rc",
                        "htop Snapshot abgeschlossen via {source} (rc={rc}).",
                    ).format(source=source, rc=rc)
                )

            self.after(0, show)

        threading.Thread(target=job, daemon=True).start()

    def vscode_async(self):
        launch_dir = os.environ.get("VSCODE_PROJECT_PATH", "").strip()
        if not launch_dir:
            repo_root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
            for candidate in (repo_root, os.getcwd(), os.path.expanduser("~")):
                if candidate and os.path.isdir(candidate):
                    launch_dir = candidate
                    break

        def _start_proc(args, note):
            try:
                proc = subprocess.Popen(
                    args,
                    cwd=launch_dir or None,
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
                self.log(
                    self.tr(
                        "vscode_log_started",
                        "VS Code gestartet ({note}, PID {pid}).",
                    ).format(note=note, pid=proc.pid)
                )
                return True
            except Exception as e:
                self.log(
                    self.tr(
                        "vscode_log_start_failed",
                        "VS Code Start fehlgeschlagen ({note}): {error}",
                    ).format(note=note, error=e)
                )
                return False

        starter = os.environ.get("VSCODE_STARTER", "").strip()
        if starter:
            starter_args = shlex.split(starter)
            allowed_starters = {"code", "codium", "code-insiders", "gtk-launch"}
            starter_bin = os.path.basename(starter_args[0]) if starter_args else ""
            if starter_args and starter_bin in allowed_starters:
                if _start_proc(starter_args, "VSCODE_STARTER"):
                    return
            else:
                self.log_debug("VSCODE_STARTER ignoriert (nicht auf Allowlist).")

        desktop_file = self._find_vscode_desktop_file()
        if desktop_file:
            desktop_id = os.path.splitext(os.path.basename(desktop_file))[0]
            if has_cmd("gtk-launch") and _start_proc(
                ["gtk-launch", desktop_id], f"gtk-launch:{desktop_id}"
            ):
                return

            desktop_exec = self._desktop_exec_command(desktop_file)
            if desktop_exec and _start_proc(
                desktop_exec, f"Exec aus {os.path.basename(desktop_file)}"
            ):
                return

        for cmd in (
            ["code", "--new-window", launch_dir] if launch_dir else ["code"],
            ["codium", "--new-window", launch_dir] if launch_dir else ["codium"],
            (
                ["code-insiders", "--new-window", launch_dir]
                if launch_dir
                else ["code-insiders"]
            ),
        ):
            if has_cmd(cmd[0]) and _start_proc(cmd, cmd[0]):
                return

        self.log(
            "VS Code Starter nicht gefunden. Setze VSCODE_STARTER oder installiere code/codium."
        )

    # ---------- Whoami / Egress info ----------
    def whoami_async(self, force=False):
        if not force and not self.auto_external_checks.get():
            return

        def job():
            rc, out = run_cmd(
                ["curl", "-s", "--max-time", "4", "https://1.1.1.1/cdn-cgi/trace"],
                timeout=6,
            )
            info = parse_cf_trace(out)

            def show():
                self.nb.select(self.tab_log)
                if rc != 0 or not info:
                    self.log("Whoami: fehlgeschlagen (curl/Netz?).")
                    if out:
                        self.log(out)
                    return
                keys = ["ip", "loc", "colo", "warp", "tls", "http"]
                nice = " | ".join([f"{k.upper()}={info.get(k,'-')}" for k in keys])
                self.log(f"Whoami (Egress): {nice}")

            self.after(0, show)

        threading.Thread(target=job, daemon=True).start()

    # ---------- Connections (PID/program -> remote) ----------
    def connections_async(self):
        if not has_cmd("ss"):
            self.log("Verbindungen: 'ss' fehlt (iproute2).")
            return

        def job():
            rc, out = run_cmd(["sudo", "-n", "ss", "-H", "-tunap"], timeout=8)
            if rc != 0 or not out.strip():
                rc, out = run_cmd(["ss", "-H", "-tunap"], timeout=8)

            conns = parse_ss_lines(out)

            def fill():
                self.nb.select(self.tab_conn)
                self._conn_rows_cache = conns

                for item in self.conn_table.get_children():
                    self.conn_table.delete(item)

                for c in conns:
                    self.conn_table.insert(
                        "",
                        "end",
                        values=(
                            c["proto"],
                            c["state"],
                            c["local"],
                            c["remote"],
                            c["pid"],
                            c["proc"],
                        ),
                    )

                self.log(
                    f"Verbindungen aktualisiert: {len(conns)} Zeilen (PID/Proc ggf. nur mit sudo sichtbar)."
                )

            self.after(0, fill)

        threading.Thread(target=job, daemon=True).start()

    def arp_async(self):
        def job():
            rc, out = run_cmd(["arp", "-a"], timeout=8)
            if rc != 0 or not (out or "").strip():
                rc, out = run_cmd(["ip", "neigh"], timeout=8)

            def show():
                self.nb.select(self.tab_log)
                self.log("---- ARP Tabelle ----")
                self.log(out if out else "Keine Daten.")

            self.after(0, show)

        threading.Thread(target=job, daemon=True).start()

    def route_async(self):
        def job():
            _, out = run_cmd(["ip", "route"], timeout=6)

            def show():
                self.nb.select(self.tab_log)
                self.log("---- Routing Tabelle ----")
                self.log(out if out else "Keine Daten.")

            self.after(0, show)

        threading.Thread(target=job, daemon=True).start()

    def apply_conn_filter(self):
        q = (self.conn_filter.get() or "").strip().lower()
        rows = self._conn_rows_cache

        for item in self.conn_table.get_children():
            self.conn_table.delete(item)

        if not q:
            for c in rows:
                self.conn_table.insert(
                    "",
                    "end",
                    values=(
                        c["proto"],
                        c["state"],
                        c["local"],
                        c["remote"],
                        c["pid"],
                        c["proc"],
                    ),
                )
            return

        for c in rows:
            blob = f"{c['proto']} {c['state']} {c['local']} {c['remote']} {c['pid']} {c['proc']}".lower()
            if q in blob:
                self.conn_table.insert(
                    "",
                    "end",
                    values=(
                        c["proto"],
                        c["state"],
                        c["local"],
                        c["remote"],
                        c["pid"],
                        c["proc"],
                    ),
                )

    def _format_connection_values(self, values) -> str:
        cols = [str(value or "").strip() for value in values]
        while len(cols) < 6:
            cols.append("")
        return (
            f"Proto: {cols[0]} | State: {cols[1]} | Local: {cols[2]} | "
            f"Remote: {cols[3]} | PID: {cols[4]} | Process: {cols[5]}"
        )

    def get_connections_analysis_text(self) -> str:
        table = self.__dict__.get("conn_table")
        if table is None:
            return ""

        selection = list(table.selection())
        target_items = selection or list(table.get_children())
        if not target_items:
            return ""

        lines = ["=== Verbindungen ===", ""]
        for item in target_items:
            try:
                values = table.item(item, "values")
            except Exception:
                values = ()
            lines.append(self._format_connection_values(values))
        return "\n".join(lines).strip()

    def copy_connections_to_clipboard(self):
        text = self.get_connections_analysis_text()
        if not text:
            try:
                self.chad_analysis_status_var.set(
                    self.tr(
                        "connections_status_nothing_to_copy",
                        "Keine Verbindungen zum Kopieren vorhanden.",
                    )
                )
            except Exception:
                pass
            return
        self.clipboard_clear()
        self.clipboard_append(text)
        try:
            self.chad_analysis_status_var.set(
                self.tr(
                    "connections_status_copied",
                    "Verbindungen in Zwischenablage kopiert.",
                )
            )
        except Exception:
            pass

    def paste_connections_filter_from_clipboard(self):
        try:
            pasted = self.clipboard_get()
        except Exception:
            pasted = ""
        if not pasted:
            return
        self.conn_filter.set(str(pasted).strip())
        entry = self.__dict__.get("conn_filter_entry")
        if entry is not None:
            try:
                entry.focus_force()
            except Exception:
                pass

    def select_all_connections(self):
        table = self.__dict__.get("conn_table")
        if table is None:
            return
        try:
            rows = list(table.get_children())
            if not rows:
                return
            table.selection_set(rows)
            table.focus(rows[0])
        except Exception:
            pass

    def analyze_connections_with_chad(self):
        text = self.get_connections_analysis_text()
        if not text:
            try:
                self.chad_analysis_status_var.set(
                    self.tr(
                        "connections_status_no_chad_data",
                        "Keine Verbindungen für Chad Analyse vorhanden.",
                    )
                )
            except Exception:
                pass
            return
        self.prepare_chad_analysis("Verbindungen Analyse", text)

    def copy_connections_as_chad_prompt(self):
        text = self.get_connections_analysis_text()
        if not text:
            try:
                self.chad_analysis_status_var.set(
                    self.tr(
                        "connections_status_no_prompt_selection",
                        "Keine Verbindungen zum Kopieren markiert.",
                    )
                )
            except Exception:
                pass
            return

        prompt = self._build_chad_analysis_prompt("Verbindungen Analyse", text)
        self.clipboard_clear()
        self.clipboard_append(prompt)
        try:
            self.chad_analysis_status_var.set(
                self.tr(
                    "chad_status_prompt_copied",
                    "Chad-Prompt in Zwischenablage kopiert.",
                )
            )
        except Exception:
            pass

    def attach_connections_context_menu(self):
        table = self.__dict__.get("conn_table")
        if table is None:
            return

        menu = tk.Menu(table, tearoff=0)
        menu.add_command(label="Kopieren", command=self.copy_connections_to_clipboard)
        menu.add_command(
            label="Einfügen",
            command=self.paste_connections_filter_from_clipboard,
        )
        menu.add_command(label="Alles auswählen", command=self.select_all_connections)
        menu.add_separator()
        menu.add_command(
            label="Auswahl an Chad Analyse senden",
            command=self.analyze_connections_with_chad,
        )
        menu.add_command(
            label="Auswahl als Chad-Prompt kopieren",
            command=self.copy_connections_as_chad_prompt,
        )

        def prepare_selection(row_id: str):
            if not row_id:
                return
            try:
                current_selection = tuple(table.selection())
            except Exception:
                current_selection = ()
            if row_id in current_selection:
                try:
                    table.focus(row_id)
                except Exception:
                    pass
                return
            try:
                table.selection_set(row_id)
                table.focus(row_id)
            except Exception:
                pass

        def show_menu(event):
            row_id = ""
            try:
                row_id = table.identify_row(event.y)
            except Exception:
                row_id = ""
            prepare_selection(row_id)
            return self._show_context_menu(menu, event)

        table.bind("<Button-3>", show_menu, add="+")
        table._netmon_context_menu = menu

    # ---------- DNS Status ----------
    def check_dnssec_status(self):
        if not has_cmd("resolvectl"):
            return "Unbekannt", "resolvectl nicht verfügbar"

        code, out = run_cmd(["resolvectl", "status"], timeout=5)
        if code != 0:
            return "Unbekannt", out.strip() or "resolvectl nicht verfügbar"

        dnssec_lines = [
            line.strip()
            for line in (out or "").splitlines()
            if "dnssec" in line.lower()
        ]
        dnssec_text = "\n".join(dnssec_lines).strip() or out
        dnssec_lower = dnssec_text.lower()

        if "dnssec setting: yes" in dnssec_lower or "dnssec=yes" in dnssec_lower:
            return "Aktiv", dnssec_text
        if "dnssec setting: allow-downgrade" in dnssec_lower:
            return "Eingeschränkt (allow-downgrade)", dnssec_text
        if "dnssec supported: yes" in dnssec_lower:
            return "Unterstützt, aber lokal nicht erzwungen", dnssec_text

        return "Nicht aktiv", dnssec_text

    def test_dnssec_validation(self):
        if has_cmd("resolvectl"):
            code, out = run_cmd(["resolvectl", "query", "dnssec-failed.org"], timeout=6)
            lower_out = (out or "").lower()
            if code != 0 and (
                "dnssec validation failed" in lower_out
                or "resolve call failed" in lower_out
                or "failed to resolve" in lower_out
                or "no appropriate name servers" in lower_out
            ):
                return (
                    True,
                    "DNSSEC funktioniert korrekt (Manipulation erkannt über systemd-resolved).",
                )
            if code == 0 and out.strip():
                return (
                    False,
                    "DNSSEC-Test fehlgeschlagen: dnssec-failed.org wurde trotzdem aufgelöst.",
                )

        if not has_cmd("dig"):
            return False, "DNSSEC-Test nicht möglich: dig ist nicht verfügbar."

        code, out = run_cmd(
            ["dig", "+dnssec", "+comments", "dnssec-failed.org"], timeout=6
        )
        lower_out = (out or "").lower()
        if code != 0 and not out.strip():
            return (
                False,
                "DNSSEC-Test fehlgeschlagen: dig konnte nicht ausgeführt werden.",
            )

        if "status: servfail" in lower_out:
            return True, "DNSSEC funktioniert korrekt (Manipulation erkannt)"

        if "status: noerror" in lower_out or "status: nxdomain" in lower_out:
            return (
                False,
                "DNSSEC-Test fehlgeschlagen: dnssec-failed.org wurde nicht geblockt.",
            )

        return False, "DNSSEC-Test uneindeutig. Resolver-Antwort bitte prüfen."

    def run_dnssec_check(self):
        self.nb.select(self.tab_dns)
        self.log(self.tr("dns_log_dnssec_started", "DNSSEC-Prüfung gestartet"))

        def job():
            status, raw = self.check_dnssec_status()
            valid, msg = self.test_dnssec_validation()
            overall = "OK" if valid else "FEHLER"
            status_line = status
            if valid and status == "Unbekannt":
                status_line = "Validierung aktiv, Resolverstatus lokal nicht abfragbar"
            elif valid and status == "Nicht aktiv":
                status_line = "Validierung aktiv, Resolverstatusmeldung bitte prüfen"

            result = (
                f"=== DNSSEC Prüfung ===\n"
                f"Gesamtbewertung: {overall}\n"
                f"DNSSEC Status: {status_line}\n"
                f"Validierung: {'OK' if valid else 'FEHLER'}\n\n"
                f"{msg}\n\n"
                f"Details:\n{raw}"
            )

            def show():
                self.nb.select(self.tab_dns)
                self.update_dns_output(result)
                self.log(
                    self.tr(
                        "dns_log_dnssec_finished",
                        "DNSSEC-Prüfung abgeschlossen",
                    )
                )

            self.after(0, show)

        threading.Thread(target=job, daemon=True).start()

    def analyze_dns_health_output(self, out: str) -> str:
        service = self._dns_control_service()
        if service is None:
            return "=== DNS Health Check ===\n\nDNS-Control-Service nicht verfügbar."
        return service.build_dns_health_report(
            out or "",
            dns_mode=self._current_dns_mode(),
        )

    def run_dns_health_check(self):
        self.nb.select(self.tab_dns)
        self.log(self.tr("dns_log_health_started", "DNS Health Check gestartet"))

        def job():
            code, out = run_cmd(["resolvectl", "status"], timeout=8)

            if code != 0:
                result = (
                    "=== DNS Health Check ===\n"
                    "Status: FEHLER\n\n"
                    "resolvectl status konnte nicht gelesen werden.\n\n"
                    f"Details:\n{out}"
                )
            else:
                result = self.analyze_dns_health_output(out)

            def show():
                self.nb.select(self.tab_dns)
                self.update_dns_output(result, append=True)
                self.log(
                    self.tr(
                        "dns_log_health_finished",
                        "DNS Health Check abgeschlossen",
                    )
                )

            self.after(0, show)

        threading.Thread(target=job, daemon=True).start()

    def get_public_ip(self):
        code, out = run_cmd(["curl", "-s", "https://api.ipify.org"], timeout=6)
        if code == 0:
            return (out or "").strip() or "unbekannt"
        return "unbekannt"

    def get_active_dns_servers(self):
        service = self._dns_control_service()
        if service is None:
            return []
        code, out = service.read_resolvectl_status(timeout=6)
        if code != 0:
            return []
        return service.parse_active_dns_servers(out)

    def _parse_active_dns_servers(self, resolvectl_output: str):
        service = self._dns_control_service()
        if service is None:
            return []
        return service.parse_active_dns_servers(resolvectl_output)

    def detect_vpn_active(self, resolvectl_output: str) -> bool:
        service = self._dns_control_service()
        if service is None:
            text = (resolvectl_output or "").lower()
            return any(
                token in text for token in ("proton", "wg", "wireguard", "tun", "vpn")
            )
        return service.detect_vpn_active(resolvectl_output)

    def _dns_leak_is_private_or_local_dns(self, dns_server: str) -> bool:
        service = self._dns_control_service()
        if service is None:
            value = (dns_server or "").strip().strip("[]")
            if not value:
                return False
            if "%" in value:
                value = value.split("%", 1)[0]
            try:
                addr = ipaddress.ip_address(value)
            except ValueError:
                return False
            return addr.is_private or addr.is_loopback or addr.is_link_local
        return service.dns_server_is_private_or_local(dns_server)

    def build_dns_leak_check_result(
        self,
        resolvectl_output: str,
        dns_servers: list[str],
        public_ip: str,
    ) -> str:
        service = self._dns_control_service()
        if service is None:
            return "=== DNS Leak Check ===\n\nDNS-Control-Service nicht verfügbar."
        return service.build_dns_leak_check_result(
            resolvectl_output=resolvectl_output,
            dns_servers=dns_servers,
            public_ip=public_ip,
        )

    def run_dns_leak_check(self, quiet: bool = False):
        if not quiet:
            self.nb.select(self.tab_dns)
            self.log(self.tr("dns_log_leak_started", "DNS Leak Check gestartet"))

        def job():
            code, out = run_cmd(["resolvectl", "status"], timeout=6)
            resolvectl_output = out or ""
            if code == 0:
                dns_servers = self._parse_active_dns_servers(resolvectl_output)
            else:
                dns_servers = []
                resolvectl_output = (
                    resolvectl_output.strip()
                    or "resolvectl status konnte nicht gelesen werden."
                )

            public_ip = self.get_public_ip()
            result = self.build_dns_leak_check_result(
                resolvectl_output=resolvectl_output,
                dns_servers=dns_servers,
                public_ip=public_ip,
            )
            self.last_dns_leak_result = result

            if code != 0:
                result = (
                    result
                    + "\n\nHinweis:\n"
                    + "- resolvectl status konnte nicht gelesen werden; VPN-Erkennung ist eingeschränkt.\n"
                    + f"\nDetails:\n{resolvectl_output}"
                )

            def show():
                if not quiet:
                    self.nb.select(self.tab_dns)
                    self.update_dns_output(result, append=True)
                    self.log(
                        self.tr(
                            "dns_log_leak_finished",
                            "DNS Leak Check abgeschlossen",
                        )
                    )

            if quiet:
                return
            self.after(0, show)

        threading.Thread(target=job, daemon=True).start()

    def find_vpn_dns_servers(self):
        service = self._dns_control_service()
        if service is None:
            return "", [], "DNS-Control-Service nicht verfügbar."
        code, out = service.read_resolvectl_status(timeout=8)
        if code != 0:
            return "", [], out
        vpn_iface, vpn_dns = service.find_vpn_dns_servers(out)
        return vpn_iface, vpn_dns, out

    def capture_dns_state(self):
        code, out = run_cmd(["resolvectl", "status"], timeout=8)
        return {
            "timestamp": datetime.now().isoformat(),
            "resolvectl_status": out if code == 0 else "",
        }

    def _show_dns_repair_message(self, title: str, text: str):
        self.update_dns_output(text, append=True)
        self.chad_analysis_status_var.set(title)

    def dns_leak_repair_async(self):
        self.nb.select(self.tab_dns)
        self.log(self.tr("dns_log_leak_repair_prepared", "DNS Leak Repair vorbereitet"))

        def find_job():
            vpn_iface, vpn_dns, raw = self.find_vpn_dns_servers()

            def ask_or_report():
                if not vpn_dns or not vpn_iface:
                    result = (
                        "=== DNS Leak Repair ===\n\n"
                        "Status: abgebrochen\n"
                        "Grund: Kein VPN-DNS gefunden.\n\n"
                        "Hinweis: Es wurden keine DNS-Server auf einem VPN-Link "
                        "(proton/wg/wireguard/tun/vpn) erkannt."
                    )
                    self._show_dns_repair_message("DNS Leak Repair", result)
                    return

                msg = self.tr(
                    "dns_leak_repair_confirm_message",
                    "DNS Leak Repair versucht, öffentliche DNS-Server zu entfernen "
                    "und VPN-DNS zu bevorzugen.\n\n"
                    "VPN-Interface: {vpn_iface}\n"
                    "Gefundene VPN-DNS:\n{vpn_dns}\n\n"
                    "Fortfahren?",
                ).format(vpn_iface=vpn_iface, vpn_dns=", ".join(vpn_dns))
                ok = messagebox.askyesno(
                    self.tr("dns_leak_repair_title", "DNS Leak Repair"),
                    msg,
                    parent=self,
                )
                if not ok:
                    self.chad_analysis_status_var.set(
                        self.tr(
                            "dns_leak_repair_status_cancelled",
                            "DNS Leak Repair abgebrochen.",
                        )
                    )
                    self.log(
                        self.tr(
                            "dns_leak_repair_log_cancelled",
                            "DNS Leak Repair abgebrochen",
                        )
                    )
                    return
                if not self.ensure_sudo_session("DNS Leak Repair"):
                    result = (
                        "=== DNS Leak Repair ===\n\n"
                        "Status: abgebrochen\n"
                        "Grund: sudo-Freigabe fehlt."
                    )
                    self._show_dns_repair_message("DNS Leak Repair", result)
                    return

                threading.Thread(
                    target=self.dns_leak_repair,
                    args=(vpn_iface, vpn_dns, raw),
                    daemon=True,
                ).start()

            self.after(0, ask_or_report)

        threading.Thread(target=find_job, daemon=True).start()

    def dns_leak_repair(self, vpn_iface: str, vpn_dns: list[str], raw: str = ""):
        self.last_dns_repair_state = self.capture_dns_state()

        code, route_out = run_cmd(["ip", "route", "show", "default"], timeout=5)
        main_iface = ""
        if code == 0:
            parts = route_out.split()
            if "dev" in parts:
                try:
                    main_iface = parts[parts.index("dev") + 1]
                except Exception:
                    main_iface = ""

        actions = []
        errors = []

        if main_iface and main_iface != vpn_iface:
            rc, out = run_cmd(
                ["sudo", "-n", "resolvectl", "revert", main_iface],
                timeout=10,
            )
            actions.append(f"resolvectl revert {main_iface}: rc={rc}")
            if rc != 0:
                errors.append(out or f"revert {main_iface} fehlgeschlagen")

        rc, out = run_cmd(
            ["sudo", "-n", "resolvectl", "dns", vpn_iface, *vpn_dns],
            timeout=10,
        )
        actions.append(f"resolvectl dns {vpn_iface} {' '.join(vpn_dns)}: rc={rc}")
        if rc != 0:
            errors.append(out or f"resolvectl dns {vpn_iface} fehlgeschlagen")

        rc, out = run_cmd(
            ["sudo", "-n", "systemctl", "restart", "systemd-resolved"],
            timeout=15,
        )
        actions.append(f"restart systemd-resolved: rc={rc}")
        if rc != 0:
            errors.append(out or "systemd-resolved restart fehlgeschlagen")

        result = (
            "=== DNS Leak Repair ===\n\n"
            f"VPN-Interface: {vpn_iface}\n"
            f"VPN-DNS: {', '.join(vpn_dns)}\n"
            f"Vorheriger DNS-Zustand gesichert: {'JA' if self.last_dns_repair_state else 'NEIN'}\n\n"
            "Aktionen:\n- " + "\n- ".join(actions)
        )

        if errors:
            result += "\n\nFehler:\n- " + "\n- ".join(str(err) for err in errors)

        def show():
            self._show_dns_repair_message("DNS Leak Repair", result)
            self.run_dns_leak_check()

        self.after(0, show)

    def dns_repair_rollback_async(self):
        self.nb.select(self.tab_dns)
        if not self.ensure_sudo_session("DNS Repair Rollback"):
            result = (
                "=== DNS Repair Rollback ===\n\n"
                "Status: abgebrochen\n"
                "Grund: sudo-Freigabe fehlt."
            )
            self._show_dns_repair_message("DNS Repair Rollback", result)
            return
        threading.Thread(target=self.dns_repair_rollback, daemon=True).start()

    def dns_repair_rollback(self):
        rc, out = run_cmd(
            ["sudo", "-n", "systemctl", "restart", "systemd-resolved"], timeout=15
        )

        captured = getattr(self, "last_dns_repair_state", None)
        captured_ts = (
            captured.get("timestamp", "nicht vorhanden")
            if captured
            else "nicht vorhanden"
        )

        result = (
            "=== DNS Repair Rollback ===\n\n"
            "Phase-1-Rollback: systemd-resolved wurde neu gestartet.\n"
            "Falls manuelle DNS-Werte gesetzt waren, bitte resolvectl status prüfen.\n\n"
            f"Gesicherter Zustand: {captured_ts}\n"
            f"rc={rc}\n{out}"
        )

        def show():
            self._show_dns_repair_message("DNS Repair Rollback", result)
            self.run_dns_leak_check()

        self.after(0, show)

    def fix_dot_leak_conflict_async(self):
        self.nb.select(self.tab_dns)

        conflict, content = self.detect_netmon_dot_dropin_conflict()
        if not content:
            result = (
                "=== DoT Leak-Konflikt beheben ===\n\n"
                "Status: nicht erforderlich\n"
                "Grund: Keine NetMon DoT-Drop-in-Datei gefunden."
            )
            self._show_dns_repair_message("DoT Leak-Konflikt", result)
            return

        if not conflict:
            result = (
                "=== DoT Leak-Konflikt beheben ===\n\n"
                "Status: nicht erforderlich\n"
                "Grund: NetMon DoT-Drop-in vorhanden, aber keine globale "
                "DNS-Erzwingung erkannt."
            )
            self._show_dns_repair_message("DoT Leak-Konflikt", result)
            return

        ok = messagebox.askyesno(
            self.tr("dot_leak_fix_confirm_title", "DoT Leak-Konflikt beheben"),
            self.tr(
                "dot_leak_fix_confirm_message",
                "Die NetMon DoT-Drop-in-Datei erzwingt globale DNS-Server.\n\n"
                "Sie wird entfernt, damit VPN-DNS nicht überschrieben wird.\n\n"
                "Fortfahren?",
            ),
            parent=self,
        )
        if not ok:
            self.chad_analysis_status_var.set(
                self.tr(
                    "dot_leak_fix_status_cancelled",
                    "DoT Leak-Konflikt Fix abgebrochen.",
                )
            )
            self.log(
                self.tr(
                    "dot_leak_fix_log_cancelled",
                    "DoT Leak-Konflikt Fix abgebrochen",
                )
            )
            return

        if not self.ensure_sudo_session("DoT Leak-Konflikt beheben"):
            result = (
                "=== DoT Leak-Konflikt beheben ===\n\n"
                "Status: abgebrochen\n"
                "Grund: sudo-Freigabe fehlt."
            )
            self._show_dns_repair_message("DoT Leak-Konflikt", result)
            return

        threading.Thread(target=self.fix_dot_leak_conflict, daemon=True).start()

    def fix_dot_leak_conflict(self):
        rc_rm, out_rm = run_cmd(["sudo", "-n", "rm", DOT_DROPIN_PATH], timeout=10)
        rc_restart, out_restart = run_cmd(
            ["sudo", "-n", "systemctl", "restart", "systemd-resolved"],
            timeout=15,
        )

        result = (
            "=== DoT Leak-Konflikt beheben ===\n\n"
            f"Datei: {DOT_DROPIN_PATH}\n"
            f"Entfernen rc={rc_rm}\n{out_rm}\n\n"
            f"Restart systemd-resolved rc={rc_restart}\n{out_restart}\n\n"
            "Bitte VPN neu verbinden und DNS Leak Check erneut ausführen."
        )

        if rc_rm == 0 and rc_restart == 0:
            result += "\n\nStatus: NetMon DoT-Drop-in entfernt."
        elif rc_rm != 0:
            result += "\n\nStatus: Entfernen fehlgeschlagen."
        else:
            result += "\n\nStatus: Datei entfernt, Restart prüfen."

        def show():
            self._show_dns_repair_message("DoT Leak-Konflikt behoben", result)
            self.show_dot_status_async()

        self.after(0, show)

    def resolved_status_async(self):
        self.nb.select(self.tab_dns)
        self.log("DNS-Status aktualisiert")

        def job():
            lines = ["=== DNS Status (resolv*) ==="]

            if has_cmd("resolvectl"):
                rc, out = run_cmd(
                    ["systemctl", "is-active", "systemd-resolved"], timeout=4
                )
                lines.append(
                    f"systemd-resolved: {out.strip() or 'unknown'} (rc={rc})\n"
                )

                _, out = run_cmd(["resolvectl", "dns"], timeout=6)
                lines.append("---- resolvectl dns ----\n" + (out or "") + "\n")

                _, out = run_cmd(["resolvectl", "domain"], timeout=6)
                lines.append("---- resolvectl domain ----\n" + (out or "") + "\n")

                _, out = run_cmd(["resolvectl", "status"], timeout=10)
                lines.append("---- resolvectl status ----\n" + (out or "") + "\n")
            else:
                lines.append(
                    "Hinweis: resolvectl nicht vorhanden → systemd-resolved ist nicht aktiv/benutzt.\n"
                )

            text = "\n".join(lines)

            def show():
                self.nb.select(self.tab_dns)
                self.update_dns_output(text)

            self.after(0, show)

        threading.Thread(target=job, daemon=True).start()

    # ---------- External IP/Country ----------
    def ip_country_async(self, force=False):
        if self.ip_country_running:
            return
        if not force and (time.time() - self.last_ip_ts) < 60:
            return
        self.ip_country_running = True
        auto_external_checks = bool(self.auto_external_checks.get())

        def job():
            try:
                if not force and not auto_external_checks:
                    self.last_ip_ts = time.time()
                    return

                _, ip = run_cmd(
                    ["curl", "-s", "--max-time", "3", "ifconfig.me"], timeout=5
                )
                ip = (ip or "").strip()
                c, _reason = self._fetch_country_from_ipinfo(timeout_s=3)

                if ip:
                    self.public_ip = ip
                self.country = c if c else "UNKNOWN"

                self.last_ip_ts = time.time()
                self.after(
                    0, lambda: self.log(f"IP: {self.public_ip} | Land: {self.country}")
                )
            finally:
                self.ip_country_running = False

        threading.Thread(target=job, daemon=True).start()

    def _expected_country_from_profile(self, profile):
        txt = (profile or "").strip().lower()
        if not txt:
            return None

        # Normalize separators to spaces and split into tokens.
        txt = txt.replace("_", " ").replace("-", " ")
        tokens = [t for t in re.split(r"\s+", txt) if t]
        if not tokens:
            return None

        # Direkt bekannte Profilnamen mit Hint
        if "proton-de" in profile.lower() or "proton de" in txt or "deutschland" in txt:
            return "DE"
        if "proton" in profile.lower() and "de" in tokens:
            return "DE"

        # ISO country as single token.
        for token in tokens:
            if token in ("nl", "ie", "de", "ch", "us"):
                return token.upper()

        # Länderbezeichnungen
        if any(tok in tokens for tok in ("niederlande", "netherlands")):
            return "NL"
        if any(tok in tokens for tok in ("irland", "ireland")):
            return "IE"
        if any(tok in tokens for tok in ("deutschland", "germany")):
            return "DE"
        if any(tok in tokens for tok in ("schweiz", "switzerland")):
            return "CH"
        if any(tok in tokens for tok in ("usa", "united", "states", "us")):
            # us token handled above; fallback.
            return "US"

        # Produktspezifische Zuordnung
        if "vpn_ch" in txt or "proton" in txt:
            return "CH"

        return None

    def _is_valid_country_code(self, value):
        value = (value or "").strip().upper()
        return bool(re.fullmatch(r"[A-Z]{2}", value))

    def _parse_ipinfo_country_response(self, raw_body, http_code):
        text = (raw_body or "").strip()
        if http_code == 429:
            return None, "rate_limit"
        if not text:
            return None, "empty"

        if text.startswith("{"):
            try:
                data = json.loads(text)
            except Exception:
                return None, "invalid_json"
            err_parts = [
                str(data.get("error") or "").strip(),
                str(data.get("message") or "").strip(),
                str(data.get("status") or "").strip(),
            ]
            err_text = " ".join(part for part in err_parts if part).lower()
            if "rate" in err_text and "limit" in err_text:
                return None, "rate_limit"
            return None, "api_error"

        candidate = text.splitlines()[0].strip().upper()
        if self._is_valid_country_code(candidate):
            return candidate, ""
        return None, "invalid_country"

    def _fetch_country_from_ipinfo(self, timeout_s=5):
        rc_ct, out_ct = run_cmd(
            [
                "curl",
                "-sS",
                "--max-time",
                str(timeout_s),
                "-w",
                "\n%{http_code}",
                "ipinfo.io/country",
            ],
            timeout=timeout_s + 3,
        )
        if rc_ct != 0:
            return None, "request_failed"

        payload = (out_ct or "").rstrip("\n")
        if "\n" in payload:
            body, last_line = payload.rsplit("\n", 1)
        else:
            body, last_line = payload, ""

        try:
            http_code = int((last_line or "").strip())
        except Exception:
            http_code = 0

        country, reason = self._parse_ipinfo_country_response(body, http_code)
        return country, reason

    def _fetch_public_ip_country(self):
        ip = getattr(self, "public_ip", "") or ""
        country = getattr(self, "country", "") or ""
        country_reason = ""

        rc_ip, out_ip = run_cmd(
            [
                "curl",
                "-s",
                "--max-time",
                "5",
                "ifconfig.me",
            ]
        )
        if rc_ip == 0:
            ip = (out_ip or "").strip()

        fetched_country, country_reason = self._fetch_country_from_ipinfo(timeout_s=5)
        if fetched_country:
            country = fetched_country
        else:
            country = "UNKNOWN"

        if ip:
            self.public_ip = ip
        self.country = country

        return ip, country, country_reason

    def verify_vpn_exit_country(self, base_profile):
        base_profile = self._base_wg_profile_name(base_profile)
        expected = self._expected_country_from_profile(base_profile)
        if not expected:
            self.log_safe(
                f"VPN Verifikation: Profil '{base_profile}' konnte nicht heuristisch einem Land zugeordnet werden."
            )
            return None

        self.log_safe(
            f"VPN Verifikation: Profil '{base_profile}' erwartet Land '{expected}'."
        )

        ip, country, country_reason = self._fetch_public_ip_country()

        if not ip and not country:
            self.log_safe(
                "VPN Verifikation: Öffentliche IP/Land konnte nicht ermittelt werden."
            )
            return False

        self.log_safe(
            f"VPN Verifikation: Profil '{base_profile}' aktiv, öffentliche IP {ip or '–'}, Land {country or '–'}."
        )

        normalized_country = (country or "").strip().upper()
        if not self._is_valid_country_code(normalized_country):
            if country_reason == "rate_limit":
                self.log_safe("Geo-IP-Verifikation fehlgeschlagen (Rate Limit).")
            else:
                self.log_safe(
                    "Geo-IP-Verifikation fehlgeschlagen (kein gültiger Ländercode)."
                )
            return None

        if normalized_country != expected.strip().upper():
            self.log_safe(
                f"Warnung: Profilname '{base_profile}' passt nicht zum realen Exit-Land '{normalized_country}'."
            )
            return False

        self.log_safe("VPN Verifikation: Exit-Land stimmt mit Profil überein.")
        return True

    def _check_wireguard_profile_for_multiple_endpoints(self, profile, path=None):
        profile = self._base_wg_profile_name(profile)
        if not path:
            path = self._wg_profile_path(profile)

        try:
            if not os.path.isfile(path):
                return
            with open(path, "r", encoding="utf-8") as f:
                lines = f.readlines()
        except Exception as e:
            self.log_safe(
                f"Warnung: Profil '{profile}' konnte nicht geprüft werden: {e}"
            )
            return

        in_peer = False
        endpoint_count = 0
        for line in lines:
            s = line.strip()
            if not s or s.startswith("#") or s.startswith(";"):
                continue
            if s.lower() == "[peer]":
                in_peer = True
                continue
            if s.startswith("[") and s.endswith("]") and s.lower() != "[peer]":
                in_peer = False
                continue
            if in_peer and re.match(r"^endpoint\s*=", s, re.IGNORECASE):
                endpoint_count += 1

        if endpoint_count > 1:
            self.log_safe(
                f"Warnung: Profil '{profile}' enthält mehrere aktive Endpoint-Zeilen (count={endpoint_count}). Profil ist möglicherweise fehlerhaft oder mehrdeutig."
            )

    # ---------- Refresh ----------
    def _update_dns_consistency_light(self):
        if "l_dnscons" not in self.__dict__:
            return
        self.update_dns_guard_indicator()

    def _update_dns_consistency_light_safe(self):
        if threading.current_thread() is threading.main_thread():
            self.update_dns_guard_indicator()
            return
        self.after(0, self.update_dns_guard_indicator)

    def refresh(self):
        inet = self.check_internet()
        inet_state = self._update_internet_indicator_state(bool(inet))
        vpn_profile = self.active_wg_profile()

        fw = self.check_ufw_active()
        ssh_rule = self.check_ssh_rule_active() if fw else None
        ks = self.killswitch_is_on()

        self.update_ipv6_controls_state()
        if hasattr(self, "vpn_profile_combo"):
            self.vpn_profile_combo["values"] = self.wg_profiles
        self.update_primary_vpn_form_controls()
        self.sync_primary_vpn_form_selection()
        self.sync_vpn_profile_form_selection()
        self.update_dns_header()
        self.update_wg_rates()

        if inet_state == "ok":
            self.set_light(self.l_inet, "green")
        elif inet_state == "warning":
            self.set_light(self.l_inet, "yellow")
        elif inet_state == "transition":
            self.set_light(self.l_inet, "yellow")
        else:
            self.set_light(self.l_inet, "red")
        self.set_light(self.l_vpn, "green" if vpn_profile else "gray")
        self.set_light(self.l_fw, "green" if fw else "red")
        if ssh_rule is True:
            self.set_light(self.l_ssh, "green")
        elif ssh_rule is False:
            self.set_light(self.l_ssh, "red")
        else:
            self.set_light(self.l_ssh, "gray")
        self.set_light(self.l_ks, "green" if ks else "red")
        active_profile_for_ipv6 = (vpn_profile or "").strip()
        active_iface_for_ipv6 = (self.get_wg_iface() or "").strip()
        if not active_profile_for_ipv6:
            ipv6_light = "gray"
        else:
            ipv6_light = (
                "green"
                if self.ipv6_runtime_available_on_iface(active_iface_for_ipv6)
                else "red"
            )
        self.set_light(self.l_ipv6, ipv6_light)

        self._update_dns_consistency_light()
        self._update_dot_light()

        self.update_vpn_slot_buttons()
        self.refresh_vpn_slots_table()

        self.ip_country_async(force=False)
        self.whoami_async(force=False)

        wg_if = self.get_wg_iface()
        default_if = self.get_default_iface() or "–"
        route_txt = f"über {wg_if}" if wg_if else f"über {default_if}"
        if not vpn_profile:
            v6txt = "IPv6: n/a"
        else:
            wg_v6_runtime = self.ipv6_runtime_available_on_iface(wg_if)
            if wg_v6_runtime:
                v6txt = "IPv6: OK"
            elif self.profile_supports_ipv6(vpn_profile) and self.ipv6_is_disabled():
                v6txt = "IPv6: Profil erwartet IPv6 (AUS)"
            else:
                v6txt = "IPv6: Profil nur IPv4"

        parts = [
            (
                "Internet OK"
                if inet_state == "ok"
                else (
                    "Internet TRANSITION"
                    if inet_state == "transition"
                    else "Internet WARN" if inet_state == "warning" else "Internet DOWN"
                )
            ),
            f"VPN: {vpn_profile}" if vpn_profile else "VPN: aus",
            "UFW: aktiv" if fw else "UFW: aus",
            f"Killswitch: {'AN' if ks else 'AUS'}",
            v6txt,
            f"Traffic: {route_txt}",
            f"RX {human_rate(self.wg_rx_rate)} / TX {human_rate(self.wg_tx_rate)}",
            f"IP {self.public_ip} ({self.country})",
            f"Stealth: {'AN' if self.isp_stealth_enabled.get() else 'AUS'}",
            f"NetMon v{NETMON_APP_VERSION}",
        ]
        self.status_text.config(text=" | ".join(parts))

    def auto_refresh(self):
        self.refresh()
        self.hdd_refresh_async()
        self.after(3000, self.auto_refresh)


if __name__ == "__main__":
    import traceback as _bootstrap_traceback

    def get_startup_error_log_path():
        base = os.environ.get("XDG_STATE_HOME", "").strip()
        if not base:
            base = os.path.expanduser("~/.local/state")

        state_dir = os.path.join(base, "netmon")
        os.makedirs(state_dir, exist_ok=True)

        return os.path.join(state_dir, "startup_error.log")

    def _report_bootstrap_failure(exc: BaseException):
        log_path = get_startup_error_log_path()
        try:
            with open(log_path, "a", encoding="utf-8") as fh:
                fh.write("\n" + "=" * 80 + "\n")
                fh.write("NetMon app.py bootstrap failure\n\n")
                _bootstrap_traceback.print_exc(file=fh)
                fh.write("=" * 80 + "\n")
        except Exception:
            pass

        try:
            import tkinter as _tk
            from tkinter import messagebox as _messagebox

            root = _tk.Tk()
            root.withdraw()
            _messagebox.showerror(
                "NetMon Startfehler",
                "NetMon konnte nicht gestartet werden.\n\n"
                f"Fehler: {exc}\n\n"
                f"Details: {log_path}",
                parent=root,
            )
            root.destroy()
        except Exception:
            pass

    try:
        try:
            from netmon.main import main
        except ModuleNotFoundError:
            from main import main

        sys.exit(main())
    except SystemExit:
        raise
    except BaseException as exc:
        _report_bootstrap_failure(exc)
        raise
